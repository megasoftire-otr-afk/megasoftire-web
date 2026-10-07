# MegaSoftire Web 2026 - corrección INSP + fechas dd/mm/aaaa (31/08/2026)
import datetime as dt
import os
import hashlib
import base64
import textwrap
import io
import flet as ft
from database import init_db, query, execute, authenticate

EVENTS = {
    'INST':'Instalación','DINS':'Desinstalación','REPA':'Reparación',
    'INSP':'Inspección','INSC':'Inspección de cierre','INVE':'Inversión',
    'ROT':'Rotación','BAJA':'Desechar / baja'
}

MODULES = [
    ('Movimiento de neumáticos', ft.Icons.SWAP_HORIZ),
    ('Neumáticos en servicio', ft.Icons.DIRECTIONS_CAR),
    ('Programa de mantenimiento', ft.Icons.BUILD_CIRCLE_OUTLINED),
    ('Neumáticos en Stand By', ft.Icons.INVENTORY_2_OUTLINED),
    ('Neumáticos de Baja', ft.Icons.DELETE_FOREVER_OUTLINED),
    ('Análisis de operación', ft.Icons.ANALYTICS_OUTLINED),
    ('Tablas y reportes', ft.Icons.ASSESSMENT_OUTLINED),
    ('Administración', ft.Icons.ADMIN_PANEL_SETTINGS_OUTLINED),
]

BG = '#E9EEF4'
NAV_BG = '#102A43'
NAV_ACCENT = '#1E5AA8'
CARD_BG = '#FFFFFF'
TEXT_MAIN = '#1B263B'
TEXT_MUTED = '#66788A'

# ============================================================
# ETAPA 1 - IDENTIFICACIÓN DE UNIDAD DE RENDIMIENTO
# El EQUIPO define si la operación se controla por HORAS o KM.
# No modifica todavía fórmulas, KPI ni reportes.
# ============================================================
PERFORMANCE_HOURS = 'HORAS'
PERFORMANCE_KILOMETERS = 'KILÓMETROS'


def normalize_performance_type(value):
    raw = str(value or '').strip().upper()
    if raw in ('HORA', 'HORAS', 'HR', 'H'):
        return PERFORMANCE_HOURS
    if raw in ('KM', 'KMS', 'KILOMETRO', 'KILOMETROS', 'KILÓMETRO', 'KILÓMETROS'):
        return PERFORMANCE_KILOMETERS
    return None


def get_equipment_performance_type(equipment_id, operation_id=None):
    """Obtiene HORAS/KILÓMETROS del equipo sin romper la aplicación si aún
    no existe el dato en PostgreSQL."""
    if equipment_id in (None, '', 0, '0'):
        return None
    try:
        if operation_id is None:
            rows = query(
                'SELECT performance_type FROM equipment WHERE id=?',
                (int(equipment_id),)
            )
        else:
            rows = query(
                'SELECT performance_type FROM equipment '
                'WHERE id=? AND operation_id=?',
                (int(equipment_id), int(operation_id))
            )
        if not rows:
            return None
        return normalize_performance_type(rows[0]['performance_type'])
    except Exception:
        return None


def performance_measure_label(performance_type):
    if performance_type == PERFORMANCE_HOURS:
        return 'Horómetro'
    if performance_type == PERFORMANCE_KILOMETERS:
        return 'Odómetro'
    return 'Medidor'


def performance_unit_label(performance_type):
    if performance_type == PERFORMANCE_HOURS:
        return 'HORAS'
    if performance_type == PERFORMANCE_KILOMETERS:
        return 'KILÓMETROS'
    return 'NO DEFINIDO'

def performance_fields_state(performance_type):
    """Retorna qué medidor debe utilizar el equipo."""
    if performance_type == PERFORMANCE_HOURS:
        return 'HOROMETRO'
    if performance_type == PERFORMANCE_KILOMETERS:
        return 'ODOMETRO'
    return None


# V05 MULTITALLER - capa de integración sobre la Web original aprobada.
# Esta versión se usa SOLO en megasoftire-multitaller-prueba.
ACTIVE_OPERATION_ID = 1  # operación inicial; la selección de sesión se gestiona en main()



MASTER_TIRE_UPDATES_20260902 = [
    ('1345', '08251Y10267', '01/01/2026', 3850.0, 'Goodyear', '18.00-25', 'SMO-5D', 'L-5S', 'Tire SOL', 100.0, 84.0, 84.0, 15.0, 2300.0, 'Convencional', 'Nueva'),
    ('1350', '08251Y10009', '01/01/2026', 3850.0, 'Goodyear', '18.00-25', 'SMO-5D', 'L-5S', 'Tire SOL', 100.0, 84.0, 84.0, 15.0, 2300.0, 'Convencional', 'Nueva'),
    ('1351', '08251Y10104', '01/01/2026', 3850.0, 'Goodyear', '18.00-25', 'SMO-5D', 'L-5S', 'Tire SOL', 100.0, 84.0, 84.0, 15.0, 2300.0, 'Convencional', 'Nueva'),
    ('1352', '08251Y10759', '01/01/2026', 3850.0, 'Goodyear', '18.00-25', 'SMO-5D', 'L-5S', 'Tire SOL', 100.0, 84.0, 84.0, 15.0, 2300.0, 'Convencional', 'Nueva'),
    ('1376', '12251Y10060', '01/01/2026', 3850.0, 'Goodyear', '18.00-25', 'SMO-5D', 'L-5S', 'Tire SOL', 100.0, 84.0, 84.0, 15.0, 2300.0, 'Convencional', 'Nueva'),
    ('1362', '12251Y10747', '01/01/2026', 3850.0, 'Goodyear', '18.00-25', 'SMO-5D', 'L-5S', 'Tire SOL', 100.0, 84.0, 84.0, 15.0, 2300.0, 'Convencional', 'Nueva'),
    ('1378', '12251Y10332', '01/01/2026', 3850.0, 'Goodyear', '18.00-25', 'SMO-5D', 'L-5S', 'Tire SOL', 100.0, 84.0, 84.0, 15.0, 2300.0, 'Convencional', 'Nueva'),
    ('1367', '12251Y10070', '01/01/2026', 3850.0, 'Goodyear', '18.00-25', 'SMO-5D', 'L-5S', 'Tire SOL', 100.0, 84.0, 84.0, 15.0, 2300.0, 'Convencional', 'Nueva'),
    ('1363', 'HC3MVC493', '01/01/2026', 7800.0, 'Yokohama', '29.5-29', 'Y524', 'L-5', 'Tire SOL', 100.0, 104.0, 104.0, 15.0, 2800.0, 'Convencional', 'Nueva'),
    ('1364', 'HC7MVC174', '01/01/2026', 7800.0, 'Yokohama', '29.5-29', 'Y524', 'L-5', 'Tire SOL', 100.0, 104.0, 104.0, 15.0, 2800.0, 'Convencional', 'Nueva'),
    ('1365', 'HC4MVC666', '01/01/2026', 7800.0, 'Yokohama', '29.5-29', 'Y524', 'L-5', 'Tire SOL', 100.0, 104.0, 104.0, 15.0, 2800.0, 'Convencional', 'Nueva'),
    ('1366', 'HC7MVC176', '01/01/2026', 7800.0, 'Yokohama', '29.5-29', 'Y524', 'L-5', 'Tire SOL', 100.0, 104.0, 104.0, 15.0, 2800.0, 'Convencional', 'Nueva'),
    ('1374', 'AE1HVC1933', '01/01/2026', 7800.0, 'Yokohama', '29.5-29', 'Y524', 'L-5', 'Tire SOL', 100.0, 104.0, 104.0, 15.0, 2800.0, 'Convencional', 'Nueva'),
    ('1375', 'HC6MVC992', '01/01/2026', 7800.0, 'Yokohama', '29.5-29', 'Y524', 'L-5', 'Tire SOL', 100.0, 104.0, 104.0, 15.0, 2800.0, 'Convencional', 'Nueva'),
    ('1354', 'XY2AVC992', '01/01/2026', 7800.0, 'Yokohama', '29.5-29', 'Y524', 'L-5', 'Tire SOL', 100.0, 104.0, 104.0, 15.0, 2800.0, 'Convencional', 'Nueva'),
    ('1377', 'HC3MVC631', '01/01/2026', 7800.0, 'Yokohama', '29.5-29', 'Y524', 'L-5', 'Tire SOL', 100.0, 104.0, 104.0, 15.0, 2800.0, 'Convencional', 'Nueva'),
    ('1346', 'XY3AVC183', '01/01/2026', 7800.0, 'Yokohama', '29.5-29', 'Y524', 'L-5', 'Tire SOL', 100.0, 104.0, 104.0, 15.0, 2800.0, 'Convencional', 'Nueva'),
    ('1357', 'AE4HVC653', '01/01/2026', 7800.0, 'Yokohama', '29.5-29', 'Y524', 'L-5', 'Tire SOL', 100.0, 104.0, 104.0, 15.0, 2800.0, 'Convencional', 'Nueva'),
    ('1320', 'AE3HVC510', '01/01/2026', 7800.0, 'Yokohama', '29.5-29', 'Y524', 'L-5', 'Tire SOL', 100.0, 104.0, 104.0, 15.0, 2800.0, 'Convencional', 'Nueva'),
    ('1321', 'AE4HVC655', '01/01/2026', 7800.0, 'Yokohama', '29.5-29', 'Y524', 'L-5', 'Tire SOL', 100.0, 104.0, 104.0, 15.0, 2800.0, 'Convencional', 'Nueva'),
    ('1359', 'AEIHVC194', '01/01/2026', 7800.0, 'Yokohama', '29.5-29', 'Y524', 'L-5', 'Tire SOL', 100.0, 104.0, 104.0, 15.0, 2800.0, 'Convencional', 'Nueva'),
    ('1355', 'AU7MVC799', '01/01/2026', 7800.0, 'Yokohama', '29.5-29', 'Y524', 'L-5', 'Tire SOL', 100.0, 104.0, 104.0, 15.0, 2800.0, 'Convencional', 'Nueva'),
    ('1316', 'AH6GVC949', '01/01/2026', 7800.0, 'Yokohama', '29.5-29', 'Y524', 'L-5', 'Tire SOL', 100.0, 104.0, 104.0, 15.0, 2800.0, 'Convencional', 'Nueva'),
    ('1317', 'XH2GVC017', '01/01/2026', 7800.0, 'Yokohama', '29.5-29', 'Y524', 'L-5', 'Tire SOL', 100.0, 104.0, 104.0, 15.0, 2800.0, 'Convencional', 'Nueva'),
    ('1370', 'XYZAVC991', '01/01/2026', 7800.0, 'Yokohama', '29.5-29', 'Y524', 'L-5', 'Tire SOL', 100.0, 104.0, 104.0, 15.0, 2800.0, 'Convencional', 'Nueva'),
    ('1349', 'XY1AVC833', '01/01/2026', 7800.0, 'Yokohama', '29.5-29', 'Y524', 'L-5', 'Tire SOL', 100.0, 104.0, 104.0, 15.0, 2800.0, 'Convencional', 'Nueva'),
    ('1329', 'AE1HVC196', '01/01/2026', 7800.0, 'Yokohama', '29.5-29', 'Y524', 'L-5', 'Tire SOL', 100.0, 104.0, 104.0, 15.0, 2800.0, 'Convencional', 'Nueva'),
    ('1344', 'XY3AVC184', '01/01/2026', 7800.0, 'Yokohama', '29.5-29', 'Y524', 'L-5', 'Tire SOL', 100.0, 104.0, 104.0, 15.0, 2800.0, 'Convencional', 'Nueva'),
    ('1353', 'XY1AVC835', '01/01/2026', 7800.0, 'Yokohama', '29.5-29', 'Y524', 'L-5', 'Tire SOL', 100.0, 104.0, 104.0, 15.0, 2800.0, 'Convencional', 'Nueva'),
    ('1358', 'AEOHVC048', '01/01/2026', 7800.0, 'Yokohama', '29.5-29', 'Y524', 'L-5', 'Tire SOL', 100.0, 104.0, 104.0, 15.0, 2800.0, 'Convencional', 'Nueva'),
    ('1327', 'HE9HVC270', '01/01/2026', 7800.0, 'Yokohama', '29.5-29', 'Y524', 'L-5', 'Tire SOL', 100.0, 104.0, 104.0, 15.0, 2800.0, 'Convencional', 'Nueva'),
    ('1328', 'AE1HVC195', '01/01/2026', 7800.0, 'Yokohama', '29.5-29', 'Y524', 'L-5', 'Tire SOL', 100.0, 104.0, 104.0, 15.0, 2800.0, 'Convencional', 'Nueva'),
    ('1371', 'AE2HVC375', '01/01/2026', 7800.0, 'Yokohama', '29.5-29', 'Y524', 'L-5', 'Tire SOL', 100.0, 104.0, 104.0, 15.0, 2800.0, 'Convencional', 'Nueva'),
    ('1372', 'CH9MVC663', '01/01/2026', 7800.0, 'Yokohama', '29.5-29', 'Y524', 'L-5', 'Tire SOL', 100.0, 104.0, 104.0, 15.0, 2800.0, 'Convencional', 'Nueva'),
    ('1324', 'AE2HVC373', '01/01/2026', 7800.0, 'Yokohama', '29.5-29', 'Y524', 'L-5', 'Tire SOL', 100.0, 104.0, 104.0, 15.0, 2800.0, 'Convencional', 'Nueva'),
    ('1373', 'HC5MYC850', '01/01/2026', 7800.0, 'Yokohama', '29.5-29', 'Y524', 'L-5', 'Tire SOL', 100.0, 104.0, 104.0, 15.0, 2800.0, 'Convencional', 'Nueva'),
]

NFU_BAJAS_20260908 = [('1308', '10241Y10617', '2025-10-02', 3850.0, 'GY', '18.00-25', 'SMO 5D', 'L-5S', 'SOLTRAK', 95.0, 84.0, 103.0, 15.0, 'Convencional', 'Nueva', '2026-05-04', 'CAT 47', 'P3', 4060.0, 22.0, 24.0, 'DR', '', 'DESGASTE REGULAR', 2502.0), ('1309', '08251Y10369', '2025-10-09', 3850.0, 'GY', '18.00-25', 'SMO 5D', 'L-5S', 'SOLTRAK', 94.0, 84.0, 103.0, 15.0, 'Convencional', 'Nueva', '2026-03-01', 'CAT 47', 'P4', 4060.0, 34.0, 36.0, 'CPB', '', 'CORTE PASANTE', 2534.0), ('1297', 'AH8GVC327', '2025-09-06', 7800.0, 'YK', '29.5-29', 'Y-524', 'L-5', 'TIRE SOL', 94.0, 103.0, 103.0, 15.0, 'Convencional', 'Nueva', '2026-01-06', 'CAT 43', 'P1', 14776.0, 22.0, 26.0, 'DR', '', 'DESGASTE REGULAR', 12915.0), ('1298', 'XS6YVC157', '2025-09-06', 7800.0, 'YK', '29.5-29', 'Y-524', 'L-5', 'TIRE SOL', 94.0, 103.0, 103.0, 15.0, 'Convencional', 'Nueva', '2026-02-21', 'CAT 43', 'P2', 15124.0, 25.0, 28.0, 'DR', '', 'DESGASTE REGULAR', 12915.0), ('1322', 'AE2HVC374', '2025-12-12', 7800.0, 'YK', '29.5-29', 'Y-524', 'L-5', 'TIRE SOL', 94.0, 103.0, 103.0, 15.0, 'Convencional', 'Nueva', '2026-03-01', 'CAT 44', 'P1', 15166.0, 80.0, 84.0, 'CPB', '', 'CORTE PASANTE', 13602.0), ('1331', 'AE4HVC656', '2026-01-17', 7800.0, 'YK', '29.5-29', 'Y-524', 'L-5', 'TIRE SOL', 94.0, 103.0, 103.0, 15.0, 'Convencional', 'Nueva', '2026-03-01', 'CAT 44', 'P2', 15210.0, 91.0, 92.0, 'CPB', '', 'CORTE PASANTE', 14087.0), ('1332', 'AE2HVC372', '2026-01-30', 7800.0, 'YK', '29.5-29', 'Y-524', 'L-5', 'TIRE SOL', 94.0, 103.0, 103.0, 15.0, 'Convencional', 'Nueva', '2026-04-30', 'CAT 51', 'P1', 3707.0, 29.0, 35.0, 'CPB', '', 'CORTE PASANTE', 2466.0), ('1333', 'AE3HVC508', '2026-01-30', 7800.0, 'YK', '29.5-29', 'Y-524', 'L-5', 'TIRE SOL', 94.0, 103.0, 103.0, 15.0, 'Convencional', 'Nueva', '2026-04-30', 'CAT 51', 'P2', 3707.0, 24.0, 27.0, 'CPB', '', 'CORTE PASANTE', 2466.0), ('1425', 'XY0AVC151', '2025-12-26', 7800.0, 'YK', '29.5-29', 'Y-524', 'L-5', 'TIRE SOL', 94.0, 103.0, 103.0, 15.0, 'Convencional', 'Nueva', '2026-05-03', 'CAT 51', 'P4', 3759.0, 88.0, 90.0, 'CPB', '', 'CORTE PASANTE', 2029.0)]


def main(page: ft.Page):
    init_db()

    # V05: estructura multitaller sin alterar el diseño original.
    execute('''CREATE TABLE IF NOT EXISTS operations(
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        code TEXT NOT NULL UNIQUE,
        name TEXT NOT NULL,
        client TEXT,
        location TEXT,
        active INTEGER NOT NULL DEFAULT 1,
        created_at TEXT DEFAULT CURRENT_TIMESTAMP
    )''')
    for _table in ('users','equipment','tires','occurrences'):
        _cols = {r['name'] for r in query(f'PRAGMA table_info({_table})')}
        if 'operation_id' not in _cols:
            execute(f'ALTER TABLE {_table} ADD COLUMN operation_id INTEGER')
    # VXX: Tapa Válvula es un dato propio de cada evento. Se agrega de forma
    # retrocompatible y se inicializa desde Observaciones para no perder la lógica histórica.
    _occ_cols = {r['name'] for r in query("PRAGMA table_info(occurrences)")}
    if 'valve_cap' not in _occ_cols:
        execute('ALTER TABLE occurrences ADD COLUMN valve_cap TEXT')
    # Migración histórica: solo completa valores vacíos. Nunca sobreescribe
    # una selección SI/NO que el usuario ya haya registrado explícitamente.
    # Se hace en Python para evitar que el adaptador de base de datos
    # interprete los caracteres '%' del LIKE como marcadores de formato.
    _occ_rows = query("SELECT id, notes, valve_cap FROM occurrences")
    for _occ in _occ_rows:
        _current_cap = str(_occ['valve_cap'] or '').strip().upper()
        if _current_cap:
            continue
        _note = str(_occ['notes'] or '').upper()
        _new_cap = 'SI' if ('TAPA' in _note or 'VALVULA' in _note or 'VÁLVULA' in _note) else 'NO'
        execute("UPDATE occurrences SET valve_cap=? WHERE id=?", (_new_cap, _occ['id']))
    execute("INSERT OR IGNORE INTO operations(id,code,name,client,location,active) VALUES(1,'CL','CERRO LINDO','NEXA RESOURCES PERU S.A.A.','Cerro Lindo',1)")
    for _table in ('users','equipment','tires','occurrences'):
        execute(f'UPDATE {_table} SET operation_id=1 WHERE operation_id IS NULL')

    # V06: operación activa por sesión. Solo el Administrador General puede cambiarla.
    # REGLA MAESTRA MULTITALLER:
    #   - Cada registro nuevo/consultado/modificado debe pertenecer a la operación activa.
    #   - Un usuario de operación queda SIEMPRE amarrado al operation_id de su usuario.
    #   - El Administrador General puede cambiar la operación desde el Panel principal.
    #   - La carga masiva Excel NO decide el taller: hereda el operation_id activo.
    operation_state = {'id': 1}

    def _is_general_admin():
        role = str((session.get('user') or {}).get('role') or '').upper()
        return role in ('ADMIN', 'ADMINISTRADOR GENERAL', 'SUPERADMIN')

    def active_operation_id():
        user = session.get('user') or {}
        if user and not _is_general_admin():
            try:
                assigned = int(user.get('operation_id') or 0)
            except (TypeError, ValueError):
                assigned = 0
            if assigned:
                return assigned
        return int(operation_state.get('id') or 1)

    def active_operation_row():
        op_id = active_operation_id()
        rows = query('SELECT id,code,name,client,location FROM operations WHERE id=? AND active=1', (op_id,))
        if rows:
            return rows[0]
        # Solo como respaldo de migración: Cerro Lindo es la operación inicial.
        return {'id':1,'code':'CL','name':'CERRO LINDO','client':'NEXA RESOURCES PERU S.A.A.','location':'Cerro Lindo'}

    # Campos maestros necesarios para la consulta operativa y el registro maestro.
    startup_cols = {r['name'] for r in query("PRAGMA table_info(tires)")}
    for col_name, col_type in [
        ('entry_date', 'TEXT'),
        ('cost_usd', 'REAL'),
        ('compound', 'TEXT'),
        ('supplier', 'TEXT'),
        ('new_tread_outer', 'REAL'),
        ('new_tread_inner', 'REAL'),
        ('construction_type', 'TEXT'),
        ('tire_condition', 'TEXT'),
        ('retirement_tread', 'REAL'),
        ('projected_life_target', 'REAL'),
        ('installation_meter', 'REAL'),
    ]:
        if col_name not in startup_cols:
            execute(f'ALTER TABLE tires ADD COLUMN {col_name} {col_type}')

    # Actualización única del maestro de 36 neumáticos.
    # Se actualiza por Código: no crea duplicados y no toca estado, equipo,
    # posición, horómetro, profundidades actuales ni historial de eventos.
    master_migration_key = 'master_tires_20260902_v1'
    master_done = query('SELECT value FROM app_meta WHERE key=?', (master_migration_key,))
    if not master_done:
        for (
            tire_code, serial, entry_date, cost_usd, brand, size, design,
            tra, supplier, pressure, tread_ext, tread_int, retirement_tread,
            life_target, construction_type, tire_condition
        ) in MASTER_TIRE_UPDATES_20260902:
            execute(
                '''
                UPDATE tires
                SET serial=?,
                    entry_date=?,
                    cost_usd=?,
                    brand=?,
                    size=?,
                    design=?,
                    compound=?,
                    supplier=?,
                    recommended_pressure=?,
                    new_tread=?,
                    new_tread_outer=?,
                    new_tread_inner=?,
                    retirement_tread=?,
                    projected_life_target=?,
                    projected_life=?,
                    construction_type=?,
                    tire_condition=?
                WHERE code=? AND operation_id=1
                ''',
                (
                    serial, entry_date, cost_usd, brand, size, design, tra,
                    supplier, pressure, tread_ext, tread_ext, tread_int,
                    retirement_tread, life_target, life_target,
                    construction_type, tire_condition, tire_code
                )
            )
        execute(
            'INSERT OR REPLACE INTO app_meta(key,value) VALUES(?,?)',
            (master_migration_key, '36 neumáticos actualizados por código - 02/09/2026')
        )

    nfu_migration_key='nfu_bajas_20260908_v1'
    if not query('SELECT value FROM app_meta WHERE key=?',(nfu_migration_key,)):
        for rec in NFU_BAJAS_20260908:
            (code,serial,entry_date,cost,brand,size,design,tra,supplier,pressure,new_ext,new_int,retirement,construction,condition,baja_date,equipment_label,position,baja_meter,rtd_ext,rtd_int,reason,location,notes,install_meter)=rec
            ex=query('SELECT id FROM tires WHERE code=? AND operation_id=1',(code,))
            if ex:
                tid=int(ex[0]['id'])
                execute("UPDATE tires SET serial=?,entry_date=?,cost_usd=?,brand=?,size=?,design=?,compound=?,supplier=?,recommended_pressure=?,new_tread=?,new_tread_outer=?,new_tread_inner=?,retirement_tread=?,construction_type=?,tire_condition=?,installation_meter=?,tread_outer=?,tread_inner=?,current_meter=?,status='BAJA',equipment_id=NULL,position=NULL,operation_id=1 WHERE id=?",(serial,entry_date,cost,brand,size,design,tra,supplier,pressure,max(new_ext,new_int),new_ext,new_int,retirement,construction,condition,install_meter,rtd_ext,rtd_int,baja_meter,tid))
            else:
                execute("INSERT INTO tires(code,serial,brand,size,design,new_tread,recommended_pressure,tread_inner,tread_outer,status,current_meter,entry_date,cost_usd,compound,supplier,new_tread_outer,new_tread_inner,construction_type,tire_condition,retirement_tread,installation_meter,operation_id) VALUES(?,?,?,?,?,?,?,?,?,'BAJA',?,?,?,?,?,?,?,?,?,?,?,1)",(code,serial,brand,size,design,max(new_ext,new_int),pressure,rtd_int,rtd_ext,baja_meter,entry_date,cost,tra,supplier,new_ext,new_int,construction,condition,retirement,install_meter))
                tid=int(query('SELECT id FROM tires WHERE code=? AND operation_id=1',(code,))[0]['id'])
            digits=''.join(ch for ch in str(equipment_label) if ch.isdigit())
            eq=None
            for cand in [equipment_label,f'SC-{digits}',f'SC {digits}',f'CAT-{digits}']:
                q=query('SELECT id FROM equipment WHERE UPPER(TRIM(code))=UPPER(TRIM(?)) AND operation_id=1 LIMIT 1',(cand,))
                if q: eq=int(q[0]['id']); break
            ob=query("SELECT id FROM occurrences WHERE tire_id=? AND operation_id=1 AND event_code='BAJA' ORDER BY id DESC LIMIT 1",(tid,))
            if ob:
                execute('UPDATE occurrences SET event_date=?,equipment_id=?,position=?,meter=?,tread_outer=?,tread_inner=?,reason=?,location=?,notes=?,operation_id=1 WHERE id=?',(baja_date,eq,position,baja_meter,rtd_ext,rtd_int,reason,location,notes,int(ob[0]['id'])))
            else:
                execute("INSERT INTO occurrences(tire_id,event_code,event_date,equipment_id,position,meter,tread_inner,tread_outer,pressure,pressure_condition,reason,location,notes,operation_id) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,1)",(tid,'BAJA',baja_date,eq,position,baja_meter,rtd_int,rtd_ext,None,'FRIO',reason,location,notes))
        execute('INSERT OR REPLACE INTO app_meta(key,value) VALUES(?,?)',(nfu_migration_key,'9 NFU/BAJA - 08/09/2026'))

    # V19 · reparación de pertenencia histórica de CERRO LINDO.
    # Las 9 NFU/BAJA son parte del inventario original NEXA/Cerro Lindo.
    # Versiones previas podían haberlas creado después del UPDATE inicial de operation_id.
    _nfu_codes = tuple(str(r[0]) for r in NFU_BAJAS_20260908)
    if _nfu_codes:
        _marks = ','.join('?' for _ in _nfu_codes)
        execute(f"UPDATE tires SET operation_id=1 WHERE code IN ({_marks}) AND operation_id IS NULL", _nfu_codes)
        _nfu_ids = query(f"SELECT id FROM tires WHERE code IN ({_marks}) AND operation_id=1", _nfu_codes)
        if _nfu_ids:
            _ids = tuple(int(r['id']) for r in _nfu_ids)
            _imarks = ','.join('?' for _ in _ids)
            execute(f"UPDATE occurrences SET operation_id=1 WHERE tire_id IN ({_imarks}) AND operation_id IS NULL", _ids)

    # Completa únicamente ocurrencias históricas huérfanas cuyos neumáticos ya pertenecen
    # inequívocamente a CERRO LINDO. No reasigna registros de otros talleres.
    execute("""UPDATE occurrences SET operation_id=1
               WHERE operation_id IS NULL
                 AND tire_id IN (SELECT id FROM tires WHERE operation_id=1)""")

    page.title = 'MegaSoftire Web 2026'
    page.padding = 0
    page.bgcolor = BG
    page.theme_mode = ft.ThemeMode.LIGHT
    page.theme = ft.Theme(color_scheme_seed=ft.Colors.BLUE_800, use_material3=True)

    session = {'user': None}
    app_host = ft.Container(expand=True)

    def snack(msg, err=False):
        page.snack_bar = ft.SnackBar(
            ft.Text(msg),
            bgcolor=ft.Colors.RED_700 if err else ft.Colors.GREEN_700
        )
        page.snack_bar.open = True
        page.update()

    def num(v):
        try:
            return float(v) if v not in (None, '') else None
        except Exception:
            return None

    def format_date(value):
        """Muestra fechas en dd/mm/aaaa sin alterar el valor almacenado.

        Acepta seriales de Excel (p. ej. 46212), fechas ISO y fechas ya
        formateadas. Esto corrige la visualización de datos importados de NEXA.
        """
        if value in (None, ''):
            return ''
        s = str(value).strip()
        # Serial de Excel: usa origen 1899-12-30 (compatibilidad Excel/LibreOffice).
        try:
            f = float(s)
            if f.is_integer() and 1 <= f <= 100000:
                d = dt.date(1899, 12, 30) + dt.timedelta(days=int(f))
                return d.strftime('%d/%m/%Y')
        except Exception:
            pass
        # Fechas textuales comunes.
        for fmt_in in ('%Y-%m-%d', '%Y/%m/%d', '%d/%m/%Y', '%d-%m-%Y'):
            try:
                d = dt.datetime.strptime(s[:10], fmt_in)
                return d.strftime('%d/%m/%Y')
            except Exception:
                pass
        return s

    def card(content, padding=18, width=None, height=None):
        return ft.Container(
            content=content,
            bgcolor=CARD_BG,
            border=ft.Border.all(1, '#E4EAF0'),
            border_radius=14,
            padding=padding,
            width=width,
            height=height,
            shadow=ft.BoxShadow(blur_radius=12, color='#12000000', offset=ft.Offset(0, 3)),
        )

    def metric_card(title, value, icon, subtitle=''):
        return card(
            ft.Column([
                ft.Row([
                    ft.Container(
                        width=40, height=40, border_radius=10,
                        bgcolor='#EAF2FF', alignment=ft.Alignment.CENTER,
                        content=ft.Icon(icon, color=NAV_ACCENT, size=21)
                    ),
                    ft.Container(expand=True),
                    ft.Text(str(value), size=28, weight=ft.FontWeight.BOLD, color=TEXT_MAIN),
                ]),
                ft.Text(title, size=13, weight=ft.FontWeight.W_600, color=TEXT_MAIN),
                ft.Text(subtitle, size=11, color=TEXT_MUTED),
            ], spacing=7), width=205
        )

    content = ft.Container(expand=True, padding=24)
    nav = None

    def page_title(title, subtitle=''):
        return ft.Column([
            ft.Text(title, size=27, weight=ft.FontWeight.BOLD, color=TEXT_MAIN),
            ft.Text(subtitle, size=12, color=TEXT_MUTED) if subtitle else ft.Container(height=0),
        ], spacing=2)

    def dashboard():
        op_id = active_operation_id()
        op = active_operation_row()

        def count(where='1=1', params=()):
            return query(f'SELECT COUNT(*) n FROM tires WHERE operation_id=? AND ({where})', (op_id, *params))[0]['n']

        # Contadores del Panel por operación activa.
        # BAJA se toma también del evento BAJA para conservar los históricos migrados
        # cuyo registro maestro pudo quedar sin operation_id durante versiones anteriores.
        service = count("status='SERVICIO'")
        standby = count("status='STAND-BY'")
        repair = count("status='REPARACIÓN'")
        baja = query("""SELECT COUNT(DISTINCT t.id) n
                        FROM tires t
                        WHERE t.status='BAJA'
                          AND (t.operation_id=? OR EXISTS (
                              SELECT 1 FROM occurrences o
                              WHERE o.tire_id=t.id AND o.operation_id=? AND o.event_code='BAJA'
                          ))""", (op_id, op_id))[0]['n']
        total = query("""SELECT COUNT(DISTINCT t.id) n
                         FROM tires t
                         WHERE t.operation_id=?
                            OR EXISTS (
                                SELECT 1 FROM occurrences o
                                WHERE o.tire_id=t.id AND o.operation_id=? AND o.event_code='BAJA'
                            )""", (op_id, op_id))[0]['n']
        equip = query('SELECT COUNT(*) n FROM equipment WHERE active=1 AND operation_id=?', (op_id,))[0]['n']


        # Selector visible únicamente en el Panel principal para Administrador General.
        operations = query('SELECT id,name,client FROM operations WHERE active=1 ORDER BY name')
        # V07: el valor real del selector es operation_id. El nombre es solo texto visible.
        # Esto evita que Flet conserve visualmente la nueva opción mientras el estado interno
        # todavía apunta a la operación anterior.
        valid_operation_ids = {int(r['id']) for r in operations}
        operation_selector = ft.Dropdown(
            label='Operación activa',
            value=str(op_id),
            width=360,
            options=[ft.DropdownOption(key=str(r['id']), text=str(r['name'])) for r in operations],
        )

        def change_operation(e):
            # Flet 1.0: Dropdown dispara on_select (no on_change).
            # El valor seleccionado se obtiene de e.control.value.
            raw = getattr(getattr(e, 'control', None), 'value', None)
            try:
                selected_id = int(str(raw))
            except (TypeError, ValueError):
                return
            if selected_id not in valid_operation_ids:
                return
            if selected_id == active_operation_id():
                return
            operation_state['id'] = selected_id
            # Reconstruye el Panel principal completo con el operation_id elegido.
            # Las demás vistas toman active_operation_id() al abrirse, por lo que
            # todo ingreso/consulta/movimiento queda inmediatamente en el taller seleccionado.
            dashboard()

        operation_selector.on_select = change_operation
        admin_general = _is_general_admin()
        selector_block = card(ft.Column([
            ft.Row([
                ft.Icon(ft.Icons.BUSINESS_OUTLINED, color=NAV_ACCENT, size=22),
                ft.Text('SELECCIÓN DE OPERACIÓN', size=14, weight=ft.FontWeight.BOLD, color=TEXT_MAIN),
            ], spacing=8),
            ft.Text('Vista como Administrador General. Seleccione la operación que desea administrar.', size=12, color=TEXT_MUTED),
            operation_selector,
            ft.Text(f"Cliente: {op['client'] or '-'}", size=12, color=TEXT_MUTED),
        ], spacing=8), width=560) if admin_general else ft.Container(height=0)


        # V16: administración multitaller desde el Panel principal (solo Administrador General).
        workshop_name = ft.TextField(label='Nombre del taller / operación', width=280)
        workshop_code = ft.TextField(label='Código', width=150)
        workshop_client = ft.TextField(label='Cliente', width=280)
        workshop_location = ft.TextField(label='Ubicación', width=220)

        def create_workshop(e):
            name=(workshop_name.value or '').strip().upper()
            code=(workshop_code.value or '').strip().upper()
            client=(workshop_client.value or '').strip()
            location=(workshop_location.value or '').strip()
            if not name or not code:
                snack('Ingrese código y nombre del taller.', True); return
            if query('SELECT id FROM operations WHERE UPPER(code)=UPPER(?) OR UPPER(name)=UPPER(?) LIMIT 1',(code,name)):
                snack('Ya existe una operación con ese código o nombre.', True); return
            execute('INSERT INTO operations(code,name,client,location,active) VALUES(?,?,?,?,1)',(code,name,client,location))
            snack(f'Taller {name} creado correctamente.')
            dashboard()

        op_options=[ft.DropdownOption(key=str(r['id']), text=str(r['name'])) for r in operations]
        user_fullname=ft.TextField(label='Nombre completo', width=280)
        user_username=ft.TextField(label='Usuario', width=200)
        user_password=ft.TextField(label='Contraseña', password=True, can_reveal_password=True, width=220)
        user_role=ft.Dropdown(label='Rol', value='USUARIO', width=250, options=[
            ft.DropdownOption(key='ADMINISTRADOR DE OPERACION', text='Administrador de Operación'),
            ft.DropdownOption(key='USUARIO', text='Usuario / Técnico'),
        ])
        user_operation=ft.Dropdown(label='Taller asignado', value=str(op_id), width=300, options=op_options)

        def create_user(e):
            fullname=(user_fullname.value or '').strip()
            username=(user_username.value or '').strip()
            password=user_password.value or ''
            role=(user_role.value or 'USUARIO').strip().upper()
            try: assigned_op=int(str(user_operation.value))
            except Exception: assigned_op=0
            if not fullname or not username or not password or not assigned_op:
                snack('Complete nombre, usuario, contraseña, rol y taller.', True); return
            if len(password) < 6:
                snack('La contraseña debe tener al menos 6 caracteres.', True); return
            if query('SELECT id FROM users WHERE LOWER(username)=LOWER(?) LIMIT 1',(username,)):
                snack('Ese nombre de usuario ya existe.', True); return
            if not query('SELECT id FROM operations WHERE id=? AND active=1',(assigned_op,)):
                snack('Seleccione un taller activo.', True); return
            pwd_hash=hashlib.sha256(password.encode('utf-8')).hexdigest()
            execute('INSERT INTO users(username,password_hash,full_name,role,active,operation_id) VALUES(?,?,?,?,1,?)',
                    (username,pwd_hash,fullname,role,assigned_op))
            snack(f'Usuario {username} creado y asignado al taller.')
            dashboard()

        workshops_rows=query('SELECT id,code,name,client,location,active FROM operations WHERE active=1 ORDER BY name')
        users_rows=query("""SELECT u.id,u.username,u.full_name,u.role,u.active,u.operation_id,o.name operation_name
                           FROM users u LEFT JOIN operations o ON o.id=u.operation_id
                           WHERE u.active=1 ORDER BY o.name,u.username""")

        def close_admin_dialog(dlg=None):
            page.pop_dialog(); page.update()

        def edit_workshop(row):
            f_code=ft.TextField(label='Código',value=str(row['code'] or ''),width=180)
            f_name=ft.TextField(label='Taller / operación',value=str(row['name'] or ''),width=300)
            f_client=ft.TextField(label='Cliente',value=str(row['client'] or ''),width=300)
            f_location=ft.TextField(label='Ubicación',value=str(row['location'] or ''),width=250)
            def save_edit(e=None):
                code=(f_code.value or '').strip().upper(); name=(f_name.value or '').strip().upper()
                if not code or not name: snack('Código y nombre son obligatorios.',True); return
                if query('SELECT id FROM operations WHERE id<>? AND (UPPER(code)=UPPER(?) OR UPPER(name)=UPPER(?)) LIMIT 1',(row['id'],code,name)):
                    snack('Ya existe otro taller con ese código o nombre.',True); return
                execute('UPDATE operations SET code=?,name=?,client=?,location=? WHERE id=?',(code,name,(f_client.value or '').strip(),(f_location.value or '').strip(),row['id']))
                page.pop_dialog(); snack('Taller actualizado correctamente.'); dashboard()
            dlg=ft.AlertDialog(modal=True,title=ft.Text('Editar taller / operación'),content=ft.Column([f_code,f_name,f_client,f_location],tight=True,spacing=10),actions=[ft.TextButton('Cancelar',on_click=lambda e: close_admin_dialog(dlg)),ft.Button('Guardar',icon=ft.Icons.SAVE,on_click=save_edit)])
            page.show_dialog(dlg); page.update()

        def delete_workshop(row):
            wid=int(row['id'])
            if wid==1: snack('CERRO LINDO es la operación base y no puede eliminarse.',True); return
            counts={t:query(f'SELECT COUNT(*) n FROM {t} WHERE operation_id=?',(wid,))[0]['n'] for t in ('equipment','tires','occurrences','users')}
            has_data=bool(counts['equipment'] or counts['tires'] or counts['occurrences'])
            msg=('Este taller contiene equipos, neumáticos o movimientos. Por seguridad no se borrará la información; el taller será DESACTIVADO.' if has_data else 'Este taller no contiene información operativa. Se eliminará de la lista de talleres.')
            def confirm(e=None):
                if has_data:
                    execute('UPDATE operations SET active=0 WHERE id=?',(wid,)); execute('UPDATE users SET active=0 WHERE operation_id=?',(wid,))
                else:
                    execute('DELETE FROM users WHERE operation_id=?',(wid,)); execute('DELETE FROM operations WHERE id=?',(wid,))
                if active_operation_id()==wid: operation_state['id']=1
                page.pop_dialog(); snack('Taller desactivado.' if has_data else 'Taller eliminado.'); dashboard()
            dlg=ft.AlertDialog(modal=True,title=ft.Text('Eliminar taller / operación'),content=ft.Text(msg),actions=[ft.TextButton('Cancelar',on_click=lambda e: close_admin_dialog(dlg)),ft.Button('Eliminar',icon=ft.Icons.DELETE_OUTLINE,on_click=confirm)])
            page.show_dialog(dlg); page.update()

        def edit_user(row):
            f_name=ft.TextField(label='Nombre completo',value=str(row['full_name'] or ''),width=280)
            f_username=ft.TextField(label='Usuario',value=str(row['username'] or ''),width=220)
            f_password=ft.TextField(label='Nueva contraseña (opcional)',password=True,can_reveal_password=True,width=260)
            is_general=str(row['role'] or '').upper() in ('ADMIN','ADMINISTRADOR GENERAL','SUPERADMIN')
            rv=str(row['role'] or 'USUARIO').upper()
            f_role=ft.Dropdown(label='Rol',value=('USUARIO' if is_general else rv),width=260,options=[ft.DropdownOption(key='ADMINISTRADOR DE OPERACION',text='Administrador de Operación'),ft.DropdownOption(key='USUARIO',text='Usuario / Técnico')],disabled=is_general)
            f_op=ft.Dropdown(label='Taller asignado',value=str(row['operation_id'] or op_id),width=300,options=[ft.DropdownOption(key=str(r['id']),text=str(r['name'])) for r in workshops_rows],disabled=is_general)
            def save_edit(e=None):
                name=(f_name.value or '').strip(); username=(f_username.value or '').strip(); pwd=f_password.value or ''
                if not name or not username: snack('Nombre y usuario son obligatorios.',True); return
                if query('SELECT id FROM users WHERE id<>? AND LOWER(username)=LOWER(?) LIMIT 1',(row['id'],username)): snack('Ese nombre de usuario ya existe.',True); return
                if pwd and len(pwd)<6: snack('La contraseña debe tener al menos 6 caracteres.',True); return
                if is_general:
                    execute('UPDATE users SET full_name=?,username=? WHERE id=?',(name,username,row['id']))
                else:
                    try: assigned=int(str(f_op.value))
                    except Exception: assigned=0
                    if not assigned or not query('SELECT id FROM operations WHERE id=? AND active=1',(assigned,)): snack('Seleccione un taller activo.',True); return
                    execute('UPDATE users SET full_name=?,username=?,role=?,operation_id=? WHERE id=?',(name,username,str(f_role.value or 'USUARIO').upper(),assigned,row['id']))
                if pwd: execute('UPDATE users SET password_hash=? WHERE id=?',(hashlib.sha256(pwd.encode('utf-8')).hexdigest(),row['id']))
                page.pop_dialog(); snack('Usuario actualizado correctamente.'); dashboard()
            dlg=ft.AlertDialog(modal=True,title=ft.Text('Editar usuario'),content=ft.Column([f_name,f_username,f_password,f_role,f_op],tight=True,spacing=10),actions=[ft.TextButton('Cancelar',on_click=lambda e: close_admin_dialog(dlg)),ft.Button('Guardar',icon=ft.Icons.SAVE,on_click=save_edit)])
            page.show_dialog(dlg); page.update()

        def delete_user(row):
            current_id=session.get('user',{}).get('id'); is_general=str(row['role'] or '').upper() in ('ADMIN','ADMINISTRADOR GENERAL','SUPERADMIN')
            if current_id and int(row['id'])==int(current_id): snack('No puede eliminar el usuario con el que inició sesión.',True); return
            if is_general: snack('El Administrador General principal no puede eliminarse desde este panel.',True); return
            def confirm(e=None):
                execute('UPDATE users SET active=0 WHERE id=?',(row['id'],)); page.pop_dialog(); snack('Usuario eliminado de la lista activa.'); dashboard()
            dlg=ft.AlertDialog(modal=True,title=ft.Text('Eliminar usuario'),content=ft.Text(f"¿Desea eliminar/desactivar al usuario {row['username']}?"),actions=[ft.TextButton('Cancelar',on_click=lambda e: close_admin_dialog(dlg)),ft.Button('Eliminar',icon=ft.Icons.DELETE_OUTLINE,on_click=confirm)])
            page.show_dialog(dlg); page.update()

        workshop_table=ft.DataTable(columns=[ft.DataColumn(ft.Text(x)) for x in ['Código','Taller / operación','Cliente','Ubicación','Acciones']],rows=[ft.DataRow(cells=[ft.DataCell(ft.Text(str(r['code'] or ''))),ft.DataCell(ft.Text(str(r['name'] or ''))),ft.DataCell(ft.Text(str(r['client'] or ''))),ft.DataCell(ft.Text(str(r['location'] or ''))),ft.DataCell(ft.Row([ft.IconButton(icon=ft.Icons.EDIT_OUTLINED,tooltip='Editar',on_click=lambda e,r=r: edit_workshop(r)),ft.IconButton(icon=ft.Icons.CLOSE,tooltip='Eliminar',on_click=lambda e,r=r: delete_workshop(r))],spacing=2))]) for r in workshops_rows])
        user_table=ft.DataTable(columns=[ft.DataColumn(ft.Text(x)) for x in ['Nombre','Usuario','Rol','Taller asignado','Acciones']],rows=[ft.DataRow(cells=[ft.DataCell(ft.Text(str(r['full_name'] or ''))),ft.DataCell(ft.Text(str(r['username'] or ''))),ft.DataCell(ft.Text('Administrador General' if str(r['role'] or '').upper()=='ADMIN' else str(r['role'] or ''))),ft.DataCell(ft.Text(str(r['operation_name'] or 'GENERAL'))),ft.DataCell(ft.Row([ft.IconButton(icon=ft.Icons.EDIT_OUTLINED,tooltip='Editar',on_click=lambda e,r=r: edit_user(r)),ft.IconButton(icon=ft.Icons.CLOSE,tooltip='Eliminar',on_click=lambda e,r=r: delete_user(r))],spacing=2))]) for r in users_rows])

        management_block = ft.Column([
            card(ft.Column([ft.Row([ft.Icon(ft.Icons.ADD_BUSINESS_OUTLINED,color=NAV_ACCENT),ft.Text('AGREGAR TALLER / OPERACIÓN',size=14,weight=ft.FontWeight.BOLD,color=TEXT_MAIN)],spacing=8),ft.Text('Cree nuevas operaciones sin modificar la información de los talleres existentes.',size=11,color=TEXT_MUTED),ft.Row([workshop_code,workshop_name,workshop_client,workshop_location],wrap=True,spacing=10),ft.Button('AGREGAR TALLER',icon=ft.Icons.ADD,on_click=create_workshop),ft.Divider(),ft.Text('TALLERES / OPERACIONES REGISTRADAS',size=13,weight=ft.FontWeight.BOLD,color=TEXT_MAIN),ft.Row([workshop_table],scroll=ft.ScrollMode.AUTO)],spacing=10)),
            card(ft.Column([ft.Row([ft.Icon(ft.Icons.PERSON_ADD_ALT_1_OUTLINED,color=NAV_ACCENT),ft.Text('CREAR USUARIO POR TALLER',size=14,weight=ft.FontWeight.BOLD,color=TEXT_MAIN)],spacing=8),ft.Text('Cada usuario queda vinculado a una sola operación y verá únicamente la información de ese taller.',size=11,color=TEXT_MUTED),ft.Row([user_fullname,user_username,user_password,user_role,user_operation],wrap=True,spacing=10),ft.Button('CREAR USUARIO',icon=ft.Icons.PERSON_ADD,on_click=create_user),ft.Divider(),ft.Text('USUARIOS REGISTRADOS',size=13,weight=ft.FontWeight.BOLD,color=TEXT_MAIN),ft.Row([user_table],scroll=ft.ScrollMode.AUTO)],spacing=10)),
        ],spacing=12) if admin_general else ft.Container(height=0)


        content.content = ft.Column([
            page_title('Panel principal', f"Vista general de la operación de neumáticos OTR · {op['name']}"),
            selector_block,
            management_block,
            ft.Row([
                metric_card('Neumáticos', total, ft.Icons.TIRE_REPAIR, 'Maestro total'),
                metric_card('En servicio', service, ft.Icons.CHECK_CIRCLE_OUTLINE, 'Actualmente instalados'),
                metric_card('Stand-by', standby, ft.Icons.PAUSE_CIRCLE_OUTLINE, 'Disponibles / retén'),
                metric_card('En reparación', repair, ft.Icons.HANDYMAN_OUTLINED, 'Pendientes de retorno'),
                metric_card('Baja', baja, ft.Icons.CANCEL_OUTLINED, 'Fuera de servicio'),
                metric_card('Equipos activos', equip, ft.Icons.PRECISION_MANUFACTURING_OUTLINED, 'Flota registrada'),
            ], wrap=True, spacing=12, run_spacing=12),
        ], scroll=ft.ScrollMode.AUTO, spacing=16)
        page.update()

    def equipment_view():
        op_id = active_operation_id()

        # PostgreSQL/SQLite compatible: no se crean tablas auxiliares desde la vista.
        # performance_type debe existir en equipment; el SQL de migración lo agrega.
        def distinct_options(field):
            rows = query(
                f"SELECT DISTINCT {field} AS value "
                "FROM equipment "
                "WHERE operation_id=? AND "
                f"{field} IS NOT NULL AND TRIM({field})<>'' "
                f"ORDER BY {field}",
                (op_id,)
            )
            return [ft.dropdown.Option(str(r['value'])) for r in rows]

        def make_field(label, field, width=190):
            options = distinct_options(field)
            return ft.Dropdown(
                label=label,
                width=width,
                editable=True,
                enable_filter=True,
                enable_search=True,
                options=options
            )

        code = ft.TextField(label='Código de equipo *', width=190)
        brand = make_field('Marca', 'brand')
        model = make_field('Modelo', 'model')
        location = make_field('Ubicación', 'location')
        kind = make_field('Tipo', 'vehicle_type')
        motor = make_field('Tipo de motor', 'motor_type')

        performance = ft.Dropdown(
            label='Rendimiento *',
            width=190,
            value=None,
            options=[
                ft.dropdown.Option('HORAS'),
                ft.dropdown.Option('KILÓMETROS'),
            ]
        )

        search = ft.TextField(
            label='Buscar equipo',
            prefix_icon=ft.Icons.SEARCH,
            width=280
        )

        performance_filter = ft.Dropdown(
            label='Filtrar rendimiento',
            width=190,
            value='TODOS',
            options=[
                ft.dropdown.Option('TODOS', 'Todos'),
                ft.dropdown.Option('HORAS', 'Horas'),
                ft.dropdown.Option('KILÓMETROS', 'Kilómetros'),
            ],
        )

        table = ft.DataTable(
            columns=[
                ft.DataColumn(ft.Text(x))
                for x in [
                    'Código',
                    'Marca / Modelo',
                    'Tipo',
                    'Ubicación',
                    'Tipo de motor',
                    'Rendimiento'
                ]
            ],
            rows=[]
        )

        def refresh():
            term = (search.value or '').strip()
            perf = (performance_filter.value or 'TODOS').strip().upper()

            where = ['operation_id=?']
            params = [op_id]

            if term:
                where.append(
                    '(code LIKE ? OR brand LIKE ? OR model LIKE ?)'
                )
                params.extend([
                    f'%{term}%',
                    f'%{term}%',
                    f'%{term}%'
                ])

            if perf != 'TODOS':
                where.append("UPPER(COALESCE(performance_type,''))=?")
                params.append(perf)

            rows = query(
                'SELECT code,brand,model,vehicle_type,location,motor_type,'
                'performance_type '
                'FROM equipment WHERE ' + ' AND '.join(where) +
                ' ORDER BY code',
                tuple(params)
            )

            table.rows = [
                ft.DataRow(
                    cells=[
                        ft.DataCell(ft.Text(str(v or '')))
                        for v in [
                            r['code'],
                            f"{r['brand'] or ''} {r['model'] or ''}".strip(),
                            r['vehicle_type'],
                            r['location'],
                            r['motor_type'],
                            r['performance_type'],
                        ]
                    ]
                )
                for r in rows
            ]
            page.update()

        def refresh_fields():
            for control, field in [
                (brand, 'brand'),
                (model, 'model'),
                (location, 'location'),
                (kind, 'vehicle_type'),
                (motor, 'motor_type'),
            ]:
                control.options = distinct_options(field)

        search.on_change = lambda e: refresh()
        performance_filter.on_select = lambda e: refresh()

        def save(e):
            if not (code.value or '').strip():
                return snack('Ingrese el código del equipo.', True)

            performance_value = (performance.value or '').strip().upper()
            if performance_value not in ('HORAS', 'KILÓMETROS'):
                return snack(
                    'Seleccione el rendimiento del equipo: HORAS o KILÓMETROS.',
                    True
                )

            try:
                execute(
                    'INSERT INTO equipment('
                    'code,brand,model,location,vehicle_type,motor_type,'
                    'performance_type,operation_id'
                    ') VALUES(?,?,?,?,?,?,?,?)',
                    (
                        code.value.strip(),
                        (brand.value or '').strip(),
                        (model.value or '').strip(),
                        (location.value or '').strip(),
                        (kind.value or '').strip(),
                        (motor.value or '').strip(),
                        performance_value,
                        op_id
                    )
                )

                code.value = ''
                brand.value = None
                model.value = None
                location.value = None
                kind.value = None
                motor.value = None
                performance.value = None

                refresh_fields()
                snack('Equipo registrado correctamente.')
                refresh()

            except Exception as ex:
                snack(str(ex), True)

        refresh()

        content.content = ft.Column([
            page_title(
                '8.1 EQUIPOS',
                'Administración de equipos · Registro y consulta de la flota'
            ),
            card(ft.Column([
                ft.Text(
                    'Nuevo equipo',
                    size=17,
                    weight=ft.FontWeight.BOLD,
                    color=TEXT_MAIN
                ),
                ft.Row(
                    [
                        code,
                        brand,
                        model,
                        location,
                        kind,
                        motor,
                        performance
                    ],
                    wrap=True
                ),
                ft.Button(
                    'Registrar equipo',
                    icon=ft.Icons.SAVE,
                    on_click=save
                )
            ])),
        ], scroll=ft.ScrollMode.AUTO, spacing=16)

        page.update()


    def tires_view(status_filter=None, prefill_code=None):
        op_id = active_operation_id()
        # Registro maestro de neumáticos.
        existing_cols = {r['name'] for r in query("PRAGMA table_info(tires)")}
        extra_cols = [
            ('entry_date', 'TEXT'),
            ('cost_usd', 'REAL'),
            ('compound', 'TEXT'),
            ('supplier', 'TEXT'),
            ('new_tread_outer', 'REAL'),
            ('new_tread_inner', 'REAL'),
            ('construction_type', 'TEXT'),
            ('tire_condition', 'TEXT'),
            ('retirement_tread', 'REAL'),
            ('projected_life_target', 'REAL'),
        ]
        for col_name, col_type in extra_cols:
            if col_name not in existing_cols:
                execute(f'ALTER TABLE tires ADD COLUMN {col_name} {col_type}')

        # Catálogos dinámicos del Registro Maestro.
        # Permiten seleccionar un valor existente o DIGITAR uno nuevo.
        execute("""
            CREATE TABLE IF NOT EXISTS tire_catalogs(
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                category TEXT NOT NULL,
                value TEXT COLLATE NOCASE NOT NULL,
                UNIQUE(category, value)
            )
        """)

        catalog_defaults = {
            'brand': ['GoodYear','Bridgestone','Michelin','Yokohama','Techking','Maxam'],
            'supplier': ['Soltrak','Nuema','PTS','Renova','Tire SOL','J.CH.','Pimentel'],
        }
        for category, values in catalog_defaults.items():
            for value in values:
                execute(
                    'INSERT OR IGNORE INTO tire_catalogs(category,value) VALUES(?,?)',
                    (category, value)
                )

        # Recuperar también valores ya existentes en los neumáticos históricos.
        historic_catalog_fields = {
            'brand': 'brand',
            'size': 'size',
            'design': 'design',
            'compound': 'compound',
            'supplier': 'supplier',
        }
        for category, field_name in historic_catalog_fields.items():
            for row in query(
                f"SELECT DISTINCT {field_name} value FROM tires "
                f"WHERE {field_name} IS NOT NULL AND TRIM({field_name})<>''"
            ):
                execute(
                    'INSERT OR IGNORE INTO tire_catalogs(category,value) VALUES(?,?)',
                    (category, str(row['value']).strip())
                )

        def catalog_options(category):
            return [
                ft.dropdown.Option(str(r['value']))
                for r in query(
                    'SELECT value FROM tire_catalogs WHERE category=? ORDER BY value COLLATE NOCASE',
                    (category,)
                )
            ]

        def make_catalog_field(label, category, width):
            # Un solo cuadro: permite escribir y al mismo tiempo filtra/autocompleta
            # con los valores ya guardados en el catálogo.
            return ft.Dropdown(
                label=f'{label} *',
                width=width,
                editable=True,
                enable_filter=True,
                enable_search=True,
                options=catalog_options(category)
            )

        def normalize_catalog_key(value):
            # Para detectar equivalencias como:
            # GoodYear / GOODYEAR / GOOD YEAR
            return ''.join(str(value or '').strip().lower().split())

        def catalog_value(dropdown):
            typed = (getattr(dropdown, 'text', None) or '').strip()
            selected = (dropdown.value or '').strip()
            raw = typed if typed else selected
            if not raw:
                return ''

            key = normalize_catalog_key(raw)
            for option in dropdown.options or []:
                option_key = getattr(option, 'key', None)
                option_text = getattr(option, 'text', None)
                candidate = str(option_key or option_text or '').strip()
                if candidate and normalize_catalog_key(candidate) == key:
                    return candidate
            return raw

        def save_catalog_value(category, value):
            clean = (value or '').strip()
            if not clean:
                return clean

            wanted_key = normalize_catalog_key(clean)
            for row in query(
                'SELECT value FROM tire_catalogs WHERE category=?',
                (category,)
            ):
                existing_value = str(row['value']).strip()
                if normalize_catalog_key(existing_value) == wanted_key:
                    return existing_value

            execute(
                'INSERT OR IGNORE INTO tire_catalogs(category,value) VALUES(?,?)',
                (category, clean)
            )
            return clean

        FIELD_W = 220

        # Carga masiva Excel: usa exactamente los mismos campos del formulario
        # Nuevo neumático. El taller/operación se toma de la sesión activa y el
        # sistema fuerza el estado inicial STAND-BY, sin equipo ni posición.
        BULK_COLUMNS = [
            'Código *', 'Serie Fab. *', 'Fecha de ingreso *', 'Costo $ *',
            'Marca *', 'Medida *', 'Diseño *', 'Clasificación TRA *',
            'Proveedor *', 'Presión recomendada *', 'Profundidad nueva EXT *',
            'Profundidad nueva INT *', 'Profundidad de retiro (mm) *',
            'Proyección de vida (h) *', 'Tipo de construcción *', 'Condición *'
        ]

        def _bulk_template_bytes():
            try:
                from openpyxl import Workbook
                from openpyxl.styles import Font, PatternFill, Alignment
                from openpyxl.worksheet.datavalidation import DataValidation
                wb = Workbook()
                ws = wb.active
                ws.title = 'REGISTRO_NEUMATICOS'
                ws.append(BULK_COLUMNS)
                for cell in ws[1]:
                    cell.font = Font(bold=True, color='FFFFFF')
                    cell.fill = PatternFill('solid', fgColor='1E5AA8')
                    cell.alignment = Alignment(horizontal='center')
                widths = [18,22,18,14,18,16,20,24,18,22,24,24,28,24,24,18]
                for i, width in enumerate(widths, 1):
                    ws.column_dimensions[chr(64+i)].width = width
                ws.freeze_panes = 'A2'
                ws.auto_filter.ref = 'A1:P1'
                dv_con = DataValidation(type='list', formula1='"Radial,Convencional"', allow_blank=False)
                dv_cond = DataValidation(type='list', formula1='"Nueva,Reencauchada"', allow_blank=False)
                ws.add_data_validation(dv_con); ws.add_data_validation(dv_cond)
                dv_con.add('O2:O5000'); dv_cond.add('P2:P5000')
                ws2 = wb.create_sheet('INSTRUCCIONES')
                instructions = [
                    ['CARGA MASIVA DE NEUMÁTICOS - MEGASOFTIRE'],
                    ['Complete únicamente la hoja REGISTRO_NEUMATICOS.'],
                    ['Los campos deben coincidir con el formulario Nuevo neumático. Un ejemplo de llenado aparece en esta hoja de instrucciones.'],
                    ['Ejemplo: 10001 | SERIE-001 | 18/09/2026 | 5000 | Yokohama | 29.5R25 | RB31 | L-3 | Proveedor | 90 | 105 | 105 | 15 | 3000 | Convencional | Nueva'],
                    ['No agregue taller, equipo, posición ni situación: el sistema usa el taller activo y registra automáticamente STAND-BY.'],
                    ['Regla principal: debe existir Código O Serie Fab.; el otro puede quedar en blanco. Los campos sin información se dejan en blanco.'],
                    ['Antes de guardar, MegaSoftire mostrará una validación y un resumen de los registros.'],
                ]
                for row in instructions: ws2.append(row)
                ws2.column_dimensions['A'].width = 110
                ws2['A1'].font = Font(bold=True, size=14, color='1E5AA8')
                buf = io.BytesIO(); wb.save(buf); return buf.getvalue()
            except Exception as ex:
                raise RuntimeError(f'No se pudo crear la plantilla Excel: {ex}')

        # Picker persistente: en Flet 1.x los servicios deben conservar una
        # referencia viva para reutilizarse durante toda la vista.
        bulk_template_picker = ft.FilePicker()

        async def download_bulk_template(e=None):
            try:
                template_bytes = _bulk_template_bytes()
                if not template_bytes:
                    raise RuntimeError('La plantilla Excel se generó vacía.')
                await bulk_template_picker.save_file(
                    file_name='PLANTILLA_CARGA_MASIVA_NEUMATICOS.xlsx',
                    file_type=ft.FilePickerFileType.CUSTOM,
                    allowed_extensions=['xlsx'],
                    src_bytes=template_bytes,
                )
                # En Web, save_file() descarga los bytes entregados al navegador.
                snack('Plantilla Excel preparada para descarga.')
            except Exception as ex:
                snack(f'No se pudo descargar la plantilla: {ex}', True)

        def _clean_header(v):
            return ' '.join(str(v or '').replace('*','').strip().split()).lower()

        def _cell_text(v):
            if v is None:
                return ''
            if isinstance(v, float) and v.is_integer():
                return str(int(v))
            return str(v).strip()

        def _bulk_num(v, label, rownum, positive=False, nonnegative=False):
            raw=_cell_text(v).replace(',','.')
            # Los campos sin información se conservan en blanco.
            if not raw:
                return None
            try:
                n=float(raw)
            except Exception:
                raise ValueError(f'Fila {rownum}: {label} no es numérico.')
            if positive and n <= 0:
                raise ValueError(f'Fila {rownum}: {label} debe ser mayor que 0.')
            if nonnegative and n < 0:
                raise ValueError(f'Fila {rownum}: {label} no puede ser negativo.')
            return n

        def _read_bulk_excel(data):
            try:
                from openpyxl import load_workbook
                wb=load_workbook(io.BytesIO(data), data_only=True, read_only=True)
                ws=wb['REGISTRO_NEUMATICOS'] if 'REGISTRO_NEUMATICOS' in wb.sheetnames else wb.active
                rows=list(ws.iter_rows(values_only=True))
                if not rows:
                    raise ValueError('El archivo Excel está vacío.')
                headers=[_clean_header(v) for v in rows[0]]
                wanted=[_clean_header(v) for v in BULK_COLUMNS]
                positions={h:i for i,h in enumerate(headers) if h}
                missing=[BULK_COLUMNS[i] for i,h in enumerate(wanted) if h not in positions]
                if missing:
                    raise ValueError('Faltan columnas obligatorias: ' + ', '.join(missing))
                data_rows=[]
                for excel_row, raw in enumerate(rows[1:], start=2):
                    if not any(v not in (None,'') for v in raw):
                        continue
                    def val(label):
                        idx=positions[_clean_header(label)]
                        return raw[idx] if idx < len(raw) else None
                    code=_cell_text(val('Código *'))
                    serial=_cell_text(val('Serie Fab. *'))
                    # Regla maestra: basta Código O Serie Fab.; no se exigen ambos.
                    if not code and not serial:
                        raise ValueError(f'Fila {excel_row}: debe ingresar Código o Serie Fab.')
                    date_value=val('Fecha de ingreso *')
                    if isinstance(date_value, (dt.datetime, dt.date)):
                        date_iso=date_value.strftime('%Y-%m-%d')
                    else:
                        date_iso=normalize_date(_cell_text(date_value))
                    if not date_iso: raise ValueError(f'Fila {excel_row}: Fecha de ingreso inválida.')
                    cost=_bulk_num(val('Costo $ *'),'Costo $',excel_row,nonnegative=True)
                    pressure_value=_bulk_num(val('Presión recomendada *'),'Presión recomendada',excel_row,positive=True)
                    new_ext=_bulk_num(val('Profundidad nueva EXT *'),'Profundidad nueva EXT',excel_row,positive=True)
                    new_int=_bulk_num(val('Profundidad nueva INT *'),'Profundidad nueva INT',excel_row,positive=True)
                    retirement=_bulk_num(val('Profundidad de retiro (mm) *'),'Profundidad de retiro',excel_row,nonnegative=True)
                    life=_bulk_num(val('Proyección de vida (h) *'),'Proyección de vida',excel_row,positive=True)
                    if retirement is not None and new_ext is not None and new_int is not None and retirement >= min(new_ext,new_int):
                        raise ValueError(f'Fila {excel_row}: la profundidad de retiro debe ser menor que la profundidad nueva.')
                    construction=_cell_text(val('Tipo de construcción *')) or None
                    condition_value=_cell_text(val('Condición *')) or None
                    if construction and construction.lower() not in ('radial','convencional'):
                        raise ValueError(f'Fila {excel_row}: Tipo de construcción debe ser Radial o Convencional.')
                    if condition_value and condition_value.lower() not in ('nueva','reencauchada'):
                        raise ValueError(f'Fila {excel_row}: Condición debe ser Nueva o Reencauchada.')
                    required_text = {
                        'Marca *': _cell_text(val('Marca *')),
                        'Medida *': _cell_text(val('Medida *')),
                        'Diseño *': _cell_text(val('Diseño *')),
                        'Clasificación TRA *': _cell_text(val('Clasificación TRA *')),
                        'Proveedor *': _cell_text(val('Proveedor *')),
                    }
                    # Los demás campos pueden quedar en blanco si el Excel no tiene información.
                    data_rows.append({
                        'code':code,'serial':serial,'entry_date':date_iso,'cost':cost,
                        'brand':required_text['Marca *'],'size':required_text['Medida *'],
                        'design':required_text['Diseño *'],'compound':required_text['Clasificación TRA *'],
                        'supplier':required_text['Proveedor *'],'pressure':pressure_value,
                        'new_ext':new_ext,'new_int':new_int,'retirement':retirement,'life':life,
                        'construction':construction,'condition':condition_value,
                    })
                if not data_rows:
                    raise ValueError('No se encontraron registros para importar.')
                codes=[r['code'].strip().lower() for r in data_rows if r['code'].strip()]
                serials=[r['serial'].strip().lower() for r in data_rows if r['serial'].strip()]
                if len(codes)!=len(set(codes)):
                    raise ValueError('El archivo contiene códigos neumático duplicados.')
                if len(serials)!=len(set(serials)):
                    raise ValueError('El archivo contiene series de fábrica duplicadas.')
                existing_codes={str(r['code']).strip().lower() for r in query('SELECT code FROM tires WHERE operation_id=?',(op_id,))}
                existing_serials={str(r['serial']).strip().lower() for r in query("SELECT serial FROM tires WHERE operation_id=? AND serial IS NOT NULL AND TRIM(serial)<>''",(op_id,))}
                dup_codes=[r['code'] for r in data_rows if r['code'].strip() and r['code'].strip().lower() in existing_codes]
                dup_serials=[r['serial'] for r in data_rows if r['serial'].strip() and r['serial'].strip().lower() in existing_serials]
                if dup_codes:
                    raise ValueError('Códigos ya existentes en el taller activo: ' + ', '.join(dup_codes[:10]) + ('...' if len(dup_codes)>10 else ''))
                if dup_serials:
                    raise ValueError('Series ya existentes en el taller activo: ' + ', '.join(dup_serials[:10]) + ('...' if len(dup_serials)>10 else ''))
                return data_rows
            except ValueError:
                raise
            except Exception as ex:
                raise ValueError(f'No se pudo leer el Excel: {ex}')

        async def process_bulk_file(f):
            try:
                if not f or not getattr(f, 'bytes', None):
                    raise ValueError('No se pudo leer el contenido del archivo seleccionado.')
                rows=_read_bulk_excel(f.bytes)
                preview='\n'.join([f"{i+1}. {r['code']} | {r['serial']} | {r['brand']} | {r['size']} | {r['design']}" for i,r in enumerate(rows[:8])])
                if len(rows)>8: preview += f"\n... y {len(rows)-8} registro(s) más."
                op=active_operation_row()
                def confirm_bulk(e2=None):
                    try:
                        from database import connect
                        with connect() as con:
                            for r in rows:
                                con.execute(
                                    '''INSERT INTO tires(
                                        code,serial,brand,size,design,new_tread,recommended_pressure,
                                        tread_inner,tread_outer,entry_date,cost_usd,compound,supplier,
                                        new_tread_outer,new_tread_inner,construction_type,tire_condition,
                                        retirement_tread,projected_life_target,projected_life,status,
                                        equipment_id,position,current_meter,operation_id
                                    ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)''',
                                    (r['code'] or None,r['serial'] or None,r['brand'] or None,r['size'] or None,r['design'] or None,(max(v for v in (r['new_ext'],r['new_int']) if v is not None) if any(v is not None for v in (r['new_ext'],r['new_int'])) else None),
                                     r['pressure'],r['new_int'],r['new_ext'],r['entry_date'] or None,r['cost'],r['compound'] or None,r['supplier'] or None,
                                     r['new_ext'],r['new_int'],r['construction'],r['condition'],r['retirement'],r['life'],r['life'],
                                     'STAND-BY',None,None,None,op_id)
                                )
                                for category,key in [('brand','brand'),('size','size'),('design','design'),('compound','compound'),('supplier','supplier')]:
                                    if r[key]:
                                        con.execute('INSERT OR IGNORE INTO tire_catalogs(category,value) VALUES(?,?)',(category,r[key]))
                        page.pop_dialog(); snack(f"Carga masiva completada: {len(rows)} neumático(s) registrados en STAND-BY."); refresh()
                    except Exception as ex:
                        snack(f'No se pudo completar la carga. No se guardó ningún registro: {ex}', True)
                dlg=ft.AlertDialog(
                    modal=True,
                    title=ft.Text(f'Confirmar carga masiva · {op["name"]}'),
                    content=ft.Column([
                        ft.Text(f'Se encontraron {len(rows)} neumático(s) válidos.'),
                        ft.Text('Todos quedarán en STAND-BY, sin equipo y sin posición.',weight=ft.FontWeight.BOLD,color=NAV_ACCENT),
                        ft.Text(preview,size=11,font_family='monospace'),
                    ],tight=True,scroll=ft.ScrollMode.AUTO),
                    actions=[ft.TextButton('Cancelar',on_click=lambda e2: (page.pop_dialog(),page.update())),ft.Button('Confirmar carga',icon=ft.Icons.UPLOAD,on_click=confirm_bulk)]
                )
                page.show_dialog(dlg); page.update()
            except Exception as ex:
                snack(str(ex),True)

        # En Web, el selector de archivos debe abrirse dentro del gesto original
        # del navegador. Por eso Cargar Excel usa la acción cliente PickFiles y
        # recibe el archivo mediante on_result. Esto evita que Chrome/Firefox
        # bloqueen el selector y funciona también en navegadores más estrictos.
        async def on_bulk_file_result(e):
            try:
                files = e.files or []
                if not files:
                    return
                await process_bulk_file(files[0])
            except Exception as ex:
                snack(str(ex), True)

        bulk_file_picker = ft.FilePicker(on_result=on_bulk_file_result)
        bulk_pick_action = ft.PickFiles(
            bulk_file_picker,
            dialog_title='Seleccionar Excel de neumáticos',
            file_type=ft.FilePickerFileType.CUSTOM,
            allowed_extensions=['xlsx'],
            allow_multiple=False,
            with_data=True,
        )

        code=ft.TextField(label='Código *',width=FIELD_W,value=(str(prefill_code) if prefill_code else ''))
        serial=ft.TextField(label='Serie Fab. *',width=FIELD_W)
        entry_date=ft.TextField(label='Fecha de ingreso *',value=dt.date.today().strftime('%d/%m/%Y'),width=FIELD_W)
        cost_usd=ft.TextField(label='Costo $ *',width=FIELD_W)

        brand = make_catalog_field('Marca', 'brand', FIELD_W)
        size = make_catalog_field('Medida', 'size', FIELD_W)
        design = make_catalog_field('Diseño', 'design', FIELD_W)
        compound = make_catalog_field('Clasificación TRA', 'compound', FIELD_W)
        supplier = make_catalog_field('Proveedor', 'supplier', FIELD_W)

        pressure=ft.TextField(label='Presión recomendada *',width=FIELD_W)
        tread_outer_new=ft.TextField(label='Profundidad nueva EXT *',width=FIELD_W)
        tread_inner_new=ft.TextField(label='Profundidad nueva INT *',width=FIELD_W)
        retirement_tread=ft.TextField(label='Profundidad de retiro (mm) *',width=FIELD_W)
        projected_life_target=ft.TextField(label='Proyección de vida (h) *',width=FIELD_W)

        construction=ft.Dropdown(
            label='Tipo de construcción *', width=FIELD_W,
            options=[ft.dropdown.Option('Radial'),ft.dropdown.Option('Convencional')]
        )
        condition=ft.Dropdown(
            label='Condición *', width=FIELD_W,
            options=[ft.dropdown.Option('Nueva'),ft.dropdown.Option('Reencauchada')]
        )

        search=ft.TextField(label='Buscar neumático',prefix_icon=ft.Icons.SEARCH,width=260)
        eq_options=[ft.dropdown.Option('', 'Todos los equipos')]
        eq_options += [ft.dropdown.Option(str(r['id']), r['code']) for r in query('SELECT id,code FROM equipment WHERE active=1 AND operation_id=? ORDER BY code',(op_id,))]
        eq_filter=ft.Dropdown(label='Equipo',width=180,value='',options=eq_options)
        summary=ft.Text('',size=12,color=TEXT_MUTED)

        # Listado maestro: se construye con filas explícitas para que los
        # encabezados sean visibles aun cuando no haya registros y para evitar
        # el bloque gris que estaba mostrando DataTable en la vista web.
        master_columns = [
            ('Código', 90),
            ('Serie Fab.', 125),
            ('Fecha ingreso', 115),
            ('Costo $', 90),
            ('Marca', 120),
            ('Medida', 105),
            ('Diseño', 115),
            ('Clasificación TRA', 135),
            ('Proveedor', 120),
            ('Presión rec.', 105),
            ('Prof. nueva EXT', 120),
            ('Prof. nueva INT', 120),
            ('Prof. retiro', 105),
            ('Proy. vida (h)', 115),
            ('Tipo construcción', 140),
            ('Condición', 120),
        ]

        def master_cell(value, width, header=False):
            return ft.Container(
                content=ft.Text(
                    value,
                    size=12,
                    weight=ft.FontWeight.BOLD if header else ft.FontWeight.NORMAL,
                    color=TEXT_MAIN,
                    no_wrap=True,
                ),
                width=width,
                padding=ft.Padding(left=6, top=8, right=6, bottom=8),
            )

        master_header = ft.Row(
            [master_cell(label, width, True) for label, width in master_columns],
            spacing=0,
        )
        master_rows = ft.Column(spacing=0)
        master_table = ft.Column([
            ft.Container(
                content=master_header,
                bgcolor='#EEF2F7',
                border=ft.Border(bottom=ft.BorderSide(1, '#D5DCE5')),
            ),
            master_rows,
        ], spacing=0)

        history=ft.DataTable(columns=[ft.DataColumn(ft.Text(x)) for x in [
            'Fecha','Evento','Neumático','Serie','Equipo','Pos.','Lectura','Cocada I/E','Presión','Ubicación'
        ]],rows=[])

        def fmt(v):
            if v is None: return ''
            if isinstance(v,float) and v.is_integer(): return str(int(v))
            return str(v)

        def refresh(e=None):
            sql='SELECT t.*,e.code equipment_code FROM tires t LEFT JOIN equipment e ON e.id=t.equipment_id'
            clauses=['t.operation_id=?']; params=[op_id]
            if status_filter:
                clauses.append('t.status=?'); params.append(status_filter)
            term=(search.value or '').strip()
            if term:
                clauses.append('(t.code LIKE ? OR t.serial LIKE ? OR t.brand LIKE ? OR e.code LIKE ?)')
                params += [f'%{term}%',f'%{term}%',f'%{term}%',f'%{term}%']
            if clauses:
                sql += ' WHERE ' + ' AND '.join(clauses)
            sql += " ORDER BY CAST(COALESCE(NULLIF(t.position,''),'999') AS INTEGER),t.code"
            rows=query(sql,tuple(params))

            master_rows.controls=[]
            for idx, r in enumerate(rows):
                values = [
                    r['code'],
                    r['serial'],
                    format_date(r['entry_date']) if r['entry_date'] else None,
                    r['cost_usd'],
                    r['brand'],
                    r['size'],
                    r['design'],
                    r['compound'],
                    r['supplier'],
                    r['recommended_pressure'],
                    r['new_tread_outer'],
                    r['new_tread_inner'],
                    r['retirement_tread'],
                    r['projected_life_target'],
                    r['construction_type'],
                    r['tire_condition'],
                ]

                display_values = [
                    '—' if v is None or str(v).strip() == '' else fmt(v)
                    for v in values
                ]

                master_rows.controls.append(
                    ft.Container(
                        content=ft.Row(
                            [
                                master_cell(display_values[i], master_columns[i][1])
                                for i in range(len(master_columns))
                            ],
                            spacing=0,
                        ),
                        bgcolor='#FFFFFF' if idx % 2 == 0 else '#F8FAFC',
                        border=ft.Border(bottom=ft.BorderSide(1, '#E5E9EF')),
                    )
                )

            summary.value=f'{len(rows)} neumático(s) registrado(s)'

            hist_sql=("SELECT o.event_date,o.event_code,t.code tire_code,t.serial,e.code equipment_code,"
                      "o.position,o.meter,o.tread_inner,o.tread_outer,o.pressure,o.location "
                      "FROM occurrences o JOIN tires t ON t.id=o.tire_id "
                      "LEFT JOIN equipment e ON e.id=o.equipment_id")
            hp=[op_id]; hc=['o.operation_id=?']
            if eq_filter.value:
                hc.append('o.equipment_id=?'); hp.append(int(eq_filter.value))
            if status_filter and not eq_filter.value:
                hc.append('t.status=?'); hp.append(status_filter)
            if term:
                hc.append('(t.code LIKE ? OR t.serial LIKE ? OR e.code LIKE ?)')
                hp += [f'%{term}%',f'%{term}%',f'%{term}%']
            if hc:
                hist_sql += ' WHERE ' + ' AND '.join(hc)
            hist_sql += ' ORDER BY o.event_date DESC,o.id DESC LIMIT 80'
            hrows=query(hist_sql,tuple(hp))

            history.rows=[ft.DataRow(cells=[ft.DataCell(ft.Text(fmt(v))) for v in [
                format_date(r['event_date']),r['event_code'],r['tire_code'],r['serial'],r['equipment_code'],r['position'],r['meter'],
                f"{fmt(r['tread_inner'])}/{fmt(r['tread_outer'])}",r['pressure'],r['location']
            ]]) for r in hrows]
            page.update()

        search.on_change=refresh
        eq_filter.on_change=refresh

        # Regla maestra de identificación: basta Código O Serie Fab.
        # Todos los demás campos pueden quedar en blanco cuando la información
        # no está disponible; se validan únicamente cuando se proporciona un valor.
        required_controls = []

        def normalize_date(value):
            raw=(value or '').strip()
            for date_fmt in ('%d/%m/%Y','%d-%m-%Y','%Y-%m-%d','%Y/%m/%d'):
                try:
                    return dt.datetime.strptime(raw[:10],date_fmt).strftime('%Y-%m-%d')
                except Exception:
                    pass
            return None

        def clear_form():
            for ctrl in [code,serial,cost_usd,pressure,
                         tread_outer_new,tread_inner_new,retirement_tread,projected_life_target]:
                ctrl.value=''
            entry_date.value=dt.date.today().strftime('%d/%m/%Y')

            for dropdown in [brand, size, design, compound, supplier]:
                dropdown.value=None
                try:
                    dropdown.text=''
                except Exception:
                    pass

            construction.value=None
            condition.value=None

        def save(e):
            code_value=(code.value or '').strip()
            serial_value=(serial.value or '').strip()
            if not code_value and not serial_value:
                return snack('Debe ingresar Código o Serie Fab. (al menos uno).', True)

            brand_value = catalog_value(brand)
            size_value = catalog_value(size)
            design_value = catalog_value(design)
            compound_value = catalog_value(compound)
            supplier_value = catalog_value(supplier)

            # Marca, medida, diseño, TRA y proveedor pueden quedar en blanco
            # cuando la información no está disponible. La regla maestra es
            # que exista Código O Serie Fab.; los demás campos se validan
            # únicamente cuando contienen un valor.

            date_iso=normalize_date(entry_date.value)
            if not date_iso:
                return snack('Fecha de ingreso inválida. Use dd/mm/aaaa.', True)

            cost=num(cost_usd.value)
            rec_pressure=num(pressure.value)
            new_ext=num(tread_outer_new.value)
            new_int=num(tread_inner_new.value)
            retirement=num(retirement_tread.value)
            life_target=num(projected_life_target.value)

            if cost is not None and cost < 0:
                return snack('Costo $ inválido.', True)
            if rec_pressure is not None and rec_pressure <= 0:
                return snack('Presión recomendada inválida.', True)
            if new_ext is not None and new_ext <= 0:
                return snack('Profundidad nueva EXT inválida.', True)
            if new_int is not None and new_int <= 0:
                return snack('Profundidad nueva INT inválida.', True)
            if retirement is not None and retirement < 0:
                return snack('Profundidad de retiro inválida.', True)
            if retirement is not None and new_ext is not None and new_int is not None and retirement >= min(float(new_ext), float(new_int)):
                return snack('La profundidad de retiro debe ser menor que la profundidad nueva.', True)
            if life_target is not None and life_target <= 0:
                return snack('Proyección de vida inválida.', True)

            new_tread_ref=max(float(new_ext), float(new_int)) if new_ext is not None and new_int is not None else None

            # Si se digitó un valor nuevo, se incorpora al catálogo y queda
            # disponible automáticamente para los siguientes registros.
            brand_value = save_catalog_value('brand', brand_value)
            size_value = save_catalog_value('size', size_value)
            design_value = save_catalog_value('design', design_value)
            compound_value = save_catalog_value('compound', compound_value)
            supplier_value = save_catalog_value('supplier', supplier_value)

            def refresh_catalog_dropdowns():
                for category, dropdown in [
                    ('brand', brand),
                    ('size', size),
                    ('design', design),
                    ('compound', compound),
                    ('supplier', supplier),
                ]:
                    dropdown.options = catalog_options(category)

            try:
                execute(
                    """INSERT INTO tires(
                           code,serial,brand,size,design,new_tread,recommended_pressure,
                           tread_inner,tread_outer,entry_date,cost_usd,compound,supplier,
                           new_tread_outer,new_tread_inner,construction_type,tire_condition,
                           retirement_tread,projected_life_target,projected_life,operation_id
                       ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                    (
                        code_value or None,
                        serial_value or None,
                        brand_value,
                        size_value,
                        design_value,
                        new_tread_ref,
                        rec_pressure,
                        new_int,
                        new_ext,
                        date_iso,
                        cost,
                        compound_value,
                        supplier_value,
                        new_ext,
                        new_int,
                        construction.value,
                        condition.value,
                        retirement,
                        life_target,
                        life_target,
                        op_id,
                    )
                )
                refresh_catalog_dropdowns()
                clear_form()
                snack('Neumático registrado correctamente.')
                refresh()
            except Exception as ex:
                snack(str(ex),True)

        refresh()
        title='8.2 NEUMÁTICOS' if not status_filter else f'Neumáticos: {status_filter}'
        subtitle='Registro maestro de neumáticos · Consulta y estado actual de cada neumático'
        blocks=[page_title(title,subtitle)]

        if not status_filter:
            blocks.append(card(ft.Column([
                ft.Text('Nuevo neumático',size=17,weight=ft.FontWeight.BOLD,color=TEXT_MAIN),
                ft.Text('Código o Serie Fab. es obligatorio. Los demás campos pueden quedar en blanco.',size=11,color=TEXT_MUTED),

                ft.Row([code,serial,entry_date,cost_usd],wrap=True,spacing=10,run_spacing=10),
                ft.Row([brand,size,design,compound],wrap=True,spacing=10,run_spacing=10),
                ft.Row([supplier,pressure,tread_outer_new,tread_inner_new],wrap=True,spacing=10,run_spacing=10),
                ft.Row([retirement_tread,projected_life_target,construction,condition],wrap=True,spacing=10,run_spacing=10),

                ft.Row([ft.Button('Registrar neumático',icon=ft.Icons.SAVE,on_click=save), ft.Button('Cargar Excel',icon=ft.Icons.UPLOAD_FILE,action=bulk_pick_action), ft.Button('Plantilla Excel',icon=ft.Icons.DESCRIPTION,url='/descargar-plantilla-excel')],wrap=True,spacing=8)
            ])))

        if status_filter:
            blocks.append(card(ft.Column([
                ft.Row([
                    ft.Text('Listado de neumáticos registrados',size=17,weight=ft.FontWeight.BOLD,color=TEXT_MAIN),
                    ft.Container(expand=True),search
                ],wrap=True),
                summary,
                ft.Row(
                    [ft.Container(content=master_table, width=1880)],
                    scroll=ft.ScrollMode.ALWAYS
                )
            ])))

        # El historial se conserva para las vistas operativas filtradas,
        # pero no se muestra dentro del Registro maestro de neumáticos.
        if status_filter:
            blocks.append(card(ft.Column([
                ft.Text('Historial operativo',size=17,weight=ft.FontWeight.BOLD,color=TEXT_MAIN),
                ft.Text('Filtra por equipo para ver juntos todos los movimientos de sus neumáticos.',size=11,color=TEXT_MUTED),
                ft.Row([history],scroll=ft.ScrollMode.AUTO)
            ])))

        content.content=ft.Column(blocks,scroll=ft.ScrollMode.AUTO,spacing=16)
        page.update()


    def service_menu_view():
        """Portada del Módulo 2 con el mismo patrón visual de tarjetas de los módulos 3-8."""
        def access_card(num, icon, accent, soft, title, desc, action):
            return ft.Container(
                width=292, height=260, bgcolor=soft,
                border=ft.Border.all(1, accent + '55'), border_radius=14, padding=18,
                on_click=lambda e: action(), ink=True,
                content=ft.Column([
                    ft.Row([
                        ft.Container(width=42, height=42, bgcolor=accent, border_radius=10,
                                     alignment=ft.Alignment.CENTER,
                                     content=ft.Text(num, color=ft.Colors.WHITE, size=12, weight=ft.FontWeight.BOLD)),
                        ft.Container(expand=True),
                        ft.Container(width=54, height=54, bgcolor=ft.Colors.WHITE, border_radius=12,
                                     alignment=ft.Alignment.CENTER,
                                     content=ft.Icon(icon, color=accent, size=29)),
                    ]),
                    ft.Container(height=5),
                    ft.Text(title, size=16, weight=ft.FontWeight.BOLD, color=TEXT_MAIN,
                            text_align=ft.TextAlign.CENTER),
                    ft.Text(desc, size=11, color=TEXT_MUTED, text_align=ft.TextAlign.CENTER),
                    ft.Container(expand=True),
                    ft.Container(height=38, bgcolor=accent, border_radius=9, alignment=ft.Alignment.CENTER,
                                 content=ft.Row([
                                     ft.Icon(ft.Icons.BAR_CHART, color=ft.Colors.WHITE, size=17),
                                     ft.Text('VER REPORTE', color=ft.Colors.WHITE, size=11, weight=ft.FontWeight.BOLD),
                                     ft.Icon(ft.Icons.CHEVRON_RIGHT, color=ft.Colors.WHITE, size=17),
                                 ], alignment=ft.MainAxisAlignment.CENTER, spacing=6)),
                ], spacing=10)
            )

        content.content=ft.Column([
            page_title('2. NEUMÁTICOS EN SERVICIO','Consulta técnica de neumáticos actualmente instalados'),
            card(ft.Column([
                ft.Row([
                    ft.Container(width=48,height=48,bgcolor='#EAF2FF',border_radius=24,
                                 alignment=ft.Alignment.CENTER,
                                 content=ft.Icon(ft.Icons.DIRECTIONS_CAR,color=NAV_ACCENT,size=26)),
                    ft.Column([
                        ft.Text('NEUMÁTICOS EN SERVICIO',size=17,weight=ft.FontWeight.BOLD,color=TEXT_MAIN),
                        ft.Text('Seleccione un reporte para visualizar el detalle.',size=11,color=TEXT_MUTED),
                    ],spacing=2)
                ],spacing=12),
                ft.Row([
                    access_card('2.1',ft.Icons.ASSESSMENT_OUTLINED,'#1565C0','#EEF5FF',
                                'REPORTE GENERAL','Vista general actual de todos los neumáticos instalados.',service_general_view),
                    access_card('2.2',ft.Icons.FACT_CHECK_OUTLINED,'#138A3D','#EEFAF2',
                                'REPORTE POR EQUIPO','Información técnica detallada de un solo equipo y sus posiciones.',service_equipment_report_view),
                    ft.Container(width=292,height=260), ft.Container(width=292,height=260),
                ],spacing=14,wrap=True)
            ],spacing=18),padding=18),
        ],scroll=ft.ScrollMode.AUTO,spacing=16)
        page.update()

    def service_general_view():
        """Consulta operativa de neumáticos actualmente instalados, aislada por operación."""
        op_id = active_operation_id()
        eq_rows = query("""
            SELECT DISTINCT e.id,e.code,e.brand,e.model,e.location,e.vehicle_type,e.tire_size
            FROM equipment e
            JOIN tires t ON t.equipment_id=e.id
            WHERE e.active=1 AND t.status='SERVICIO'
              AND e.operation_id=? AND t.operation_id=?
            ORDER BY e.code
        """, (op_id, op_id))
        ALL='__ALL__'
        eq_filter = ft.Dropdown(
            label='Equipo en servicio',
            width=220,
            value=ALL,
            options=[ft.dropdown.Option(key=ALL, text='Todos los equipos')] +
                    [ft.dropdown.Option(key=str(r['id']), text=r['code']) for r in eq_rows]
        )
        equipment_by_code = {str(r['code']).strip().upper(): str(r['id']) for r in eq_rows}
        equipment_ids = {str(r['id']) for r in eq_rows}
        equipment_state = {'id': ALL}
        equipment_rows_state = {'rows': None}
        mode_state = {'mode': None}  # None | 'search' | 'equipment'
        tire_filter = ft.Dropdown(
            label='Neumático',
            width=280,
            value=ALL,
            options=[ft.dropdown.Option(key=ALL, text='Todos los neumáticos')]
        )
        search = ft.TextField(
            label='Buscar código / serie',
            prefix_icon=ft.Icons.SEARCH,
            width=260
        )
        eq_info = ft.Text('', size=12, color=TEXT_MUTED)
        summary = ft.Text('', size=12, color=TEXT_MUTED)

        metrics = ft.Row([], wrap=True, spacing=12, run_spacing=12)
        position_grid = ft.Row([], wrap=True, spacing=12, run_spacing=12)

        # Dashboard visual tipo Power BI. Se alimenta exclusivamente de los
        # mismos datos calculados para la tabla de neumáticos en servicio.
        rem_chart_body = ft.Column([], spacing=7)
        hours_chart_body = ft.Column([], spacing=7)
        brand_pie_body = ft.Column([], spacing=8, horizontal_alignment=ft.CrossAxisAlignment.CENTER)
        pressure_pie_body = ft.Column([], spacing=8, horizontal_alignment=ft.CrossAxisAlignment.CENTER)
        valve_pie_body = ft.Column([], spacing=8, horizontal_alignment=ft.CrossAxisAlignment.CENTER)
        fleet_body = ft.Column([], spacing=8, horizontal_alignment=ft.CrossAxisAlignment.CENTER)

        def dashboard_bar(label, value, max_value, suffix='', decimals=1):
            try:
                val = float(value)
            except Exception:
                val = 0.0
            try:
                mx = max(float(max_value), 0.0001)
            except Exception:
                mx = 1.0
            ratio = max(0.0, min(1.0, val / mx))
            return ft.Row([
                ft.Text(label, width=72, size=10, weight=ft.FontWeight.BOLD, color=TEXT_MAIN),
                ft.Stack([
                    ft.Container(width=260, height=14, bgcolor='#E9EEF5', border_radius=7),
                    ft.Container(width=max(3, 260 * ratio), height=14, bgcolor=NAV_ACCENT, border_radius=7),
                ], width=260, height=14),
                ft.Text(f'{val:.{decimals}f}{suffix}', width=72, size=10, text_align=ft.TextAlign.RIGHT, color=TEXT_MAIN),
            ], spacing=8, vertical_alignment=ft.CrossAxisAlignment.CENTER)

        def _soft_dotted_grid(width=1200, segments=120):
            """Cuadrícula horizontal punteada suave, sin dependencias nuevas."""
            return ft.Row(
                [ft.Container(width=4, height=1, bgcolor='#D8E1EB') for _ in range(segments)],
                spacing=5,
                width=width,
                height=1,
            )

        def grouped_remanente_chart(groups):
            """Remanente: equipos en X, P1-P4 agrupadas y eje Y fijo 0-100%."""
            pos_order = ['P1', 'P2', 'P3', 'P4']
            if not any(v is not None for vals in groups.values() for v in vals.values()):
                return ft.Text('Sin datos suficientes para graficar.', size=11, color=TEXT_MUTED)

            chart_h = 150
            bar_w = 10
            ticks = [100, 80, 60, 40, 20, 0]
            y_axis = ft.Column(
                [ft.Text(f'{t}%', size=8.5, color=TEXT_MUTED) for t in ticks],
                height=chart_h,
                alignment=ft.MainAxisAlignment.SPACE_BETWEEN,
                horizontal_alignment=ft.CrossAxisAlignment.END,
            )

            group_blocks = []
            group_items = list(groups.items())
            for idx, (eq, pos_values) in enumerate(group_items):
                bars = []
                for pos in pos_order:
                    val = pos_values.get(pos)
                    if val is None:
                        bar = ft.Container(width=bar_w, height=2, bgcolor='#CBD5E1', border_radius=2)
                    else:
                        v = max(0.0, min(100.0, float(val)))
                        h = max(3, chart_h * v / 100.0)
                        bar = ft.Container(width=bar_w, height=h, bgcolor=NAV_ACCENT, border_radius=2)
                    bars.append(ft.Column([
                        ft.Container(height=chart_h, alignment=ft.Alignment.BOTTOM_CENTER, content=bar),
                        ft.Text(pos, size=7.2, weight=ft.FontWeight.BOLD, color=TEXT_MUTED,
                                text_align=ft.TextAlign.CENTER),
                    ], spacing=2, horizontal_alignment=ft.CrossAxisAlignment.CENTER))

                group_blocks.append(ft.Container(
                    width=76,
                    content=ft.Column([
                        ft.Row(bars, spacing=4, alignment=ft.MainAxisAlignment.CENTER,
                               vertical_alignment=ft.CrossAxisAlignment.END),
                        ft.Text(eq, size=9, weight=ft.FontWeight.BOLD,
                                color=TEXT_MAIN, text_align=ft.TextAlign.CENTER),
                    ], spacing=3, horizontal_alignment=ft.CrossAxisAlignment.CENTER),
                ))
                if idx < len(group_items) - 1:
                    group_blocks.append(ft.Container(
                        width=12,
                        height=chart_h + 28,
                        alignment=ft.Alignment.BOTTOM_CENTER,
                        content=ft.Text('|', size=12, weight=ft.FontWeight.BOLD, color=TEXT_MUTED),
                    ))

            # Separación compacta entre equipos; la barra vertical queda centrada en el espacio blanco.
            plot = ft.Stack([
                ft.Column(
                    [_soft_dotted_grid() for _ in ticks],
                    height=chart_h,
                    alignment=ft.MainAxisAlignment.SPACE_BETWEEN,
                    spacing=0,
                ),
                ft.Row(
                    group_blocks,
                    spacing=4,
                    alignment=ft.MainAxisAlignment.START,
                    vertical_alignment=ft.CrossAxisAlignment.END,
                ),
            ], height=chart_h + 28)

            return ft.Row([
                ft.Column([
                    ft.Text('% Rem.', size=8.5, color=TEXT_MUTED),
                    y_axis,
                    ft.Container(height=24),
                ], spacing=1, horizontal_alignment=ft.CrossAxisAlignment.END),
                ft.Container(expand=True, content=plot),
            ], spacing=8, vertical_alignment=ft.CrossAxisAlignment.START)

        def grouped_performance_chart(groups, performance_type):
            """Acumulado + restante proyectado, separado por unidad de rendimiento."""
            import math
            pos_order = ['P1', 'P2', 'P3', 'P4']
            unit = performance_unit_label(performance_type)
            if not groups:
                return ft.Text(f'Sin datos suficientes para graficar en {unit}.', size=11, color=TEXT_MUTED)
            all_totals = []
            for vals in groups.values():
                for data in vals.values():
                    if data:
                        all_totals.append(max(0.0, float(data.get('worked') or 0.0)) + max(0.0, float(data.get('remaining') or 0.0)))
            if not all_totals:
                return ft.Text(f'Sin datos suficientes para graficar en {unit}.', size=11, color=TEXT_MUTED)
            step = 1000 if performance_type == PERFORMANCE_HOURS else 10000
            max_total = max(all_totals)
            if max_total <= step and max_total > 0:
                step = max(1, int(math.ceil(max_total / 4.0)))
            y_max = max(step * 4, (int(math.floor(max_total / step)) + 1) * step)
            ticks = list(range(y_max, -1, -step))
            chart_h, bar_w = 150, 10
            y_axis = ft.Column([ft.Text(f'{t:,.0f}', size=8.5, color=TEXT_MUTED) for t in ticks], height=chart_h,
                               alignment=ft.MainAxisAlignment.SPACE_BETWEEN, horizontal_alignment=ft.CrossAxisAlignment.END)
            group_blocks=[]
            items=list(groups.items())
            for idx,(label,pos_values) in enumerate(items):
                bars=[]
                for pos in pos_order:
                    data=pos_values.get(pos)
                    if not data:
                        stacked=ft.Column([ft.Container(width=bar_w,height=2,bgcolor='#CBD5E1',border_radius=2)],spacing=0,height=chart_h,alignment=ft.MainAxisAlignment.END)
                    else:
                        worked=max(0.0,float(data.get('worked') or 0.0)); remaining=max(0.0,float(data.get('remaining') or 0.0))
                        wh=chart_h*min(worked,y_max)/y_max; rh=chart_h*min(remaining,max(0.0,y_max-worked))/y_max
                        seg=[]
                        if remaining>0: seg.append(ft.Container(width=bar_w,height=max(2,rh),bgcolor='#F59E0B',border_radius=ft.BorderRadius.only(top_left=2,top_right=2)))
                        if worked>0: seg.append(ft.Container(width=bar_w,height=max(3,wh),bgcolor=NAV_ACCENT,border_radius=(ft.BorderRadius.only(bottom_left=2,bottom_right=2) if remaining>0 else 2)))
                        stacked=ft.Column(seg,spacing=0,height=chart_h,alignment=ft.MainAxisAlignment.END,horizontal_alignment=ft.CrossAxisAlignment.CENTER)
                    bars.append(ft.Column([ft.Container(height=chart_h,alignment=ft.Alignment.BOTTOM_CENTER,content=stacked),ft.Text(pos,size=7.2,weight=ft.FontWeight.BOLD,color=TEXT_MUTED)],spacing=2,horizontal_alignment=ft.CrossAxisAlignment.CENTER))
                group_blocks.append(ft.Container(width=92,content=ft.Column([ft.Row(bars,spacing=4,alignment=ft.MainAxisAlignment.CENTER,vertical_alignment=ft.CrossAxisAlignment.END),ft.Text(label,size=8.5,weight=ft.FontWeight.BOLD,color=TEXT_MAIN,text_align=ft.TextAlign.CENTER,max_lines=2,overflow=ft.TextOverflow.ELLIPSIS)],spacing=3,horizontal_alignment=ft.CrossAxisAlignment.CENTER)))
                if idx<len(items)-1: group_blocks.append(ft.Container(width=10,height=chart_h+28,alignment=ft.Alignment.BOTTOM_CENTER,content=ft.Text('|',size=12,weight=ft.FontWeight.BOLD,color=TEXT_MUTED)))
            plot=ft.Stack([ft.Column([_soft_dotted_grid() for _ in ticks],height=chart_h,alignment=ft.MainAxisAlignment.SPACE_BETWEEN,spacing=0),ft.Row(group_blocks,spacing=4,alignment=ft.MainAxisAlignment.START,vertical_alignment=ft.CrossAxisAlignment.END)],height=chart_h+28)
            legend=ft.Row([ft.Row([ft.Container(width=10,height=10,bgcolor=NAV_ACCENT,border_radius=2),ft.Text(f'Acumulado ({unit})',size=9,color=TEXT_MUTED)],spacing=5),ft.Row([ft.Container(width=10,height=10,bgcolor='#F59E0B',border_radius=2),ft.Text(f'Restante proyectado ({unit})',size=9,color=TEXT_MUTED)],spacing=5)],spacing=16)
            return ft.Column([ft.Row([ft.Text(f'{unit} por modelo de equipo',size=9,weight=ft.FontWeight.BOLD,color=TEXT_MAIN),ft.Container(expand=True),ft.Text('P1 · P2 · P3 · P4',size=8.5,color=TEXT_MUTED)]),legend,ft.Row([ft.Column([ft.Text(unit,size=8.5,color=TEXT_MUTED),y_axis,ft.Container(height=24)],spacing=1,horizontal_alignment=ft.CrossAxisAlignment.END),ft.Container(content=ft.Row([plot],scroll=ft.ScrollMode.ALWAYS),expand=True)],spacing=8,vertical_alignment=ft.CrossAxisAlignment.START)],spacing=4)

        def donut_svg_chart(items, total=None, palette=None, legend_lines=None, center_label='Total'):
            """Dona grande centrada + leyenda inferior en 2 columnas con alto dinámico."""
            if not items:
                return ft.Text('Sin datos suficientes para graficar.', size=11, color=TEXT_MUTED)

            import base64, math

            if total is None:
                total = sum(count for _label, count in items)

            if palette is None:
                palette = ['#1D4ED8', '#16A34A', '#EA580C', '#A855F7', '#DC2626',
                           '#0D9488', '#CA8A04', '#475569']

            positive = [(str(label), int(count)) for label, count in items if int(count) > 0]

            # ============================================================
            # 1. CALCULAR AUTOMÁTICAMENTE LA LEYENDA
            # ============================================================
            # Dos columnas. El sistema determina cuántas filas necesita.
            legend_count = len(positive)
            legend_rows = max(1, math.ceil(legend_count / 2))

            # Altura por fila y márgenes de la leyenda.
            LEGEND_ROW_H = 20
            LEGEND_TOP_BOTTOM = 12
            legend_height = legend_rows * LEGEND_ROW_H + LEGEND_TOP_BOTTOM

            # ============================================================
            # 2. CALCULAR AUTOMÁTICAMENTE EL ALTO TOTAL
            # ============================================================
            # La dona conserva un tamaño grande y centrado.
            chart_width = 480
            donut_height = 290
            separator_height = 4

            # El alto se adapta a la cantidad real de elementos.
            total_height = donut_height + separator_height + legend_height

            cx, cy = chart_width / 2, 132
            radius, stroke = 108, 44
            circumference = 2 * math.pi * radius

            # ============================================================
            # 3. DIBUJAR LA DONA
            # ============================================================
            circles = []
            offset = 0.0

            for idx, (_label, count) in enumerate(positive):
                color = palette[idx % len(palette)]
                dash = circumference * (count / total) if total else 0.0
                gap = max(0.0, circumference - dash)

                circles.append(
                    f'<circle cx="{cx}" cy="{cy}" r="{radius}" fill="none" '
                    f'stroke="{color}" stroke-width="{stroke}" '
                    f'stroke-dasharray="{dash:.3f} {gap:.3f}" '
                    f'stroke-dashoffset="{-offset:.3f}" '
                    f'transform="rotate(-90 {cx} {cy})" />'
                )
                offset += dash

            svg = (
                f'<svg xmlns="http://www.w3.org/2000/svg" '
                f'width="{chart_width}" height="{donut_height}" '
                f'viewBox="0 0 {chart_width} {donut_height}">'
                f'<circle cx="{cx}" cy="{cy}" r="{radius}" fill="none" '
                f'stroke="#E2E8F0" stroke-width="{stroke}" />'
                + ''.join(circles) +
                f'<text x="{cx}" y="{cy-2}" text-anchor="middle" '
                f'font-family="Arial" font-size="30" font-weight="700" '
                f'fill="#172033">{total}</text>'
                f'<text x="{cx}" y="{cy+20}" text-anchor="middle" '
                f'font-family="Arial" font-size="10" fill="#64748B">'
                f'{center_label}</text>'
                '</svg>'
            )

            svg_b64 = base64.b64encode(svg.encode('utf-8')).decode('ascii')

            chart = ft.Container(
                width=chart_width,
                height=donut_height,
                alignment=ft.Alignment(0, 0),
                content=ft.Image(
                    src='data:image/svg+xml;base64,' + svg_b64,
                    width=chart_width,
                    height=donut_height,
                    fit=ft.BoxFit.CONTAIN,
                ),
            )

            # ============================================================
            # 4. LEYENDA INFERIOR EN DOS COLUMNAS
            # ============================================================
            legend_entries = []

            for idx, (label, count) in enumerate(positive):
                color = palette[idx % len(palette)]
                pct_value = (count / total * 100.0) if total else 0.0

                legend_entries.append(
                    ft.Row(
                        [
                            ft.Container(
                                width=10,
                                height=10,
                                bgcolor=color,
                                border_radius=2,
                            ),
                            ft.Text(
                                f'{label}: {count} ({pct_value:.1f}%)',
                                size=10,
                                color=TEXT_MAIN,
                                no_wrap=True,
                                overflow=ft.TextOverflow.ELLIPSIS,
                            ),
                        ],
                        spacing=6,
                        vertical_alignment=ft.CrossAxisAlignment.CENTER,
                    )
                )

            # Repartición automática en dos columnas.
            mid = math.ceil(len(legend_entries) / 2)
            left_entries = legend_entries[:mid]
            right_entries = legend_entries[mid:]

            while len(left_entries) < legend_rows:
                left_entries.append(ft.Container(height=LEGEND_ROW_H))

            while len(right_entries) < legend_rows:
                right_entries.append(ft.Container(height=LEGEND_ROW_H))

            legend = ft.Container(
                width=chart_width,
                height=legend_height,
                padding=ft.Padding(left=8, top=2, right=8, bottom=8),
                content=ft.Row(
                    [
                        ft.Column(
                            left_entries,
                            spacing=0,
                            expand=True,
                            horizontal_alignment=ft.CrossAxisAlignment.START,
                        ),
                        ft.Column(
                            right_entries,
                            spacing=0,
                            expand=True,
                            horizontal_alignment=ft.CrossAxisAlignment.START,
                        ),
                    ],
                    spacing=18,
                    vertical_alignment=ft.CrossAxisAlignment.START,
                ),
            )

            controls = [chart, ft.Container(height=separator_height), legend]

            if legend_lines:
                extra_height = len(legend_lines) * 18 + 24
                total_height += extra_height

                controls.append(
                    ft.Container(
                        width=chart_width,
                        padding=ft.Padding(left=8, top=2, right=8, bottom=4),
                        content=ft.Column(
                            [
                                ft.Text(
                                    'LEYENDA / CRITERIO',
                                    size=9,
                                    weight=ft.FontWeight.BOLD,
                                    color=TEXT_MAIN,
                                ),
                                *[
                                    ft.Text(line, size=9, color=TEXT_MAIN)
                                    for line in legend_lines
                                ],
                            ],
                            spacing=3,
                        ),
                    )
                )

            # ============================================================
            # 5. EL CONTENEDOR EXTERIOR TAMBIÉN CRECE
            # ============================================================
            return ft.Container(
                width=chart_width,
                height=total_height,
                alignment=ft.Alignment(0, 0),
                content=ft.Column(
                    controls,
                    spacing=0,
                    horizontal_alignment=ft.CrossAxisAlignment.CENTER,
                ),
            )

        def horizontal_bar_chart(items, title, subtitle='', max_items=20):
            """Barras horizontales con altura adaptable."""
            if not items:
                return ft.Text('Sin datos suficientes para graficar.', size=11, color=TEXT_MUTED)

            import math
            ordered = sorted(
                [(str(label), int(count)) for label, count in items if int(count) > 0],
                key=lambda x: (-x[1], x[0])
            )[:max_items]

            total = sum(c for _, c in ordered) or 1
            max_value = max(c for _, c in ordered) or 1
            palette = ['#1D4ED8','#16A34A','#EA580C','#A855F7',
                       '#DC2626','#0D9488','#CA8A04','#475569']

            rows = []
            for i, (label, count) in enumerate(ordered):
                color = palette[i % len(palette)]
                rows.append(
                    ft.Row([
                        ft.Container(
                            width=125,
                            content=ft.Text(label, size=10, color=TEXT_MAIN,
                                            no_wrap=True,
                                            overflow=ft.TextOverflow.ELLIPSIS),
                        ),
                        ft.Container(
                            width=max(35, 270 * count / max_value),
                            height=22,
                            bgcolor=color,
                            border_radius=4,
                            alignment=ft.Alignment(1, 0),
                            padding=ft.Padding(left=5, top=0, right=6, bottom=0),
                            content=ft.Text(str(count), size=10,
                                            weight=ft.FontWeight.BOLD,
                                            color='#FFFFFF'),
                        ),
                    ], spacing=6)
                )

            legend_items = [
                ft.Text(f'{label}: {count} ({count / total * 100:.1f}%)',
                        size=9, color=TEXT_MUTED, no_wrap=True,
                        overflow=ft.TextOverflow.ELLIPSIS)
                for label, count in ordered
            ]

            mid = math.ceil(len(legend_items) / 2)
            left, right = legend_items[:mid], legend_items[mid:]
            rows_needed = max(len(left), len(right))
            while len(left) < rows_needed:
                left.append(ft.Container(height=16))
            while len(right) < rows_needed:
                right.append(ft.Container(height=16))

            legend = ft.Row([
                ft.Column(left, spacing=2, expand=True),
                ft.Column(right, spacing=2, expand=True),
            ], spacing=12)

            dynamic_height = 300

            return ft.Container(
                height=300,
                content=ft.Column([
                    ft.Text(title, size=13, weight=ft.FontWeight.BOLD, color=TEXT_MAIN),
                    ft.Text(subtitle, size=9, color=TEXT_MUTED),
                    ft.Container(height=4),
                    ft.Column(rows, spacing=5),
                    ft.Divider(height=8, color='#E2E8F0'),
                    legend,
                ], spacing=3),
            )

        def brand_bar_chart(brand_counts):
            ordered = sorted(brand_counts.items(), key=lambda x: (-x[1], x[0]))
            return horizontal_bar_chart(
                ordered,
                'DISTRIBUCIÓN POR MARCA',
                'Cantidad de neumáticos actualmente en servicio',
            )

        def fleet_vehicle_chart(operation_id):
            """Flota completa del taller activo: característica + modelo + cantidad."""
            rows = query(
                """
                SELECT
                    COALESCE(NULLIF(TRIM(vehicle_type), ''), 'SIN CARACTERÍSTICA') AS characteristic,
                    COALESCE(NULLIF(TRIM(model), ''), 'SIN MODELO') AS model,
                    COUNT(*) AS qty
                FROM equipment
                WHERE operation_id=?
                GROUP BY 1,2
                ORDER BY 1,2
                """,
                (operation_id,)
            )

            items = [
                (f"{str(r['characteristic']).strip()} · {str(r['model']).strip()}", int(r['qty']))
                for r in rows
            ]

            return horizontal_bar_chart(
                items,
                'FLOTA VEHICULAR',
                'Equipos del taller activo por característica y modelo',
                max_items=30,
            )

        def pressure_pie_chart(counts):
            items = [
                ('± 5 psi (OK)', counts.get('green', 0)),
                ('> 5 a 10 psi', counts.get('orange', 0)),
                ('> 10 psi', counts.get('red', 0)),
                ('Sobrepresión > 20%', counts.get('purple', 0)),
            ]
            return donut_svg_chart(
                items,
                palette=['#16A34A', '#F59E0B', '#DC2626', '#7C3AED'],
            )

        def valve_status_chart(ops_rows):
            """Gráfico SI/NO de tapa de válvula para el panel 2.1.
            Eje X = SI / NO; altura = cantidad; debajo se muestran las posiciones.
            """
            yes_pos=[]
            no_pos=[]
            for r, od in ops_rows:
                pos = str(r['position'] or '—').strip()
                if str(od.get('valve_cap') or '').strip().upper() == 'SI':
                    yes_pos.append(pos)
                else:
                    no_pos.append(pos)

            values=[('SÍ',len(yes_pos),'#28B7E6'),('NO',len(no_pos),'#EF4444')]
            max_count=max([v[1] for v in values] + [1])
            chart_w=290
            chart_h=175
            base_y=122
            plot_top=20
            bar_w=58
            xs=[72, 170]
            layers=[ft.Container(left=0,top=plot_top,width=chart_w,height=1,bgcolor='#D5DDE6')]
            # Ejes y guías ligeras
            for tick in range(max_count+1):
                y=base_y-(tick/max_count)*88 if max_count else base_y
                layers.append(ft.Container(left=32,top=y,width=220,height=1,bgcolor='#E5EAF0'))
                layers.append(ft.Container(left=10,top=y-7,width=18,height=14,alignment=ft.Alignment.CENTER,
                                           content=ft.Text(str(tick),size=8,color=TEXT_MUTED,text_align=ft.TextAlign.CENTER)))
            for x,(label,count,color) in zip(xs, values):
                bh=(count/max_count)*88 if max_count else 0
                y=base_y-bh
                layers.append(ft.Container(left=x,top=y,width=bar_w,height=max(bh,2),bgcolor=color,
                                           border_radius=ft.BorderRadius(7,7,0,0),
                                           alignment=ft.Alignment.TOP_CENTER,
                                           content=ft.Container(padding=ft.Padding(left=2,top=4,right=2,bottom=2),
                                                                content=ft.Text(str(count),size=13,weight=ft.FontWeight.BOLD,color=ft.Colors.WHITE,text_align=ft.TextAlign.CENTER))))
                layers.append(ft.Container(left=x-5,top=base_y+7,width=bar_w+10,
                                           alignment=ft.Alignment.CENTER,
                                           content=ft.Text(label,size=11,weight=ft.FontWeight.BOLD,color=TEXT_MAIN,text_align=ft.TextAlign.CENTER)))

            plot=ft.Stack(controls=layers,width=chart_w,height=chart_h)
            yes_text=', '.join(yes_pos) if yes_pos else 'Ninguna'
            no_text=', '.join(no_pos) if no_pos else 'Ninguna'
            return ft.Column([
                ft.Row([plot],alignment=ft.MainAxisAlignment.CENTER),
                ft.Text(f'SÍ: {len(yes_pos)} neumático(s) · Posiciones: {yes_text}',size=8.5,color='#15803D',weight=ft.FontWeight.BOLD,text_align=ft.TextAlign.CENTER),
                ft.Text(f'NO: {len(no_pos)} neumático(s) · Posiciones: {no_text}',size=8.5,color='#B91C1C',weight=ft.FontWeight.BOLD,text_align=ft.TextAlign.CENTER),
            ],spacing=3,horizontal_alignment=ft.CrossAxisAlignment.CENTER)

        def dashboard_card(title, subtitle, body, height=None):
            return ft.Container(
                expand=True,
                height=height,
                bgcolor=CARD_BG,
                border=ft.Border.all(1, '#E2E8F0'),
                border_radius=14,
                padding=16,
                content=ft.Column([
                    ft.Text(title, size=14, weight=ft.FontWeight.BOLD, color=TEXT_MAIN),
                    ft.Text(subtitle, size=10, color=TEXT_MUTED),
                    ft.Divider(height=10, color='#E2E8F0'),
                    body,
                ], spacing=5),
            )

        # Dashboard 2.1: tres columnas.
        dashboard = ft.Row([
            ft.Container(
                expand=True,
                content=dashboard_card(
                    'DISTRIBUCIÓN POR MARCA',
                    'Cantidad de neumáticos actualmente en servicio',
                    brand_pie_body,
                         height=300,
                ),
            ),
            ft.Container(
                expand=True,
                content=dashboard_card(
                    'FLOTA VEHICULAR',
                    'Flota completa del taller activo',
                    fleet_body,
                         height=300,
                ),
            ),
            ft.Container(
                expand=True,
                content=ft.Column([
                    dashboard_card(
                        'EVALUACIÓN DE PRESIÓN (PSI)',
                        'Diferencia entre presión actual y recomendada',
                        pressure_pie_body,
                         height=300,
                    ),
                    dashboard_card(
                        'EVALUACIÓN DE EVALUACIÓN DE TAPA VÁLVULASS',
                        'Estado de tapa de válvula en neumáticos en servicio',
                        valve_pie_body,
                    ),
                ], spacing=12),
            ),
        ], spacing=12, vertical_alignment=ft.CrossAxisAlignment.START)

        # Tabla técnica compacta: encabezado fijo + desplazamiento vertical interno.
        # El contenedor horizontal usa ScrollMode.ALWAYS para mantener disponible
        # la barra horizontal durante todo el recorrido de la tabla.
        service_columns = [
            'EQUIPO',
            'POSICIÓN',
            'CÓDIGO',
            'SERIE',
            'MARCA',
            'MEDIDA',
            'DISEÑO',
            'RTD\nINICIAL',
            'RTD\nEXT',
            'RTD\nINT',
            'DIFERENCIA\nRTD',
            'VIDA ÚTIL\n(%)',
            'FECHA DE\nINSPECCIÓN',
            'HORÓMETRO DE\nINSPECCIÓN',
            'HORAS DE\nRODADO',
            'Hs/mm',
            'HORAS\nPROYECTADAS',
            'HORAS DE RODADO +\nPROYECCIÓN',
        ]
        service_widths = [78,58,70,96,78,80,88,72,62,62,80,132,98,106,86,68,100,190]
        service_total_width = sum(service_widths)

        def service_cell(value, width, header=False, bold=False, bgcolor=None, content=None):
            return ft.Container(
                width=width,
                height=44 if header else 34,
                bgcolor=('#173B5E' if header else bgcolor),
                padding=ft.Padding(left=4, top=0, right=4, bottom=0),
                alignment=ft.Alignment(0, 0),
                content=(content if content is not None else ft.Text(
                    str(value),
                    size=8.7 if header else 9.5,
                    color='#FFFFFF' if header else TEXT_MAIN,
                    weight=ft.FontWeight.BOLD if header or bold else None,
                    text_align=ft.TextAlign.CENTER,
                    max_lines=2 if header else 1,
                )),
                border=ft.Border(
                    bottom=ft.BorderSide(1, '#C9D5E2' if header else '#D7DEE8'),
                    right=ft.BorderSide(1, '#5E7891' if header else '#DCE5ED'),
                ),
            )

        def life_graph_cell(percent, width, bgcolor):
            if percent is None:
                return service_cell('—', width, bgcolor=bgcolor)
            p = max(0.0, min(100.0, float(percent)))
            track_w = max(42.0, width - 48.0)
            fill_w = max(2.0, track_w * p / 100.0)
            graph = ft.Row([
                ft.Stack([
                    ft.Container(width=track_w, height=13, bgcolor='#E5E7EB', border_radius=2),
                    ft.Container(width=fill_w, height=13, bgcolor='#22A06B', border_radius=2),
                ], width=track_w, height=13),
                ft.Text(f'{p:.0f}%', width=34, size=10, weight=ft.FontWeight.BOLD,
                        text_align=ft.TextAlign.RIGHT, color=TEXT_MAIN),
            ], spacing=5, alignment=ft.MainAxisAlignment.CENTER,
               vertical_alignment=ft.CrossAxisAlignment.CENTER)
            return service_cell('', width, bgcolor=bgcolor, content=graph)

        def hours_projection_graph_cell(worked, projected, width, bgcolor):
            if worked is None and projected is None:
                return service_cell('—', width, bgcolor=bgcolor)
            w = max(0.0, float(worked or 0.0))
            p = max(0.0, float(projected or 0.0))
            total = w + p
            track_w = max(80.0, width - 18.0)
            blue_w = (track_w * w / total) if total > 0 else 0.0
            orange_w = max(0.0, track_w - blue_w) if total > 0 else 0.0
            graph = ft.Column([
                ft.Stack([
                    ft.Container(width=track_w, height=15, bgcolor='#E5E7EB', border_radius=2),
                    ft.Row([
                        ft.Container(width=max(2.0, blue_w) if w > 0 else 0, height=15,
                                     bgcolor=NAV_ACCENT,
                                     border_radius=ft.BorderRadius.only(top_left=2,bottom_left=2)),
                        ft.Container(width=max(2.0, orange_w) if p > 0 else 0, height=15,
                                     bgcolor='#F47C20',
                                     border_radius=ft.BorderRadius.only(top_right=2,bottom_right=2)),
                    ], spacing=0),
                ], width=track_w, height=15),
                ft.Text(f'{w:,.0f} + {p:,.0f} = {total:,.0f} h', size=8.2, color=TEXT_MUTED,
                        text_align=ft.TextAlign.CENTER),
            ], spacing=1, horizontal_alignment=ft.CrossAxisAlignment.CENTER)
            return service_cell('', width, bgcolor=bgcolor, content=graph)

        service_header = ft.Row(
            [service_cell(label, width, header=True) for label, width in zip(service_columns, service_widths)],
            spacing=0,
        )
        # Tabla técnica con viewport fijo: la barra horizontal queda visible desde
        # el momento en que se abre 2.1, sin necesidad de llegar al final de las filas.
        service_body = ft.Column([], spacing=0, scroll=ft.ScrollMode.ALWAYS)
        service_table_inner = ft.Column([
            service_header,
            ft.Container(height=420, content=service_body),
        ], spacing=0, width=service_total_width)
        service_table_view = ft.Container(
            height=480,
            content=ft.Row(
                [service_table_inner],
                scroll=ft.ScrollMode.ALWAYS,
                spacing=0,
                vertical_alignment=ft.CrossAxisAlignment.START,
            ),
        )
        # Historial operativo retirado de este módulo; los datos históricos se conservan en la base.


        def fmt(v, dec=0):
            if v in (None, ''):
                return ''
            try:
                f = float(v)
                if dec == 0 and f.is_integer():
                    return str(int(f))
                return f'{f:.{dec}f}'
            except Exception:
                return str(v)

        def pct(v):
            return '' if v is None else f'{v:.1f}%'

        visible_tire_ids = []
        export_rows_state = {'rows': []}

        def selected_tire_id():
            try:
                return int(tire_filter.value) if tire_filter.value not in (None, '', ALL) else None
            except Exception:
                return None

        def effective_tire_id(allow_search_auto=True):
            """Devuelve el neumático activo según el modo de consulta.

            Modo búsqueda: solo auto-selecciona cuando el texto deja un único resultado.
            Modo equipo: solo acepta un neumático elegido manualmente.
            Sin modo activo: no habilita eventos.
            """
            mode = mode_state['mode']
            if mode == 'equipment':
                return selected_tire_id()
            if mode == 'search':
                term = (search.value or '').strip()
                if allow_search_auto and term and len(visible_tire_ids) == 1:
                    return visible_tire_ids[0]
                return None
            return None

        # V10: caché por refresco para evitar 3 consultas SQLite por cada neumático.
        # Con 33 neumáticos, la vista pasaba de ~100 consultas a solo 3 consultas masivas.
        operational_cache_state = {'inst': {}, 'last': {}, 'insp': {}}

        def tire_operational_data(r):
            tid = int(r['id'])
            eid = r['equipment_id']
            inst_row = operational_cache_state['inst'].get((tid, eid))
            if inst_row is None and eid is None:
                inst_row = operational_cache_state['inst'].get((tid, None))
            inst_meter = inst_row['meter'] if inst_row else None
            inst_date = inst_row['event_date'] if inst_row else ''
            last_row = operational_cache_state['last'].get(tid)
            inspection_row = operational_cache_state['insp'].get(tid) or last_row

            current_meter = r['current_meter']
            worked = None
            if current_meter is not None and inst_meter is not None:
                worked = max(0, float(current_meter) - float(inst_meter))

            vals = [v for v in (r['tread_inner'], r['tread_outer'])
                    if isinstance(v, (int, float))]
            min_tread = min(vals) if vals else None
            new_tread = r['new_tread']
            wear = None
            rem = None
            hpmm = None
            if min_tread is not None and new_tread not in (None, 0):
                wear = max(0, float(new_tread) - float(min_tread))
                rem = max(0, min(100, float(min_tread) / float(new_tread) * 100))
                if worked is not None and wear > 0:
                    hpmm = worked / wear

            last_pressure = inspection_row['pressure'] if inspection_row and inspection_row['pressure'] is not None else None
            # Tapa Válvula representa el estado del ÚLTIMO evento registrado.
            valve_source = last_row or inspection_row
            valve_raw = str(valve_source['valve_cap'] or '').strip().upper() if valve_source and 'valve_cap' in valve_source.keys() else ''
            if valve_raw in ('SI','SÍ'):
                valve_cap = 'SI'
            elif valve_raw == 'NO':
                valve_cap = 'NO'
            else:
                note_text = str(valve_source['notes'] or '').upper() if valve_source and 'notes' in valve_source.keys() else ''
                valve_cap = 'SI' if ('TAPA' in note_text or 'VALVULA' in note_text or 'VÁLVULA' in note_text) else 'NO'
            return {
                'inst_meter': inst_meter, 'inst_date': inst_date, 'worked': worked,
                'min_tread': min_tread, 'wear': wear, 'rem': rem, 'hpmm': hpmm,
                'last_pressure': last_pressure,
                'last_event': last_row['event_code'] if last_row else '',
                'last_date': last_row['event_date'] if last_row else '',
                'inspection_date': inspection_row['event_date'] if inspection_row else '',
                'inspection_meter': inspection_row['meter'] if inspection_row else None,
                'valve_cap': valve_cap,
            }

        def goto_movement(event_code=None):
            tid = effective_tire_id()
            if not tid:
                return snack('Seleccione un neumático para continuar.', True)
            session['movement_tire_id'] = str(tid)
            session['movement_event'] = event_code
            nav.selected_index = 1
            select(1)

        # Accesos directos a todos los eventos desde la vista de neumáticos en servicio.
        # INST permanece bloqueado porque todo neumático mostrado aquí ya está instalado.
        event_icons = {
            'INST': ft.Icons.ADD_CIRCLE_OUTLINE,
            'INSP': ft.Icons.CHECK_CIRCLE_OUTLINE,
            'INSC': ft.Icons.FACT_CHECK_OUTLINED,
            'ROT': ft.Icons.SYNC_ALT,
            'INVE': ft.Icons.SWAP_HORIZ,
            'DINS': ft.Icons.REMOVE_CIRCLE_OUTLINE,
            'REPA': ft.Icons.HANDYMAN_OUTLINED,
            'BAJA': ft.Icons.DELETE_OUTLINE,
        }
        event_buttons = {}
        for code in ['INST','INSP','INSC','ROT','INVE','DINS','REPA','BAJA']:
            event_buttons[code] = ft.OutlinedButton(
                code,
                icon=event_icons[code],
                tooltip=EVENTS.get(code, code),
                on_click=lambda e, ec=code: goto_movement(ec),
                disabled=True
            )

        event_buttons['INST'].tooltip = 'Instalación bloqueada: el neumático ya está EN SERVICIO'
        event_buttons['ROT'].tooltip = 'Rotación bloqueada: funcionalidad pendiente de definición'

        def refresh_tire_options(rows):
            current = tire_filter.value
            opts = [ft.dropdown.Option(key=ALL, text='Todos los neumáticos')]
            for r in rows:
                opts.append(ft.dropdown.Option(key=str(r['id']), text=f"{r['code']} | P{r['position'] or '-'} | {r['serial'] or 's/serie'}"))
            tire_filter.options = opts
            valid = set([ALL] + [str(r['id']) for r in rows])
            if current not in valid:
                tire_filter.value = ALL

        def refresh(e=None):
            selected_value = str(tire_filter.value or ALL)

            # El selector Neumático depende EXCLUSIVAMENTE del equipo seleccionado.
            # Si on_equipment_change ya construyó la lista, reutilizamos exactamente
            # esos registros y no volvemos a mezclarlos con el buscador.
            if equipment_rows_state['rows'] is not None:
                option_rows = list(equipment_rows_state['rows'])
            else:
                options_sql = """
                    SELECT t.*,e.code equipment_code,e.brand equipment_brand,e.model equipment_model,
                           e.location equipment_location,e.vehicle_type,e.tire_size equipment_tire_size
                    FROM tires t
                    LEFT JOIN equipment e ON e.id=t.equipment_id
                    WHERE t.status='SERVICIO' AND t.operation_id=?
                """
                options_params = [op_id]
                if equipment_state['id'] not in (None, '', ALL):
                    options_sql += ' AND t.equipment_id=?'
                    options_params.append(int(equipment_state['id']))
                options_sql += " ORDER BY COALESCE(e.code,''), CAST(COALESCE(NULLIF(t.position,''),'999') AS INTEGER), t.code"
                option_rows = query(options_sql, tuple(options_params))

            refresh_tire_options(option_rows)

            valid_ids = {str(r['id']) for r in option_rows}
            if selected_value not in ('', ALL) and selected_value in valid_ids:
                tire_filter.value = selected_value
                tid = int(selected_value)
            else:
                if tire_filter.value in (None, '') or str(tire_filter.value) not in valid_ids | {ALL}:
                    tire_filter.value = ALL
                tid = selected_tire_id()

            # La consulta visible sí responde a búsqueda + equipo + neumático.
            rows = list(option_rows)
            term = (search.value or '').strip()
            if term:
                term_l = term.lower()
                rows = [r for r in rows if
                        term_l in str(r['code'] or '').lower() or
                        term_l in str(r['serial'] or '').lower() or
                        term_l in str(r['brand'] or '').lower() or
                        term_l in str(r['design'] or '').lower()]
            if tid is not None:
                rows = [r for r in rows if int(r['id']) == int(tid)]

            visible_tire_ids.clear()
            visible_tire_ids.extend([int(r['id']) for r in rows])

            # Habilitar eventos solo si:
            # 1) se seleccionó manualmente un neumático, o
            # 2) Buscar código / serie dejó un único resultado.
            # Elegir únicamente un equipo NO activa eventos.
            active_tid = effective_tire_id(allow_search_auto=True)
            can_open = active_tid is not None and str(active_tid) in valid_ids
            for code, btn in event_buttons.items():
                # INST: bloqueado porque el neumático ya está en servicio.
                # ROT: bloqueado hasta definir su funcionalidad.
                btn.disabled = (not can_open) or code in ('INST', 'ROT')

            # V10: precarga masiva de ocurrencias para los neumáticos visibles.
            # Evita el patrón N+1 y mejora notablemente el tiempo de apertura/refresco.
            operational_cache_state['inst'].clear()
            operational_cache_state['last'].clear()
            operational_cache_state['insp'].clear()
            row_ids = [int(r['id']) for r in rows]
            if row_ids:
                marks = ','.join('?' for _ in row_ids)
                # Optimización crítica: NO cargar todo el historial de occurrences.
                # El reporte general solo necesita el último evento, la última INSP/INSC
                # y la última INST de cada neumático visible. Con muchos movimientos
                # históricos, traer todas las ocurrencias hacía que Argentum tardara
                # demasiado en abrir el reporte.
                latest_rows = query(
                    f"""SELECT id,tire_id,equipment_id,event_date,event_code,meter,
                               tread_inner,tread_outer,pressure,pressure_condition,
                               location,reason,notes,valve_cap
                        FROM (
                            SELECT o.*, ROW_NUMBER() OVER (
                                PARTITION BY o.tire_id
                                ORDER BY o.id DESC
                            ) rn
                            FROM occurrences o
                            WHERE o.operation_id=? AND o.tire_id IN ({marks})
                        ) q
                        WHERE rn=1""",
                    (op_id, *row_ids)
                )
                for o in latest_rows:
                    operational_cache_state['last'][int(o['tire_id'])] = o

                insp_rows = query(
                    f"""SELECT id,tire_id,equipment_id,event_date,event_code,meter,
                               tread_inner,tread_outer,pressure,pressure_condition,
                               location,reason,notes,valve_cap
                        FROM (
                            SELECT o.*, ROW_NUMBER() OVER (
                                PARTITION BY o.tire_id
                                ORDER BY o.id DESC
                            ) rn
                            FROM occurrences o
                            WHERE o.operation_id=?
                              AND o.tire_id IN ({marks})
                              AND o.event_code IN ('INSP','INSC')
                        ) q
                        WHERE rn=1""",
                    (op_id, *row_ids)
                )
                for o in insp_rows:
                    operational_cache_state['insp'][int(o['tire_id'])] = o

                inst_rows = query(
                    f"""SELECT id,tire_id,equipment_id,event_date,event_code,meter
                        FROM (
                            SELECT o.*, ROW_NUMBER() OVER (
                                PARTITION BY o.tire_id,o.equipment_id
                                ORDER BY o.id DESC
                            ) rn
                            FROM occurrences o
                            WHERE o.operation_id=?
                              AND o.event_code='INST'
                              AND o.tire_id IN ({marks})
                        ) q
                        WHERE rn=1""",
                    (op_id, *row_ids)
                )
                for o in inst_rows:
                    operational_cache_state['inst'][(int(o['tire_id']), o['equipment_id'])] = o

            ops = [(r, tire_operational_data(r)) for r in rows]
            service_body.controls = []
            export_rows_state['rows'] = []
            rem_values = []
            worked_values = []

            previous_equipment = None
            equipment_group_index = -1
            equipment_group_colors = ['#F4F7FA', '#EAF1F7']
            equipment_tire_counts = {}
            for rr in rows:
                if rr['equipment_id'] is not None:
                    equipment_tire_counts[rr['equipment_id']] = equipment_tire_counts.get(rr['equipment_id'], 0) + 1

            # Cada equipo se presenta como un acordeón. Al abrir 2.1 solo se
            # muestran los encabezados de los equipos; los neumáticos quedan
            # ocultos hasta pulsar el botón deslizable del equipo.
            current_group = None
            current_children = None

            def make_equipment_header(equipment_code, group_detail, tire_count, children_container):
                expanded = {'value': False}

                def toggle(_e):
                    expanded['value'] = not expanded['value']
                    children_container.visible = expanded['value']
                    toggle_icon.icon = (
                        ft.Icons.KEYBOARD_ARROW_UP if expanded['value']
                        else ft.Icons.KEYBOARD_ARROW_DOWN
                    )
                    page.update()

                toggle_icon = ft.IconButton(
                    icon=ft.Icons.KEYBOARD_ARROW_DOWN,
                    icon_size=22,
                    icon_color='#173B5E',
                    tooltip='Mostrar / ocultar neumáticos',
                    on_click=toggle,
                )

                return ft.Container(
                    height=42,
                    bgcolor='#DCE7F2',
                    padding=ft.Padding(left=8, top=0, right=8, bottom=0),
                    alignment=ft.Alignment.CENTER_LEFT,
                    content=ft.Row([
                        toggle_icon,
                        ft.Icon(ft.Icons.PRECISION_MANUFACTURING_OUTLINED, size=17, color='#173B5E'),
                        ft.Text(
                            f'EQUIPO: {equipment_code}  ·  {group_detail}',
                            size=10.5, weight=ft.FontWeight.BOLD, color='#173B5E', expand=True,
                        ),
                        ft.Container(
                            bgcolor='#173B5E', border_radius=12,
                            padding=ft.Padding(left=10, top=4, right=10, bottom=4),
                            content=ft.Text(
                                f'{tire_count} NEUMÁTICO' + ('S' if tire_count != 1 else '') +
                                ' CARGADO' + ('S' if tire_count != 1 else ''),
                                size=9.5, weight=ft.FontWeight.BOLD, color='#FFFFFF'
                            )
                        ),
                    ], spacing=6),
                )

            for r, od in ops:
                if od['rem'] is not None:
                    rem_values.append(od['rem'])
                if od['worked'] is not None:
                    worked_values.append(od['worked'])

                equipment_code = r['equipment_code'] or '—'

                if equipment_code != previous_equipment:
                    equipment_group_index += 1
                    if current_group is not None:
                        service_body.controls.append(current_group)

                    if previous_equipment is not None:
                        service_body.controls.append(ft.Container(height=5, bgcolor='#DCE6EF'))

                    tire_count = equipment_tire_counts.get(r['equipment_id'], 0)
                    vehicle_type = str(r['vehicle_type'] or '').strip()
                    equipment_model = str(r['equipment_model'] or '').strip()
                    group_detail = ' · '.join(x for x in (vehicle_type, equipment_model) if x) or 'Sin tipo/modelo'

                    current_children = ft.Column(
                        [],
                        spacing=0,
                        visible=False,
                    )
                    header = make_equipment_header(
                        equipment_code, group_detail, tire_count, current_children
                    )
                    current_group = ft.Column(
                        [header, current_children],
                        spacing=0,
                    )

                row_bgcolor = '#F4F7FA' if (equipment_group_index % 2 == 0) else '#EAF1F7'
                previous_equipment = equipment_code

                original_outer = r['new_tread_outer'] if 'new_tread_outer' in r.keys() else None
                original_inner = r['new_tread_inner'] if 'new_tread_inner' in r.keys() else None
                if original_outer is None:
                    original_outer = r['new_tread']
                if original_inner is None:
                    original_inner = r['new_tread']

                original_vals = [v for v in (original_outer, original_inner) if isinstance(v, (int, float))]
                current_vals = [v for v in (r['tread_outer'], r['tread_inner']) if isinstance(v, (int, float))]
                original_min = min(original_vals) if original_vals else None
                current_min = min(current_vals) if current_vals else None
                rem_pct = None
                wear_mm = None
                hs_mm = None
                if original_min not in (None, 0) and current_min is not None:
                    rem_pct = max(0, min(100, float(current_min) / float(original_min) * 100))
                    wear_mm = max(0, float(original_min) - float(current_min))
                    if od['worked'] is not None and wear_mm > 0:
                        hs_mm = float(od['worked']) / wear_mm

                cost_hour = None
                if r['cost_usd'] is not None and od['worked'] is not None and float(od['worked']) > 0:
                    cost_hour = float(r['cost_usd']) / float(od['worked'])

                projected_hours = None
                retirement_tread = r['retirement_tread'] if 'retirement_tread' in r.keys() else None
                if (hs_mm is not None and current_min is not None and retirement_tread is not None):
                    remaining_mm = max(0.0, float(current_min) - float(retirement_tread))
                    projected_hours = remaining_mm * float(hs_mm)

                rtd_ext = r['tread_outer'] if isinstance(r['tread_outer'], (int, float)) else None
                rtd_int = r['tread_inner'] if isinstance(r['tread_inner'], (int, float)) else None
                rtd_diff = abs(float(rtd_ext) - float(rtd_int)) if rtd_ext is not None and rtd_int is not None else None

                row_values = [
                    equipment_code,
                    f"P{r['position']}" if r['position'] not in (None, '') else '—',
                    r['code'] or '—',
                    r['serial'] or '—',
                    r['brand'] or '—',
                    r['size'] or r['equipment_tire_size'] or '—',
                    r['design'] or '—',
                    fmt(original_min) if original_min is not None else '—',
                    fmt(rtd_ext) if rtd_ext is not None else '—',
                    fmt(rtd_int) if rtd_int is not None else '—',
                    f"{rtd_diff:.1f}" if rtd_diff is not None else '—',
                    None,
                    format_date(od['inspection_date']) or '—',
                    fmt(od['inspection_meter']) or '—',
                    fmt(od['worked']) or '—',
                    f"{hs_mm:.2f}" if hs_mm is not None else '—',
                    f"{projected_hours:,.0f}" if projected_hours is not None else '—',
                    None,
                ]

                row_cells = []
                for idx, v in enumerate(row_values):
                    if idx == 11:
                        row_cells.append(life_graph_cell(rem_pct, service_widths[idx], row_bgcolor))
                    elif idx == 17:
                        row_cells.append(hours_projection_graph_cell(
                            od['worked'], projected_hours, service_widths[idx], row_bgcolor
                        ))
                    else:
                        row_cells.append(service_cell(
                            v, service_widths[idx], bold=idx in (0,2), bgcolor=row_bgcolor
                        ))
                current_children.controls.append(ft.Row(row_cells, spacing=0))

                export_rows_state['rows'].append({
                    'EQUIPO': equipment_code,
                    'POSICIÓN': f"P{r['position']}" if r['position'] not in (None, '') else '—',
                    'CÓDIGO': r['code'] or '—',
                    'SERIE': r['serial'] or '—',
                    'MARCA': r['brand'] or '—',
                    'MEDIDA': r['size'] or r['equipment_tire_size'] or '—',
                    'DISEÑO': r['design'] or '—',
                    'RTD INICIAL': original_min,
                    'RTD EXT': rtd_ext,
                    'RTD INT': rtd_int,
                    'DIFERENCIA RTD': rtd_diff,
                    'VIDA ÚTIL (%)': rem_pct,
                    'FECHA DE INSPECCIÓN': format_date(od['inspection_date']) or '—',
                    'HORÓMETRO DE INSPECCIÓN': od['inspection_meter'],
                    'HORAS DE RODADO': od['worked'],
                    'Hs/mm': hs_mm,
                    'HORAS PROYECTADAS': projected_hours,
                    'HORAS DE RODADO + PROYECCIÓN': (
                        float(od['worked'] or 0) + float(projected_hours or 0)
                    ) if (od['worked'] is not None or projected_hours is not None) else None,
                })

            if current_group is not None:
                service_body.controls.append(current_group)

            total_service = len(rows)
            eq_count = len({r['equipment_id'] for r in rows if r['equipment_id'] is not None})
            avg_rem = sum(rem_values) / len(rem_values) if rem_values else None
            current_meter = max([float(r['current_meter']) for r in rows if r['current_meter'] is not None], default=None)
            metrics.controls = [
                metric_card('En servicio', total_service, ft.Icons.TIRE_REPAIR, 'Neumáticos del filtro actual'),
                metric_card('Equipos con neumáticos', eq_count, ft.Icons.PRECISION_MANUFACTURING_OUTLINED, 'Flota del filtro actual'),
                metric_card('Remanente promedio', pct(avg_rem) if avg_rem is not None else '—', ft.Icons.ASSESSMENT_OUTLINED, 'Sobre profundidad nueva'),
            ]

            # Los gráficos REMANENTE POR POSICIÓN y ACUMULADO + RESTANTE
            # quedan ocultos en esta reestructuración. No se recalculan para
            # reducir el tiempo de carga de talleres con muchos equipos.
            rem_chart_body.controls = []
            hours_chart_body.controls = []

            brand_counts = {}
            for r, _od in ops:
                brand = str(r['brand'] or 'Sin marca').strip() or 'Sin marca'
                brand_counts[brand] = brand_counts.get(brand, 0) + 1
            brand_pie_body.controls = [brand_bar_chart(brand_counts)]
            fleet_body.controls = [fleet_vehicle_chart(op_id)]

            # Distribución de presión respecto a la presión recomendada.
            # Sobrepresión es una categoría exclusiva: presión actual > 120% de la recomendada.
            # Los demás neumáticos conservan el criterio previo por diferencia absoluta:
            # <=5 psi verde; >5 y <=10 psi naranja; >10 psi rojo.
            pressure_counts = {'green': 0, 'orange': 0, 'red': 0, 'purple': 0}
            for r, od in ops:
                rec = r['recommended_pressure']
                act = od['last_pressure']
                if rec is None or act is None:
                    continue
                try:
                    rec_f = float(rec)
                    act_f = float(act)
                    if rec_f <= 0:
                        continue
                    diff = abs(act_f - rec_f)
                except Exception:
                    continue
                if act_f > rec_f * 1.20:
                    pressure_counts['purple'] += 1
                elif diff <= 5:
                    pressure_counts['green'] += 1
                elif diff <= 10:
                    pressure_counts['orange'] += 1
                else:
                    pressure_counts['red'] += 1
            pressure_pie_body.controls = [pressure_pie_chart(pressure_counts)]

            valve_counts = {'yes': 0, 'no': 0}
            for _r, od in ops:
                if str(od['valve_cap'] or '').strip().upper() == 'SI':
                    valve_counts['yes'] += 1
                else:
                    valve_counts['no'] += 1
            valve_pie_body.controls = [valve_status_chart(ops)]

            if equipment_state['id'] not in (None, '', ALL):
                er = query('SELECT * FROM equipment WHERE id=? AND operation_id=?', (int(equipment_state['id']), op_id))
                if er:
                    q = er[0]
                    eq_info.value = (
                        f"{q['code']} · {(q['brand'] or '').strip()} {(q['model'] or '').strip()} · "
                        f"{q['vehicle_type'] or ''} · Ubicación: {q['location'] or ''} · "
                        f"Medida: {q['tire_size'] or ''}"
                    )
                else:
                    eq_info.value = ''
            else:
                eq_info.value = 'Vista consolidada de todos los equipos con neumáticos en servicio.'

            position_grid.controls = []
            for r, od in ops:
                position_grid.controls.append(
                    card(ft.Column([
                        ft.Row([
                            ft.Container(
                                width=42, height=42, border_radius=21,
                                bgcolor='#EAF2FF',
                                alignment=ft.Alignment.CENTER,
                                content=ft.Text(f"P{r['position'] or '-'}", weight=ft.FontWeight.BOLD, color=NAV_ACCENT)
                            ),
                            ft.Column([
                                ft.Text(r['code'], size=18, weight=ft.FontWeight.BOLD, color=TEXT_MAIN),
                                ft.Text(r['serial'] or 'Sin serie', size=11, color=TEXT_MUTED),
                            ], spacing=1),
                        ]),
                        ft.Text(f"{r['brand'] or ''} · {r['size'] or ''} · {r['design'] or ''}", size=11, color=TEXT_MUTED),
                        ft.Row([
                            ft.Column([ft.Text('Cocada', size=10, color=TEXT_MUTED), ft.Text(f"{fmt(r['tread_inner'])}/{fmt(r['tread_outer'])} mm", weight=ft.FontWeight.BOLD)]),
                            ft.Column([ft.Text('Remanente', size=10, color=TEXT_MUTED), ft.Text(pct(od['rem']) or '—', weight=ft.FontWeight.BOLD)]),
                            ft.Column([ft.Text('Horas trab.', size=10, color=TEXT_MUTED), ft.Text(fmt(od['worked']) or '—', weight=ft.FontWeight.BOLD)]),
                        ], alignment=ft.MainAxisAlignment.SPACE_BETWEEN),
                        ft.Text(f"Último evento: {od['last_event'] or '—'} · {format_date(od['last_date']) or '—'}", size=10, color=TEXT_MUTED),
                    ], spacing=9), width=300)
                )

            summary.value = f"{len(rows)} neumático(s) mostrado(s)"
            page.update()

        def build_excel_bytes():
            '''Genera un XLSX real con los registros visibles del 2.1.'''
            import io, zipfile, html
            headers=[h.replace('\n',' ') for h in service_columns]
            keys=['EQUIPO','POSICIÓN','CÓDIGO','SERIE','MARCA','MEDIDA','DISEÑO','RTD INICIAL','RTD EXT','RTD INT','DIFERENCIA RTD','VIDA ÚTIL (%)','FECHA DE INSPECCIÓN','HORÓMETRO DE INSPECCIÓN','HORAS DE RODADO','Hs/mm','HORAS PROYECTADAS','HORAS DE RODADO + PROYECCIÓN']
            rows=export_rows_state['rows']
            def xlcol(n):
                out=''
                while n:
                    n,r=divmod(n-1,26); out=chr(65+r)+out
                return out
            def esc(v): return html.escape(str('' if v is None else v),quote=False)
            def cell(ref,v,style):
                if isinstance(v,(int,float)) and not isinstance(v,bool):
                    return f'<c r="{ref}" s="{style}"><v>{v}</v></c>'
                return f'<c r="{ref}" s="{style}" t="inlineStr"><is><t>{esc(v)}</t></is></c>'
            xml_rows=['<row r="1" ht="34" customHeight="1">']
            for i,h in enumerate(headers,1): xml_rows.append(cell(f'{xlcol(i)}1',h,1))
            xml_rows.append('</row>')
            for ri,row in enumerate(rows,2):
                xml_rows.append(f'<row r="{ri}" ht="22" customHeight="1">')
                for ci,key in enumerate(keys,1):
                    style=2 if ri%2==0 else 3
                    if key=='VIDA ÚTIL (%)': style=4 if ri%2==0 else 5
                    xml_rows.append(cell(f'{xlcol(ci)}{ri}',row.get(key) if row.get(key) is not None else '—',style))
                xml_rows.append('</row>')
            widths=[12,10,12,16,13,13,14,12,10,10,14,15,17,18,15,12,17,24]
            cols=''.join(f'<col min="{i}" max="{i}" width="{w}" customWidth="1"/>' for i,w in enumerate(widths,1))
            last=xlcol(len(headers)); lastrow=max(1,len(rows)+1)
            sheet='<?xml version="1.0" encoding="UTF-8" standalone="yes"?><worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetViews><sheetView workbookViewId="0"><pane ySplit="1" topLeftCell="A2" activePane="bottomLeft" state="frozen"/></sheetView></sheetViews><cols>'+cols+'</cols><sheetData>'+''.join(xml_rows)+'</sheetData><autoFilter ref="A1:'+last+str(lastrow)+'"/></worksheet>'
            styles='''<?xml version="1.0" encoding="UTF-8" standalone="yes"?><styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><fonts count="3"><font><sz val="10"/><name val="Calibri"/></font><font><b/><color rgb="FFFFFFFF"/><sz val="10"/><name val="Calibri"/></font><font><b/><color rgb="FF137A4B"/><sz val="10"/><name val="Calibri"/></font></fonts><fills count="5"><fill><patternFill patternType="none"/></fill><fill><patternFill patternType="gray125"/></fill><fill><patternFill patternType="solid"><fgColor rgb="FF173B5E"/></patternFill></fill><fill><patternFill patternType="solid"><fgColor rgb="FFF4F7FA"/></patternFill></fill><fill><patternFill patternType="solid"><fgColor rgb="FFEAF1F7"/></patternFill></fill></fills><borders count="2"><border/><border><left style="thin"><color rgb="FFD7E0E8"/></left><right style="thin"><color rgb="FFD7E0E8"/></right><top style="thin"><color rgb="FFD7E0E8"/></top><bottom style="thin"><color rgb="FFD7E0E8"/></bottom></border></borders><cellXfs count="6"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/><xf numFmtId="0" fontId="1" fillId="2" borderId="1" applyAlignment="1"><alignment horizontal="center" vertical="center" wrapText="1"/></xf><xf numFmtId="0" fontId="0" fillId="3" borderId="1" applyAlignment="1"><alignment horizontal="center" vertical="center"/></xf><xf numFmtId="0" fontId="0" fillId="4" borderId="1" applyAlignment="1"><alignment horizontal="center" vertical="center"/></xf><xf numFmtId="0" fontId="2" fillId="3" borderId="1" applyAlignment="1"><alignment horizontal="center" vertical="center"/></xf><xf numFmtId="0" fontId="2" fillId="4" borderId="1" applyAlignment="1"><alignment horizontal="center" vertical="center"/></xf></cellXfs></styleSheet>'''
            ct='''<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/><Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/><Override PartName="/xl/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/></Types>'''
            rel='''<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/></Relationships>'''
            wb='''<?xml version="1.0" encoding="UTF-8" standalone="yes"?><workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><sheets><sheet name="Reporte 2.1" sheetId="1" r:id="rId1"/></sheets></workbook>'''
            wbr='''<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/><Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/></Relationships>'''
            bio=io.BytesIO()
            with zipfile.ZipFile(bio,'w',zipfile.ZIP_DEFLATED) as z:
                z.writestr('[Content_Types].xml',ct); z.writestr('_rels/.rels',rel); z.writestr('xl/workbook.xml',wb); z.writestr('xl/_rels/workbook.xml.rels',wbr); z.writestr('xl/styles.xml',styles); z.writestr('xl/worksheets/sheet1.xml',sheet)
            return bio.getvalue()

        def build_service_pdf_bytes():
            '''PDF A3 apaisado con las 18 columnas visibles.'''
            headers=[h.replace('\n',' ') for h in service_columns]
            keys=['EQUIPO','POSICIÓN','CÓDIGO','SERIE','MARCA','MEDIDA','DISEÑO','RTD INICIAL','RTD EXT','RTD INT','DIFERENCIA RTD','VIDA ÚTIL (%)','FECHA DE INSPECCIÓN','HORÓMETRO DE INSPECCIÓN','HORAS DE RODADO','Hs/mm','HORAS PROYECTADAS','HORAS DE RODADO + PROYECCIÓN']
            rows=export_rows_state['rows']; pw,ph=1191,842; margin=24; usable=pw-2*margin
            weights=[6,5,6,8,7,7,7,6,5,5,7,8,8,9,7,6,8,11]; sc=usable/sum(weights); widths=[w*sc for w in weights]
            def esc(v):
                raw=str('—' if v is None else v).replace('–','-').replace('—','-').encode('cp1252','replace').decode('latin-1')
                return raw.replace('\\','\\\\').replace('(','\\(').replace(')','\\)')
            def val(v,k):
                if v is None:return '—'
                if isinstance(v,(int,float)):
                    if k=='VIDA ÚTIL (%)':return f'{v:.0f}%'
                    if k in ('DIFERENCIA RTD','Hs/mm'):return f'{v:.2f}'
                    return f'{v:,.0f}'
                return str(v)
            pages=[]; cmds=[]; y=ph-margin
            def text(x,yy,t,size=5.6,bold=False,color='0.10 0.16 0.22'):
                cmds.append(f'BT {"/F2" if bold else "/F1"} {size} Tf {color} rg {x:.1f} {yy:.1f} Td ({esc(t)}) Tj ET')
            def box(x,yy,w,h,fill):
                cmds.append(f'{fill} rg {x:.1f} {yy:.1f} {w:.1f} {h:.1f} re f'); cmds.append(f'0.78 0.83 0.88 RG .35 w {x:.1f} {yy:.1f} {w:.1f} {h:.1f} re S')
            def header():
                nonlocal y
                text(margin,ph-26,'MEGASOFTIRE - 2.1 REPORTE GENERAL DE NEUMÁTICOS EN SERVICIO',12,True,'0.09 0.23 0.37'); y=ph-48; h=31; x=margin
                for hd,w in zip(headers,widths):
                    box(x,y-h,w,h,'0.09 0.23 0.37'); short=hd if len(hd)<=max(6,int(w/4.5)) else hd[:max(4,int(w/4.5)-1)]+'...'; text(x+2,y-18,short,5.2,True,'1 1 1'); x+=w
                y-=h
            def newpage():
                nonlocal cmds,y
                if cmds: pages.append('\n'.join(cmds))
                cmds=[]; header()
            newpage()
            for ri,row in enumerate(rows):
                h=20
                if y-h<margin+10:newpage()
                x=margin; fill='0.96 0.97 0.98' if ri%2==0 else '0.92 0.95 0.97'
                for k,w in zip(keys,widths):
                    box(x,y-h,w,h,fill); t=val(row.get(k),k); mc=max(3,int(w/4.7)); t=t if len(t)<=mc else t[:max(1,mc-1)]+'...'; text(x+2,y-13,t,5.4,k in ('EQUIPO','CÓDIGO','VIDA ÚTIL (%)'),'0.05 0.42 0.24' if k=='VIDA ÚTIL (%)' else '0.10 0.16 0.22'); x+=w
                y-=h
            if cmds:pages.append('\n'.join(cmds))
            objs=[b'<< /Type /Catalog /Pages 2 0 R >>',b'',b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>',b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica-Bold /Encoding /WinAnsiEncoding >>']; pids=[]
            for st in pages:
                bs=st.encode('latin-1','replace'); cid=len(objs)+1; objs.append(b'<< /Length '+str(len(bs)).encode()+b' >>\nstream\n'+bs+b'\nendstream'); pid=len(objs)+1; pids.append(pid); objs.append(f'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 {pw} {ph}] /Resources << /Font << /F1 3 0 R /F2 4 0 R >> >> /Contents {cid} 0 R >>'.encode())
            objs[1]=f'<< /Type /Pages /Kids [{" ".join(f"{p} 0 R" for p in pids)}] /Count {len(pids)} >>'.encode(); out=bytearray(b'%PDF-1.4\n%\xe2\xe3\xcf\xd3\n'); offs=[0]
            for i,o in enumerate(objs,1):offs.append(len(out));out.extend(f'{i} 0 obj\n'.encode());out.extend(o);out.extend(b'\nendobj\n')
            xr=len(out);out.extend(f'xref\n0 {len(objs)+1}\n'.encode());out.extend(b'0000000000 65535 f \n')
            for off in offs[1:]:out.extend(f'{off:010d} 00000 n \n'.encode())
            out.extend(f'trailer\n<< /Size {len(objs)+1} /Root 1 0 R >>\nstartxref\n{xr}\n%%EOF'.encode());return bytes(out)

        async def download_service_excel(e):
            try:
                if not export_rows_state['rows']: return snack('No hay datos visibles para exportar.',True)
                name='MegaSoftire_2.1_Reporte_General.xlsx'; await ft.FilePicker().save_file(file_name=name,file_type=ft.FilePickerFileType.CUSTOM,allowed_extensions=['xlsx'],src_bytes=build_excel_bytes()); snack(f'Excel preparado: {name}')
            except Exception as ex: snack(f'No se pudo descargar Excel: {ex}',True)

        async def download_service_pdf(e):
            try:
                if not export_rows_state['rows']: return snack('No hay datos visibles para exportar.',True)
                name='MegaSoftire_2.1_Reporte_General.pdf'; await ft.FilePicker().save_file(file_name=name,file_type=ft.FilePickerFileType.CUSTOM,allowed_extensions=['pdf'],src_bytes=build_service_pdf_bytes()); snack(f'PDF preparado: {name}')
            except Exception as ex: snack(f'No se pudo descargar PDF: {ex}',True)

        def resolve_equipment_from_event(e):
            candidates = []
            ctrl = getattr(e, 'control', None)
            if ctrl is not None:
                candidates.append(getattr(ctrl, 'value', None))
            candidates.append(getattr(e, 'data', None))

            for raw in candidates:
                if raw in (None, ''):
                    continue
                val = str(raw).strip()
                if val == ALL or val.lower() == 'todos los equipos':
                    return ALL
                if val in equipment_ids:
                    return val
                eid = equipment_by_code.get(val.upper())
                if eid:
                    return eid
            return ALL

        def on_equipment_change(e):
            """Cambio de equipo totalmente independiente del buscador.

            Flujo:
            equipo -> limpiar neumático -> bloquear eventos -> consultar solo ese
            equipment_id -> reconstruir selector -> refrescar vista.
            """
            eid = resolve_equipment_from_event(e)
            equipment_state['id'] = eid
            eq_filter.value = eid

            # Nunca conservar un neumático seleccionado del equipo anterior.
            tire_filter.value = ALL
            visible_tire_ids.clear()
            for code, btn in event_buttons.items():
                btn.disabled = True

            sql = """
                SELECT t.*,e.code equipment_code,e.brand equipment_brand,e.model equipment_model,
                       e.location equipment_location,e.vehicle_type,e.tire_size equipment_tire_size
                FROM tires t
                LEFT JOIN equipment e ON e.id=t.equipment_id
                WHERE t.status='SERVICIO' AND t.operation_id=?
            """
            params = [op_id]
            if eid != ALL:
                sql += ' AND t.equipment_id=?'
                params.append(int(eid))
            sql += " ORDER BY CAST(COALESCE(NULLIF(t.position,''),'999') AS INTEGER),t.code"
            option_rows = query(sql, tuple(params))

            # Guardamos exactamente la lista del equipo; refresh no la recalcula
            # a partir del buscador ni de un neumático anterior.
            equipment_rows_state['rows'] = list(option_rows)
            tire_filter.options = [ft.dropdown.Option(key=ALL, text='Todos los neumáticos')] + [
                ft.dropdown.Option(
                    key=str(r['id']),
                    text=f"{r['code']} | P{r['position'] or '-'} | {r['serial'] or 's/serie'}"
                ) for r in option_rows
            ]
            tire_filter.value = ALL
            refresh()

        def set_mode(mode):
            """Activa un único camino de consulta y bloquea el otro."""
            mode_state['mode'] = mode
            if mode == 'search':
                eq_filter.disabled = True
                tire_filter.disabled = True
                equipment_state['id'] = ALL
                equipment_rows_state['rows'] = None
                eq_filter.value = ALL
                tire_filter.value = ALL
            elif mode == 'equipment':
                search.disabled = True
                search.value = ''
                tire_filter.disabled = False
            else:
                search.disabled = False
                eq_filter.disabled = False
                tire_filter.disabled = False

        def reset_query_modes():
            """ESC: vuelve al estado inicial y elimina cualquier selección activa."""
            mode_state['mode'] = None
            search.value = ''
            search.disabled = False
            eq_filter.disabled = False
            eq_filter.value = ALL
            equipment_state['id'] = ALL
            equipment_rows_state['rows'] = None
            tire_filter.disabled = False
            tire_filter.value = ALL
            visible_tire_ids.clear()
            for code, btn in event_buttons.items():
                btn.disabled = True
            refresh()

        def on_search_focus(e):
            if not search.disabled:
                set_mode('search')
                refresh()

        def on_search_change(e):
            if mode_state['mode'] != 'search':
                set_mode('search')
            refresh(e)

        def on_equipment_focus(e):
            if not eq_filter.disabled:
                set_mode('equipment')
                page.update()

        # Conservamos la lógica de filtrado por equipo, pero activando primero
        # su modo exclusivo y eliminando cualquier búsqueda previa.
        original_on_equipment_change = on_equipment_change
        def on_equipment_change_mode(e):
            set_mode('equipment')
            original_on_equipment_change(e)

        def on_tire_change(e):
            if mode_state['mode'] != 'equipment':
                return
            selected = getattr(e, 'data', None)
            if selected not in (None, ''):
                tire_filter.value = str(selected)
            elif getattr(e, 'control', None) is not None:
                tire_filter.value = str(e.control.value or ALL)
            refresh(e)

        def on_service_keyboard(e):
            key = str(getattr(e, 'key', '') or '').upper()
            if key in ('ESCAPE', 'ESC'):
                reset_query_modes()

        search.on_focus = on_search_focus
        search.on_change = on_search_change
        eq_filter.on_focus = on_equipment_focus
        eq_filter.on_change = on_equipment_change_mode
        tire_filter.on_change = on_tire_change
        page.on_keyboard_event = on_service_keyboard

        refresh()

        content.content = ft.Column([
            page_title(
                'Neumáticos en servicio',
                'Estado actual, posiciones y horas de trabajo por equipo'
            ),
            metrics,
            dashboard,
            card(ft.Column([
                ft.Row([
                    ft.Text('Detalle técnico de neumáticos instalados', size=17, weight=ft.FontWeight.BOLD, color=TEXT_MAIN),
                    ft.Container(expand=True),
                    summary,
                    ft.IconButton(icon=ft.Icons.TABLE_VIEW_OUTLINED,icon_color='#168A55',tooltip='Descargar Excel',on_click=download_service_excel),
                    ft.IconButton(icon=ft.Icons.PICTURE_AS_PDF_OUTLINED,icon_color='#C0392B',tooltip='Descargar PDF',on_click=download_service_pdf),
                ], spacing=4),
                ft.Text(
                    'Una fila por neumático instalado. Cada equipo muestra P1, P2, P3 y P4 en orden, con separación antes del siguiente equipo.',
                    size=10, color=TEXT_MUTED
                ),
                service_table_view
            ]))        ], scroll=ft.ScrollMode.AUTO, spacing=16)
        page.update()

    def service_equipment_report_view():
        op_id = active_operation_id()
        """2.2 Reporte por equipo: ficha y detalle de todas las posiciones instaladas."""
        eqs=query("SELECT id,code FROM equipment WHERE active=1 AND operation_id=? ORDER BY code", (op_id,))
        selector=ft.Dropdown(label='Seleccione el equipo',width=250,
            options=[ft.dropdown.Option(key=str(r['id']),text=r['code']) for r in eqs])
        info=ft.Column([],spacing=8)
        table_area=ft.Column([],spacing=10)
        rem_chart_22=ft.Column([],spacing=6)
        hours_chart_22=ft.Column([],spacing=6)
        pressure_chart_22=ft.Column([],spacing=6)
        valve_chart_22=ft.Column([],spacing=6)
        pressure_history_22=ft.Column([],spacing=6)

        def metric(title,value,subtitle,accent='#1565C0'):
            return ft.Container(width=220,height=92,bgcolor=ft.Colors.WHITE,border_radius=12,
                border=ft.Border.all(1,'#D8E1EB'),padding=12,
                content=ft.Column([ft.Text(title,size=10,color=TEXT_MUTED,weight=ft.FontWeight.BOLD),
                                   ft.Text(str(value),size=22,color=accent,weight=ft.FontWeight.BOLD),
                                   ft.Text(subtitle,size=9,color=TEXT_MUTED)],spacing=2))

        def refresh(e=None):
            if not selector.value:
                info.controls=[ft.Text('Seleccione un equipo para generar el reporte.',color=TEXT_MUTED)]
                table_area.controls=[]; rem_chart_22.controls=[]; hours_chart_22.controls=[]; pressure_chart_22.controls=[]; valve_chart_22.controls=[]; page.update(); return
            eid=int(selector.value)
            eq_rows=query('SELECT * FROM equipment WHERE id=? AND operation_id=?',(eid,op_id));
            if not eq_rows:
                return
            eq=eq_rows[0]
            rows=query("""SELECT t.* FROM tires t WHERE t.equipment_id=? AND t.operation_id=? AND t.status='SERVICIO'
                          ORDER BY CAST(COALESCE(NULLIF(t.position,''),'999') AS INTEGER),t.code""",(eid,op_id))
            info.controls=[ft.Row([
                ft.Column([ft.Text('Código',size=9,color=TEXT_MUTED),ft.Text(eq['code'] or '—',weight=ft.FontWeight.BOLD)],width=125),
                ft.Column([ft.Text('Marca',size=9,color=TEXT_MUTED),ft.Text(eq['brand'] or '—',weight=ft.FontWeight.BOLD)],width=145),
                ft.Column([ft.Text('Modelo',size=9,color=TEXT_MUTED),ft.Text(eq['model'] or '—',weight=ft.FontWeight.BOLD)],width=145),
                ft.Column([ft.Text('Tipo',size=9,color=TEXT_MUTED),ft.Text(eq['vehicle_type'] or '—',weight=ft.FontWeight.BOLD)],width=125),
                ft.Column([ft.Text('Motor',size=9,color=TEXT_MUTED),ft.Text((eq['motor_type'] if 'motor_type' in eq.keys() else None) or '—',weight=ft.FontWeight.BOLD)],width=125),
                ft.Column([ft.Text('Ubicación',size=9,color=TEXT_MUTED),ft.Text(eq['location'] or '—',weight=ft.FontWeight.BOLD)],width=145),
            ],wrap=True,spacing=10)]
            data=[]; rems=[]; costs=[]; pressures_ok=0
            rem_by_pos={}; hours_by_pos={}; pressure_by_pos={}; valve_by_pos={}
            for r in rows:
                occ=query("""SELECT * FROM occurrences WHERE tire_id=? AND operation_id=? ORDER BY id DESC LIMIT 1""",(r['id'],op_id))
                o=occ[0] if occ else None
                insp=query("""SELECT * FROM occurrences WHERE tire_id=? AND operation_id=? AND event_code IN ('INSP','INSC') ORDER BY id DESC LIMIT 1""",(r['id'],op_id))
                io=insp[0] if insp else None
                # Horas acumuladas: usar exactamente el mismo criterio de 2.1.
                # Se toma el horómetro del último evento INST del neumático en el
                # equipo actual y se resta del horómetro actual guardado en tires.
                inst=query("""SELECT meter FROM occurrences
                              WHERE tire_id=? AND operation_id=? AND event_code='INST'
                                AND (? IS NULL OR equipment_id=?)
                              ORDER BY event_date DESC,id DESC LIMIT 1""",
                           (r['id'],op_id,eid,eid))
                inst_meter=inst[0]['meter'] if inst else None
                current_meter=r['current_meter']
                worked=None
                if current_meter is not None and inst_meter is not None:
                    try:
                        worked=max(0.0,float(current_meter)-float(inst_meter))
                    except Exception:
                        worked=None
                newvals=[x for x in (r['new_tread_outer'],r['new_tread_inner'],r['new_tread']) if isinstance(x,(int,float))]
                curvals=[x for x in (r['tread_outer'],r['tread_inner']) if isinstance(x,(int,float))]
                rem=(min(curvals)/min(newvals)*100) if curvals and newvals and min(newvals)>0 else None
                if rem is not None: rems.append(rem)
                cph=(float(r['cost_usd'])/worked) if r['cost_usd'] is not None and worked and worked>0 else None
                if cph is not None: costs.append(cph)
                pactual=io['pressure'] if io and io['pressure'] is not None else None
                prec=r['recommended_pressure']
                if pactual is not None and prec not in (None,0) and abs(float(pactual)-float(prec))/float(prec)<=0.05: pressures_ok+=1
                poskey=f"P{r['position']}" if r['position'] else '—'
                # Los gráficos deben representar TODAS las posiciones cargadas.
                # No se limita a P1-P4: algunos equipos tienen P5, P6, ... P11, P12, etc.
                if poskey != '—':
                    if rem is not None: rem_by_pos[poskey]=rem
                    remaining_hours=0.0
                    if worked is not None and newvals and curvals:
                        original_min=min(newvals); current_min=min(curvals)
                        wear_mm=max(0.0,float(original_min)-float(current_min))
                        retirement=r['retirement_tread'] if 'retirement_tread' in r.keys() else None
                        if wear_mm>0 and retirement is not None:
                            remaining_hours=max(0.0,float(current_min)-float(retirement))*(float(worked)/wear_mm)
                    if worked is not None: hours_by_pos[poskey]={'worked':float(worked),'remaining':float(remaining_hours)}
                    pressure_by_pos[poskey]={'actual':float(pactual) if pactual is not None else None,'recommended':float(prec) if prec is not None else None}
                    valve_source = o or io
                    valve_raw=str(valve_source['valve_cap'] or '').strip().upper() if valve_source and 'valve_cap' in valve_source.keys() else ''
                    if valve_raw in ('SI','SÍ'):
                        valve_by_pos[poskey]='SI'
                    elif valve_raw == 'NO':
                        valve_by_pos[poskey]='NO'
                    else:
                        note_text=str(valve_source['notes'] or '').upper() if valve_source and 'notes' in valve_source.keys() else ''
                        valve_by_pos[poskey]='SI' if ('TAPA' in note_text or 'VALVULA' in note_text or 'VÁLVULA' in note_text) else 'NO'
                data.append([
                    f"P{r['position']}" if r['position'] else '—',r['code'] or '—',r['serial'] or '—',r['brand'] or '—',
                    r['size'] or '—',r['design'] or '—',f"{worked:.0f}" if worked is not None else '—',
                    f"$ {cph:.2f}/h" if cph is not None else '—',
                    f"{min(newvals):.0f}" if newvals else '—',
                    f"{r['tread_outer'] or 0:g}/{r['tread_inner'] or 0:g}" if curvals else '—',
                    f"{rem:.1f}%" if rem is not None else '—',f"{pactual:g}" if pactual is not None else '—',
                    f"{prec:g}" if prec is not None else '—',r['tire_condition'] or '—',
                    (o['event_code'] if o else '—'),format_date(io['event_date']) if io else '—'
                ])
            headers=['POS.','CÓDIGO','SERIE','MARCA','MEDIDA','DISEÑO','HRS ACUM.','COSTO X HORA','COCADA ORIG.','COCADA EXT/INT','% REM.','PSI ACT.','PSI REC.','CONDICIÓN','ÚLT. EVENTO','FECHA ÚLT. INSP.']
            widths=[55,75,105,90,80,85,80,95,85,100,70,65,65,90,85,100]
            head=ft.Row([ft.Container(width=widths[i],height=44,bgcolor=NAV_BG,padding=6,alignment=ft.Alignment.CENTER,
                                     content=ft.Text(h,size=8,color=ft.Colors.WHITE,weight=ft.FontWeight.BOLD,text_align=ft.TextAlign.CENTER)) for i,h in enumerate(headers)],spacing=0)
            body=[head]
            for ri,row in enumerate(data):
                bg='#F7FAFD' if ri%2==0 else '#FFFFFF'
                body.append(ft.Row([ft.Container(width=widths[i],height=36,bgcolor=bg,padding=6,
                    alignment=ft.Alignment.CENTER_LEFT,content=ft.Text(str(v),size=8.5,color=TEXT_MAIN)) for i,v in enumerate(row)],spacing=0))
            table_area.controls=[ft.Row([ft.Column(body,spacing=1)],scroll=ft.ScrollMode.AUTO)]
            def position_sort_key(pos):
                try:
                    return (0, int(str(pos).lstrip('P')))
                except Exception:
                    return (1, str(pos))

            def ordered_positions(values):
                return sorted(values.keys(), key=position_sort_key)

            def _dash_line(width=1, color='#D5DDE6', dash=5, gap=4):
                parts=[]
                for _ in range(18):
                    parts.append(ft.Container(width=dash,height=width,bgcolor=color))
                    parts.append(ft.Container(width=gap,height=width))
                return ft.Row(parts,spacing=0)

            def bar_chart(values, y_max, suffix='', secondary=None, secondary_label=None):
                """Remanente por posición: barras horizontales dinámicas, escala 0-100% y valor visible."""
                positions=ordered_positions(values)
                if not positions:
                    return ft.Text('Sin datos suficientes para graficar.',size=11,color=TEXT_MUTED)

                # Escala fija 0-100% para que la comparación entre posiciones sea directa.
                max_value=float(y_max or 100)
                if max_value <= 0:
                    max_value=100

                # Escala superior alineada con el área de barras.
                scale_labels=list(range(0,101,10))
                # La escala coincide exactamente con los extremos de la pista.
                # 0% se alinea con el inicio de la barra y 100% con el final.
                scale_row=ft.Row([
                    ft.Container(width=34),
                    ft.Container(
                        expand=True,
                        content=ft.Row(
                            [ft.Text(f'{v}{suffix}',size=8,color=TEXT_MUTED) for v in scale_labels],
                            spacing=0,
                            alignment=ft.MainAxisAlignment.SPACE_BETWEEN
                        )
                    )
                ],spacing=8)

                # Área de barras. Cada barra usa exactamente el valor calculado
                # para la posición correspondiente; no hay valores fijos P1-P4.
                rows=[]
                for pos in positions:
                    try:
                        v=float(values.get(pos) or 0)
                    except (TypeError,ValueError):
                        v=0.0
                    v=max(0.0,min(v,max_value))
                    pct=(v/max_value) if max_value else 0.0

                    fill=ft.Container(
                        width=0,
                        height=26,
                        bgcolor=NAV_ACCENT,
                        border_radius=5,
                        content=ft.Text(
                            f'{v:.1f}{suffix}',
                            size=10,
                            weight=ft.FontWeight.BOLD,
                            color=ft.Colors.WHITE,
                            text_align=ft.TextAlign.RIGHT
                        ),
                        padding=ft.Padding(left=6,top=4,right=8,bottom=4)
                    )

                    # La barra se coloca sobre una pista gris. La anchura se
                    # calcula en proporción al ancho disponible del contenedor.
                    bar_track=ft.Container(
                        expand=True,
                        height=26,
                        bgcolor='#E7EDF4',
                        border_radius=5,
                        content=ft.Row(
                            [fill],
                            spacing=0
                        )
                    )

                    # Actualizar el ancho relativo después de tener el espacio
                    # disponible mediante una barra flexible: la fila interior
                    # usa una proporción basada en Expanded.
                    units=int(round(pct*100))
                    units=max(0,min(100,units))
                    fill.expand=units if units > 0 else False
                    remaining_units=100-units

                    # La parte vacía completa la pista para que la barra siempre
                    # conserve la proporción 0-100%, incluso al cambiar de equipo.
                    if remaining_units > 0:
                        bar_track.content=ft.Row(
                            [fill,
                             ft.Container(expand=remaining_units,height=26)],
                            spacing=0
                        )
                    else:
                        bar_track.content=ft.Row([fill],spacing=0)

                    rows.append(
                        ft.Row([
                            ft.Container(
                                width=34,
                                content=ft.Text(pos,size=9,weight=ft.FontWeight.BOLD,color=TEXT_MAIN)
                            ),
                            bar_track
                        ],spacing=8)
                    )

                # Líneas verticales punteadas como referencia visual 20/40/60/80/100.
                # Se mantienen discretas para no interferir con los valores.
                return ft.Column(
                    [scale_row, ft.Column(rows,spacing=7)],
                    spacing=6
                )

            def stacked_hours_chart(vals):
                """Horas acumuladas + proyección con leyenda y eje X oculto."""
                positions=ordered_positions(vals)
                if not positions:
                    return ft.Text('Sin datos suficientes para graficar.',size=11,color=TEXT_MUTED)

                totals=[
                    max(0,float((d or {}).get('worked',0))) +
                    max(0,float((d or {}).get('remaining',0)))
                    for d in vals.values() if d
                ]
                max_total=max(totals or [0])
                xmax=max(2000, int((max_total + 1999)//2000)*2000)

                plot_w=510
                label_w=38
                total_w=58
                rows=[]

                for pos in positions:
                    d=vals.get(pos,{}) or {}
                    w=max(0,float(d.get('worked',0)))
                    r=max(0,float(d.get('remaining',0)))
                    total=w+r

                    ww=(plot_w*w/xmax) if w else 0
                    rw=(plot_w*r/xmax) if r else 0

                    if w and ww < 42:
                        ww=42
                    if r and rw < 42:
                        rw=42

                    if ww+rw > plot_w and total > 0:
                        scale=plot_w/(ww+rw)
                        ww*=scale
                        rw*=scale

                    blue = ft.Container(
                        width=ww, height=26,
                        bgcolor=NAV_ACCENT, border_radius=4,
                        alignment=ft.Alignment.CENTER,
                        content=ft.Text(
                            f'{w:,.0f} h', size=9,
                            color=ft.Colors.WHITE,
                            weight=ft.FontWeight.BOLD,
                            text_align=ft.TextAlign.CENTER
                        )
                    ) if w else ft.Container(width=0)

                    orange = ft.Container(
                        width=rw, height=26,
                        bgcolor='#F59E0B', border_radius=4,
                        alignment=ft.Alignment.CENTER,
                        content=ft.Text(
                            f'{r:,.0f} h', size=9,
                            color=ft.Colors.WHITE,
                            weight=ft.FontWeight.BOLD,
                            text_align=ft.TextAlign.CENTER
                        )
                    ) if r else ft.Container(width=0)

                    rows.append(
                        ft.Row([
                            ft.Container(
                                width=label_w, height=26,
                                alignment=ft.Alignment.CENTER_RIGHT,
                                content=ft.Text(
                                    pos, size=9, weight=ft.FontWeight.BOLD,
                                    text_align=ft.TextAlign.RIGHT
                                )
                            ),
                            blue,
                            orange,
                            ft.Container(
                                width=total_w, height=26,
                                alignment=ft.Alignment.CENTER_LEFT,
                                content=ft.Text(
                                    f'{total:,.0f} h', size=9,
                                    color=TEXT_MAIN,
                                    weight=ft.FontWeight.BOLD
                                )
                            )
                        ], spacing=2)
                    )

                # Leyenda solicitada: cuadrito azul + Horas acumuladas,
                # cuadrito naranja + Proyección.
                legend = ft.Row([
                    ft.Row([
                        ft.Container(width=10, height=10, bgcolor=NAV_ACCENT, border_radius=1),
                        ft.Text('Horas acumuladas', size=9, color=TEXT_MUTED)
                    ], spacing=4),
                    ft.Row([
                        ft.Container(width=10, height=10, bgcolor='#F59E0B', border_radius=1),
                        ft.Text('Proyección', size=9, color=TEXT_MUTED)
                    ], spacing=4)
                ], spacing=14)

                # Eje X: se conserva la escala dinámica, pero se ocultan
                # completamente los valores numéricos.
                axis = ft.Container(
                    height=10,
                    content=ft.Row([
                        ft.Container(width=label_w),
                        ft.Container(width=plot_w),
                        ft.Container(width=total_w)
                    ], spacing=2)
                )

                # Espaciador superior para alinear exactamente el inicio
                # de P1/P2/P3/P4 con las barras del gráfico Remanente por posición.
                alignment_spacer = ft.Container(height=1)

                return ft.Column(
                    [legend, alignment_spacer] + rows + [axis],
                    spacing=6
                )

            def pressure_chart(vals):
                """Presión actual por posición: barras horizontales con altura dinámica."""
                positions=ordered_positions(vals)
                if not positions:
                    return ft.Text('Sin datos suficientes para graficar.',size=11,color=TEXT_MUTED)

                rows=[]
                all_values=[]
                for p in positions:
                    item=vals.get(p) or {}
                    actual=item.get('actual')
                    recommended=item.get('recommended')
                    try:
                        actual=float(actual) if actual is not None else None
                    except (TypeError,ValueError):
                        actual=None
                    try:
                        recommended=float(recommended) if recommended is not None else None
                    except (TypeError,ValueError):
                        recommended=None
                    if actual is not None:
                        all_values.append(actual)
                    if recommended is not None:
                        all_values.append(recommended)
                    rows.append((p,actual,recommended))

                max_psi=max(120.0, ((max(all_values)+9)//10)*10 if all_values else 120.0)
                label_w=34
                track_w=390
                bar_h=20
                row_gap=9

                legend=ft.Row([
                    ft.Row([ft.Container(width=10,height=10,bgcolor='#2E9B45',border_radius=3),
                            ft.Text('Actual',size=8.5,color=TEXT_MAIN)],spacing=4),
                    ft.Row([ft.Container(width=3,height=12,bgcolor='#DC2626',border_radius=1),
                            ft.Text('Recomendada',size=8.5,color=TEXT_MAIN)],spacing=4),
                ],alignment=ft.MainAxisAlignment.CENTER,spacing=18)

                scale=ft.Row([
                    ft.Container(width=label_w),
                    ft.Container(width=track_w,content=ft.Row([
                        ft.Text('0',size=7,color=TEXT_MUTED,width=16,text_align=ft.TextAlign.LEFT),
                        ft.Container(expand=True),
                        ft.Text(f'{max_psi:.0f}',size=7,color=TEXT_MUTED,width=30,text_align=ft.TextAlign.RIGHT),
                    ],spacing=0)),
                ],spacing=4)

                visual_rows=[]
                for pos,actual,recommended in rows:
                    actual_ratio=max(0,min(1,(actual/max_psi) if actual is not None else 0))
                    rec_ratio=max(0,min(1,(recommended/max_psi) if recommended is not None else 0))
                    fill_w=max(0,(track_w-2)*actual_ratio)
                    rec_x=max(0,(track_w-2)*rec_ratio)

                    bar_layers=[
                        ft.Container(left=0,top=0,width=track_w,height=bar_h,
                                     bgcolor='#EEF2F6',border_radius=6,
                                     border=ft.Border.all(1,'#D6DEE7')),
                    ]
                    if fill_w>0:
                        bar_layers.append(ft.Container(left=1,top=1,width=max(2,fill_w),height=bar_h-2,
                                                       bgcolor='#2E9B45',border_radius=5,
                                                       alignment=ft.Alignment.CENTER_RIGHT,
                                                       padding=ft.Padding(left=4,top=1,right=5,bottom=1),
                                                       content=ft.Text(f'{actual:.0f}',size=8,weight=ft.FontWeight.BOLD,
                                                                       color=ft.Colors.WHITE,text_align=ft.TextAlign.RIGHT)))
                    if recommended is not None:
                        bar_layers.append(ft.Container(left=rec_x,top=-2,width=3,height=bar_h+7,
                                                       bgcolor='#DC2626',border_radius=1,
                                                       tooltip=f'Presión recomendada: {recommended:.0f} PSI'))

                    track=ft.Stack(controls=bar_layers,width=track_w,height=bar_h+7)
                    visual_rows.append(
                        ft.Row([
                            ft.Container(width=label_w,alignment=ft.Alignment.CENTER_LEFT,
                                         content=ft.Text(pos,size=8.5,weight=ft.FontWeight.BOLD,color=TEXT_MAIN)),
                            track,
                        ],spacing=4,vertical_alignment=ft.CrossAxisAlignment.CENTER)
                    )

                # Altura del gráfico = contenido real + márgenes, sin una altura fija.
                dynamic_h = 76 + len(visual_rows) * (bar_h + 7 + row_gap)
                return ft.Container(
                    height=dynamic_h,
                    content=ft.Column([
                        legend,
                        scale,
                        ft.Column(visual_rows,spacing=row_gap)
                    ],spacing=4)
                )

            def valve_chart(vals):
                """Dona de tapa válvulas: dona + leyenda/datos a la derecha."""
                positions=ordered_positions(vals)
                if not positions:
                    return ft.Text('Sin datos suficientes para graficar.',size=11,color=TEXT_MUTED)

                import base64, math
                yes=[p for p in positions if str(vals.get(p,'NO')).upper()=='SI']
                no=[p for p in positions if str(vals.get(p,'NO')).upper()!='SI']
                total=len(positions)

                # Dona 25% más grande respecto de la versión anterior.
                chart_w=250
                chart_h=210
                cx,cy=125,105
                radius=81.25
                stroke=37.5
                circumference=2*math.pi*radius

                values=[('SÍ',len(yes),'#28B7E6'),('NO',len(no),'#EF4444')]
                circles=[]
                offset=0.0
                for label,count,color in values:
                    if count<=0:
                        continue
                    dash=circumference*(count/total)
                    gap=max(0.0,circumference-dash)
                    circles.append(
                        f'<circle cx="{cx}" cy="{cy}" r="{radius}" fill="none" '
                        f'stroke="{color}" stroke-width="{stroke}" '
                        f'stroke-dasharray="{dash:.3f} {gap:.3f}" '
                        f'stroke-dashoffset="{-offset:.3f}" '
                        f'transform="rotate(-90 {cx} {cy})" />'
                    )
                    offset += dash

                svg=(
                    f'<svg xmlns="http://www.w3.org/2000/svg" width="{chart_w}" height="{chart_h}" viewBox="0 0 {chart_w} {chart_h}">'
                    f'<circle cx="{cx}" cy="{cy}" r="{radius}" fill="none" stroke="#E2E8F0" stroke-width="{stroke}" />'
                    + ''.join(circles) +
                    f'<text x="{cx}" y="{cy-2}" text-anchor="middle" font-family="Arial" font-size="24" font-weight="700" fill="#172033">{total}</text>'
                    f'<text x="{cx}" y="{cy+16}" text-anchor="middle" font-family="Arial" font-size="9" fill="#64748B">Neumáticos</text>'
                    '</svg>'
                )
                b64=base64.b64encode(svg.encode('utf-8')).decode('ascii')

                legend=ft.Column([
                    ft.Text('RESUMEN',size=10,weight=ft.FontWeight.BOLD,color=TEXT_MAIN),
                    ft.Row([
                        ft.Container(width=10,height=10,bgcolor='#28B7E6',border_radius=5),
                        ft.Text(f'SÍ ({len(yes)})',size=10,color=TEXT_MAIN,weight=ft.FontWeight.BOLD)
                    ],spacing=5),
                    ft.Row([
                        ft.Container(width=10,height=10,bgcolor='#EF4444',border_radius=5),
                        ft.Text(f'NO ({len(no)})',size=10,color=TEXT_MAIN,weight=ft.FontWeight.BOLD)
                    ],spacing=5),
                    ft.Divider(height=1,color='#D8E1EB'),
                    ft.Text(f'TOTAL: {total} neumáticos',size=10,color=TEXT_MAIN,weight=ft.FontWeight.BOLD),
                    ft.Text('SÍ: '+(', '.join(yes) if yes else 'Ninguna'),
                            size=8.5,color='#15803D',weight=ft.FontWeight.BOLD,no_wrap=False),
                    ft.Text('NO: '+(', '.join(no) if no else 'Ninguna'),
                            size=8.5,color='#B91C1C',weight=ft.FontWeight.BOLD,no_wrap=False),
                ],spacing=6)

                return ft.Row([
                    ft.Container(width=chart_w,height=chart_h,alignment=ft.Alignment.CENTER,
                                 content=ft.Image(src='data:image/svg+xml;base64,'+b64,width=chart_w,height=chart_h,fit=ft.BoxFit.CONTAIN)),
                    ft.Container(expand=1,content=legend,padding=ft.Padding(left=4,top=14,right=2,bottom=0))
                ],spacing=4,vertical_alignment=ft.CrossAxisAlignment.START)

            rem_chart_22.controls=[bar_chart(rem_by_pos,100,'%')]
            hours_chart_22.controls=[stacked_hours_chart(hours_by_pos)]
            def pressure_history_chart(rows_for_equipment, equipment_id):
                """Historial de presión en barras horizontales, agrupado por código.
                La longitud de cada barra es proporcional al PSI real sobre una
                escala fija de 0 a 120 PSI. Cada código conserva un tono de azul.
                """
                groups=[]
                all_values=[]
                total_events=0

                for r in rows_for_equipment:
                    code=str(r['code'] or '').strip() or f"ID {r['id']}"
                    pos=f"P{r['position']}" if r['position'] else '—'
                    try:
                        recommended=float(r['recommended_pressure']) if r['recommended_pressure'] is not None else None
                    except (TypeError,ValueError):
                        recommended=None

                    tire_events=query("""
                        SELECT id,event_code,event_date,pressure
                        FROM occurrences
                        WHERE tire_id=? AND operation_id=?
                          AND pressure IS NOT NULL
                        ORDER BY event_date ASC,id ASC
                    """,(r['id'],op_id))

                    points=[]
                    for n,ev in enumerate(tire_events,1):
                        try:
                            psi=float(ev['pressure'])
                        except (TypeError,ValueError):
                            continue
                        all_values.append(psi)
                        total_events += 1
                        points.append({
                            'n':n,
                            'psi':psi,
                            'recommended':recommended,
                            'event':str(ev['event_code'] or ''),
                            'date':format_date(ev['event_date']) if ev['event_date'] else '—'
                        })

                    groups.append((code,pos,points))

                if not groups or not any(points for _,_,points in groups):
                    return ft.Text(
                        'Sin historial de presión disponible para los neumáticos del equipo seleccionado.',
                        size=11,color=TEXT_MUTED
                    )

                # Escala fija solicitada: 0–120 PSI.
                max_psi = 120.0

                # Dimensiones de la pista, ajustadas al ancho de la tarjeta.
                label_w = 46
                track_w = 355
                bar_h = 10
                event_gap = 2
                group_gap = 19

                # Eje X: marcas cada 20 PSI hasta 120 PSI.
                tick_values = [0,20,40,60,80,100,120]
                tick_items = []
                for tv in tick_values:
                    align = ft.TextAlign.LEFT if tv == 0 else (
                        ft.TextAlign.RIGHT if tv == 120 else ft.TextAlign.CENTER
                    )
                    tick_items.append(
                        ft.Container(
                            expand=1,
                            content=ft.Text(
                                str(tv),size=10,color=ft.Colors.BLACK,
                                weight=ft.FontWeight.BOLD,
                                text_align=align,no_wrap=True
                            )
                        )
                    )

                scale = ft.Row([
                    ft.Container(width=label_w),
                    ft.Container(
                        width=track_w,
                        content=ft.Row(tick_items,spacing=0)
                    ),
                ],spacing=4)

                blue_palette=[
                    '#0B84B7','#119ED0','#1AA8D8','#249FCE','#2BA7D5',
                    '#39AED9','#45B4DE','#2E8FBE','#1789BA','#3A9BC8',
                    '#0E76A8','#4AA9D2'
                ]
                code_colors={
                    code:blue_palette[i % len(blue_palette)]
                    for i,(code,_,_) in enumerate(groups)
                }

                code_groups=[]
                for code,pos,points in groups:
                    if not points:
                        continue

                    rows_for_code=[]
                    for pnt in points:
                        # La longitud de la barra representa el PSI real:
                        # 120 PSI = 100% del ancho disponible.
                        psi = max(0.0, min(max_psi, float(pnt['psi'])))
                        ratio = psi / max_psi
                        fill_w = max(2, (track_w - 2) * ratio)
                        pct = ratio * 100.0

                        rec = pnt.get('recommended')
                        fill_color = code_colors.get(code,'#16A9D8')
                        class_label = f'Neumático {code}'

                        # Pista + barra proporcional + marcas verticales de 20 PSI.
                        tick_lines=[]
                        for tv in (20,40,60,80,100):
                            x = max(0, min(track_w-1, (track_w-1) * (tv/max_psi)))
                            tick_lines.append(
                                ft.Container(
                                    left=x,
                                    top=0,
                                    width=1,
                                    height=bar_h,
                                    bgcolor='#D6DEE7'
                                )
                            )

                        # El valor de PSI se coloca al final de la barra,
                        # reemplazando la etiqueta de porcentaje.
                        # Para valores muy cercanos a 120 PSI se mantiene visible
                        # dentro del área de la tarjeta.
                        psi_left = min(track_w - 32, max(3, fill_w + 4))
                        psi_label = ft.Container(
                            left=psi_left,
                            top=-1,
                            width=32,
                            height=bar_h + 2,
                            alignment=ft.Alignment.CENTER_LEFT,
                            content=ft.Text(
                                f"{pnt['psi']:.0f} PSI",
                                size=10.5,
                                weight=ft.FontWeight.BOLD,
                                color='#0B3D66',
                                no_wrap=True,
                                text_align=ft.TextAlign.LEFT,
                            )
                        )

                        bar_track = ft.Stack(
                            controls=[
                                ft.Container(
                                    width=track_w,
                                    height=bar_h,
                                    bgcolor='#EEF2F6',
                                    border_radius=4,
                                    border=ft.Border.all(1,'#D6DEE7')
                                ),
                                *tick_lines,
                                ft.Container(
                                    width=fill_w,
                                    height=max(2,bar_h-2),
                                    left=1,
                                    top=1,
                                    bgcolor=fill_color,
                                    border_radius=3,
                                ),
                                psi_label,
                            ],
                            width=track_w,
                            height=bar_h
                        )

                        row=ft.Row(
                            [bar_track],
                            spacing=4,
                            vertical_alignment=ft.CrossAxisAlignment.CENTER
                        )
                        row.tooltip=(
                            f"{code} · {pos} · {pnt['date']} · {pnt['psi']:.0f} PSI · {pct:.0f}% · {class_label}"
                            + (f" · Recomendada {rec:.0f} PSI" if rec is not None else '')
                        )
                        rows_for_code.append(row)

                    # Código del neumático centrado verticalmente respecto a
                    # todas las barras de su propio grupo.
                    group_rows_height = len(rows_for_code) * bar_h + max(0, len(rows_for_code)-1) * event_gap
                    centered_code = ft.Container(
                        width=label_w,
                        height=group_rows_height,
                        alignment=ft.Alignment.CENTER,
                        content=ft.Text(
                            code,
                            size=10.5,
                            weight=ft.FontWeight.BOLD,
                            color=ft.Colors.BLACK,
                            no_wrap=True,
                            overflow=ft.TextOverflow.ELLIPSIS,
                            text_align=ft.TextAlign.LEFT,
                        )
                    )

                    group_content = ft.Row(
                        [
                            centered_code,
                            ft.Column(rows_for_code,spacing=event_gap),
                        ],
                        spacing=4,
                        vertical_alignment=ft.CrossAxisAlignment.CENTER,
                    )

                    code_groups.append(
                        ft.Container(
                            padding=ft.Padding(left=0,top=0,right=0,bottom=group_gap),
                            content=group_content
                        )
                    )

                dynamic_h = 44 + total_events * (bar_h + event_gap + 1) + len(code_groups) * group_gap
                return ft.Container(
                    height=dynamic_h,
                    content=ft.Column([
                        scale,
                        ft.Column(code_groups,spacing=0)
                    ],spacing=4)
                )

            pressure_chart_22.controls=[pressure_chart(pressure_by_pos)]
            valve_chart_22.controls=[valve_chart(valve_by_pos)]
            pressure_history_22.controls=[pressure_history_chart(rows,eid)]

            # Altura dinámica únicamente para los dos gráficos de presión.
            # Presión actual crece según la cantidad de posiciones.
            # Historial crece según la cantidad de registros de presión.
            pressure_card_height = max(240, 145 + len(pressure_by_pos) * 40)

            history_event_count = 0
            history_group_count = 0
            for _trow in rows:
                _hist = query(
                    """SELECT id FROM occurrences
                       WHERE tire_id=? AND operation_id=? AND pressure IS NOT NULL""",
                    (_trow['id'],op_id)
                )
                if _hist:
                    history_group_count += 1
                    history_event_count += len(_hist)

            # Altura dinámica del historial: se ajusta al número real de barras
            # y a los grupos de códigos, evitando un espacio blanco excesivo.
            # Se deja un margen superior/inferior para título, escala y separación.
            history_card_height = max(
                270,
                118 + history_event_count * 13 + max(0, history_group_count - 1) * 19
            )

            graph_row2_left.height = pressure_card_height
            graph_row2_mid.height = history_card_height

            # Altura dinámica REAL de la FILA 01 (Remanente por posición + Horas acumuladas).
            # Cada posición necesita su propia fila visible en ambos gráficos.
            # La altura crece aproximadamente 33 px por posición para evitar
            # que equipos con muchas posiciones queden recortados.
            graph_row1_height = max(260, 90 + len(rows) * 33)
            graph_row1_left.height = graph_row1_height
            graph_row1_right.height = graph_row1_height

            graph_row1_left.content = card(ft.Column([
                ft.Text('REMANENTE POR POSICIÓN (%)',size=14,weight=ft.FontWeight.BOLD,color=TEXT_MAIN),
                ft.Text('Eje X: % remanente · Eje Y: posiciones instaladas',size=9,color=TEXT_MUTED),
                ft.Container(expand=True, content=rem_chart_22)
            ],spacing=8),padding=14)

            graph_row1_right.content = card(ft.Column([
                ft.Text('HORAS ACUMULADAS + PROYECCION',size=14,weight=ft.FontWeight.BOLD,color=TEXT_MAIN),
                ft.Container(expand=True, content=hours_chart_22)
            ],spacing=8),padding=14)

            graph_row2_left.content = card(ft.Column([
                ft.Text('PRESIÓN ACTUAL VS. PRESIÓN RECOMENDADA (PSI)',size=13,weight=ft.FontWeight.BOLD,color=TEXT_MAIN),
                ft.Text('Eje X: PSI · Eje Y: posiciones instaladas',size=9,color=TEXT_MUTED),
                ft.Container(expand=True, content=pressure_chart_22)
            ],spacing=8),padding=14)

            graph_row2_mid.content = card(ft.Column([
                ft.Text('HISTORIAL DE PRESIÓN POR NEUMÁTICO',size=13,weight=ft.FontWeight.BOLD,color=TEXT_MAIN),
                ft.Text('Eje X: PSI · Eje Y: código del neumático',size=9,color=TEXT_MUTED),
                ft.Container(expand=True, content=pressure_history_22)
            ],spacing=8),padding=14)

            graph_row2_right.content = card(ft.Column([
                ft.Text('ESTADO DE EVALUACIÓN DE TAPA VÁLVULAS',size=13,weight=ft.FontWeight.BOLD,color=TEXT_MAIN),
                ft.Text('Por posición instalada: SÍ / NO',size=9,color=TEXT_MUTED),
                ft.Container(expand=True, content=valve_chart_22)
            ],spacing=8),padding=14)

            # La tarjeta de Tapa Válvulas queda 25% más alta que su altura base.
            graph_row2_right.height = 338

            page.update()
        selector.on_change=refresh

        # Contenedores de gráficos con altura actualizable desde refresh().
        # Esto evita depender de una variable `rows` que solo existe dentro de refresh.
        graph_row1_left = ft.Container(expand=1)
        graph_row1_right = ft.Container(expand=1)
        graph_row2_left = ft.Container(expand=1)
        graph_row2_mid = ft.Container(expand=1)
        graph_row2_right = ft.Container(expand=1)

        # Alturas iniciales. refresh() las recalcula según el número real de neumáticos.
        initial_graph_height = 260
        initial_graph_height_row2 = 270
        for _c in (graph_row1_left, graph_row1_right):
            _c.height = initial_graph_height
        for _c in (graph_row2_left, graph_row2_mid, graph_row2_right):
            _c.height = initial_graph_height_row2

        content.content=ft.Column([
            page_title('2.2 REPORTE POR EQUIPO','Información detallada de un solo equipo'),
            ft.Row([ft.OutlinedButton('VOLVER A NEUMÁTICOS EN SERVICIO',icon=ft.Icons.ARROW_BACK,on_click=lambda e:service_menu_view())]),
            card(ft.Column([
                ft.Row([selector,ft.FilledButton('GENERAR REPORTE',icon=ft.Icons.SEARCH,on_click=refresh)],spacing=12),
                ft.Divider(height=1,color='#D8E1EB'),
                ft.Text('DATOS DEL EQUIPO',size=14,weight=ft.FontWeight.BOLD,color=TEXT_MAIN),info,
            ],spacing=12),padding=16),
            card(ft.Column([ft.Text('NEUMÁTICOS DEL EQUIPO',size=15,weight=ft.FontWeight.BOLD,color=TEXT_MAIN),table_area],spacing=10),padding=14),
            ft.Row([
                graph_row1_left,
                graph_row1_right,
            ],spacing=12,vertical_alignment=ft.CrossAxisAlignment.START),
            # FILA 2: tres gráficos alineados y simétricos.
            # Cada tarjeta ocupa exactamente 1/3 del ancho disponible y comparte
            # la misma altura, evitando que el historial quede debajo de los otros.
            ft.Row([
                graph_row2_left,
                graph_row2_mid,
                graph_row2_right,
            ],spacing=12,vertical_alignment=ft.CrossAxisAlignment.START),
        ],scroll=ft.ScrollMode.AUTO,spacing=14)
        refresh()
        page.update()

    def movement_view():
        op_id = active_operation_id()
        tire=ft.Dropdown(label='Neumático *',width=310,options=[ft.dropdown.Option(str(r['id']),f"{r['code']} | {r['serial'] or 's/serie'}") for r in query('SELECT id,code,serial FROM tires WHERE operation_id=? ORDER BY code',(op_id,))])
        # ROT se mantiene fuera del selector hasta definir su funcionalidad.
        event=ft.Dropdown(
            label='Evento *',
            width=265,
            options=[ft.dropdown.Option(k,f'{k} - {v}') for k,v in EVENTS.items() if k != 'ROT']
        )
        date=ft.TextField(label='Fecha',value=dt.date.today().strftime('%d/%m/%Y'),width=260,dense=True,bgcolor='#FFFFFF')
        equip=ft.Dropdown(label='Equipo',width=155,dense=True,bgcolor='#FFFFFF',options=[ft.dropdown.Option(str(r['id']),r['code']) for r in query('SELECT id,code FROM equipment WHERE active=1 AND operation_id=? ORDER BY code',(op_id,))])
        performance_info=ft.Text(
            'Rendimiento: —',
            size=10,
            color=TEXT_MUTED,
            weight=ft.FontWeight.BOLD
        )
        pos=ft.TextField(label='Pos.',width=95,dense=True,bgcolor='#FFFFFF')
        # Regla fundamental: ambos campos existen, pero solo uno puede utilizarse
        # según el rendimiento del EQUIPO seleccionado. Ambos alimentan el mismo
        # campo `meter` de occurrences para conservar el esquema actual.
        horometer=ft.TextField(label='Horómetro',width=260,dense=True,bgcolor='#FFFFFF',disabled=True)
        odometer=ft.TextField(label='Odómetro',width=260,dense=True,bgcolor='#FFFFFF',disabled=True)
        meter=horometer  # Compatibilidad interna: el valor operativo se obtiene mediante active_meter_value().
        ti=ft.TextField(label='INT',width=125,dense=True,bgcolor='#FFFFFF')
        to=ft.TextField(label='EXT',width=125,dense=True,bgcolor='#FFFFFF')
        press=ft.TextField(label='Psi',width=125,dense=True,bgcolor='#FFFFFF')
        cond=ft.Dropdown(
            label='Cond.', width=125, value='FRIO', dense=True, bgcolor='#FFFFFF',
            options=[ft.dropdown.Option('FRIO','FRIO'), ft.dropdown.Option('CALIENTE','CALIENTE')]
        )
        BAJA_REASONS = [
            ('GAST','Desgaste regular'),
            ('CORL','Corte lateral'),
            ('CORB','Corte en banda'),
            ('SEPA','Separación banda'),
            ('PSIB','Baja presión'),
            ('RTEQ','Retiro de equipo'),
            ('EXPL','Exposición de lona / alambre'),
            ('REEN','Reencauche'),
            ('DIB','Daño irregular'),
            ('CGH','Corte / golpe hombro'),
            ('DTB','Daño talón'),
            ('IRB','Irregularidad banda'),
            ('IBR','Impacto banda rodamiento'),
        ]
        REPA_REASONS = [
            ('ARO','Aro'),
            ('PERF','Perforación'),
            ('CORB','Corte banda'),
            ('CORL','Corte lateral'),
            ('VALV','Válvula'),
            ('ORIN','O-ring'),
            ('OTRO','Otro'),
        ]
        reason=ft.Dropdown(label='Motivo',width=260,dense=True,bgcolor='#FFFFFF',options=[])
        loc=ft.TextField(label='Lugar',width=260,dense=True,bgcolor='#FFFFFF')
        notes=ft.TextField(label='Observaciones',multiline=True,min_lines=1,max_lines=2,width=260,dense=True,bgcolor='#FFFFFF')
        valve_cap=ft.Dropdown(
            label='Tapa Válvula *', width=260, dense=True, bgcolor='#FFFFFF', value='NO',
            options=[ft.DropdownOption(key='SI', text='SI'), ft.DropdownOption(key='NO', text='NO')]
        )
        ref=ft.Text('',size=11,color=TEXT_MUTED)
        pre_tire = session.pop('movement_tire_id', None)
        pre_event = session.pop('movement_event', None)
        if pre_tire:
            # V22 MULTITALLER: nunca reutilizar una selección proveniente de otra operación.
            _pre_ok = query('SELECT id FROM tires WHERE id=? AND operation_id=?', (int(pre_tire), op_id)) if str(pre_tire).isdigit() else []
            tire.value = pre_tire if _pre_ok else None
        if pre_event:
            event.value = pre_event

        hist=ft.DataTable(
            columns=[ft.DataColumn(ft.Text(x)) for x in [
                'Fecha','Evento','Equipo','Pos.','Lectura','Cocada E/I','Presión','Condición','Ubicación','Acción'
            ]],
            rows=[]
        )

        def fmt(v):
            if v is None:
                return ''
            if isinstance(v,float) and v.is_integer():
                return str(int(v))
            return str(v)

        def select_all_on_focus(e):
            c=e.control
            try:
                value=str(c.value or '')
                c.selection=ft.TextSelection(base_offset=0, extent_offset=len(value))
                c.update()
            except Exception:
                pass

        for ctrl in [date,pos,meter,ti,to,press,loc,notes]:
            ctrl.on_focus=select_all_on_focus

        def current_tire():
            if not tire.value:
                return None
            rows=query(
                'SELECT t.*,e.code equipment_code,e.location equipment_location '
                'FROM tires t LEFT JOIN equipment e ON e.id=t.equipment_id AND e.operation_id=t.operation_id WHERE t.id=? AND t.operation_id=?',
                (int(tire.value), op_id)
            )
            return rows[0] if rows else None

        def selected_performance_type():
            return get_equipment_performance_type(equip.value, op_id) if equip.value else None

        def active_meter_control():
            perf = selected_performance_type()
            return horometer if perf == PERFORMANCE_HOURS else odometer if perf == PERFORMANCE_KILOMETERS else None

        def active_meter_value():
            ctrl = active_meter_control()
            return ctrl.value if ctrl is not None else ''

        def set_active_meter_value(value):
            perf = selected_performance_type()
            if perf == PERFORMANCE_HOURS:
                horometer.value = fmt(value) if value not in (None, '') else ''
                odometer.value = ''
            elif perf == PERFORMANCE_KILOMETERS:
                odometer.value = fmt(value) if value not in (None, '') else ''
                horometer.value = ''
            else:
                horometer.value = ''
                odometer.value = ''

        def apply_meter_mode(lock=False):
            perf = selected_performance_type()
            # Sin equipo no se puede determinar la unidad: ambos quedan bloqueados.
            if perf == PERFORMANCE_HOURS:
                horometer.disabled = bool(lock)
                odometer.disabled = True
                horometer.label = 'Horómetro'
                odometer.label = 'Odómetro (no aplica)'
            elif perf == PERFORMANCE_KILOMETERS:
                horometer.disabled = True
                odometer.disabled = bool(lock)
                horometer.label = 'Horómetro (no aplica)'
                odometer.label = 'Odómetro'
            else:
                horometer.disabled = True
                odometer.disabled = True
                horometer.label = 'Horómetro'
                odometer.label = 'Odómetro'
            try:
                horometer.update(); odometer.update()
            except Exception:
                pass
            return perf

        def historical_limits(tid):
            # El horómetro conserva su validación histórica por lectura máxima.
            # Las cocadas, en cambio, deben tomar como referencia el ÚLTIMO
            # evento registrado. Esto es indispensable después de una INVE,
            # porque EXT/INT cambian físicamente de lado.
            meter_rows=query(
                'SELECT MAX(meter) max_meter FROM occurrences WHERE tire_id=? AND operation_id=?',
                (tid, op_id)
            )
            last_tread=query(
                '''SELECT tread_inner,tread_outer
                   FROM occurrences
                   WHERE tire_id=? AND operation_id=?
                     AND (tread_inner IS NOT NULL OR tread_outer IS NOT NULL)
                   ORDER BY id DESC LIMIT 1''',
                (tid, op_id)
            )
            return {
                'max_meter': meter_rows[0]['max_meter'] if meter_rows else None,
                # Se conservan estas claves para no alterar el resto del módulo;
                # ahora representan la última lectura válida, no mínimos históricos.
                'min_ti': last_tread[0]['tread_inner'] if last_tread else None,
                'min_to': last_tread[0]['tread_outer'] if last_tread else None,
            }

        def parse_event_date(value):
            if value in (None, ''):
                return None
            raw=str(value).strip()
            try:
                f=float(raw)
                if f.is_integer() and 1 <= f <= 100000:
                    return dt.date(1899,12,30)+dt.timedelta(days=int(f))
            except Exception:
                pass
            for f in ('%d/%m/%Y','%d-%m-%Y','%Y/%m/%d','%Y-%m-%d'):
                try:
                    return dt.datetime.strptime(raw[:10],f).date()
                except Exception:
                    pass
            return None

        def latest_event_date(tid):
            rows=query('SELECT event_date FROM occurrences WHERE tire_id=? AND operation_id=?',(tid,op_id))
            dates=[parse_event_date(r['event_date']) for r in rows]
            dates=[d for d in dates if d is not None]
            return max(dates) if dates else None

        def recalc_numeric_state(tid):
            lim=historical_limits(tid)
            if not lim:
                return
            execute(
                '''UPDATE tires SET
                       current_meter=COALESCE(?,current_meter),
                       tread_inner=COALESCE(?,tread_inner),
                       tread_outer=COALESCE(?,tread_outer)
                   WHERE id=? AND operation_id=?''',
                (lim['max_meter'],lim['min_ti'],lim['min_to'],tid,op_id)
            )

        def refresh_equipment_performance(lock_meter=False):
            performance_type = get_equipment_performance_type(equip.value, op_id)
            performance_info.value = f'Rendimiento: {performance_unit_label(performance_type)}'
            apply_meter_mode(lock=lock_meter)
            try:
                performance_info.update()
            except Exception:
                pass
            return performance_type

        def load_current_state(e=None):
            r=current_tire()
            if not r:
                equip.value=None
                refresh_equipment_performance()
                pos.value=''
                horometer.value=''
                odometer.value=''
                ti.value=''
                to.value=''
                press.value=''
                cond.value='FRIO'
                loc.value=''
                ref.value=''
                return

            equip.value=str(r['equipment_id']) if r['equipment_id'] is not None else None
            refresh_equipment_performance()
            pos.value=fmt(r['position'])

            lim=historical_limits(int(r['id']))
            set_active_meter_value(lim['max_meter'] if lim and lim['max_meter'] is not None else r['current_meter'])
            ti.value=fmt(lim['min_ti'] if lim and lim['min_ti'] is not None else r['tread_inner'])
            to.value=fmt(lim['min_to'] if lim and lim['min_to'] is not None else r['tread_outer'])

            last=query(
                '''SELECT pressure,pressure_condition,location,valve_cap,notes
                   FROM occurrences
                   WHERE tire_id=? AND operation_id=?
                   ORDER BY id DESC LIMIT 1''',
                (int(r['id']), op_id)
            )
            last_row=last[0] if last else None
            press.value=fmt(last_row['pressure']) if last_row and last_row['pressure'] is not None else fmt(r['recommended_pressure'])

            last_cond=(last_row['pressure_condition'] if last_row else None) or 'FRIO'
            last_cond=str(last_cond).strip().upper().replace('Í','I')
            cond.value='CALIENTE' if last_cond.startswith('CAL') else 'FRIO'

            loc.value=fmt(last_row['location']) if last_row and last_row['location'] else fmt(r['equipment_location'])
            valve_raw = str(last_row['valve_cap'] or '').strip().upper() if last_row and 'valve_cap' in last_row.keys() else ''
            if valve_raw in ('SI','SÍ'):
                valve_cap.value='SI'
            elif valve_raw == 'NO':
                valve_cap.value='NO'
            else:
                note_text = str(last_row['notes'] or '').upper() if last_row and 'notes' in last_row.keys() else ''
                valve_cap.value='SI' if ('TAPA' in note_text or 'VALVULA' in note_text or 'VÁLVULA' in note_text) else 'NO'
            ref.value=(
                f"Estado actual: {r['status']} · Equipo: {r['equipment_code'] or '-'} · "
                f"Pos.: {r['position'] or '-'} · Última lectura válida: {active_meter_value() or '-'} · "
                f"Cocada E/I válida: {to.value or '-'}/{ti.value or '-'}"
            )

        def apply_event_rules(e=None):
            ec=event.value
            r=current_tire()

            # Estado editable por defecto. Cada evento aplica solo sus bloqueos propios.
            equip.disabled=False
            pos.disabled=False
            horometer.disabled=False
            odometer.disabled=False
            ti.disabled=False
            to.disabled=False
            # Tapa Válvula es editable en INST, INSP, INSC e INVE.
            # En otros eventos se conserva el último valor como referencia y queda bloqueado.
            valve_cap.disabled = ec not in ('INST','INSP','INSC','INVE')

            # MOTIVO: se utiliza únicamente en REPA y BAJA.
            # En INST, INSP, INSC, ROT, INVE y DINS permanece bloqueado.
            reason.disabled = ec not in ('REPA','BAJA')
            if ec == 'BAJA':
                reason.options=[ft.dropdown.Option(k, f'{k} - {v}') for k,v in BAJA_REASONS]
                if reason.value not in {k for k,_ in BAJA_REASONS}:
                    reason.value=None
            elif ec == 'REPA':
                reason.options=[ft.dropdown.Option(k, f'{k} - {v}') for k,v in REPA_REASONS]
                if reason.value not in {k for k,_ in REPA_REASONS}:
                    reason.value=None
            else:
                # Evita arrastrar un motivo si el usuario cambia a un evento
                # donde el campo Motivo no corresponde.
                reason.options=[]
                reason.value=None

            # INSP/INSC, BAJA y DINS conservan el equipo y la posición actuales.
            # En BAJA y DINS estos campos son solo referencia y no pueden modificarse.
            locked=ec in ('INSP','INSC','BAJA','DINS')
            equip.disabled=locked
            pos.disabled=locked
            if locked and r:
                equip.value=str(r['equipment_id']) if r['equipment_id'] is not None else None
                refresh_equipment_performance()
                pos.value=fmt(r['position'])
                if r['status'] != 'SERVICIO':
                    ref.value=(ref.value + ' · ADVERTENCIA: el neumático no figura EN SERVICIO').strip(' ·')

            # INVE - Inversión:
            # equipo, posición y horómetro permanecen iguales y bloqueados;
            # las cocadas EXT/INT se intercambian automáticamente y quedan bloqueadas.
            if ec == 'INVE' and r:
                equip.value=str(r['equipment_id']) if r['equipment_id'] is not None else None
                refresh_equipment_performance()
                pos.value=fmt(r['position'])
                last=query(
                    '''SELECT meter,tread_inner,tread_outer
                       FROM occurrences
                       WHERE tire_id=? AND operation_id=?
                       ORDER BY id DESC LIMIT 1''',
                    (int(r['id']), op_id)
                )
                last_row=last[0] if last else None
                if last_row:
                    if last_row['meter'] is not None:
                        set_active_meter_value(last_row['meter'])
                    prev_int=last_row['tread_inner']
                    prev_ext=last_row['tread_outer']
                else:
                    prev_int=r['tread_inner']
                    prev_ext=r['tread_outer']
                # Después de invertir físicamente el neumático:
                # EXT nueva = INT anterior / INT nueva = EXT anterior.
                to.value=fmt(prev_int)
                ti.value=fmt(prev_ext)
                equip.disabled=True
                pos.disabled=True
                horometer.disabled=True
                odometer.disabled=True
                to.disabled=True
                ti.disabled=True

            # La unidad la determina el equipo.
            # INSP e INSC deben permitir ingresar el horómetro/odómetro de la inspección.
            # BAJA, DINS e INVE mantienen el medidor bloqueado.
            meter_locked = ec in ('INVE','BAJA','DINS')
            apply_meter_mode(lock=meter_locked)
            page.update()

        def ask_delete(occ_id):
            if not tire.value:
                return
            tid=int(tire.value)

            # Regla de seguridad: solo se puede eliminar el último evento.
            latest=query(
                '''SELECT id,event_code
                   FROM occurrences
                   WHERE tire_id=? AND operation_id=?
                   ORDER BY id DESC LIMIT 1''',
                (tid, op_id)
            )
            if not latest or int(latest[0]['id']) != int(occ_id):
                snack('Solo se puede eliminar el último evento.', True)
                return
            if str(latest[0]['event_code'] or '').strip().upper() == 'INST':
                snack('El evento INST está bloqueado y no puede eliminarse.', True)
                return

            def close_dialog(e=None):
                page.pop_dialog()
                page.update()

            def do_delete(e=None):
                # Segunda validación justo antes de borrar: el evento debe seguir
                # siendo el último y no puede ser INST.
                check=query(
                    '''SELECT id,event_code
                       FROM occurrences
                       WHERE tire_id=? AND operation_id=?
                       ORDER BY id DESC LIMIT 1''',
                    (tid, op_id)
                )
                if not check or int(check[0]['id']) != int(occ_id):
                    page.pop_dialog()
                    snack('El evento ya no es el último. No se eliminó.', True)
                    refresh()
                    return
                if str(check[0]['event_code'] or '').strip().upper() == 'INST':
                    page.pop_dialog()
                    snack('El evento INST está bloqueado y no puede eliminarse.', True)
                    refresh()
                    return
                execute('DELETE FROM occurrences WHERE id=? AND tire_id=? AND operation_id=?',(occ_id,tid,op_id))
                recalc_numeric_state(tid)
                page.pop_dialog()
                snack('Evento eliminado correctamente.')
                refresh()

            dlg=ft.AlertDialog(
                modal=True,
                title=ft.Text('Eliminar evento'),
                content=ft.Text('¿Desea eliminar este evento del historial? Esta acción no se puede deshacer.'),
                actions=[
                    ft.TextButton('Cancelar',on_click=close_dialog),
                    ft.Button('Eliminar',icon=ft.Icons.DELETE_OUTLINE,on_click=do_delete)
                ],
                actions_alignment=ft.MainAxisAlignment.END
            )
            page.show_dialog(dlg)
            page.update()

        save_btn=ft.Button(
            'Guardar movimiento',
            icon=ft.Icons.SAVE,
            disabled=True
        )

        def form_is_valid():
            if not tire.value or not event.value:
                return False
            # Seguridad adicional: ROT no puede guardarse mientras esté bloqueado.
            if event.value == 'ROT':
                return False
            r=current_tire()
            if not r:
                return False

            # Reglas fundamentales por estado del neumático:
            # SERVICIO/instalado -> INSP, INSC, INVE, DINS y BAJA. REPA no aplica.
            # STAND-BY -> INST, INVE, REPA y BAJA.
            # ROT continúa bloqueado hasta definir su funcionalidad.
            status_norm=str(r['status'] or '').strip().upper().replace('_','-')
            installed = r['equipment_id'] is not None and str(r['position'] or '').strip() != ''
            if installed and status_norm == 'SERVICIO':
                allowed_events={'INSP','INSC','INVE','DINS','BAJA'}
            elif status_norm in ('STAND-BY','STAND BY','STANDBY'):
                allowed_events={'INST','INVE','REPA','BAJA'}
            else:
                allowed_events=set()
            if event.value not in allowed_events:
                return False
            # REPA y BAJA requieren seleccionar un motivo del catálogo correspondiente.
            if event.value in ('REPA','BAJA') and not (reason.value or '').strip():
                return False

            tid=int(tire.value)
            lim=historical_limits(tid)

            entered_date=parse_event_date(date.value)
            last_date=latest_event_date(tid)
            if entered_date is None:
                return False
            if last_date is not None and entered_date < last_date:
                return False

            new_meter=num(active_meter_value())
            if new_meter is None:
                return False
            if lim and lim['max_meter'] is not None and float(new_meter) < float(lim['max_meter']):
                return False

            new_ti=num(ti.value)
            new_to=num(to.value)
            max_new=num(r['new_tread'])
            if new_ti is not None:
                if new_ti < 0:
                    return False
                if max_new is not None and new_ti > max_new:
                    return False
                if event.value != 'INVE' and lim and lim['min_ti'] is not None and new_ti > float(lim['min_ti']):
                    return False
            if new_to is not None:
                if new_to < 0:
                    return False
                if max_new is not None and new_to > max_new:
                    return False
                if event.value != 'INVE' and lim and lim['min_to'] is not None and new_to > float(lim['min_to']):
                    return False

            if event.value in ('INSP','INSC'):
                if r['status'] != 'SERVICIO' or r['equipment_id'] is None or not r['position']:
                    return False
            if event.value == 'INST' and (not equip.value or not (pos.value or '').strip()):
                return False

            # BAJA/DINS: el remanente no puede ser menor al último valor existente.
            # Como la regla general tampoco permite aumentarlo, en BAJA y DINS el RTD
            # queda necesariamente igual al último RTD válido registrado.
            if event.value in ('BAJA','DINS'):
                if lim and lim['min_ti'] is not None:
                    if new_ti is None or float(new_ti) < float(lim['min_ti']):
                        return False
                if lim and lim['min_to'] is not None:
                    if new_to is None or float(new_to) < float(lim['min_to']):
                        return False
            return True

        def update_save_state(e=None):
            save_btn.disabled = not form_is_valid()
            try:
                refresh_new_event_preview()
            except Exception:
                pass
            page.update()

        def refresh(e=None):
            if not tire.value:
                hist.rows=[]
                load_current_state()
            else:
                rows=query(
                    '''SELECT o.*,e.code equipment_code
                       FROM occurrences o
                       LEFT JOIN equipment e ON e.id=o.equipment_id
                       WHERE o.tire_id=? AND o.operation_id=?
                       ORDER BY o.id DESC''',
                    (int(tire.value), op_id)
                )
                hist.rows=[]
                for idx, r in enumerate(rows):
                    hist.rows.append(
                        ft.DataRow(cells=[
                            ft.DataCell(ft.Text(format_date(r['event_date']))),
                            ft.DataCell(ft.Text(fmt(r['event_code']))),
                            ft.DataCell(ft.Text(fmt(r['equipment_code']))),
                            ft.DataCell(ft.Text(fmt(r['position']))),
                            ft.DataCell(ft.Text(fmt(r['meter']))),
                            ft.DataCell(ft.Text(f"{fmt(r['tread_outer'])}/{fmt(r['tread_inner'])}")),
                            ft.DataCell(ft.Text(fmt(r['pressure']))),
                            ft.DataCell(ft.Text(fmt(r['pressure_condition']))),
                            ft.DataCell(ft.Text(fmt(r['location']))),
                            ft.DataCell(ft.IconButton(
                                icon=ft.Icons.DELETE_OUTLINE,
                                tooltip=('Eliminar último evento' if idx == 0 and str(r['event_code'] or '').upper() != 'INST' else 'Bloqueado: solo se puede eliminar el último evento'),
                                disabled=(idx != 0 or str(r['event_code'] or '').upper() == 'INST'),
                                on_click=(lambda e, oid=r['id']: ask_delete(oid)) if (idx == 0 and str(r['event_code'] or '').upper() != 'INST') else None
                            ))
                        ])
                    )
                load_current_state()
                apply_event_rules()
            save_btn.disabled = not form_is_valid()
            page.update()

        def on_tire_change(e):
            refresh(e)

        def on_event_change(e):
            apply_event_rules(e)
            save_btn.disabled = not form_is_valid()
            page.update()

        tire.on_select=on_tire_change
        event.on_select=on_event_change
        date.on_change=update_save_state
        def on_equipment_change(e):
            refresh_equipment_performance()
            update_save_state(e)

        equip.on_select=on_equipment_change
        pos.on_change=update_save_state
        meter.on_change=update_save_state
        ti.on_change=update_save_state
        to.on_change=update_save_state
        press.on_change=update_save_state
        cond.on_select=update_save_state
        reason.on_select=update_save_state
        loc.on_change=update_save_state
        valve_cap.on_select=update_save_state

        def save(e):
            if not tire.value or not event.value:
                return snack('Seleccione neumático y evento.',True)

            tid=int(tire.value)
            ec=event.value
            if ec == 'ROT':
                return snack('ROT está bloqueado hasta definir su funcionalidad.', True)
            r=current_tire()
            if not r:
                return snack('No se encontró el neumático seleccionado.',True)

            # Protección de guardado: aplicar la misma matriz de eventos que los botones.
            status_norm=str(r['status'] or '').strip().upper().replace('_','-')
            installed = r['equipment_id'] is not None and str(r['position'] or '').strip() != ''
            if installed and status_norm == 'SERVICIO':
                allowed_events={'INSP','INSC','INVE','DINS','BAJA'}
            elif status_norm in ('STAND-BY','STAND BY','STANDBY'):
                allowed_events={'INST','INVE','REPA','BAJA'}
            else:
                allowed_events=set()
            if ec not in allowed_events:
                if ec == 'REPA' and installed:
                    return snack('REPA no está permitido para neumáticos instalados. Primero debe pasar a STAND-BY.', True)
                return snack(f'El evento {ec} no está permitido para el estado actual del neumático.', True)
            if ec in ('REPA','BAJA') and not (reason.value or '').strip():
                return snack(f'Seleccione un motivo para {ec}.', True)
            if ec in ('INST','INSP','INSC','INVE') and (str(valve_cap.value or '').strip().upper() not in ('SI','NO')):
                return snack('Seleccione SI o NO en Tapa Válvula.', True)
            if ec in ('INST','INSP','INSC','INVE') and (str(valve_cap.value or '').strip().upper() not in ('SI','NO')):
                return snack('Seleccione SI o NO en Tapa Válvula.', True)
            valid_reason_codes = ({k for k,_ in REPA_REASONS} if ec == 'REPA' else {k for k,_ in BAJA_REASONS}) if ec in ('REPA','BAJA') else set()
            if ec in ('REPA','BAJA') and reason.value not in valid_reason_codes:
                return snack(f'El motivo seleccionado no es válido para {ec}.', True)

            lim=historical_limits(tid)
            new_meter=num(active_meter_value())
            max_meter=lim['max_meter'] if lim else None
            if new_meter is None:
                return snack('Ingrese el Horómetro u Odómetro según el rendimiento del equipo.',True)
            if max_meter is not None and float(new_meter) < float(max_meter):
                return snack(
                    f'Lectura inválida: {fmt(new_meter)} es menor que la última lectura válida {fmt(max_meter)}.',
                    True
                )

            new_ti=num(ti.value)
            new_to=num(to.value)
            max_new=num(r['new_tread'])

            if new_ti is not None:
                if new_ti < 0:
                    return snack('La cocada interior no puede ser negativa.',True)
                if max_new is not None and new_ti > max_new:
                    return snack(f'Cocada interior inválida: no puede superar la profundidad nueva ({fmt(max_new)} mm).',True)
                if ec != 'INVE' and lim and lim['min_ti'] is not None and new_ti > float(lim['min_ti']):
                    return snack(
                        f'Cocada interior inválida: {fmt(new_ti)} mm es mayor que la última cocada válida {fmt(lim["min_ti"])} mm.',
                        True
                    )

            if new_to is not None:
                if new_to < 0:
                    return snack('La cocada exterior no puede ser negativa.',True)
                if max_new is not None and new_to > max_new:
                    return snack(f'Cocada exterior inválida: no puede superar la profundidad nueva ({fmt(max_new)} mm).',True)
                if ec != 'INVE' and lim and lim['min_to'] is not None and new_to > float(lim['min_to']):
                    return snack(
                        f'Cocada exterior inválida: {fmt(new_to)} mm es mayor que la última cocada válida {fmt(lim["min_to"])} mm.',
                        True
                    )

            # Regla específica BAJA/DINS: no se permite reducir el remanente respecto
            # del último RTD válido. El botón Guardar ya queda deshabilitado por
            # form_is_valid(); esta validación adicional evita cualquier bypass.
            if ec in ('BAJA','DINS'):
                if lim and lim['min_ti'] is not None:
                    if new_ti is None or float(new_ti) < float(lim['min_ti']):
                        return snack(
                            f'Remanente interior inválido: {fmt(new_ti)} mm no puede ser menor que el existente {fmt(lim["min_ti"])} mm.',
                            True
                        )
                if lim and lim['min_to'] is not None:
                    if new_to is None or float(new_to) < float(lim['min_to']):
                        return snack(
                            f'Remanente exterior inválido: {fmt(new_to)} mm no puede ser menor que el existente {fmt(lim["min_to"])} mm.',
                            True
                        )

            raw_date=(date.value or '').strip()
            event_date=raw_date
            date_ok=False
            for _fmt in ('%d/%m/%Y','%d-%m-%Y','%Y/%m/%d','%Y-%m-%d'):
                try:
                    event_date=dt.datetime.strptime(raw_date[:10],_fmt).strftime('%Y-%m-%d')
                    date_ok=True
                    break
                except Exception:
                    pass
            if not date_ok:
                return snack('Fecha inválida. Use dd/mm/aaaa.',True)

            entered_date=parse_event_date(event_date)
            last_date=latest_event_date(tid)
            if last_date is not None and entered_date is not None and entered_date < last_date:
                return snack(
                    f'Fecha inválida: {entered_date.strftime("%d/%m/%Y")} es anterior al último evento {last_date.strftime("%d/%m/%Y")}.',
                    True
                )

            if ec in ('INSP','INSC'):
                if r['status'] != 'SERVICIO' or r['equipment_id'] is None or not r['position']:
                    return snack('Para registrar una inspección el neumático debe estar instalado y EN SERVICIO.',True)
                eid=int(r['equipment_id'])
                event_pos=str(r['position'])
                equip.value=str(eid)
                pos.value=event_pos
            elif ec in ('BAJA','DINS'):
                # BAJA y DINS conservan obligatoriamente el equipo y la posición actuales.
                eid=int(r['equipment_id']) if r['equipment_id'] is not None else None
                event_pos=str(r['position'] or '')
                equip.value=str(eid) if eid is not None else None
                pos.value=event_pos
            else:
                eid=int(equip.value) if equip.value else None
                event_pos=pos.value

            # V22 MULTITALLER: un movimiento jamás puede apuntar a un equipo de otra operación.
            if eid is not None:
                _eq_ok=query('SELECT id FROM equipment WHERE id=? AND operation_id=? AND active=1',(eid,op_id))
                if not _eq_ok:
                    return snack('El equipo seleccionado no pertenece a la operación activa.',True)

            if ec=='INSC' and query(
                "SELECT id FROM occurrences WHERE tire_id=? AND operation_id=? AND event_code='INSC' AND event_date=? AND COALESCE(meter,-1)=COALESCE(?, -1)",
                (tid,op_id,event_date,new_meter)
            ):
                return snack('Ya existe una INSC con la misma fecha y lectura.',True)

            condition=(cond.value or 'FRIO').strip().upper().replace('Í','I')
            condition='CALIENTE' if condition.startswith('CAL') else 'FRIO'

            # Por seguridad, únicamente REPA y BAJA pueden grabar Motivo.
            # En cualquier otro evento se descarta un valor residual.
            reason_for_db=(reason.value or '').strip() if ec in ('REPA','BAJA') else ''
            valve_for_db=(valve_cap.value or '').strip().upper() if ec in ('INST','INSP','INSC','INVE') else (valve_cap.value or 'NO').strip().upper()
            valve_for_db='SI' if valve_for_db in ('SI','SÍ') else 'NO'

            execute(
                '''INSERT INTO occurrences(
                       tire_id,event_code,event_date,equipment_id,position,meter,
                       tread_inner,tread_outer,pressure,pressure_condition,reason,location,notes,valve_cap,operation_id
                   ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)''',
                (tid,ec,event_date,eid,event_pos,new_meter,new_ti,new_to,
                 num(press.value),condition,reason_for_db,loc.value,notes.value,valve_for_db,op_id)
            )

            if ec=='INST':
                execute(
                    "UPDATE tires SET status='SERVICIO',equipment_id=?,position=?,current_meter=?,"
                    "tread_inner=COALESCE(?,tread_inner),tread_outer=COALESCE(?,tread_outer) WHERE id=? AND operation_id=?",
                    (eid,event_pos,new_meter,new_ti,new_to,tid,op_id)
                )
            elif ec=='DINS':
                execute(
                    "UPDATE tires SET status='STAND-BY',equipment_id=NULL,position=NULL,current_meter=? WHERE id=? AND operation_id=?",
                    (new_meter,tid,op_id)
                )
            elif ec=='REPA':
                execute("UPDATE tires SET status='REPARACIÓN' WHERE id=? AND operation_id=?",(tid,op_id))
            elif ec=='BAJA':
                execute("UPDATE tires SET status='BAJA',equipment_id=NULL,position=NULL WHERE id=? AND operation_id=?",(tid,op_id))
            else:
                execute(
                    'UPDATE tires SET current_meter=COALESCE(?,current_meter),'
                    'tread_inner=COALESCE(?,tread_inner),tread_outer=COALESCE(?,tread_outer) WHERE id=? AND operation_id=?',
                    (new_meter,new_ti,new_to,tid,op_id)
                )

            snack(f'Evento {ec} registrado correctamente.')
            set_inline_event_mode(False)
            refresh()
            load_foxpro_ficha(tid)
            update_event_button_states(tid)
            save_btn.disabled = True
            page.update()

        save_btn.on_click=save

        if pre_tire:
            refresh()
        else:
            save_btn.disabled = True

        history_scroller=ft.Row(
            [ft.Container(content=hist,width=1280)],
            scroll=ft.ScrollMode.ALWAYS
        )

        # ------------------------------------------------------------------
        # CABECERA OPERATIVA - lógica visual basada en MegaSoftire FoxPro.
        # Primero se busca/selecciona el neumático; luego se muestra su ficha
        # vertical y los eventos. El formulario de movimiento queda oculto
        # hasta escoger un evento.
        # ------------------------------------------------------------------
        search_tire = ft.TextField(
            label='Buscar código / serie',
            prefix_icon=ft.Icons.SEARCH,
            width=255
        )
        search_result = ft.Dropdown(
            label='Neumático encontrado',
            width=255,
            visible=False,
            options=[]
        )

        register_missing_btn = ft.Button(
            'REGISTRAR NEUMÁTICO',
            icon=ft.Icons.ADD_CIRCLE_OUTLINE,
            visible=False
        )

        foxpro_values = {}
        foxpro_order = [
            'Código',
            'Serie',
            'Marca',
            'Medida',
            'Modelo / Diseño',
            'Clasificación TRA',
            'Costo $',
            'Estado',
            'Nro. Eventos',
            'Fecha',
            'Equipo-Posición',
            'Horómetro',
            'Odómetro',
            'Hrs Acumuladas',
            'Ext/Int - Inicial',
            'Ext/Int - Último',
            'Psi Act(F/C)-Rec',
            'Proyección Hrs',
            'Horas Acumuladas',
            'Costo x Hrs.',
            'Proyección Km',
            'Km Acumulados',
            'Costo x Km',
            'Tapa Válvula',
            'Lugar de Operación',
        ]
        for label in foxpro_order:
            foxpro_values[label] = ft.Text('—', size=13, color=TEXT_MAIN)

        header_tire = ft.Text('Seleccione un neumático', size=18, weight=ft.FontWeight.BOLD, color=TEXT_MAIN)
        header_detail = ft.Text('', size=12, color=TEXT_MUTED)

        ficha_rows = []
        vertical_labels = [
            'Nro. Eventos',
            'Fecha',
            'Equipo-Posición',
            'Horómetro',
            'Odómetro',
            'Hrs Acumuladas',
            'Ext/Int - Inicial',
            'Ext/Int - Último',
            'Psi Act(F/C)-Rec',
            'Proyección Hrs',
            'Horas Acumuladas',
            'Costo x Hrs.',
            'Proyección Km',
            'Km Acumulados',
            'Costo x Km',
            'Tapa Válvula',
            'Lugar de Operación',
            'Motivo',
            'Observaciones',
        ]
        # Cuatro columnas históricas. Se muestran en orden inverso: el evento
        # más reciente a la izquierda y el EVENTO 01 (instalación inicial) a la derecha.
        event_values = [
            {label: ft.Text('—', size=13, color=TEXT_MAIN) for label in vertical_labels}
            for _ in range(4)
        ]

        event_headers = [
            ft.Text('EVENTO 04', size=11, weight=ft.FontWeight.BOLD, color='#5E35B1'),
            ft.Text('EVENTO 03', size=11, weight=ft.FontWeight.BOLD, color='#9A4D00'),
            ft.Text('EVENTO 02', size=11, weight=ft.FontWeight.BOLD, color='#0D47A1'),
            ft.Text('EVENTO 01', size=11, weight=ft.FontWeight.BOLD, color='#1B5E20'),
        ]
        event_header_bg = ['#EFE7FF', '#FFF3E0', '#E3F2FD', '#E8F5E9']
        # Ancho uniforme de las columnas históricas. Aproximadamente 70%
        # del ancho que tenían originalmente cuando ocupaban el espacio disponible.
        HIST_EVENT_WIDTH = 175

        # Primera columna editable que aparece al seleccionar un evento.
        new_event_header = ft.Text('NUEVO EVENTO', size=11, weight=ft.FontWeight.BOLD, color='#1565C0')
        new_event_header_box = ft.Container(content=new_event_header, width=280, visible=False, bgcolor='#DDF3E4', padding=8)
        new_event_cells = {}
        new_event_text = {
            'Nro. Eventos': ft.Text('NUEVO', size=12, weight=ft.FontWeight.BOLD, color='#1565C0'),
            'Hrs Acumuladas': ft.Text('—', size=12, color=TEXT_MAIN),
            'Ext/Int - Inicial': ft.Text('—', size=12, color=TEXT_MAIN),
            'Proyección Hrs': ft.Text('—', size=12, color=TEXT_MAIN),
            'Horas Acumuladas': ft.Text('—', size=12, color=TEXT_MAIN),
            'Costo x Hrs.': ft.Text('—', size=12, color=TEXT_MAIN),
            'Proyección Km': ft.Text('—', size=12, color=TEXT_MAIN),
            'Km Acumulados': ft.Text('—', size=12, color=TEXT_MAIN),
            'Costo x Km': ft.Text('—', size=12, color=TEXT_MAIN),
            'Tapa Válvula': ft.Text('—', size=12, color=TEXT_MAIN),
        }
        inline_controls = {
            'Nro. Eventos': new_event_text['Nro. Eventos'],
            'Fecha': date,
            'Equipo-Posición': ft.Row([equip, pos, performance_info], spacing=4, tight=True, wrap=True),
            'Horómetro': horometer,
            'Odómetro': odometer,
            'Hrs Acumuladas': new_event_text['Hrs Acumuladas'],
            'Ext/Int - Inicial': new_event_text['Ext/Int - Inicial'],
            'Ext/Int - Último': ft.Row([to, ti], spacing=4, tight=True),
            'Psi Act(F/C)-Rec': ft.Row([press, cond], spacing=4, tight=True),
            'Proyección Hrs': new_event_text['Proyección Hrs'],
            'Horas Acumuladas': new_event_text['Horas Acumuladas'],
            'Costo x Hrs.': new_event_text['Costo x Hrs.'],
            'Proyección Km': new_event_text['Proyección Km'],
            'Km Acumulados': new_event_text['Km Acumulados'],
            'Costo x Km': new_event_text['Costo x Km'],
            'Tapa Válvula': valve_cap,
            'Lugar de Operación': loc,
            'Motivo': reason,
            'Observaciones': notes,
        }

        # Cabecera: NUEVO EVENTO + cuatro columnas históricas.
        ficha_rows.append(
            ft.Row([
                ft.Container(
                    content=ft.Text('DETALLE', size=11, weight=ft.FontWeight.BOLD, color='#FFFFFF'),
                    width=155, bgcolor=NAV_ACCENT, padding=8,
                    border_radius=ft.BorderRadius(top_left=6, top_right=6, bottom_left=0, bottom_right=0),
                ),
                new_event_header_box,
                *[
                    ft.Container(
                        content=event_headers[i], width=HIST_EVENT_WIDTH, bgcolor=event_header_bg[i], padding=8,
                        alignment=ft.Alignment(0, 0),
                        border=ft.Border(left=ft.BorderSide(1, '#B7C8D9'))
                    ) for i in range(4)
                ],
            ], spacing=0)
        )

        # Filas calculadas que se mantienen internamente, pero no se muestran
        # en la ficha histórica. No se eliminan cálculos ni datos.
        hidden_history_rows = {'Proyección Hrs', 'Proyección Km'}
        for label in vertical_labels:
            if label in hidden_history_rows:
                continue
            new_box = ft.Container(content=inline_controls[label], width=280, visible=False, bgcolor='#FFFFFF', padding=6, border=ft.Border(bottom=ft.BorderSide(1, '#DCE6EF')))
            new_event_cells[label] = new_box
            ficha_rows.append(
                ft.Row([
                    ft.Container(
                        content=ft.Text(label, size=11, weight=ft.FontWeight.W_600, color='#35556F'),
                        width=155, bgcolor='#EAF2F8', padding=6,
                        border=ft.Border(bottom=ft.BorderSide(1, '#DCE6EF'))
                    ),
                    new_box,
                    *[
                        ft.Container(
                            content=event_values[i][label], width=HIST_EVENT_WIDTH, padding=6,
                            bgcolor='#FFFFFF' if vertical_labels.index(label) % 2 == 0 else '#F7FAFD',
                            border=ft.Border(
                                left=ft.BorderSide(1, '#B7C8D9'),
                                bottom=ft.BorderSide(1, '#E4EBF2')
                            )
                        )
                        for i in range(4)
                    ],
                ], spacing=0)
            )

        inline_action_box = ft.Container(
            content=ft.Row([], spacing=8),
            width=280,
            visible=False,
            bgcolor='#F3F8FC',
            padding=6
        )
        ficha_rows.append(
            ft.Row([
                ft.Container(
                    content=ft.Text('Acción', size=11, weight=ft.FontWeight.W_600, color='#35556F'),
                    width=155, bgcolor='#EAF2F8', padding=6
                ),
                inline_action_box,
                ft.Container(width=HIST_EVENT_WIDTH, border=ft.Border(left=ft.BorderSide(1, '#B7C8D9'))),
                ft.Container(width=HIST_EVENT_WIDTH, border=ft.Border(left=ft.BorderSide(1, '#B7C8D9'))),
                ft.Container(width=HIST_EVENT_WIDTH, border=ft.Border(left=ft.BorderSide(1, '#B7C8D9'))),
                ft.Container(width=HIST_EVENT_WIDTH, border=ft.Border(left=ft.BorderSide(1, '#B7C8D9'))),
            ], spacing=0)
        )

        events_area = ft.Container(
            content=ft.Column(ficha_rows, spacing=0),
            bgcolor='#EDF4FA',
            padding=8,
            border_radius=8,
            border=ft.Border(
                left=ft.BorderSide(1, '#D3E0EC'), right=ft.BorderSide(1, '#D3E0EC'),
                top=ft.BorderSide(1, '#D3E0EC'), bottom=ft.BorderSide(1, '#D3E0EC')
            )
        )

        ficha_panel = card(
            ft.Column([
                ft.Row([
                    ft.Container(
                        content=ft.Column([
                            ft.Text('Consulta del neumático', size=12, weight=ft.FontWeight.W_600, color=TEXT_MUTED),
                            header_tire,
                        ], spacing=1, alignment=ft.MainAxisAlignment.CENTER),
                        expand=2,
                        height=74,
                        alignment=ft.Alignment(-1, 0),
                    ),
                    ft.Container(
                        content=ft.Column([
                            ft.Text('Marca:', size=11, weight=ft.FontWeight.W_600, color=TEXT_MUTED),
                            foxpro_values['Marca'],
                            ft.Text('Diseño:', size=11, weight=ft.FontWeight.W_600, color=TEXT_MUTED),
                            foxpro_values['Modelo / Diseño'],
                        ], spacing=2),
                        expand=1
                    ),
                    ft.Container(
                        content=ft.Column([
                            ft.Text('Medida:', size=11, weight=ft.FontWeight.W_600, color=TEXT_MUTED),
                            foxpro_values['Medida'],
                            ft.Text('Estado:', size=11, weight=ft.FontWeight.W_600, color=TEXT_MUTED),
                            foxpro_values['Estado'],
                        ], spacing=2),
                        expand=1
                    ),
                    ft.Container(
                        content=ft.Column([
                            ft.Text('Clasificación TRA:', size=11, weight=ft.FontWeight.W_600, color=TEXT_MUTED),
                            foxpro_values['Clasificación TRA'],
                            ft.Text('Costo $:', size=11, weight=ft.FontWeight.W_600, color=TEXT_MUTED),
                            foxpro_values['Costo $'],
                        ], spacing=2),
                        expand=1
                    ),
                ], spacing=18, vertical_alignment=ft.CrossAxisAlignment.CENTER),
                ft.Divider(height=4),
                events_area,
            ], spacing=7),
            width=None
        )

        # El formulario tradicional se mantiene como marcador interno, pero ya no se muestra.
        # Los mismos controles se editan directamente en la primera columna de eventos.
        movement_form = ft.Container(visible=False)
        inline_mode = {'active': False}

        cancel_inline_btn = ft.TextButton('Cancelar', icon=ft.Icons.CLOSE)
        save_btn.text = 'GUARDAR'
        save_btn.icon = ft.Icons.SAVE
        inline_action_box.content = ft.Row([save_btn, cancel_inline_btn], spacing=6, wrap=True)

        def set_inline_event_mode(active):
            inline_mode['active'] = bool(active)
            new_event_header_box.visible = bool(active)
            inline_action_box.visible = bool(active)
            for box in new_event_cells.values():
                box.visible = bool(active)
            if not active:
                event.value = None

        def refresh_new_event_preview():
            if not inline_mode['active'] or not tire.value:
                return
            r = current_tire()
            if not r:
                return
            ec = event.value or ''
            new_event_header.value = f'NUEVO: {ec}' if ec else 'NUEVO EVENTO'

            occ = query('SELECT * FROM occurrences WHERE tire_id=? AND operation_id=? ORDER BY id', (int(tire.value), op_id))
            base_inst = None
            for item in reversed(occ):
                if item['event_code'] == 'INST':
                    base_inst = item
                    break

            init_e = base_inst['tread_outer'] if base_inst and base_inst['tread_outer'] is not None else r['new_tread']
            init_i = base_inst['tread_inner'] if base_inst and base_inst['tread_inner'] is not None else r['new_tread']
            new_event_text['Ext/Int - Inicial'].value = f"{fmt(init_e)}/{fmt(init_i)}"

            event_distance = None
            try:
                m = num(active_meter_value())
                if m is not None and base_inst and base_inst['meter'] is not None:
                    event_distance = max(0, float(m) - float(base_inst['meter']))
            except Exception:
                event_distance = None
            perf = selected_performance_type()
            hrs_text = f'{event_distance:.1f}' if perf == PERFORMANCE_HOURS and event_distance is not None else '—'
            km_text = f'{event_distance:.1f}' if perf == PERFORMANCE_KILOMETERS and event_distance is not None else '—'
            new_event_text['Hrs Acumuladas'].value = hrs_text
            new_event_text['Horas Acumuladas'].value = hrs_text
            new_event_text['Km Acumulados'].value = km_text

            life_value = r['projected_life_target'] if r['projected_life_target'] is not None else r['projected_life']
            new_event_text['Proyección Hrs'].value = fmt(life_value) if perf == PERFORMANCE_HOURS else '—'
            new_event_text['Proyección Km'].value = fmt(life_value) if perf == PERFORMANCE_KILOMETERS else '—'
            try:
                c = float(r['cost_usd']) if r['cost_usd'] is not None else None
            except Exception:
                c = None
            new_event_text['Costo x Hrs.'].value = (
                f'$ {c / event_distance:.2f}/h' if perf == PERFORMANCE_HOURS and c is not None and event_distance is not None and event_distance > 0 else '—'
            )
            new_event_text['Costo x Km'].value = (
                f'$ {c / event_distance:.2f}/km' if perf == PERFORMANCE_KILOMETERS and c is not None and event_distance is not None and event_distance > 0 else '—'
            )
            if str(valve_cap.value or '').strip().upper() not in ('SI','NO'):
                valve_cap.value = 'NO'

        def cancel_inline(e=None):
            set_inline_event_mode(False)
            if tire.value:
                load_foxpro_ficha(int(tire.value))
                update_event_button_states(int(tire.value))
            page.update()

        cancel_inline_btn.on_click = cancel_inline

        def clear_foxpro_ficha():
            header_tire.value = 'Seleccione un neumático'
            header_detail.value = ''
            for label in foxpro_order:
                foxpro_values[label].value = '—'
            for i in range(3):
                for label in vertical_labels:
                    event_values[i][label].value = '—'
            set_inline_event_mode(False)
            event_headers[0].value = 'ÚLTIMO EVENTO'
            event_headers[1].value = 'PENÚLTIMO EVENTO'
            event_headers[2].value = 'ANTEPENÚLTIMO EVENTO'

        def load_foxpro_ficha(tid):
            rows = query(
                '''SELECT t.*,e.code equipment_code
                   FROM tires t LEFT JOIN equipment e ON e.id=t.equipment_id AND e.operation_id=t.operation_id
                   WHERE t.id=? AND t.operation_id=?''',
                (int(tid), op_id)
            )
            if not rows:
                clear_foxpro_ficha()
                return
            r = rows[0]
            # Orden cronológico real de los eventos. No dependemos del formato
            # textual almacenado en SQLite ni del id de inserción.
            occ_raw = query(
                '''SELECT * FROM occurrences
                   WHERE tire_id=? AND operation_id=?''',
                (int(tid), op_id)
            )

            def occurrence_sort_key(item):
                parsed = parse_event_date(item['event_date'])
                # Fechas no interpretables quedan al inicio para no desplazar
                # un evento válido reciente de la columna "Último".
                return (parsed or dt.datetime.min, int(item['id'] or 0))

            occ = sorted(occ_raw, key=occurrence_sort_key)
            last = occ[-1] if occ else None
            first = occ[0] if occ else None

            inst = [o for o in occ if o['event_code'] == 'INST']
            last_inst = inst[-1] if inst else None

            # Para la ficha operativa usamos el horómetro del evento
            # cronológicamente más reciente. Si no existe, conservamos el
            # current_meter del maestro como respaldo.
            current_meter = (
                last['meter'] if last and last['meter'] is not None
                else r['current_meter']
            )
            inst_meter = last_inst['meter'] if last_inst else None
            hrs_acum = None
            try:
                if current_meter is not None and inst_meter is not None:
                    hrs_acum = max(0, float(current_meter) - float(inst_meter))
            except Exception:
                hrs_acum = None

            initial_i = first['tread_inner'] if first and first['tread_inner'] is not None else r['new_tread']
            initial_e = first['tread_outer'] if first and first['tread_outer'] is not None else r['new_tread']
            last_i = last['tread_inner'] if last and last['tread_inner'] is not None else r['tread_inner']
            last_e = last['tread_outer'] if last and last['tread_outer'] is not None else r['tread_outer']

            actual_press = last['pressure'] if last and last['pressure'] is not None else None
            press_cond = last['pressure_condition'] if last and last['pressure_condition'] else ''
            rec_press = r['recommended_pressure']

            header_tire.value = f"{r['code']} | {r['serial'] or 's/serie'}"
            header_detail.value = (
                f"{r['brand'] or ''} · {r['size'] or ''} · {r['design'] or ''} · Estado: {r['status'] or '-'}"
            )

            foxpro_values['Código'].value = str(r['code'] or '—')
            foxpro_values['Serie'].value = str(r['serial'] or 's/serie')
            foxpro_values['Marca'].value = str(r['brand'] or '—')
            foxpro_values['Medida'].value = str(r['size'] or '—')
            foxpro_values['Modelo / Diseño'].value = str(r['design'] or '—')
            foxpro_values['Clasificación TRA'].value = str(r['compound'] or '—')
            try:
                header_cost = float(r['cost_usd']) if r['cost_usd'] is not None else None
            except Exception:
                header_cost = None
            foxpro_values['Costo $'].value = f"$ {header_cost:,.2f}" if header_cost is not None else '—'
            foxpro_values['Estado'].value = str(r['status'] or '—')
            foxpro_values['Nro. Eventos'].value = str(len(occ))
            foxpro_values['Fecha'].value = format_date(last['event_date']) if last else '—'
            foxpro_values['Equipo-Posición'].value = (
                f"{r['equipment_code'] or '-'} - P{r['position'] or '-'}"
            )
            perf = get_equipment_performance_type(r['equipment_id'], op_id) if r['equipment_id'] is not None else None
            foxpro_values['Horómetro'].value = fmt(current_meter) if perf == PERFORMANCE_HOURS else '—'
            foxpro_values['Odómetro'].value = fmt(current_meter) if perf == PERFORMANCE_KILOMETERS else '—'
            foxpro_values['Hrs Acumuladas'].value = f"{hrs_acum:.1f}" if perf == PERFORMANCE_HOURS and hrs_acum is not None else '—'
            foxpro_values['Ext/Int - Inicial'].value = f"{fmt(initial_e)}/{fmt(initial_i)}"
            foxpro_values['Ext/Int - Último'].value = f"{fmt(last_e)}/{fmt(last_i)}"
            foxpro_values['Psi Act(F/C)-Rec'].value = (
                f"{fmt(actual_press) or '-'} ({press_cond or '-'}) / {fmt(rec_press) or '-'}"
            )
            life_value = r['projected_life_target'] if r['projected_life_target'] is not None else r['projected_life']
            foxpro_values['Proyección Hrs'].value = fmt(life_value) if perf == PERFORMANCE_HOURS else '—'
            foxpro_values['Horas Acumuladas'].value = f"{hrs_acum:.1f}" if perf == PERFORMANCE_HOURS and hrs_acum is not None else '—'
            foxpro_values['Costo x Hrs.'].value = (
                f"$ {header_cost / hrs_acum:.2f}/h" if perf == PERFORMANCE_HOURS and header_cost is not None and hrs_acum is not None and hrs_acum > 0 else '—'
            )
            foxpro_values['Proyección Km'].value = fmt(life_value) if perf == PERFORMANCE_KILOMETERS else '—'
            foxpro_values['Km Acumulados'].value = f"{hrs_acum:.1f}" if perf == PERFORMANCE_KILOMETERS and hrs_acum is not None else '—'
            foxpro_values['Costo x Km'].value = (
                f"$ {header_cost / hrs_acum:.2f}/km" if perf == PERFORMANCE_KILOMETERS and header_cost is not None and hrs_acum is not None and hrs_acum > 0 else '—'
            )
            valve_raw = str(last['valve_cap'] or '').strip().upper() if last and 'valve_cap' in last.keys() else ''
            if valve_raw in ('SI','SÍ'):
                foxpro_values['Tapa Válvula'].value='SI'
            elif valve_raw == 'NO':
                foxpro_values['Tapa Válvula'].value='NO'
            else:
                note_text = str(last['notes'] or '').upper() if last and 'notes' in last.keys() else ''
                foxpro_values['Tapa Válvula'].value = 'SI' if ('TAPA' in note_text or 'VALVULA' in note_text or 'VÁLVULA' in note_text) else 'NO'
            foxpro_values['Lugar de Operación'].value = fmt(last['location']) if last and last['location'] else '—'


            # --------------------------------------------------------------
            # Visualización de los cuatro últimos eventos en paralelo.
            # Se muestran de izquierda a derecha en orden inverso. EVENTO 01
            # corresponde al primer evento registrado (normalmente INST).
            # --------------------------------------------------------------
            last_four = list(reversed(occ[-4:]))

            def fill_event_column(col_idx, target):
                values = event_values[col_idx]
                if not target:
                    for label in vertical_labels:
                        values[label].value = '—'
                    return

                # Posición ordinal real del evento dentro del historial.
                target_index = next(
                    (i for i, item in enumerate(occ) if item['id'] == target['id']),
                    None
                )
                event_number = (target_index + 1) if target_index is not None else '—'

                # Última instalación vigente hasta ese evento.
                base_inst = None
                if target_index is not None:
                    for item in reversed(occ[:target_index + 1]):
                        if item['event_code'] == 'INST':
                            base_inst = item
                            break

                event_hours = None
                try:
                    if target['meter'] is not None and base_inst and base_inst['meter'] is not None:
                        event_hours = max(0, float(target['meter']) - float(base_inst['meter']))
                except Exception:
                    event_hours = None

                init_i = (
                    base_inst['tread_inner']
                    if base_inst and base_inst['tread_inner'] is not None
                    else r['new_tread']
                )
                init_e = (
                    base_inst['tread_outer']
                    if base_inst and base_inst['tread_outer'] is not None
                    else r['new_tread']
                )
                evt_i = target['tread_inner'] if target['tread_inner'] is not None else '—'
                evt_e = target['tread_outer'] if target['tread_outer'] is not None else '—'
                evt_press = target['pressure'] if target['pressure'] is not None else None
                evt_cond = target['pressure_condition'] if target['pressure_condition'] else ''

                values['Nro. Eventos'].value = str(event_number)
                values['Fecha'].value = format_date(target['event_date'])
                values['Equipo-Posición'].value = (
                    f"{target['equipment_id'] or '-'} - P{target['position'] or '-'}"
                )

                # Mostrar código de equipo en lugar del id cuando exista.
                if target['equipment_id']:
                    eq_row = query(
                        'SELECT code FROM equipment WHERE id=? AND operation_id=?',
                        (int(target['equipment_id']), op_id)
                    )
                    if eq_row:
                        values['Equipo-Posición'].value = (
                            f"{eq_row[0]['code']} - P{target['position'] or '-'}"
                        )

                target_perf = get_equipment_performance_type(target['equipment_id'], op_id) if target['equipment_id'] is not None else None
                values['Horómetro'].value = fmt(target['meter']) if target_perf == PERFORMANCE_HOURS else '—'
                values['Odómetro'].value = fmt(target['meter']) if target_perf == PERFORMANCE_KILOMETERS else '—'
                values['Hrs Acumuladas'].value = (
                    f"{event_hours:.1f}" if target_perf == PERFORMANCE_HOURS and event_hours is not None else '—'
                )
                values['Ext/Int - Inicial'].value = f"{fmt(init_e)}/{fmt(init_i)}"
                values['Ext/Int - Último'].value = f"{fmt(evt_e)}/{fmt(evt_i)}"
                values['Psi Act(F/C)-Rec'].value = (
                    f"{fmt(evt_press) or '-'} ({evt_cond or '-'}) / {fmt(rec_press) or '-'}"
                )
                values['Proyección Hrs'].value = fmt(life_value) if target_perf == PERFORMANCE_HOURS else '—'
                values['Horas Acumuladas'].value = (
                    f"{event_hours:.1f}" if target_perf == PERFORMANCE_HOURS and event_hours is not None else '—'
                )
                values['Costo x Hrs.'].value = (
                    f"$ {header_cost / event_hours:.2f}/h" if target_perf == PERFORMANCE_HOURS and header_cost is not None and event_hours is not None and event_hours > 0 else '—'
                )
                values['Proyección Km'].value = fmt(life_value) if target_perf == PERFORMANCE_KILOMETERS else '—'
                values['Km Acumulados'].value = (
                    f"{event_hours:.1f}" if target_perf == PERFORMANCE_KILOMETERS and event_hours is not None else '—'
                )
                values['Costo x Km'].value = (
                    f"$ {header_cost / event_hours:.2f}/km" if target_perf == PERFORMANCE_KILOMETERS and header_cost is not None and event_hours is not None and event_hours > 0 else '—'
                )
                valve_raw = str(target['valve_cap'] or '').strip().upper() if 'valve_cap' in target.keys() else ''
                if valve_raw in ('SI','SÍ'):
                    values['Tapa Válvula'].value='SI'
                elif valve_raw == 'NO':
                    values['Tapa Válvula'].value='NO'
                else:
                    note_text = str(target['notes'] or '').upper() if 'notes' in target.keys() else ''
                    values['Tapa Válvula'].value = 'SI' if ('TAPA' in note_text or 'VALVULA' in note_text or 'VÁLVULA' in note_text) else 'NO'
                values['Lugar de Operación'].value = fmt(target['location']) if target['location'] else '—'
                values['Motivo'].value = fmt(target['reason']) if 'reason' in target.keys() and target['reason'] else '—'
                values['Observaciones'].value = fmt(target['notes']) if 'notes' in target.keys() and target['notes'] else '—'

            for col_idx in range(4):
                target = last_four[col_idx] if col_idx < len(last_four) else None
                fill_event_column(col_idx, target)
                if target:
                    # EVENTO 01 es el primer evento del historial; por eso usamos
                    # su ordinal real y mostramos el tipo entre paréntesis.
                    target_index = next(
                        (i for i, item in enumerate(occ) if item['id'] == target['id']),
                        None
                    )
                    event_number = (target_index + 1) if target_index is not None else None
                    if event_number is not None:
                        event_headers[col_idx].value = f"EVENTO {event_number:02d}\n({target['event_code']})"
                    else:
                        event_headers[col_idx].value = f"EVENTO --\n({target['event_code']})"
                else:
                    # Mantiene la secuencia visual 04 / 03 / 02 / 01 cuando aún
                    # no existe información suficiente para llenar las cuatro columnas.
                    event_headers[col_idx].value = f"EVENTO {4-col_idx:02d}"

        def select_operational_tire(tid):
            if not tid:
                return
            # V22: rechazo explícito de cualquier neumático que no pertenezca al taller activo.
            _owned=query('SELECT id FROM tires WHERE id=? AND operation_id=?',(int(tid),op_id))
            if not _owned:
                tire.value=None
                clear_foxpro_ficha()
                update_event_button_states(None)
                return snack('El neumático no pertenece a la operación activa.', True)
            tire.value = str(tid)
            set_inline_event_mode(False)
            movement_form.visible = False
            load_foxpro_ficha(tid)
            update_event_button_states(tid)
            refresh()
            movement_form.visible = False
            page.update()

        def go_register_missing(e=None):
            pending_code = (search_tire.value or '').strip()
            if not pending_code:
                return
            tires_view(prefill_code=pending_code)

        register_missing_btn.on_click = go_register_missing

        def do_search(e=None):
            term = (search_tire.value or '').strip()
            set_inline_event_mode(False)
            movement_form.visible = False
            register_missing_btn.visible = False
            if not term:
                search_result.visible = False
                search_result.options = []
                search_result.value = None
                tire.value = None
                clear_foxpro_ficha()
                update_event_button_states(None)
                page.update()
                return

            rows = query(
                '''SELECT id,code,serial FROM tires
                   WHERE operation_id=? AND (code LIKE ? OR serial LIKE ?)
                   ORDER BY code''',
                (op_id, f'%{term}%', f'%{term}%')
            )
            search_result.options = [
                ft.dropdown.Option(str(r['id']), f"{r['code']} | {r['serial'] or 's/serie'}")
                for r in rows
            ]
            if len(rows) == 1:
                search_result.visible = False
                search_result.value = str(rows[0]['id'])
                select_operational_tire(rows[0]['id'])
            elif len(rows) > 1:
                search_result.visible = True
                search_result.value = None
                clear_foxpro_ficha()
                update_event_button_states(None)
                page.update()
            else:
                search_result.visible = False
                search_result.value = None
                tire.value = None
                clear_foxpro_ficha()
                update_event_button_states(None)
                register_missing_btn.visible = True
                snack('Código no registrado. Puede registrar el neumático desde el botón habilitado.', True)
                page.update()

        def on_search_result(e):
            selected = getattr(e, 'data', None) or search_result.value
            if selected:
                search_result.value = str(selected)
                select_operational_tire(selected)

        search_tire.on_submit = do_search
        search_tire.on_change = do_search
        search_result.on_select = on_search_result

        event_icons_local = {
            'INST': ft.Icons.ADD_CIRCLE_OUTLINE,
            'INSP': ft.Icons.CHECK_CIRCLE_OUTLINE,
            'INSC': ft.Icons.FACT_CHECK_OUTLINED,
            'ROT': ft.Icons.SYNC_ALT,
            'INVE': ft.Icons.SWAP_HORIZ,
            'DINS': ft.Icons.REMOVE_CIRCLE_OUTLINE,
            'REPA': ft.Icons.HANDYMAN_OUTLINED,
            'BAJA': ft.Icons.DELETE_OUTLINE,
        }
        event_buttons_local = {}

        def update_event_button_states(tid=None):
            # Matriz operativa fundamental por estado:
            # - SERVICIO + instalado: INSP, INSC, INVE, DINS y BAJA.
            #   REPA queda bloqueado mientras el neumático esté instalado.
            # - STAND-BY: únicamente INST, INVE, REPA y BAJA.
            # - BAJA / REPARACIÓN / otros estados: todos bloqueados.
            # - ROT permanece bloqueado hasta definir su funcionalidad.
            if not tid:
                for code, btn in event_buttons_local.items():
                    btn.disabled = True
                return

            rows = query(
                """SELECT id,equipment_id,position,status
                   FROM tires
                   WHERE id=? AND operation_id=?""",
                (int(tid), op_id)
            )
            if not rows:
                for code, btn in event_buttons_local.items():
                    btn.disabled = True
                return

            r = rows[0]
            installed = r['equipment_id'] is not None and str(r['position'] or '').strip() != ''
            status_norm = str(r['status'] or '').strip().upper().replace('_','-')

            if installed and status_norm == 'SERVICIO':
                allowed = {'INSP','INSC','INVE','DINS','BAJA'}
            elif status_norm in ('STAND-BY','STAND BY','STANDBY'):
                allowed = {'INST','INVE','REPA','BAJA'}
            else:
                allowed = set()

            for code, btn in event_buttons_local.items():
                btn.disabled = code not in allowed

        def open_event_form(ec):
            if not tire.value:
                return snack('Primero busque y seleccione un neumático.', True)
            if ec == 'ROT':
                return snack('ROT permanece bloqueado hasta definir su funcionalidad.', True)
            event.value = ec
            movement_form.visible = False
            set_inline_event_mode(True)
            load_current_state()
            apply_event_rules()
            refresh_new_event_preview()
            update_save_state()
            page.update()

        for ec in ['INST','INSP','INSC','ROT','INVE','DINS','REPA','BAJA']:
            event_buttons_local[ec] = ft.OutlinedButton(
                ec,
                icon=event_icons_local[ec],
                tooltip=EVENTS.get(ec, ec),
                disabled=True,
                on_click=lambda e, code=ec: open_event_form(code)
            )

        if pre_tire:
            search_tire.value = str(current_tire()['code']) if current_tire() else ''
            load_foxpro_ficha(pre_tire)
            update_event_button_states(pre_tire)

        left_panel = ft.Column([
            card(ft.Column([
                ft.Text('Consulta operativa', size=17, weight=ft.FontWeight.BOLD, color=TEXT_MAIN),
                search_tire,
                search_result,
                register_missing_btn,
            ], spacing=8), width=285),
            card(
                ft.Column([
                    ft.Text('Registrar evento', size=12, weight=ft.FontWeight.W_600, color=TEXT_MUTED),
                    *[
                        ft.Container(
                            content=event_buttons_local[c],
                            width=140
                        )
                        for c in ['INST','INSP','INSC','ROT','INVE','DINS','REPA','BAJA']
                    ],
                ], spacing=8),
                width=185
            ),
        ], spacing=12)

        top_operational_area = ft.Row([
            ft.Container(content=left_panel, width=300),
            ft.Container(content=ficha_panel, expand=True),
        ], spacing=12, vertical_alignment=ft.CrossAxisAlignment.START)

        content.content=ft.Column([
            page_title('Movimiento de neumáticos','Registro operativo del ciclo de vida'),
            top_operational_area,
            card(ft.Column([
                ft.Text('Historial del neumático',size=17,weight=ft.FontWeight.BOLD,color=TEXT_MAIN),
                ft.Text('Use la barra inferior para desplazarse horizontalmente. La columna Acción permite eliminar eventos.',size=11,color=TEXT_MUTED),
                history_scroller
            ]))
        ],scroll=ft.ScrollMode.AUTO,spacing=16)
        page.update()

    def maintenance_menu_view():
        """3. Programa de mantenimiento - menú visual de accesos."""
        maintenance_meta={
            '31':{
                'num':'3.1','icon':ft.Icons.CHECK_CIRCLE_OUTLINE,'accent':'#1565C0','soft':'#EEF5FF',
                'title':'EVALUACIÓN DE\nREMANENTE',
                'desc':'Condición general según profundidad remanente RTD.',
                'action':maintenance_view,
            },
            '32':{
                'num':'3.2','icon':ft.Icons.COMPARE_ARROWS_OUTLINED,'accent':'#138A3D','soft':'#EEFAF2',
                'title':'DIFERENCIA RTD\nENTRE HOMBROS',
                'desc':'Comparación de cocada exterior e interior por neumático.',
                'action':maintenance_shoulders_view,
            },
            '33':{
                'num':'3.3','icon':ft.Icons.SWAP_HORIZ_OUTLINED,'accent':'#EF6C00','soft':'#FFF5EA',
                'title':'DIFERENCIA RTD\nMISMO EJE',
                'desc':'Evaluación Tire Mismatch entre posiciones del mismo eje.',
                'action':maintenance_axles_view,
            },
            '34':{
                'num':'3.4','icon':ft.Icons.ALIGN_HORIZONTAL_CENTER_OUTLINED,'accent':'#7B1FA2','soft':'#F8EEFC',
                'title':'DIFERENCIA ENTRE EJES\nPOR EQUIPO',
                'desc':'Comparación del RTD entre las cuatro posiciones del equipo.',
                'action':maintenance_four_positions_view,
            },
            '35':{
                'num':'3.5','icon':ft.Icons.SPEED_OUTLINED,'accent':'#C62828','soft':'#FFF0F0',
                'title':'NIVELACIÓN DE\nPRESIÓN',
                'desc':'Evaluación de presión actual y condición de tapa válvula.',
                'action':maintenance_pressure_view,
            },
            '36':{
                'num':'3.6','icon':ft.Icons.DESCRIPTION_OUTLINED,'accent':'#455A64','soft':'#F1F5F7',
                'title':'REPORTE FINAL DE\nMANTENIMIENTO',
                'desc':'Consolidado de actividades y acciones generadas por las evaluaciones 3.1–3.5.',
                'action':maintenance_final_report_view,
            },
        }

        def maintenance_access_card(key):
            m=maintenance_meta[key]
            return ft.Container(
                width=292,height=260,bgcolor=m['soft'],
                border=ft.Border.all(1,m['accent']+'55'),border_radius=14,padding=18,
                on_click=lambda e,k=key: maintenance_meta[k]['action'](),ink=True,
                content=ft.Column([
                    ft.Row([
                        ft.Container(bgcolor=m['accent'],border_radius=9,padding=ft.Padding(11,6,11,6),
                                     content=ft.Text(m['num'],size=16,weight=ft.FontWeight.BOLD,color='#FFFFFF')),
                        ft.Container(expand=True),
                        ft.Container(width=54,height=54,border_radius=14,bgcolor='#FFFFFF',alignment=ft.Alignment.CENTER,
                                     content=ft.Icon(m['icon'],size=31,color=m['accent'])),
                    ]),
                    ft.Text(m['title'],size=16,weight=ft.FontWeight.BOLD,color=TEXT_MAIN,text_align=ft.TextAlign.CENTER),
                    ft.Text(m['desc'],size=11,color=TEXT_MUTED,text_align=ft.TextAlign.CENTER),
                    ft.Container(height=4),
                    ft.Container(bgcolor=m['accent'],border_radius=9,padding=10,alignment=ft.Alignment.CENTER,
                                 content=ft.Row([
                                     ft.Icon(ft.Icons.BAR_CHART,color='#FFFFFF',size=20),
                                     ft.Text('VER REPORTE',color='#FFFFFF',weight=ft.FontWeight.BOLD,size=13),
                                     ft.Icon(ft.Icons.CHEVRON_RIGHT,color='#FFFFFF',size=20),
                                 ],alignment=ft.MainAxisAlignment.CENTER,spacing=8)),
                ],spacing=12,horizontal_alignment=ft.CrossAxisAlignment.CENTER)
            )

        empty_slot=ft.Container(width=292,height=260)
        content.content=ft.Column([
            page_title('3. PROGRAMA DE MANTENIMIENTO','Evaluación técnica y acciones preventivas de neumáticos en servicio'),
            card(ft.Column([
                ft.Row([
                    ft.Container(width=48,height=48,border_radius=24,bgcolor='#EAF2FF',alignment=ft.Alignment.CENTER,
                                 content=ft.Icon(ft.Icons.BUILD_CIRCLE_OUTLINED,color=NAV_ACCENT,size=27)),
                    ft.Column([
                        ft.Text('PROGRAMA DE MANTENIMIENTO',size=18,weight=ft.FontWeight.BOLD,color=TEXT_MAIN),
                        ft.Text('Seleccione una evaluación para visualizar el detalle.',size=11,color=TEXT_MUTED),
                    ],spacing=2),
                ],spacing=12),
                ft.Row([maintenance_access_card('31'),maintenance_access_card('32'),maintenance_access_card('33'),maintenance_access_card('34')],spacing=12),
                ft.Row([maintenance_access_card('35'),maintenance_access_card('36'),ft.Container(width=292,height=260),ft.Container(width=292,height=260)],spacing=12),
            ],spacing=16),padding=16),
        ],scroll=ft.ScrollMode.AUTO,spacing=16)
        page.update()

    def maintenance_nav_button(label, icon, on_click=None, active=False):
        """Acceso uniforme para los submódulos del Programa de mantenimiento."""
        return ft.Button(
            label,
            icon=icon,
            on_click=on_click,
            width=280,
            height=44,
            disabled=active,
            bgcolor='#D7DFEA' if active else '#EAF2FF',
            color='#5D6875' if active else '#1E5AA8',
            style=ft.ButtonStyle(
                shape=ft.RoundedRectangleBorder(radius=10),
                text_style=ft.TextStyle(size=14, weight=ft.FontWeight.BOLD),
            ),
        )

    def maintenance_nav_row(active=None):
        """Accesos estandarizados 3.1–3.6 + volver. Todos con el mismo tamaño."""
        return ft.Row([
            maintenance_nav_button('VOLVER A PROGRAMA DE MANTENIMIENTO', ft.Icons.ARROW_BACK,
                                   lambda e: maintenance_menu_view(), active=False),
            maintenance_nav_button('3.1 Evaluación de remanente', ft.Icons.CHECK_CIRCLE_OUTLINE,
                                   lambda e: maintenance_view(), active == '31'),
            maintenance_nav_button('3.2 Diferencia RTD entre hombros', ft.Icons.COMPARE_ARROWS,
                                   lambda e: maintenance_shoulders_view(), active == '32'),
            maintenance_nav_button('3.3 Diferencia RTD mismo eje', ft.Icons.COMPARE_ARROWS,
                                   lambda e: maintenance_axles_view(), active == '33'),
            maintenance_nav_button('3.4 Diferencia entre ejes por equipo', ft.Icons.COMPARE_ARROWS,
                                   lambda e: maintenance_four_positions_view(), active == '34'),
            maintenance_nav_button('3.5 Nivelación de presión', ft.Icons.SPEED,
                                   lambda e: maintenance_pressure_view(), active == '35'),
            maintenance_nav_button('3.6 Reporte final de mantenimiento', ft.Icons.DESCRIPTION_OUTLINED,
                                   lambda e: maintenance_final_report_view(), active == '36'),
        ], spacing=10, run_spacing=10, wrap=True)

    def maintenance_view():
        """Programa de mantenimiento - Prueba 01: evaluación de remanente (RTD)."""
        op_id = active_operation_id()
        rows = query("""
            SELECT
                t.code, t.serial, t.brand, t.size, t.design,
                t.tread_inner, t.tread_outer,
                e.id AS equipment_id, e.code AS equipment_code,
                e.vehicle_type, e.model AS equipment_model, t.position
            FROM tires t
            LEFT JOIN equipment e ON e.id=t.equipment_id
            WHERE t.status='SERVICIO' AND t.operation_id=?
            ORDER BY COALESCE(e.code,''), t.position, t.code
        """, (op_id,))

        def rtd_value(r):
            vals=[]
            for value in (r['tread_inner'], r['tread_outer']):
                try:
                    if value is not None and str(value).strip() != '':
                        vals.append(float(value))
                except Exception:
                    pass
            return min(vals) if vals else None

        def rtd_condition(rtd):
            # Criterio general conservado para los indicadores globales.
            if rtd is None or rtd < 0:
                return 'SIN LECTURA'
            if rtd > 30:
                return 'BUEN ESTADO'
            if rtd > 20:
                return 'PRÓXIMO CAMBIO'
            return 'CAMBIO URGENTE'

        def rtd_condition_by_category(category, rtd):
            # Semáforo específico para remanente mínimo por tipo de equipo.
            # El valor menor representa mayor proximidad al retiro.
            if rtd is None or rtd < 0:
                return 'SIN LECTURA'
            thresholds = {
                # Cambio urgente: 0–20 / 0–10 / 0–4.
                # Próximo cambio: 21–30 / 11–15 / 5–6.
                # Buen estado: >30 / >15 / >6.
                'SCOOP': (20.0, 30.0),
                'VOLQUETES': (10.0, 15.0),
                'CAMIONETAS': (4.0, 6.0),
                'OTROS': (10.0, 15.0),
            }
            emergency_max, preventive_max = thresholds.get(str(category).upper(), (10.0, 15.0))
            if rtd > preventive_max:
                return 'BUEN ESTADO'
            if rtd > emergency_max:
                return 'PRÓXIMO CAMBIO'
            return 'CAMBIO URGENTE'

        evaluated=[]
        for r in rows:
            rtd=rtd_value(r)
            evaluated.append((r, rtd, rtd_condition(rtd)))

        total=len(rows)
        equipment_count=len({r['equipment_id'] for r in rows if r['equipment_id'] is not None})
        counts={
            'BUEN ESTADO': 0,
            'PRÓXIMO CAMBIO': 0,
            'CAMBIO URGENTE': 0,
            'SIN LECTURA': 0,
        }
        for _,_,condition in evaluated:
            counts[condition]=counts.get(condition,0)+1

        evaluated_total=total-counts['SIN LECTURA']
        good_pct=(counts['BUEN ESTADO']/evaluated_total*100) if evaluated_total else 0
        attention=counts['PRÓXIMO CAMBIO']+counts['CAMBIO URGENTE']
        attention_pct=(attention/evaluated_total*100) if evaluated_total else 0

        def pct(value, base):
            return (value/base*100) if base else 0

        def top_metric(title, value, subtitle, value_color=TEXT_MAIN):
            return ft.Container(
                expand=1,
                height=108,
                bgcolor=CARD_BG,
                border=ft.Border.all(1, '#DDE5ED'),
                border_radius=10,
                padding=14,
                content=ft.Column([
                    ft.Text(title, size=12, weight=ft.FontWeight.BOLD, color=TEXT_MAIN,
                            text_align=ft.TextAlign.CENTER),
                    ft.Text(str(value), size=29, weight=ft.FontWeight.BOLD, color=value_color,
                            text_align=ft.TextAlign.CENTER),
                    ft.Text(subtitle, size=10, color=TEXT_MUTED, text_align=ft.TextAlign.CENTER),
                ], spacing=3, horizontal_alignment=ft.CrossAxisAlignment.CENTER)
            )

        criteria_rows = ft.Column([
            ft.Container(
                bgcolor='#F5F7FA', padding=8,
                content=ft.Row([
                    ft.Text('CONDICIÓN', width=150, size=11, weight=ft.FontWeight.BOLD, color=TEXT_MAIN),
                    ft.Text('CRITERIO RTD', width=150, size=11, weight=ft.FontWeight.BOLD, color=TEXT_MAIN),
                    ft.Text('INTERPRETACIÓN', expand=True, size=11, weight=ft.FontWeight.BOLD, color=TEXT_MAIN),
                ])
            ),
            ft.Container(
                border=ft.Border(bottom=ft.BorderSide(1,'#E6EBF0')), padding=8,
                content=ft.Row([
                    ft.Container(width=8, height=28, bgcolor='#2E9B45', border_radius=4),
                    ft.Text('BUEN ESTADO', width=135, size=11, weight=ft.FontWeight.BOLD, color='#2E9B45'),
                    ft.Text('Mayor a 30 mm', width=150, size=11, color=TEXT_MAIN),
                    ft.Text('Continuar en servicio.', expand=True, size=11, color=TEXT_MUTED),
                ], spacing=8)
            ),
            ft.Container(
                border=ft.Border(bottom=ft.BorderSide(1,'#E6EBF0')), padding=8,
                content=ft.Row([
                    ft.Container(width=8, height=28, bgcolor='#F2A900', border_radius=4),
                    ft.Text('PRÓXIMO CAMBIO', width=135, size=11, weight=ft.FontWeight.BOLD, color='#C98600'),
                    ft.Text('20 a 30 mm', width=150, size=11, color=TEXT_MAIN),
                    ft.Text('Programar seguimiento / próximo cambio.', expand=True, size=11, color=TEXT_MUTED),
                ], spacing=8)
            ),
            ft.Container(
                padding=8,
                content=ft.Row([
                    ft.Container(width=8, height=28, bgcolor='#D92D20', border_radius=4),
                    ft.Text('CAMBIO URGENTE', width=135, size=11, weight=ft.FontWeight.BOLD, color='#D92D20'),
                    ft.Text('0 a 20 mm', width=150, size=11, color=TEXT_MAIN),
                    ft.Text('Programar cambio con prioridad.', expand=True, size=11, color=TEXT_MUTED),
                ], spacing=8)
            ),
        ], spacing=0)

        def rtd_donut_chart():
            """Dona ampliada y centrada, con leyenda a la derecha usando exactamente los mismos colores."""
            import base64, math
            palette={
                'Buen Estado':'#2E9B45',
                'Próximo Cambio':'#F2A900',
                'Cambio Urgente':'#D92D20',
            }
            items=[
                ('Buen Estado', counts['BUEN ESTADO']),
                ('Próximo Cambio', counts['PRÓXIMO CAMBIO']),
                ('Cambio Urgente', counts['CAMBIO URGENTE']),
            ]
            total_chart=sum(n for _,n in items)
            cx=cy=100; radius=61; stroke=27
            circumference=2*math.pi*radius
            offset=0.0; circles=[]; legend=[]
            for label,n in items:
                color=palette[label]
                dash=circumference*(n/total_chart) if total_chart else 0
                gap=max(0.0,circumference-dash)
                if n>0:
                    circles.append(
                        f'<circle cx="{cx}" cy="{cy}" r="{radius}" fill="none" stroke="{color}" '
                        f'stroke-width="{stroke}" stroke-dasharray="{dash:.3f} {gap:.3f}" '
                        f'stroke-dashoffset="{-offset:.3f}" transform="rotate(-90 {cx} {cy})" />'
                    )
                offset += dash
                pc=(n/total_chart*100) if total_chart else 0
                legend.append(ft.Row([
                    ft.Container(width=10,height=10,bgcolor=color,border_radius=2),
                    ft.Text(f'{label}: {n} ({pc:.1f}%)',size=8.8,color=TEXT_MAIN),
                ],spacing=6))
            svg=(
                '<svg xmlns="http://www.w3.org/2000/svg" width="200" height="200" viewBox="0 0 200 200">'
                '<circle cx="100" cy="100" r="61" fill="none" stroke="#E2E8F0" stroke-width="27" />'
                + ''.join(circles) +
                f'<text x="100" y="98" text-anchor="middle" font-family="Arial" font-size="27" font-weight="700" fill="#172033">{total_chart}</text>'
                '<text x="100" y="118" text-anchor="middle" font-family="Arial" font-size="10" fill="#64748B">neumáticos</text>'
                '</svg>'
            )
            src='data:image/svg+xml;base64,'+base64.b64encode(svg.encode('utf-8')).decode('ascii')
            # Distribución fija para asegurar que la leyenda quede siempre a la derecha
            # y la dona permanezca realmente centrada dentro de su zona.
            donut_box=ft.Container(
                width=255,
                height=215,
                alignment=ft.Alignment.CENTER,
                content=ft.Image(src=src,width=220,height=220,fit=ft.BoxFit.CONTAIN),
            )
            legend_box=ft.Container(
                width=190,
                height=215,
                alignment=ft.Alignment.CENTER_LEFT,
                content=ft.Column(legend,spacing=9,horizontal_alignment=ft.CrossAxisAlignment.START),
            )
            return ft.Container(
                width=455,
                height=215,
                alignment=ft.Alignment.CENTER,
                content=ft.Row([
                    donut_box,
                    legend_box,
                ],spacing=10,vertical_alignment=ft.CrossAxisAlignment.CENTER)
            )

        # Resumen por tipo de vehículo, basado en los neumáticos actualmente en servicio.
        vehicle_summary={}
        for r,_,_ in evaluated:
            vehicle=(str(r['vehicle_type']).strip().upper() if r['vehicle_type'] else 'SIN CLASIFICAR')
            vehicle_summary[vehicle]=vehicle_summary.get(vehicle,0)+1
        vehicle_table=ft.DataTable(
            heading_row_height=34,
            data_row_min_height=30,
            data_row_max_height=30,
            column_spacing=26,
            columns=[
                ft.DataColumn(ft.Text('Tipo de Vehículo', size=11, weight=ft.FontWeight.BOLD)),
                ft.DataColumn(ft.Text('N° Neumáticos', size=11, weight=ft.FontWeight.BOLD), numeric=True),
                ft.DataColumn(ft.Text('% del Total', size=11, weight=ft.FontWeight.BOLD), numeric=True),
            ],
            rows=[
                ft.DataRow(cells=[
                    ft.DataCell(ft.Text(vehicle, size=10)),
                    ft.DataCell(ft.Text(str(n), size=10)),
                    ft.DataCell(ft.Text(f'{pct(n,total):.1f}%', size=10)),
                ])
                for vehicle,n in sorted(vehicle_summary.items(), key=lambda kv:(-kv[1],kv[0]))
            ] + ([ft.DataRow(cells=[
                ft.DataCell(ft.Text('TOTAL', size=10, weight=ft.FontWeight.BOLD)),
                ft.DataCell(ft.Text(str(total), size=10, weight=ft.FontWeight.BOLD)),
                ft.DataCell(ft.Text('100.0%' if total else '0.0%', size=10, weight=ft.FontWeight.BOLD)),
            ])] if total else [])
        )

        # Resumen específico de SCOOP por modelo registrado en Administración de equipos.
        scoop_summary={}
        for r,_,_ in evaluated:
            vehicle=(str(r['vehicle_type']).strip().upper() if r['vehicle_type'] else '')
            if 'SCOOP' in vehicle:
                model=(str(r['equipment_model']).strip().upper() if r['equipment_model'] else 'SIN MODELO')
                scoop_summary[model]=scoop_summary.get(model,0)+1
        scoop_total=sum(scoop_summary.values())
        scoop_table=ft.DataTable(
            heading_row_height=34, data_row_min_height=30, data_row_max_height=30, column_spacing=18,
            columns=[
                ft.DataColumn(ft.Text('Modelo Scoop',size=11,weight=ft.FontWeight.BOLD)),
                ft.DataColumn(ft.Text('N° Neumáticos',size=11,weight=ft.FontWeight.BOLD),numeric=True),
                ft.DataColumn(ft.Text('% del Total',size=11,weight=ft.FontWeight.BOLD),numeric=True),
            ],
            rows=[ft.DataRow(cells=[
                ft.DataCell(ft.Text(model,size=10)),
                ft.DataCell(ft.Text(str(n),size=10)),
                ft.DataCell(ft.Text(f'{pct(n,scoop_total):.1f}%',size=10)),
            ]) for model,n in sorted(scoop_summary.items())] + ([ft.DataRow(cells=[
                ft.DataCell(ft.Text('TOTAL',size=10,weight=ft.FontWeight.BOLD)),
                ft.DataCell(ft.Text(str(scoop_total),size=10,weight=ft.FontWeight.BOLD)),
                ft.DataCell(ft.Text('100.0%' if scoop_total else '0.0%',size=10,weight=ft.FontWeight.BOLD)),
            ])] if scoop_total else [])
        )

        # Condición general desglosada por tipo de equipo.
        # Cada neumático se coloca en la fila que corresponde a su condición RTD
        # y en la columna de su tipo de equipo. Los umbrales son los específicos
        # de cada categoría definidos en rtd_condition_by_category().
        def classify_equipment(vehicle_type):
            kind = str(vehicle_type or '').strip().upper()
            kind = (kind.replace('Á','A').replace('É','E').replace('Í','I')
                         .replace('Ó','O').replace('Ú','U').replace('Ü','U'))
            if 'VOLQUETE' in kind:
                return 'VOLQUETES'
            if 'CAMIONETA' in kind or 'CAMION' in kind:
                return 'CAMIONETAS'
            if 'SCOOP' in kind:
                return 'SCOOP'
            return 'OTROS'

        condition_by_category = {}
        category_order = ['VOLQUETES', 'CAMIONETAS', 'SCOOP', 'OTROS']
        present_categories = []
        for r, rtd, _ in evaluated:
            category = classify_equipment(r.get('vehicle_type'))
            if category not in present_categories:
                present_categories.append(category)
            category_counts = condition_by_category.setdefault(category, {
                'BUEN ESTADO': 0,
                'PRÓXIMO CAMBIO': 0,
                'CAMBIO URGENTE': 0,
                'SIN LECTURA': 0,
            })
            category_condition = rtd_condition_by_category(category, rtd)
            category_counts[category_condition] = category_counts.get(category_condition, 0) + 1

        present_categories = [c for c in category_order if c in present_categories]
        # Si existiera una categoría no contemplada, se conserva al final.
        present_categories += [c for c in condition_by_category if c not in present_categories]

        # CONSOLIDADO GENERAL:
        # Los indicadores superiores deben ser exactamente la suma de todas
        # las tablas por tipo de equipo. No se utiliza el criterio general
        # rtd_condition(), porque cada categoría tiene sus propios umbrales.
        counts = {
            'BUEN ESTADO': sum(condition_by_category.get(cat, {}).get('BUEN ESTADO', 0)
                               for cat in present_categories),
            'PRÓXIMO CAMBIO': sum(condition_by_category.get(cat, {}).get('PRÓXIMO CAMBIO', 0)
                                  for cat in present_categories),
            'CAMBIO URGENTE': sum(condition_by_category.get(cat, {}).get('CAMBIO URGENTE', 0)
                                  for cat in present_categories),
            'SIN LECTURA': sum(condition_by_category.get(cat, {}).get('SIN LECTURA', 0)
                               for cat in present_categories),
        }
        evaluated_total = (
            counts['BUEN ESTADO']
            + counts['PRÓXIMO CAMBIO']
            + counts['CAMBIO URGENTE']
        )
        good_pct = (counts['BUEN ESTADO'] / evaluated_total * 100) if evaluated_total else 0

        condition_rows=[]
        for label,text_color,row_color in [
            ('BUEN ESTADO','#176B2C','#DDF3E3'),
            ('PRÓXIMO CAMBIO','#8A5A00','#FFF0C2'),
            ('CAMBIO URGENTE','#A61B12','#FAD9D6'),
        ]:
            category_values = [condition_by_category.get(cat, {}).get(label, 0) for cat in present_categories]
            row_total = sum(category_values)
            cells = [
                ft.DataCell(ft.Text(label.title(), size=10, weight=ft.FontWeight.BOLD, color=text_color)),
            ]
            cells += [
                ft.DataCell(ft.Text(str(value), size=10, weight=ft.FontWeight.BOLD, color=text_color))
                for value in category_values
            ]
            cells += [
                ft.DataCell(ft.Text(str(row_total), size=10, weight=ft.FontWeight.BOLD, color=text_color)),
                ft.DataCell(ft.Text(f'{pct(row_total,evaluated_total):.1f}%', size=10, weight=ft.FontWeight.BOLD, color=text_color)),
            ]
            condition_rows.append(ft.DataRow(color=row_color, cells=cells))

        # Se muestran lecturas faltantes solo si existen.
        no_reading_values = [condition_by_category.get(cat, {}).get('SIN LECTURA', 0) for cat in present_categories]
        no_reading_total = sum(no_reading_values)
        if no_reading_total:
            cells = [ft.DataCell(ft.Text('Sin lectura', size=10, weight=ft.FontWeight.BOLD, color=TEXT_MUTED))]
            cells += [ft.DataCell(ft.Text(str(value), size=10, color=TEXT_MUTED)) for value in no_reading_values]
            cells += [
                ft.DataCell(ft.Text(str(no_reading_total), size=10, color=TEXT_MUTED)),
                ft.DataCell(ft.Text(f'{pct(no_reading_total,total):.1f}%', size=10, color=TEXT_MUTED)),
            ]
            condition_rows.append(ft.DataRow(cells=cells))

        # Totales por tipo de equipo y total general.
        category_totals = [sum(condition_by_category.get(cat, {}).values()) for cat in present_categories]
        cells = [ft.DataCell(ft.Text('TOTAL', size=10, weight=ft.FontWeight.BOLD, color='#FFFFFF'))]
        cells += [ft.DataCell(ft.Text(str(value), size=10, weight=ft.FontWeight.BOLD, color='#FFFFFF')) for value in category_totals]
        cells += [
            ft.DataCell(ft.Text(str(total), size=10, weight=ft.FontWeight.BOLD, color='#FFFFFF')),
            ft.DataCell(ft.Text('100.0%' if total else '0.0%', size=10, weight=ft.FontWeight.BOLD, color='#FFFFFF')),
        ]
        condition_rows.append(ft.DataRow(color='#000000', cells=cells))

        condition_columns = [
            ft.DataColumn(ft.Text('Condición', size=11, weight=ft.FontWeight.BOLD, color='#FFFFFF')),
        ]
        condition_columns += [
            ft.DataColumn(ft.Text(category, size=10, weight=ft.FontWeight.BOLD, color='#FFFFFF'), numeric=True)
            for category in present_categories
        ]
        condition_columns += [
            ft.DataColumn(ft.Text('Total', size=11, weight=ft.FontWeight.BOLD, color='#FFFFFF'), numeric=True),
            ft.DataColumn(ft.Text('% del Total', size=10, weight=ft.FontWeight.BOLD, color='#FFFFFF'), numeric=True),
        ]

        condition_table=ft.DataTable(
            heading_row_height=34,
            heading_row_color='#000000',
            data_row_min_height=30,
            data_row_max_height=30,
            column_spacing=22,
            columns=condition_columns,
            rows=condition_rows
        )

        # Tablas independientes por tipo de equipo. Cada tabla conserva el
        # mismo formato de la tabla general, pero solo cuantifica los neumáticos
        # de su propia categoría y se inserta en la fila correspondiente.
        def build_condition_table_category(category):
            data = condition_by_category.get(category, {})
            category_total = sum(data.get(k, 0) for k in (
                'BUEN ESTADO', 'PRÓXIMO CAMBIO', 'CAMBIO URGENTE', 'SIN LECTURA'
            ))
            if category_total <= 0:
                return None

            rows_cat = []
            for label, text_color, row_color in [
                ('BUEN ESTADO', '#176B2C', '#DDF3E3'),
                ('PRÓXIMO CAMBIO', '#8A5A00', '#FFF0C2'),
                ('CAMBIO URGENTE', '#A61B12', '#FAD9D6'),
            ]:
                value = data.get(label, 0)
                rows_cat.append(ft.DataRow(
                    color=row_color,
                    cells=[
                        ft.DataCell(ft.Text(label.title(), size=9.5, weight=ft.FontWeight.BOLD, color=text_color)),
                        ft.DataCell(ft.Text(str(value), size=9.5, weight=ft.FontWeight.BOLD, color=text_color)),
                        ft.DataCell(ft.Text(f'{pct(value, category_total):.1f}%', size=9.5, weight=ft.FontWeight.BOLD, color=text_color)),
                    ]
                ))

            no_reading = data.get('SIN LECTURA', 0)
            if no_reading:
                rows_cat.append(ft.DataRow(cells=[
                    ft.DataCell(ft.Text('Sin lectura', size=9.5, color=TEXT_MUTED)),
                    ft.DataCell(ft.Text(str(no_reading), size=9.5, color=TEXT_MUTED)),
                    ft.DataCell(ft.Text(f'{pct(no_reading, category_total):.1f}%', size=9.5, color=TEXT_MUTED)),
                ]))

            rows_cat.append(ft.DataRow(
                color='#000000',
                cells=[
                    ft.DataCell(ft.Text('TOTAL', size=9.5, weight=ft.FontWeight.BOLD, color='#FFFFFF')),
                    ft.DataCell(ft.Text(str(category_total), size=9.5, weight=ft.FontWeight.BOLD, color='#FFFFFF')),
                    ft.DataCell(ft.Text('100.0%', size=9.5, weight=ft.FontWeight.BOLD, color='#FFFFFF')),
                ]
            ))

            return ft.DataTable(
                heading_row_height=30,
                heading_row_color='#000000',
                data_row_min_height=28,
                data_row_max_height=28,
                column_spacing=14,
                columns=[
                    ft.DataColumn(ft.Text('Condición', size=9.5, weight=ft.FontWeight.BOLD, color='#FFFFFF')),
                    ft.DataColumn(ft.Text('Cantidad', size=9.5, weight=ft.FontWeight.BOLD, color='#FFFFFF'), numeric=True),
                    ft.DataColumn(ft.Text('% Total', size=9.5, weight=ft.FontWeight.BOLD, color='#FFFFFF'), numeric=True),
                ],
                rows=rows_cat,
            )

        # Tercera fila: matriz dinámica invertida.
        # Filas = posiciones de neumáticos. Columnas = equipos agrupados por tipo.
        # Cada celda usa el menor RTD entre EXT e INT y se representa con un círculo
        # cuyo color corresponde a la condición RTD aprobada.
        equipment_positions = {}
        equipment_categories = {}
        all_positions = set()

        for r, rtd, condition in evaluated:
            equipment = (str(r['equipment_code']).strip().upper() if r['equipment_code'] else 'SIN EQUIPO')
            position_raw = str(r['position']).strip().upper() if r['position'] is not None else ''
            if not position_raw:
                continue
            if position_raw.isdigit():
                position = f'P{position_raw}'
            elif position_raw.startswith('P') and position_raw[1:].isdigit():
                position = f"P{int(position_raw[1:])}"
            else:
                position = position_raw

            all_positions.add(position)
            equipment_positions.setdefault(equipment, {})
            equipment_categories[equipment] = classify_equipment(r['vehicle_type'])
            current = equipment_positions[equipment].get(position)
            if rtd is not None and (current is None or rtd < current):
                equipment_positions[equipment][position] = rtd
            elif position not in equipment_positions[equipment]:
                equipment_positions[equipment][position] = None

        def position_sort_key(pos):
            text = str(pos)
            if text.startswith('P') and text[1:].isdigit():
                return (0, int(text[1:]))
            if text.isdigit():
                return (0, int(text))
            return (1, text)

        ordered_positions = sorted(all_positions, key=position_sort_key)
        category_order = ['VOLQUETES', 'CAMIONETAS', 'SCOOP', 'OTROS']
        grouped_equipment = {category: [] for category in category_order}
        for equipment in sorted(equipment_positions.keys()):
            grouped_equipment.setdefault(equipment_categories.get(equipment, 'OTROS'), []).append(equipment)
        grouped_equipment = {k: v for k, v in grouped_equipment.items() if v}

        status_colors = {
            'BUEN ESTADO': '#2E9B45',
            'PRÓXIMO CAMBIO': '#F2A900',
            'CAMBIO URGENTE': '#D92D20',
            'SIN LECTURA': '#AEB9C4',
        }

        def rtd_circle(rtd, category):
            condition = rtd_condition_by_category(category, rtd)
            if rtd is None or condition == 'SIN LECTURA':
                label = '—'
                color = status_colors['SIN LECTURA']
            else:
                label = f'{rtd:.0f}' if abs(rtd - round(rtd)) < 0.05 else f'{rtd:.1f}'
                color = status_colors[condition]
            return ft.Container(
                width=30,
                height=30,
                bgcolor=color,
                border_radius=15,
                alignment=ft.Alignment.CENTER,
                content=ft.Text(
                    label, size=10.5, weight=ft.FontWeight.BOLD, color='#FFFFFF',
                    text_align=ft.TextAlign.CENTER
                ),
            )

        # Tercera fila: matriz de remanente mínimo separada dinámicamente por
        # tipo de equipo, igual que la fila 4. Cada tipo genera su propia tarjeta.
        POSITION_COL_W = 70
        EQUIPMENT_COL_W = 40

        def build_remanente_matrix_category(category):
            equipments = grouped_equipment.get(category, [])
            if not equipments:
                return None

            # Solo posiciones que realmente aparecen en los equipos de esta categoría.
            category_positions = set()
            for equipment in equipments:
                category_positions.update(equipment_positions.get(equipment, {}).keys())
            ordered_category_positions = sorted(category_positions, key=position_sort_key)

            header = ft.Row([
                ft.Container(
                    width=POSITION_COL_W, height=82, bgcolor='#172B3A',
                    alignment=ft.Alignment.CENTER,
                    content=ft.Text('POSICIÓN', size=10.5, weight=ft.FontWeight.BOLD, color='#FFFFFF')
                )
            ], spacing=2)

            for equipment in equipments:
                header.controls.append(
                    ft.Container(
                        width=EQUIPMENT_COL_W, height=82,
                        bgcolor='#F0F4F8',
                        border=ft.Border(bottom=ft.BorderSide(1, '#D9E2EA')),
                        alignment=ft.Alignment.CENTER,
                        content=ft.Text(
                            equipment, size=8.5, weight=ft.FontWeight.BOLD,
                            color=TEXT_MAIN, text_align=ft.TextAlign.CENTER,
                            rotate=ft.Rotate(angle=-1.5708),
                        )
                    )
                )

            rows = []
            for pos in ordered_category_positions:
                row_controls = [
                    ft.Container(
                        width=POSITION_COL_W, height=36, bgcolor='#F5F7FA',
                        alignment=ft.Alignment.CENTER,
                        border=ft.Border(right=ft.BorderSide(1, '#D9E2EA')),
                        content=ft.Text(pos, size=10, weight=ft.FontWeight.BOLD, color=TEXT_MAIN)
                    )
                ]
                for equipment in equipments:
                    row_controls.append(
                        ft.Container(
                            width=EQUIPMENT_COL_W, height=36,
                            alignment=ft.Alignment.CENTER,
                            content=rtd_circle(equipment_positions[equipment].get(pos), category)
                        )
                    )
                rows.append(ft.Row(row_controls, spacing=2))

            body = ft.Column(rows, spacing=1) if rows else ft.Container(
                height=58, alignment=ft.Alignment.CENTER,
                content=ft.Text('Sin posiciones de neumáticos para mostrar.', size=10, color=TEXT_MUTED)
            )

            matrix = ft.Row([
                ft.Column([header, body], spacing=1, tight=True)
            ], scroll=ft.ScrollMode.AUTO)

            # Ancho mínimo basado en la cantidad real de equipos; permite scroll
            # si un tipo de equipo tiene una flota grande.
            content_width = POSITION_COL_W + (EQUIPMENT_COL_W * len(equipments)) + (2 * len(equipments)) + 20
            return ft.Container(
                width=max(330, min(760, content_width)),
                bgcolor=CARD_BG,
                border=ft.Border.all(1, '#DDE5ED'),
                border_radius=10,
                padding=10,
                content=ft.Column([
                    ft.Text(
                        f'REMANENTE MÍNIMO · {category}', size=12,
                        weight=ft.FontWeight.BOLD, color=TEXT_MAIN,
                        text_align=ft.TextAlign.CENTER
                    ),
                    ft.Text(
                        'Menor RTD entre EXT e INT.', size=9, color=TEXT_MUTED,
                        text_align=ft.TextAlign.CENTER
                    ),
                    matrix,
                ], spacing=4, horizontal_alignment=ft.CrossAxisAlignment.CENTER)
            )

        # ------------------------------------------------------------------
        # FILAS 3+ DINÁMICAS: cada tipo de equipo agrupa su matriz de
        # remanente mínimo y su gráfico de cuantificación en la MISMA fila.
        # Orden: VOLQUETES, CAMIONETAS, SCOOP, OTROS.
        # ------------------------------------------------------------------
        # Leyenda de colores oculta por solicitud del usuario.
        # Se usa un contenedor vacío en lugar de None para evitar que Flet
        # reciba un control nulo dentro de la columna.
        matrix_legend = ft.Container(height=0)

        # Cuantificación dinámica del remanente mínimo por milímetro,
        # segregada por tipo de equipo. Solo valores con cantidad > 0.
        equipment_categories = {}
        for r, _rtd, _condition in evaluated:
            category = classify_equipment(r['vehicle_type'])
            equipment_categories.setdefault(category, [])
            equipment_categories[category].append((r, _rtd, _condition))

        category_order_chart = ['VOLQUETES', 'CAMIONETAS', 'SCOOP', 'OTROS']
        chart_categories = [
            category for category in category_order_chart
            if equipment_categories.get(category)
        ]

        def build_remanente_chart(category):
            category_rows = equipment_categories.get(category, [])
            remanente_por_mm = {}
            for _r, rtd, _condition in category_rows:
                if rtd is None:
                    continue
                try:
                    mm = int(round(float(rtd)))
                except (TypeError, ValueError):
                    continue
                if mm < 1:
                    continue
                remanente_por_mm[mm] = remanente_por_mm.get(mm, 0) + 1

            items = sorted(
                ((mm, cantidad) for mm, cantidad in remanente_por_mm.items() if cantidad > 0),
                key=lambda item: item[0]
            )
            chart_max_count = max((cantidad for _, cantidad in items), default=1)
            # Escala Y en múltiplos de 5, comenzando siempre en 0.
            y_max = max(5, ((chart_max_count + 4) // 5) * 5)
            BAR_AREA_H = 185
            BAR_W = 18
            BAR_GAP = 8

            chart_bars = []
            for mm, cantidad in items:
                condition = rtd_condition_by_category(category, mm)
                bar_color = status_colors.get(condition, status_colors['SIN LECTURA'])
                bar_h = max(8, int((cantidad / chart_max_count) * (BAR_AREA_H - 10)))
                chart_bars.append(
                    ft.Container(
                        width=BAR_W + BAR_GAP,
                        height=BAR_AREA_H + 50,
                        alignment=ft.Alignment.BOTTOM_CENTER,
                        content=ft.Column([
                            ft.Container(
                                width=BAR_W,
                                height=BAR_AREA_H,
                                alignment=ft.Alignment.BOTTOM_CENTER,
                                content=ft.Container(
                                    width=BAR_W,
                                    height=bar_h,
                                    bgcolor=bar_color,
                                    border_radius=ft.BorderRadius(top_left=3, top_right=3, bottom_left=0, bottom_right=0),
                                    alignment=ft.Alignment.TOP_CENTER,
                                    content=ft.Container(
                                        margin=ft.Margin(top=-20, left=0, right=0, bottom=0),
                                        content=ft.Text(
                                            str(cantidad), size=9.5, weight=ft.FontWeight.BOLD,
                                            color=TEXT_MAIN, text_align=ft.TextAlign.CENTER
                                        )
                                    )
                                )
                            ),
                            ft.Text(str(mm), size=8.5, weight=ft.FontWeight.BOLD,
                                    color=TEXT_MAIN, text_align=ft.TextAlign.CENTER, no_wrap=True),
                        ], spacing=4, horizontal_alignment=ft.CrossAxisAlignment.CENTER)
                    )
                )

            if not chart_bars:
                return ft.Container(
                    height=BAR_AREA_H + 60,
                    alignment=ft.Alignment.CENTER,
                    content=ft.Text('No hay lecturas RTD válidas para cuantificar.', size=10, color=TEXT_MUTED)
                )

            # Escala Y uniforme: 0, 5, 10, 15... hasta y_max.
            y_ticks = list(range(y_max, -1, -5))
            # El área de barras incluye 50 px adicionales para las etiquetas X.
            # El eje Y debe terminar exactamente en la base de las barras,
            # no en la parte inferior de las etiquetas X.
            y_scale_ticks = ft.Column(
                [ft.Text(str(tick), size=9, color=TEXT_MUTED) for tick in y_ticks],
                height=BAR_AREA_H,
                alignment=ft.MainAxisAlignment.SPACE_BETWEEN
            )
            y_scale = ft.Container(
                width=22,
                height=BAR_AREA_H + 50,
                alignment=ft.Alignment.TOP_LEFT,
                content=y_scale_ticks,
            )

            plot_row = ft.Row([
                y_scale,
                ft.Container(width=1, height=BAR_AREA_H, bgcolor='#CBD5E1'),
                ft.Container(
                    content=ft.Row(
                        chart_bars,
                        spacing=0,
                        vertical_alignment=ft.CrossAxisAlignment.END,
                        scroll=ft.ScrollMode.AUTO,
                    ),
                    expand=True,
                ),
            ], spacing=8, vertical_alignment=ft.CrossAxisAlignment.END)

            return plot_row

        # Construye una fila completa por categoría: izquierda = matriz,
        # derecha = cuantificación. Si una categoría no existe en el taller,
        # simplemente no se genera esa fila.
        paired_category_rows = []
        PAIR_LEFT_W = 540
        PAIR_RIGHT_W = 540
        PAIR_CONDITION_W = 360

        for category in category_order:
            matrix = build_remanente_matrix_category(category)
            if matrix is None:
                continue

            chart_plot = build_remanente_chart(category)

            # El título y subtítulo ya forman parte de la matriz construida
            # por build_remanente_matrix_category(). No repetirlos en el panel.
            matrix_panel = ft.Container(
                width=PAIR_LEFT_W,
                bgcolor=CARD_BG,
                border=ft.Border.all(1, '#DDE5ED'),
                border_radius=10,
                padding=12,
                content=matrix,
            )

            chart_panel = ft.Container(
                width=PAIR_RIGHT_W,
                bgcolor=CARD_BG,
                border=ft.Border.all(1, '#DDE5ED'),
                border_radius=10,
                padding=12,
                content=ft.Column([
                    ft.Text(
                        f'CUANTIFICACIÓN DEL REMANENTE · {category}', size=12,
                        weight=ft.FontWeight.BOLD, color=TEXT_MAIN,
                        text_align=ft.TextAlign.CENTER
                    ),
                    ft.Text(
                        'Cantidad de neumáticos por cada milímetro de remanente mínimo.',
                        size=9, color=TEXT_MUTED, text_align=ft.TextAlign.CENTER
                    ),
                    chart_plot,
                ], spacing=4, horizontal_alignment=ft.CrossAxisAlignment.CENTER)
            )

            category_condition_table = build_condition_table_category(category)
            condition_panel = ft.Container(
                width=PAIR_RIGHT_W,
                bgcolor=CARD_BG,
                border=ft.Border.all(1, '#DDE5ED'),
                border_radius=10,
                padding=12,
                content=ft.Column([
                    ft.Text(
                        f'CONDICIÓN RTD · {category}', size=12,
                        weight=ft.FontWeight.BOLD, color=TEXT_MAIN,
                        text_align=ft.TextAlign.CENTER
                    ),
                    ft.Text(
                        'Distribución de neumáticos según su condición.',
                        size=9, color=TEXT_MUTED, text_align=ft.TextAlign.CENTER
                    ),
                    ft.Row([category_condition_table], scroll=ft.ScrollMode.AUTO, alignment=ft.MainAxisAlignment.CENTER),
                ], spacing=4, horizontal_alignment=ft.CrossAxisAlignment.CENTER)
            )

            # Cada tipo de equipo queda en una fila independiente: la matriz
            # va a la izquierda y, a la derecha, la cuantificación arriba y
            # la tabla de condición RTD inmediatamente debajo.
            right_column = ft.Column([
                chart_panel,
                ft.Container(height=12),
                condition_panel,
            ], spacing=0, horizontal_alignment=ft.CrossAxisAlignment.CENTER)

            paired_category_rows.append(
                ft.Row(
                    [matrix_panel, right_column],
                    spacing=12,
                    vertical_alignment=ft.CrossAxisAlignment.START,
                    scroll=ft.ScrollMode.AUTO,
                )
            )

        if paired_category_rows:
            paired_rows_with_separators = []
            for idx, category_row in enumerate(paired_category_rows):
                paired_rows_with_separators.append(category_row)
                if idx < len(paired_category_rows) - 1:
                    paired_rows_with_separators.append(
                        ft.Container(
                            width=1100,
                            height=1,
                            bgcolor='#CBD5E1',
                            margin=ft.Margin(top=10, bottom=10, left=0, right=0),
                        )
                    )
            paired_rows_content = ft.Column(
                paired_rows_with_separators,
                spacing=0,
                horizontal_alignment=ft.CrossAxisAlignment.CENTER,
            )
        else:
            paired_rows_content = ft.Container(
                height=180,
                alignment=ft.Alignment.CENTER,
                content=ft.Text('No hay posiciones de neumáticos en servicio para mostrar.', size=10, color=TEXT_MUTED)
            )

        paired_remanente_section = card(ft.Column([
            ft.Text(
                'REMANENTE MÍNIMO + CUANTIFICACIÓN POR TIPO DE EQUIPO',
                size=16, weight=ft.FontWeight.BOLD,
                color=TEXT_MAIN, text_align=ft.TextAlign.CENTER
            ),
            ft.Text(
                'Cada tipo de equipo reúne su matriz de posiciones y su cuantificación de remanente en la misma fila.',
                size=10, color=TEXT_MUTED, text_align=ft.TextAlign.CENTER
            ),
            ft.Container(height=2),
            paired_rows_content,
            matrix_legend,
        ], spacing=8, horizontal_alignment=ft.CrossAxisAlignment.CENTER), padding=14)

        content.content=ft.Column([
            page_title('3. Programa de mantenimiento · 3.1 Evaluación de Remanente (RTD)',
                       'Evaluación automática de neumáticos en servicio según profundidad remanente'),
            maintenance_nav_row('31'),
            ft.Row([
                top_metric('NEUMÁTICOS EN SERVICIO', total, 'Total actualmente instalado', '#C81D2A'),
                top_metric('EQUIPOS EN SERVICIO', equipment_count, 'Equipos con neumáticos instalados'),
                top_metric('BUEN ESTADO', f'{good_pct:.1f}%',
                           f'{counts["BUEN ESTADO"]} neumáticos', '#2E9B45'),
                top_metric('PRÓXIMO CAMBIO', f'{pct(counts["PRÓXIMO CAMBIO"], evaluated_total):.1f}%',
                           f'{counts["PRÓXIMO CAMBIO"]} neumáticos', '#F2A900'),
                top_metric('CAMBIO URGENTE', f'{pct(counts["CAMBIO URGENTE"], evaluated_total):.1f}%',
                           f'{counts["CAMBIO URGENTE"]} neumáticos', '#C81D2A'),
            ], spacing=12, vertical_alignment=ft.CrossAxisAlignment.START),
            ft.Container(height=2),
            paired_remanente_section,
        ], scroll=ft.ScrollMode.AUTO, spacing=16)
        page.update()


    def maintenance_shoulders_view():
        """3.2 Diferencia de RTD entre hombros (EXT vs INT) del mismo neumático."""
        op_id = active_operation_id()
        rows = query("""
            SELECT
                t.id, t.code, t.serial, t.brand, t.size, t.design,
                t.tread_inner, t.tread_outer,
                e.id AS equipment_id, e.code AS equipment_code,
                t.position
            FROM tires t
            LEFT JOIN equipment e ON e.id=t.equipment_id
            WHERE t.status='SERVICIO' AND t.operation_id=?
            ORDER BY COALESCE(e.code,''), t.position, t.code
        """, (op_id,))

        def latest_rtd(r):
            # Se prioriza la última inspección por ID para evitar problemas con
            # fechas históricas guardadas en formatos mixtos.
            insp = query("""
                SELECT tread_inner, tread_outer
                FROM occurrences
                WHERE tire_id=? AND operation_id=? AND event_code IN ('INSP','INSC')
                ORDER BY id DESC LIMIT 1
            """, (r['id'], op_id))
            inner = insp[0]['tread_inner'] if insp and insp[0]['tread_inner'] is not None else r['tread_inner']
            outer = insp[0]['tread_outer'] if insp and insp[0]['tread_outer'] is not None else r['tread_outer']
            try:
                inner = float(inner) if inner is not None and str(inner).strip() != '' else None
            except Exception:
                inner = None
            try:
                outer = float(outer) if outer is not None and str(outer).strip() != '' else None
            except Exception:
                outer = None
            return outer, inner

        def shoulder_condition(outer, inner):
            if outer is None or inner is None:
                return 'SIN LECTURA', None
            diff = inner - outer
            # Regla aprobada 3.2:
            # Emergencia: RTD INT - RTD EXT >= 10 mm
            # Preventivo: 7 <= diferencia < 10 mm
            # Normal: cualquier otra condición (incluye EXT >= INT)
            if diff >= 10:
                return 'EMERGENCIA', diff
            if diff >= 7:
                return 'PREVENTIVO', diff
            return 'NORMAL', diff

        evaluated=[]
        counts={'NORMAL':0, 'PREVENTIVO':0, 'EMERGENCIA':0, 'SIN LECTURA':0}
        for r in rows:
            outer, inner = latest_rtd(r)
            condition, diff = shoulder_condition(outer, inner)
            counts[condition] = counts.get(condition, 0) + 1
            evaluated.append((r, outer, inner, diff, condition))

        total=len(rows)
        evaluated_total=total-counts['SIN LECTURA']

        def pct(n, base):
            return (n/base*100) if base else 0

        # Tarjeta KPI local para 3.2. Se define dentro de esta vista para no
        # depender del helper local de 3.1.
        def top_metric(title, value, subtitle, value_color=TEXT_MAIN):
            return ft.Container(
                width=250,
                height=108,
                bgcolor=CARD_BG,
                border=ft.Border.all(1, '#DDE5ED'),
                border_radius=10,
                padding=14,
                content=ft.Column([
                    ft.Text(title, size=12, weight=ft.FontWeight.BOLD, color=TEXT_MAIN,
                            text_align=ft.TextAlign.CENTER),
                    ft.Text(str(value), size=29, weight=ft.FontWeight.BOLD, color=value_color,
                            text_align=ft.TextAlign.CENTER),
                    ft.Text(subtitle, size=10, color=TEXT_MUTED, text_align=ft.TextAlign.CENTER),
                ], spacing=3, horizontal_alignment=ft.CrossAxisAlignment.CENTER)
            )

        # Semáforo sin encabezado independiente de "Criterio": el criterio se
        # muestra al costado del nombre de cada condición, como pidió el usuario.
        condition_rows=[]
        row_specs=[
            ('NORMAL', '#176B2C', '#DDF3E3', '< 7 mm  o  EXT ≥ INT'),
            ('PREVENTIVO', '#A66000', '#FFF0C2', '≥ 7 y < 10 mm'),
            ('EMERGENCIA', '#A61B12', '#FAD9D6', '≥ 10 mm'),
        ]
        for label, text_color, row_color, criterion in row_specs:
            n=counts[label]
            condition_rows.append(ft.DataRow(color=row_color, cells=[
                ft.DataCell(ft.Row([
                    ft.Container(width=10, height=10, bgcolor={
                        'NORMAL':'#2E9B45','PREVENTIVO':'#F2A900','EMERGENCIA':'#D92D20'
                    }[label], border_radius=5),
                    ft.Text(label.title(), width=100, size=10.5, weight=ft.FontWeight.BOLD, color=text_color),
                    ft.Text(criterion, size=10, color=TEXT_MAIN),
                ], spacing=8)),
                ft.DataCell(ft.Text(str(n), size=11, weight=ft.FontWeight.BOLD, color=text_color)),
                ft.DataCell(ft.Text(f'{pct(n,evaluated_total):.1f}%', size=11, weight=ft.FontWeight.BOLD, color=text_color)),
            ]))
        if counts['SIN LECTURA']:
            condition_rows.append(ft.DataRow(cells=[
                ft.DataCell(ft.Text('Sin lectura', size=10, weight=ft.FontWeight.BOLD, color=TEXT_MUTED)),
                ft.DataCell(ft.Text(str(counts['SIN LECTURA']), size=10)),
                ft.DataCell(ft.Text(f'{pct(counts["SIN LECTURA"],total):.1f}%', size=10)),
            ]))
        condition_rows.append(ft.DataRow(color='#000000', cells=[
            ft.DataCell(ft.Text('TOTAL', size=10.5, weight=ft.FontWeight.BOLD, color='#FFFFFF')),
            ft.DataCell(ft.Text(str(total), size=10.5, weight=ft.FontWeight.BOLD, color='#FFFFFF')),
            ft.DataCell(ft.Text('100.0%' if total else '0.0%', size=10.5, weight=ft.FontWeight.BOLD, color='#FFFFFF')),
        ]))

        condition_table=ft.DataTable(
            heading_row_height=34,
            heading_row_color='#000000',
            data_row_min_height=38,
            data_row_max_height=38,
            column_spacing=28,
            columns=[
                ft.DataColumn(ft.Text('Condición', size=11, weight=ft.FontWeight.BOLD, color='#FFFFFF')),
                ft.DataColumn(ft.Text('Cantidad', size=11, weight=ft.FontWeight.BOLD, color='#FFFFFF'), numeric=True),
                ft.DataColumn(ft.Text('% del Total', size=11, weight=ft.FontWeight.BOLD, color='#FFFFFF'), numeric=True),
            ],
            rows=condition_rows,
        )

        def shoulders_donut_chart():
            import base64, math
            items=[
                ('Normal', counts['NORMAL'], '#2E9B45'),
                ('Preventivo', counts['PREVENTIVO'], '#F2A900'),
                ('Emergencia', counts['EMERGENCIA'], '#D92D20'),
            ]
            total_chart=sum(n for _,n,_ in items)
            cx=cy=82; radius=49; stroke=24
            circumference=2*math.pi*radius
            offset=0.0; circles=[]; legend=[]
            for label,n,color in items:
                dash=circumference*(n/total_chart) if total_chart else 0
                gap=max(0.0,circumference-dash)
                if n>0:
                    circles.append(
                        f'<circle cx="{cx}" cy="{cy}" r="{radius}" fill="none" stroke="{color}" '
                        f'stroke-width="{stroke}" stroke-dasharray="{dash:.3f} {gap:.3f}" '
                        f'stroke-dashoffset="{-offset:.3f}" transform="rotate(-90 {cx} {cy})" />'
                    )
                offset += dash
                pc=(n/total_chart*100) if total_chart else 0
                legend.append(ft.Row([
                    ft.Container(width=10,height=10,bgcolor=color,border_radius=5),
                    ft.Text(f'{label}: {n} ({pc:.1f}%)',size=10,color=TEXT_MAIN),
                ],spacing=6))
            svg=(
                '<svg xmlns="http://www.w3.org/2000/svg" width="164" height="164" viewBox="0 0 164 164">'
                '<circle cx="82" cy="82" r="49" fill="none" stroke="#E2E8F0" stroke-width="24" />'
                + ''.join(circles) +
                f'<text x="82" y="80" text-anchor="middle" font-family="Arial" font-size="25" font-weight="700" fill="#172033">{total_chart}</text>'
                '<text x="82" y="99" text-anchor="middle" font-family="Arial" font-size="9" fill="#64748B">neumáticos</text>'
                '</svg>'
            )
            src='data:image/svg+xml;base64,'+base64.b64encode(svg.encode('utf-8')).decode('ascii')
            return ft.Row([
                ft.Image(src=src,width=164,height=164,fit=ft.BoxFit.CONTAIN),
                ft.Column(legend,spacing=9),
            ],spacing=12,vertical_alignment=ft.CrossAxisAlignment.CENTER)

        content.content=ft.Column([
            page_title('3. Programa de mantenimiento · 3.2 Diferencia de RTD entre hombros',
                       'Evaluación del desgaste entre hombro exterior (EXT) e interior (INT) del mismo neumático'),
            maintenance_nav_row('32'),
            ft.Row([
                top_metric('NEUMÁTICOS EN SERVICIO', total, 'Total actualmente instalado', '#C81D2A'),
                top_metric('EQUIPOS EN SERVICIO', len({r['equipment_id'] for r in rows if r['equipment_id'] is not None}), 'Equipos con neumáticos instalados'),
                top_metric('NEUMÁTICOS EN BUEN ESTADO', f'{pct(counts["NORMAL"] + counts["PREVENTIVO"], evaluated_total):.1f}%',
                           f'{counts["NORMAL"] + counts["PREVENTIVO"]} neumáticos · Normal + Preventivo', '#2E9B45'),
                top_metric('NEUMÁTICOS QUE REQUIEREN INVERSIÓN', f'{pct(counts["EMERGENCIA"], evaluated_total):.1f}%',
                           f'{counts["EMERGENCIA"]} neumáticos · Solo Emergencia', '#C81D2A'),
            ], wrap=True, spacing=12, run_spacing=12),
            ft.Row([
                ft.Container(expand=1, content=card(ft.Column([
                    ft.Text('SEMÁFORO DE CONDICIÓN', size=14, weight=ft.FontWeight.BOLD, color=TEXT_MAIN),
                    ft.Row([condition_table], scroll=ft.ScrollMode.AUTO),
                    ft.Text('Diferencia evaluada = RTD INT − RTD EXT.', size=9.5, italic=True, color=TEXT_MUTED),
                ], spacing=8))),
                ft.Container(expand=1, content=card(ft.Column([
                    ft.Text('DISTRIBUCIÓN DE NEUMÁTICOS', size=14, weight=ft.FontWeight.BOLD, color=TEXT_MAIN),
                    shoulders_donut_chart(),
                    ft.Text('Porcentajes calculados sobre neumáticos con lecturas EXT/INT disponibles.',
                            size=9.5, italic=True, color=TEXT_MUTED),
                ], spacing=8))),
            ], spacing=12, vertical_alignment=ft.CrossAxisAlignment.START),
        ], scroll=ft.ScrollMode.AUTO, spacing=16)
        page.update()


    def maintenance_axles_view():
        """3.3 Diferencia de RTD entre neumáticos del mismo eje (P1-P2 y P3-P4)."""
        op_id = active_operation_id()
        rows = query("""
            SELECT
                t.id, t.code, t.serial, t.brand, t.size, t.design,
                t.tread_inner, t.tread_outer,
                e.id AS equipment_id, e.code AS equipment_code,
                e.vehicle_type, e.model AS equipment_model,
                t.position
            FROM tires t
            LEFT JOIN equipment e ON e.id=t.equipment_id
            WHERE t.status='SERVICIO' AND t.operation_id=?
            ORDER BY COALESCE(e.code,''), t.position, t.code
        """, (op_id,))

        def latest_rtd(r):
            # Igual que 3.2: se toma la última inspección por ID. De este modo
            # cualquier nueva INSP/INSC actualiza inmediatamente la evaluación.
            insp = query("""
                SELECT tread_inner, tread_outer
                FROM occurrences
                WHERE tire_id=? AND operation_id=? AND event_code IN ('INSP','INSC')
                ORDER BY id DESC LIMIT 1
            """, (r['id'], op_id))
            inner = insp[0]['tread_inner'] if insp and insp[0]['tread_inner'] is not None else r['tread_inner']
            outer = insp[0]['tread_outer'] if insp and insp[0]['tread_outer'] is not None else r['tread_outer']
            try:
                inner = float(inner) if inner is not None and str(inner).strip() != '' else None
            except Exception:
                inner = None
            try:
                outer = float(outer) if outer is not None and str(outer).strip() != '' else None
            except Exception:
                outer = None
            return outer, inner

        def avg_rtd(r):
            outer, inner = latest_rtd(r)
            if outer is None or inner is None:
                return None
            return (outer + inner) / 2.0

        def normalize_position(value):
            s = str(value or '').strip().upper().replace(' ', '')
            if s in ('1', 'P01', 'POS1', 'POS01'):
                return 'P1'
            if s in ('2', 'P02', 'POS2', 'POS02'):
                return 'P2'
            if s in ('3', 'P03', 'POS3', 'POS03'):
                return 'P3'
            if s in ('4', 'P04', 'POS4', 'POS04'):
                return 'P4'
            return s

        def axle_condition(diff):
            # Regla aprobada para 3.3 (sin porcentajes):
            # Normal:      diferencia < 5 mm
            # Preventivo:  diferencia >= 5 y <= 7.5 mm
            # Emergencia:  diferencia > 7.5 mm
            if diff is None:
                return 'SIN LECTURA'
            if diff > 7.5:
                return 'EMERGENCIA'
            if diff >= 5.0:
                return 'PREVENTIVO'
            return 'NORMAL'

        # Agrupar neumáticos por equipo y posición.
        equipment_map = {}
        for r in rows:
            eq_id = r['equipment_id']
            if eq_id is None:
                continue
            pos = normalize_position(r['position'])
            if pos not in ('P1','P2','P3','P4'):
                continue
            bucket = equipment_map.setdefault(eq_id, {
                'code': r['equipment_code'] or f'Equipo {eq_id}',
                'model': r['equipment_model'] or '',
                'vehicle_type': r['vehicle_type'] or '',
                'positions': {}
            })
            bucket['positions'][pos] = r

        evaluated = []
        missing_axes = 0
        for eq_id, eq in sorted(equipment_map.items(), key=lambda kv: str(kv[1]['code'])):
            for axle_name, p_left, p_right in [('Eje delantero','P1','P2'), ('Eje posterior','P3','P4')]:
                left = eq['positions'].get(p_left)
                right = eq['positions'].get(p_right)
                if not left or not right:
                    missing_axes += 1
                    continue
                rtd_left = avg_rtd(left)
                rtd_right = avg_rtd(right)
                if rtd_left is None or rtd_right is None:
                    missing_axes += 1
                    continue
                diff = abs(rtd_left - rtd_right)
                condition = axle_condition(diff)
                evaluated.append({
                    'equipment_id': eq_id,
                    'equipment_code': eq['code'],
                    'model': eq['model'],
                    'vehicle_type': eq['vehicle_type'],
                    'axle': axle_name,
                    'positions': f'{p_left} – {p_right}',
                    'left_code': left['code'],
                    'right_code': right['code'],
                    'left_rtd': rtd_left,
                    'right_rtd': rtd_right,
                    'diff': diff,
                    'condition': condition,
                })

        counts = {'NORMAL':0, 'PREVENTIVO':0, 'EMERGENCIA':0}
        for item in evaluated:
            counts[item['condition']] = counts.get(item['condition'], 0) + 1

        equipment_count = len(equipment_map)
        evaluated_total = len(evaluated)

        def top_metric(title, value, subtitle, value_color=TEXT_MAIN, tint=None):
            return ft.Container(
                width=205,
                height=104,
                bgcolor=tint or CARD_BG,
                border=ft.Border.all(1, '#DDE5ED'),
                border_radius=10,
                padding=12,
                content=ft.Column([
                    ft.Text(title, size=11.2, weight=ft.FontWeight.BOLD, color=TEXT_MAIN,
                            text_align=ft.TextAlign.CENTER),
                    ft.Text(str(value), size=28, weight=ft.FontWeight.BOLD, color=value_color,
                            text_align=ft.TextAlign.CENTER),
                    ft.Text(subtitle, size=9.2, color=TEXT_MUTED, text_align=ft.TextAlign.CENTER),
                ], spacing=3, horizontal_alignment=ft.CrossAxisAlignment.CENTER)
            )

        def axles_bar_chart():
            """Gráfico horizontal tipo Power BI en SVG, expresado únicamente en mm."""
            import base64
            chart_items = evaluated
            if not chart_items:
                return ft.Container(
                    height=160,
                    alignment=ft.Alignment.CENTER,
                    content=ft.Text('No hay ejes completos con lecturas RTD disponibles.', color=TEXT_MUTED)
                )

            row_h = 42
            left_w = 240
            bar_x = 285
            bar_w = 470
            max_diff = max([15.0] + [i['diff'] for i in chart_items])
            # Redondeo de escala a múltiplos de 2.5 para mantener lectura técnica.
            import math
            max_axis = max(10.0, math.ceil(max_diff / 2.5) * 2.5)
            width = 820
            height = 62 + row_h * len(chart_items) + 42
            y0 = 48
            limit_x = bar_x + (7.5 / max_axis) * bar_w

            parts = [
                f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}">',
                '<rect width="100%" height="100%" fill="#FFFFFF"/>',
                '<text x="12" y="22" font-family="Arial" font-size="12" font-weight="700" fill="#1B263B">Equipo / Eje</text>',
                '<text x="285" y="22" font-family="Arial" font-size="12" font-weight="700" fill="#1B263B">Diferencia de RTD (mm)</text>',
                f'<line x1="{limit_x:.1f}" y1="38" x2="{limit_x:.1f}" y2="{height-34}" stroke="#D92D20" stroke-width="2" stroke-dasharray="6 5"/>',
                f'<text x="{limit_x:.1f}" y="34" text-anchor="middle" font-family="Arial" font-size="10" font-weight="700" fill="#D92D20">7.5 mm</text>',
            ]

            colors = {'NORMAL':'#2E9B45', 'PREVENTIVO':'#F2A900', 'EMERGENCIA':'#D92D20'}
            last_equipment = None
            for idx, item in enumerate(chart_items):
                y = y0 + idx * row_h
                if last_equipment is not None and item['equipment_code'] != last_equipment:
                    parts.append(f'<line x1="12" y1="{y-8}" x2="{width-14}" y2="{y-8}" stroke="#DDE5ED" stroke-width="1"/>')
                last_equipment = item['equipment_code']
                color = colors[item['condition']]
                fill_w = max(3, (item['diff'] / max_axis) * bar_w)
                parts += [
                    f'<text x="14" y="{y+14}" font-family="Arial" font-size="12" font-weight="700" fill="#1B263B">{item["equipment_code"]}</text>',
                    f'<text x="92" y="{y+14}" font-family="Arial" font-size="10.5" fill="#334155">{item["axle"]}</text>',
                    f'<text x="205" y="{y+14}" font-family="Arial" font-size="11" fill="#334155">{item["positions"]}</text>',
                    f'<rect x="{bar_x}" y="{y-2}" width="{bar_w}" height="24" rx="3" fill="#E8EEF5"/>',
                    f'<rect x="{bar_x}" y="{y-2}" width="{fill_w:.1f}" height="24" rx="3" fill="{color}"/>',
                    f'<text x="{min(bar_x+fill_w+8, width-38):.1f}" y="{y+15}" font-family="Arial" font-size="11" font-weight="700" fill="{color}">{item["diff"]:.1f}</text>',
                ]

            # Escala inferior.
            tick = 2.5
            t = 0.0
            axis_y = height - 24
            parts.append(f'<line x1="{bar_x}" y1="{axis_y}" x2="{bar_x+bar_w}" y2="{axis_y}" stroke="#94A3B8" stroke-width="1"/>')
            while t <= max_axis + 0.001:
                x = bar_x + (t / max_axis) * bar_w
                label = f'{t:g}'
                parts.append(f'<line x1="{x:.1f}" y1="{axis_y}" x2="{x:.1f}" y2="{axis_y+4}" stroke="#94A3B8"/>')
                parts.append(f'<text x="{x:.1f}" y="{axis_y+16}" text-anchor="middle" font-family="Arial" font-size="9" fill="#64748B">{label}</text>')
                t += tick
            parts.append('</svg>')
            svg = ''.join(parts)
            src = 'data:image/svg+xml;base64,' + base64.b64encode(svg.encode('utf-8')).decode('ascii')
            return ft.Image(src=src, width=820, height=height, fit=ft.BoxFit.CONTAIN)

        def axles_donut_chart():
            import base64, math
            items = [
                ('Normal', counts['NORMAL'], '#2E9B45'),
                ('Preventivo', counts['PREVENTIVO'], '#F2A900'),
                ('Emergencia', counts['EMERGENCIA'], '#D92D20'),
            ]
            total_chart = sum(n for _, n, _ in items)
            cx = cy = 82; radius = 49; stroke = 24
            circumference = 2 * math.pi * radius
            offset = 0.0; circles = []; legend = []
            for label, n, color in items:
                dash = circumference * (n / total_chart) if total_chart else 0
                gap = max(0.0, circumference - dash)
                if n > 0:
                    circles.append(
                        f'<circle cx="{cx}" cy="{cy}" r="{radius}" fill="none" stroke="{color}" '
                        f'stroke-width="{stroke}" stroke-dasharray="{dash:.3f} {gap:.3f}" '
                        f'stroke-dashoffset="{-offset:.3f}" transform="rotate(-90 {cx} {cy})" />'
                    )
                offset += dash
                legend.append(ft.Row([
                    ft.Container(width=10, height=10, bgcolor=color, border_radius=5),
                    ft.Text(f'{label}: {n}', size=10.5, color=TEXT_MAIN),
                ], spacing=7))
            svg = (
                '<svg xmlns="http://www.w3.org/2000/svg" width="164" height="164" viewBox="0 0 164 164">'
                '<circle cx="82" cy="82" r="49" fill="none" stroke="#E2E8F0" stroke-width="24" />'
                + ''.join(circles) +
                f'<text x="82" y="80" text-anchor="middle" font-family="Arial" font-size="25" font-weight="700" fill="#172033">{total_chart}</text>'
                '<text x="82" y="99" text-anchor="middle" font-family="Arial" font-size="9" fill="#64748B">ejes</text>'
                '</svg>'
            )
            src = 'data:image/svg+xml;base64,' + base64.b64encode(svg.encode('utf-8')).decode('ascii')
            return ft.Row([
                ft.Image(src=src, width=164, height=164, fit=ft.BoxFit.CONTAIN),
                ft.Column(legend, spacing=10),
            ], spacing=12, vertical_alignment=ft.CrossAxisAlignment.CENTER)

        # Semáforo de 3.3 sin porcentajes.
        criteria_table = ft.DataTable(
            heading_row_height=34,
            heading_row_color='#000000',
            data_row_min_height=38,
            data_row_max_height=38,
            column_spacing=25,
            columns=[
                ft.DataColumn(ft.Text('Estado', size=11, weight=ft.FontWeight.BOLD, color='#FFFFFF')),
                ft.DataColumn(ft.Text('Diferencia RTD', size=11, weight=ft.FontWeight.BOLD, color='#FFFFFF')),
                ft.DataColumn(ft.Text('Cantidad', size=11, weight=ft.FontWeight.BOLD, color='#FFFFFF'), numeric=True),
            ],
            rows=[
                ft.DataRow(color='#DDF3E3', cells=[
                    ft.DataCell(ft.Row([ft.Container(width=10,height=10,bgcolor='#2E9B45',border_radius=5), ft.Text('Normal', size=10.5, weight=ft.FontWeight.BOLD, color='#176B2C')], spacing=7)),
                    ft.DataCell(ft.Text('< 5 mm', size=10.5, color=TEXT_MAIN)),
                    ft.DataCell(ft.Text(str(counts['NORMAL']), size=11, weight=ft.FontWeight.BOLD, color='#176B2C')),
                ]),
                ft.DataRow(color='#FFF0C2', cells=[
                    ft.DataCell(ft.Row([ft.Container(width=10,height=10,bgcolor='#F2A900',border_radius=5), ft.Text('Preventivo', size=10.5, weight=ft.FontWeight.BOLD, color='#A66000')], spacing=7)),
                    ft.DataCell(ft.Text('5 a 7.5 mm', size=10.5, color=TEXT_MAIN)),
                    ft.DataCell(ft.Text(str(counts['PREVENTIVO']), size=11, weight=ft.FontWeight.BOLD, color='#A66000')),
                ]),
                ft.DataRow(color='#FAD9D6', cells=[
                    ft.DataCell(ft.Row([ft.Container(width=10,height=10,bgcolor='#D92D20',border_radius=5), ft.Text('Emergencia', size=10.5, weight=ft.FontWeight.BOLD, color='#A61B12')], spacing=7)),
                    ft.DataCell(ft.Text('> 7.5 mm', size=10.5, color=TEXT_MAIN)),
                    ft.DataCell(ft.Text(str(counts['EMERGENCIA']), size=11, weight=ft.FontWeight.BOLD, color='#A61B12')),
                ]),
                ft.DataRow(color='#000000', cells=[
                    ft.DataCell(ft.Text('TOTAL EVALUADO', size=10.5, weight=ft.FontWeight.BOLD, color='#FFFFFF')),
                    ft.DataCell(ft.Text('P1–P2 / P3–P4', size=10.5, color='#FFFFFF')),
                    ft.DataCell(ft.Text(str(evaluated_total), size=10.5, weight=ft.FontWeight.BOLD, color='#FFFFFF')),
                ]),
            ]
        )

        worst = max(evaluated, key=lambda i: i['diff']) if evaluated else None
        worst_panel = ft.Column([
            ft.Text('EJE CON MAYOR DIFERENCIA DE RTD', size=14, weight=ft.FontWeight.BOLD, color=TEXT_MAIN),
            ft.Divider(height=4, color='#E6EBF0'),
        ], spacing=6)
        if worst:
            worst_color = {'NORMAL':'#2E9B45','PREVENTIVO':'#F2A900','EMERGENCIA':'#D92D20'}[worst['condition']]
            worst_panel.controls.extend([
                ft.Row([ft.Text('Equipo:', width=105, size=11, color=TEXT_MUTED), ft.Text(str(worst['equipment_code']), size=12, weight=ft.FontWeight.BOLD, color=TEXT_MAIN)]),
                ft.Row([ft.Text('Eje:', width=105, size=11, color=TEXT_MUTED), ft.Text(f'{worst["positions"]} ({worst["axle"]})', size=12, weight=ft.FontWeight.BOLD, color=TEXT_MAIN)]),
                ft.Row([ft.Text('RTD promedio:', width=105, size=11, color=TEXT_MUTED), ft.Text(f'{worst["left_rtd"]:.1f} / {worst["right_rtd"]:.1f} mm', size=12, color=TEXT_MAIN)]),
                ft.Row([ft.Text('Diferencia:', width=105, size=11, color=TEXT_MUTED), ft.Text(f'{worst["diff"]:.1f} mm', size=15, weight=ft.FontWeight.BOLD, color=worst_color)]),
                ft.Row([ft.Text('Estado:', width=105, size=11, color=TEXT_MUTED), ft.Row([ft.Container(width=11,height=11,bgcolor=worst_color,border_radius=6), ft.Text(worst['condition'].title(), size=12, weight=ft.FontWeight.BOLD, color=worst_color)], spacing=7)]),
            ])
        else:
            worst_panel.controls.append(ft.Text('Sin ejes evaluables.', size=11, color=TEXT_MUTED))

        content.content = ft.Column([
            page_title('3. Programa de mantenimiento · 3.3 Diferencia de RTD en el mismo eje',
                       'Comparación del RTD promedio entre P1–P2 y P3–P4 · Vista tipo Power BI'),
            maintenance_nav_row('33'),
            ft.Row([
                top_metric('EQUIPOS EN SERVICIO', equipment_count, 'Equipos con posiciones P1–P4'),
                top_metric('EJES EVALUADOS', evaluated_total, 'P1–P2 y P3–P4 con lectura'),
                top_metric('EJES EN CONDICIÓN NORMAL', counts['NORMAL'], '< 5 mm', '#2E9B45', '#F1FAF3'),
                top_metric('EJES EN PREVENTIVO', counts['PREVENTIVO'], '5 a 7.5 mm', '#C98600', '#FFF9E8'),
                top_metric('EJES EN EMERGENCIA', counts['EMERGENCIA'], '> 7.5 mm', '#C81D2A', '#FFF1F0'),
            ], wrap=True, spacing=10, run_spacing=10),
            ft.Row([
                ft.Container(expand=2, content=card(ft.Column([
                    ft.Text('DIFERENCIA DE RTD POR EJE', size=14, weight=ft.FontWeight.BOLD, color=TEXT_MAIN),
                    ft.Text('Cada barra representa |RTD promedio neumático A − RTD promedio neumático B|.', size=9.5, color=TEXT_MUTED),
                    ft.Row([axles_bar_chart()], scroll=ft.ScrollMode.AUTO),
                    ft.Text('Línea roja de referencia: 7.5 mm.', size=9.5, italic=True, color=TEXT_MUTED),
                ], spacing=7))),
                ft.Container(expand=1, content=ft.Column([
                    card(ft.Column([
                        ft.Text('CRITERIOS DE EVALUACIÓN', size=14, weight=ft.FontWeight.BOLD, color=TEXT_MAIN),
                        ft.Row([criteria_table], scroll=ft.ScrollMode.AUTO),
                    ], spacing=8)),
                    card(ft.Column([
                        ft.Text('DISTRIBUCIÓN DE EJES POR ESTADO', size=14, weight=ft.FontWeight.BOLD, color=TEXT_MAIN),
                        axles_donut_chart(),
                        ft.Text('La torta contabiliza únicamente ejes con ambas posiciones y lecturas EXT/INT disponibles.', size=9, italic=True, color=TEXT_MUTED),
                    ], spacing=8)),
                ], spacing=12)),
            ], spacing=12, vertical_alignment=ft.CrossAxisAlignment.START),
            ft.Row([
                ft.Container(expand=1, content=card(worst_panel)),
                ft.Container(expand=1, content=card(ft.Column([
                    ft.Text('NOTAS TÉCNICAS', size=14, weight=ft.FontWeight.BOLD, color=TEXT_MAIN),
                    ft.Text('• RTD de cada neumático = promedio entre hombro EXT e INT.', size=10.5, color=TEXT_MAIN),
                    ft.Text('• Diferencia por eje = valor absoluto entre los RTD promedio de P1–P2 o P3–P4.', size=10.5, color=TEXT_MAIN),
                    ft.Text('• Los datos se toman de la última INSP/INSC registrada por neumático.', size=10.5, color=TEXT_MAIN),
                    ft.Text('• Regla 3.3: Normal < 5 mm · Preventivo 5 a 7.5 mm · Emergencia > 7.5 mm.', size=10.5, color=TEXT_MAIN),
                    ft.Text(f'• Ejes no evaluados por posición o lectura faltante: {missing_axes}.', size=10.5, color=TEXT_MUTED),
                ], spacing=7))),
            ], spacing=12, vertical_alignment=ft.CrossAxisAlignment.START),
        ], scroll=ft.ScrollMode.AUTO, spacing=16)
        page.update()


    def maintenance_four_positions_view():
        """3.4 Diferencia de RTD entre las cuatro posiciones P1-P4, sin considerar diámetro."""
        op_id = active_operation_id()
        rows = query("""
            SELECT t.id, t.code, t.tread_inner, t.tread_outer, t.recommended_pressure,
                   e.id AS equipment_id, e.code AS equipment_code, e.vehicle_type, t.position
            FROM tires t
            LEFT JOIN equipment e ON e.id=t.equipment_id
            WHERE t.status='SERVICIO' AND t.operation_id=?
            ORDER BY COALESCE(e.code,''), t.position, t.code
        """, (op_id,))

        def norm_pos(v):
            x=str(v or '').strip().upper().replace(' ','')
            return {'1':'P1','P01':'P1','POS1':'P1','POS01':'P1',
                    '2':'P2','P02':'P2','POS2':'P2','POS02':'P2',
                    '3':'P3','P03':'P3','POS3':'P3','POS03':'P3',
                    '4':'P4','P04':'P4','POS4':'P4','POS04':'P4'}.get(x,x)

        def latest_avg(r):
            z=query("""SELECT tread_inner,tread_outer FROM occurrences
                       WHERE tire_id=? AND operation_id=? AND event_code IN ('INSP','INSC')
                       ORDER BY id DESC LIMIT 1""",(r['id'],op_id))
            inn=z[0]['tread_inner'] if z and z[0]['tread_inner'] is not None else r['tread_inner']
            out=z[0]['tread_outer'] if z and z[0]['tread_outer'] is not None else r['tread_outer']
            try: inn=float(inn)
            except Exception: return None
            try: out=float(out)
            except Exception: return None
            return (inn+out)/2.0

        def state(d):
            if d>7.5: return 'EMERGENCIA'
            if d>=5.0: return 'PREVENTIVO'
            return 'NORMAL'

        eqmap={}
        for r in rows:
            if r['equipment_id'] is None: continue
            pos=norm_pos(r['position'])
            if pos not in ('P1','P2','P3','P4'): continue
            b=eqmap.setdefault(r['equipment_id'],{'code':r['equipment_code'] or f"Equipo {r['equipment_id']}",'pos':{}})
            b['pos'][pos]=r

        evaluated=[]
        for _,eq in sorted(eqmap.items(),key=lambda kv:str(kv[1]['code'])):
            if not all(p in eq['pos'] for p in ('P1','P2','P3','P4')): continue
            vals={p:latest_avg(eq['pos'][p]) for p in ('P1','P2','P3','P4')}
            if any(v is None for v in vals.values()): continue
            diff=max(vals.values())-min(vals.values())
            evaluated.append({'equipment_code':eq['code'],'vals':vals,'diff':diff,'condition':state(diff)})

        counts={'NORMAL':0,'PREVENTIVO':0,'EMERGENCIA':0}
        for x in evaluated: counts[x['condition']]+=1
        total=len(evaluated)
        worst=max(evaluated,key=lambda x:x['diff']) if evaluated else None

        def metric(title,value,subtitle,color=TEXT_MAIN,tint=None):
            return ft.Container(width=205,height=104,bgcolor=tint or CARD_BG,border=ft.Border.all(1,'#DDE5ED'),border_radius=10,padding=12,
                content=ft.Column([ft.Text(title,size=11.2,weight=ft.FontWeight.BOLD,color=TEXT_MAIN,text_align=ft.TextAlign.CENTER),
                    ft.Text(str(value),size=28,weight=ft.FontWeight.BOLD,color=color,text_align=ft.TextAlign.CENTER),
                    ft.Text(subtitle,size=9.2,color=TEXT_MUTED,text_align=ft.TextAlign.CENTER)],spacing=3,horizontal_alignment=ft.CrossAxisAlignment.CENTER))

        def bars():
            import base64, math
            if not evaluated:
                return ft.Container(height=150,alignment=ft.Alignment.CENTER,content=ft.Text('No hay equipos con P1–P4 y lecturas completas.',color=TEXT_MUTED))
            row_h=40; bar_x=405; bar_w=390; width=840; y0=46
            maxd=max([15.0]+[x['diff'] for x in evaluated]); maxa=max(10.0,math.ceil(maxd/2.5)*2.5)
            h=62+row_h*len(evaluated)+34; lx=bar_x+(7.5/maxa)*bar_w
            a=[f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{h}" viewBox="0 0 {width} {h}">','<rect width="100%" height="100%" fill="#FFFFFF"/>',
               '<text x="12" y="22" font-family="Arial" font-size="11" font-weight="700" fill="#1B263B">Equipo</text>',
               '<text x="95" y="22" font-family="Arial" font-size="10" font-weight="700" fill="#1B263B">P1</text>',
               '<text x="160" y="22" font-family="Arial" font-size="10" font-weight="700" fill="#1B263B">P2</text>',
               '<text x="225" y="22" font-family="Arial" font-size="10" font-weight="700" fill="#1B263B">P3</text>',
               '<text x="290" y="22" font-family="Arial" font-size="10" font-weight="700" fill="#1B263B">P4</text>',
               '<text x="405" y="22" font-family="Arial" font-size="11" font-weight="700" fill="#1B263B">Diferencia (mm)</text>',
               f'<line x1="{lx:.1f}" y1="32" x2="{lx:.1f}" y2="{h-30}" stroke="#D92D20" stroke-width="2" stroke-dasharray="6 5"/>',
               f'<text x="{lx:.1f}" y="30" text-anchor="middle" font-family="Arial" font-size="9" font-weight="700" fill="#D92D20">7.5</text>']
            colors={'NORMAL':'#2E9B45','PREVENTIVO':'#F2A900','EMERGENCIA':'#D92D20'}
            for i,x in enumerate(evaluated):
                y=y0+i*row_h; c=colors[x['condition']]; fw=max(3,(x['diff']/maxa)*bar_w)
                a += [f'<text x="12" y="{y+13}" font-family="Arial" font-size="11" font-weight="700" fill="#1B263B">{x["equipment_code"]}</text>']
                for j,p in enumerate(('P1','P2','P3','P4')):
                    a.append(f'<text x="{95+j*65}" y="{y+13}" font-family="Arial" font-size="10" fill="#334155">{x["vals"][p]:.1f}</text>')
                a += [f'<rect x="{bar_x}" y="{y-3}" width="{bar_w}" height="23" rx="3" fill="#E8EEF5"/>',
                      f'<rect x="{bar_x}" y="{y-3}" width="{fw:.1f}" height="23" rx="3" fill="{c}"/>',
                      f'<text x="{min(bar_x+fw+7,width-32):.1f}" y="{y+13}" font-family="Arial" font-size="10.5" font-weight="700" fill="{c}">{x["diff"]:.1f}</text>']
            a.append('</svg>')
            src='data:image/svg+xml;base64,'+base64.b64encode(''.join(a).encode()).decode()
            return ft.Image(src=src,width=840,height=h,fit=ft.BoxFit.CONTAIN)

        def donut():
            import base64, math
            its=[('Normal',counts['NORMAL'],'#2E9B45'),('Preventivo',counts['PREVENTIVO'],'#F2A900'),('Emergencia',counts['EMERGENCIA'],'#D92D20')]
            circ=2*math.pi*49; off=0; cs=[]; leg=[]
            for lab,n,c in its:
                dash=circ*(n/total) if total else 0; gap=max(0,circ-dash)
                if n: cs.append(f'<circle cx="82" cy="82" r="49" fill="none" stroke="{c}" stroke-width="24" stroke-dasharray="{dash:.3f} {gap:.3f}" stroke-dashoffset="{-off:.3f}" transform="rotate(-90 82 82)"/>')
                off+=dash
                leg.append(ft.Row([ft.Container(width=10,height=10,bgcolor=c,border_radius=5),ft.Text(f'{lab}: {n}',size=10.5,color=TEXT_MAIN)],spacing=7))
            svg='<svg xmlns="http://www.w3.org/2000/svg" width="164" height="164"><circle cx="82" cy="82" r="49" fill="none" stroke="#E2E8F0" stroke-width="24"/>'+''.join(cs)+f'<text x="82" y="80" text-anchor="middle" font-family="Arial" font-size="30" font-weight="700" fill="#172033">{total}</text><text x="82" y="99" text-anchor="middle" font-family="Arial" font-size="9" fill="#64748B">equipos</text></svg>'
            src='data:image/svg+xml;base64,'+base64.b64encode(svg.encode()).decode()
            return ft.Row([ft.Image(src=src,width=164,height=164),ft.Column(leg,spacing=10)],spacing=12)

        criteria=ft.DataTable(heading_row_height=34,heading_row_color='#000000',data_row_min_height=38,data_row_max_height=38,column_spacing=22,
            columns=[ft.DataColumn(ft.Text('Estado',size=11,weight=ft.FontWeight.BOLD,color='white')),ft.DataColumn(ft.Text('Diferencia RTD',size=11,weight=ft.FontWeight.BOLD,color='white'))],
            rows=[ft.DataRow(color='#DDF3E3',cells=[ft.DataCell(ft.Text('● Normal',color='#176B2C',weight=ft.FontWeight.BOLD)),ft.DataCell(ft.Text('< 5 mm'))]),
                  ft.DataRow(color='#FFF0C2',cells=[ft.DataCell(ft.Text('● Preventivo',color='#A66000',weight=ft.FontWeight.BOLD)),ft.DataCell(ft.Text('5 a 7.5 mm'))]),
                  ft.DataRow(color='#FAD9D6',cells=[ft.DataCell(ft.Text('● Emergencia',color='#A61B12',weight=ft.FontWeight.BOLD)),ft.DataCell(ft.Text('> 7.5 mm'))])])

        content.content=ft.Column([
            page_title('3. Programa de mantenimiento · 3.4 Diferencia entre ejes por equipo','Comparación del RTD promedio entre P1, P2, P3 y P4 · Sin considerar diámetro'),
            maintenance_nav_row('34'),
            ft.Row([metric('EQUIPOS EVALUADOS',total,'Equipos con P1–P4'),metric('EN CONDICIÓN NORMAL',counts['NORMAL'],'< 5 mm','#2E9B45','#F1FAF3'),
                    metric('EN PREVENTIVO',counts['PREVENTIVO'],'5 a 7.5 mm','#C98600','#FFF9E8'),metric('EN EMERGENCIA',counts['EMERGENCIA'],'> 7.5 mm','#C81D2A','#FFF1F0'),
                    metric('DIFERENCIA MÁXIMA',f'{worst["diff"]:.1f} mm' if worst else '—',f'Equipo: {worst["equipment_code"]}' if worst else 'Sin datos','#C81D2A')],wrap=True,spacing=10,run_spacing=10),
            ft.Row([ft.Container(expand=2,content=card(ft.Column([ft.Text('DIFERENCIA ENTRE EJES POR EQUIPO',size=14,weight=ft.FontWeight.BOLD,color=TEXT_MAIN),
                        ft.Text('Diferencia = RTD promedio mayor − RTD promedio menor. No se considera diámetro.',size=9.5,color=TEXT_MUTED),ft.Row([bars()],scroll=ft.ScrollMode.AUTO)],spacing=7))),
                    ft.Container(expand=1,content=ft.Column([card(ft.Column([ft.Text('DISTRIBUCIÓN DE EQUIPOS POR ESTADO',size=14,weight=ft.FontWeight.BOLD,color=TEXT_MAIN),donut()],spacing=8)),
                        card(ft.Column([ft.Text('CRITERIOS DE EVALUACIÓN',size=14,weight=ft.FontWeight.BOLD,color=TEXT_MAIN),ft.Row([criteria],scroll=ft.ScrollMode.AUTO)],spacing=8)),
                        card(ft.Column([ft.Text('NOTA',size=13,weight=ft.FontWeight.BOLD,color=TEXT_MAIN),ft.Text('Se usa la última INSP/INSC de cada neumático. RTD de cada posición = promedio EXT/INT. La diferencia corresponde al mayor menos el menor RTD de P1–P4.',size=10.2,color=TEXT_MAIN)],spacing=6))],spacing=12))],spacing=12,vertical_alignment=ft.CrossAxisAlignment.START)
        ],scroll=ft.ScrollMode.AUTO,spacing=16)
        page.update()

    def maintenance_pressure_view():
        """3.5 Nivelación de presión: cuadro resumen + gráfico, con los criterios del Módulo 2."""
        op_id = active_operation_id()
        rows = query("""
            SELECT t.id, t.code, t.recommended_pressure,
                   e.code AS equipment_code, t.position
            FROM tires t
            LEFT JOIN equipment e ON e.id=t.equipment_id
            WHERE t.status='SERVICIO' AND t.operation_id=?
            ORDER BY COALESCE(e.code,''), t.position, t.code
        """, (op_id,))

        def norm_pos(v):
            x=str(v or '').strip().upper().replace(' ','')
            return {'1':'P1','P01':'P1','POS1':'P1','POS01':'P1',
                    '2':'P2','P02':'P2','POS2':'P2','POS02':'P2',
                    '3':'P3','P03':'P3','POS3':'P3','POS03':'P3',
                    '4':'P4','P04':'P4','POS4':'P4','POS04':'P4'}.get(x,x)

        counts={'green':0,'orange':0,'red':0,'purple':0}
        activities=[]
        evaluated=0
        detail=[]
        for r in rows:
            z=query("""SELECT pressure FROM occurrences
                       WHERE tire_id=? AND operation_id=? AND event_code IN ('INSP','INSC')
                       ORDER BY id DESC LIMIT 1""",(r['id'], op_id))
            act=z[0]['pressure'] if z and z[0]['pressure'] is not None else None
            rec=r['recommended_pressure']
            try:
                act=float(act); rec=float(rec)
                if rec <= 0:
                    continue
            except Exception:
                continue
            evaluated += 1
            diff=abs(act-rec)
            if act > rec * 1.20:
                status='purple'; label='Sobrepresión > 20%'
            elif diff <= 5:
                status='green'; label='± 5 psi (OK)'
            elif diff <= 10:
                status='orange'; label='> 5 a 10 psi'
            else:
                status='red'; label='> 10 psi'
            counts[status] += 1
            eq=r['equipment_code'] or 'SIN EQUIPO'
            pos=norm_pos(r['position']) or 'SIN POSICIÓN'
            detail.append((eq,pos,r['code'],rec,act,diff,label,status))
            if status in ('red','purple'):
                activities.append(f'NIVELAR PRESIÓN DEL NEUMÁTICO {pos} DEL EQUIPO {eq}.')

        import base64, math
        def pressure_donut():
            items=[
                ('± 5 psi (OK)',counts['green'],'#16A34A'),
                ('> 5 a 10 psi',counts['orange'],'#F59E0B'),
                ('> 10 psi',counts['red'],'#DC2626'),
                ('Sobrepresión > 20%',counts['purple'],'#7C3AED'),
            ]
            total=evaluated
            cx=82; cy=82; radius=50; stroke=24
            circumference=2*math.pi*radius
            offset=0.0; circles=[]; legend=[]
            for label,count,color in items:
                pct=(count/total*100.0) if total else 0.0
                dash=circumference*(count/total) if total else 0.0
                gap=max(0.0,circumference-dash)
                if count>0:
                    circles.append(
                        f'<circle cx="{cx}" cy="{cy}" r="{radius}" fill="none" stroke="{color}" '
                        f'stroke-width="{stroke}" stroke-dasharray="{dash:.3f} {gap:.3f}" '
                        f'stroke-dashoffset="{-offset:.3f}" transform="rotate(-90 {cx} {cy})" />'
                    )
                offset += dash
                legend.append(ft.Row([
                    ft.Container(width=9,height=9,bgcolor=color,border_radius=2),
                    ft.Text(f'{label}: {count} ({pct:.1f}%)',size=9.5,color=TEXT_MAIN),
                ],spacing=6))
            svg=(
                '<svg xmlns="http://www.w3.org/2000/svg" width="164" height="164" viewBox="0 0 164 164">'
                '<circle cx="82" cy="82" r="50" fill="none" stroke="#E2E8F0" stroke-width="24" />'
                + ''.join(circles) +
                f'<text x="82" y="79" text-anchor="middle" font-family="Arial" font-size="23" font-weight="700" fill="#172033">{total}</text>'
                '<text x="82" y="98" text-anchor="middle" font-family="Arial" font-size="10" fill="#64748B">Total</text>'
                '</svg>'
            )
            src='data:image/svg+xml;base64,'+base64.b64encode(svg.encode('utf-8')).decode('ascii')
            return ft.Row([
                ft.Image(src=src,width=164,height=164,fit=ft.BoxFit.CONTAIN),
                ft.Column(legend,spacing=7),
            ],spacing=12,vertical_alignment=ft.CrossAxisAlignment.CENTER)

        def pct(n):
            return f'{(n/evaluated*100.0):.1f}%' if evaluated else '0.0%'

        summary_table=ft.DataTable(
            heading_row_height=36,
            data_row_min_height=38,
            data_row_max_height=38,
            column_spacing=24,
            columns=[
                ft.DataColumn(ft.Text('Condición',size=10.5,weight=ft.FontWeight.BOLD)),
                ft.DataColumn(ft.Text('Criterio',size=10.5,weight=ft.FontWeight.BOLD)),
                ft.DataColumn(ft.Text('Cantidad',size=10.5,weight=ft.FontWeight.BOLD)),
                ft.DataColumn(ft.Text('%',size=10.5,weight=ft.FontWeight.BOLD)),
            ],
            rows=[
                ft.DataRow(cells=[ft.DataCell(ft.Text('OK',size=10.5)),ft.DataCell(ft.Text('± 5 psi',size=10.5)),ft.DataCell(ft.Text(str(counts['green']),size=10.5)),ft.DataCell(ft.Text(pct(counts['green']),size=10.5))]),
                ft.DataRow(cells=[ft.DataCell(ft.Text('Preventivo',size=10.5)),ft.DataCell(ft.Text('> 5 a 10 psi',size=10.5)),ft.DataCell(ft.Text(str(counts['orange']),size=10.5)),ft.DataCell(ft.Text(pct(counts['orange']),size=10.5))]),
                ft.DataRow(cells=[ft.DataCell(ft.Text('Nivelación',size=10.5)),ft.DataCell(ft.Text('> 10 psi',size=10.5)),ft.DataCell(ft.Text(str(counts['red']),size=10.5)),ft.DataCell(ft.Text(pct(counts['red']),size=10.5))]),
                ft.DataRow(cells=[ft.DataCell(ft.Text('Crítico',size=10.5)),ft.DataCell(ft.Text('Sobrepresión > 20%',size=10.5)),ft.DataCell(ft.Text(str(counts['purple']),size=10.5)),ft.DataCell(ft.Text(pct(counts['purple']),size=10.5))]),
            ]
        )

        controls=[
            page_title('3. Programa de mantenimiento · 3.5 Nivelación de presión',
                       'Comparación de la última presión INSP/INSC contra la presión recomendada'),
            maintenance_nav_row('35'),
            ft.Row([
                ft.Container(expand=1,content=card(ft.Column([
                    ft.Text('CUADRO DE EVALUACIÓN DE PRESIÓN',size=14,weight=ft.FontWeight.BOLD,color=TEXT_MAIN),
                    ft.Text('Mismos parámetros establecidos en Neumáticos en servicio.',size=10,color=TEXT_MUTED),
                    ft.Row([summary_table],scroll=ft.ScrollMode.AUTO),
                ],spacing=8))),
                ft.Container(expand=1,content=card(ft.Column([
                    ft.Text('EVALUACIÓN DE PRESIÓN (PSI)',size=14,weight=ft.FontWeight.BOLD,color=TEXT_MAIN),
                    ft.Text('Diferencia absoluta entre presión actual y recomendada.',size=10,color=TEXT_MUTED),
                    pressure_donut(),
                ],spacing=8))),
            ],spacing=12,vertical_alignment=ft.CrossAxisAlignment.START),
        ]
        lines=[ft.Text('ACTIVIDADES DE NIVELACIÓN DE PRESIÓN',size=14,weight=ft.FontWeight.BOLD,color=TEXT_MAIN)]
        if activities:
            lines += [ft.Text(f'- {x}',size=12,color=TEXT_MAIN) for x in activities]
        else:
            lines.append(ft.Text('- SIN ACTIVIDADES DE MANTENIMIENTO PENDIENTES.',size=12,color=TEXT_MAIN))
        controls.append(card(ft.Column(lines,spacing=7),padding=20))
        content.content=ft.Column(controls,scroll=ft.ScrollMode.AUTO,spacing=16)
        page.update()

    def maintenance_final_report_view():
        """Reporte final: actividades de mantenimiento generadas solo por condiciones de emergencia."""
        op_id = active_operation_id()
        rows = query("""
            SELECT t.id, t.code, t.tread_inner, t.tread_outer, t.recommended_pressure,
                   e.id AS equipment_id, e.code AS equipment_code,
                   e.vehicle_type AS equipment_vehicle_type, t.position
            FROM tires t
            LEFT JOIN equipment e ON e.id=t.equipment_id
            WHERE t.status='SERVICIO' AND t.operation_id=?
            ORDER BY COALESCE(e.code,''), t.position, t.code
        """, (op_id,))

        def norm_pos(v):
            x=str(v or '').strip().upper().replace(' ','')
            return {'1':'P1','P01':'P1','POS1':'P1','POS01':'P1',
                    '2':'P2','P02':'P2','POS2':'P2','POS02':'P2',
                    '3':'P3','P03':'P3','POS3':'P3','POS03':'P3',
                    '4':'P4','P04':'P4','POS4':'P4','POS04':'P4'}.get(x,x)

        def latest_pair(r):
            z=query("""SELECT tread_inner,tread_outer FROM occurrences
                       WHERE tire_id=? AND operation_id=? AND event_code IN ('INSP','INSC')
                       ORDER BY id DESC LIMIT 1""",(r['id'], op_id))
            inn=z[0]['tread_inner'] if z and z[0]['tread_inner'] is not None else r['tread_inner']
            out=z[0]['tread_outer'] if z and z[0]['tread_outer'] is not None else r['tread_outer']
            try: inn=float(inn)
            except Exception: inn=None
            try: out=float(out)
            except Exception: out=None
            return out,inn

        def latest_avg(r):
            out,inn=latest_pair(r)
            if out is None or inn is None: return None
            return (out+inn)/2.0

        activities_31=[]; activities_32=[]; activities_33=[]; activities_34=[]; activities_35=[]

        def fmt_tech(v):
            try:
                n=float(v)
                return str(int(n)) if n.is_integer() else f'{n:.1f}'
            except Exception:
                return str(v or '')

        # Cada actividad guarda: (texto de la acción, dato técnico a resaltar en rojo).

        # 3.1: mismo criterio de cambio urgente de 3.1 Evaluación de Remanente,
        # pero aplicado según el tipo de equipo.
        def equipment_category(r):
            # Normaliza vehicle_type para usar los mismos criterios de 3.1.
            raw = str(
                r.get('equipment_vehicle_type')
                or r.get('vehicle_type')
                or ''
            ).strip().upper()
            if 'SCOOP' in raw or 'LHD' in raw:
                return 'SCOOP'
            if 'VOLQUETE' in raw or 'DUMPER' in raw or 'MINERO' in raw:
                return 'VOLQUETES'
            if 'CAMIONETA' in raw or 'PICKUP' in raw:
                return 'CAMIONETAS'
            return 'OTROS'

        def rtd_change_category(category, rtd):
            # Exactamente los mismos umbrales definidos en 3.1.
            thresholds = {
                'SCOOP': (20.0, 30.0),
                'VOLQUETES': (10.0, 15.0),
                'CAMIONETAS': (4.0, 6.0),
                'OTROS': (10.0, 15.0),
            }
            emergency_max, _ = thresholds.get(category, thresholds['OTROS'])
            if rtd is None or rtd < 0:
                return False
            return rtd <= emergency_max

        for r in rows:
            vals=[]
            for v in (r['tread_inner'],r['tread_outer']):
                try:
                    if v is not None and str(v).strip()!='': vals.append(float(v))
                except Exception: pass
            if vals:
                rtd_min = min(vals)
                category = equipment_category(r)
                if rtd_change_category(category, rtd_min):
                    eq=r['equipment_code'] or 'SIN EQUIPO'; pos=norm_pos(r['position']) or 'SIN POSICIÓN'
                    activities_31.append((
                        f'CAMBIO DE NEUMÁTICO DE LA {pos} DEL EQUIPO {eq}.',
                        f'{pos} RTD {fmt_tech(rtd_min)} MM'
                    ))

        # 3.2: inversión cuando RTD INT - RTD EXT >= 10 mm.
        for r in rows:
            out,inn=latest_pair(r)
            if out is not None and inn is not None and (inn-out) >= 10:
                eq=r['equipment_code'] or 'SIN EQUIPO'; pos=norm_pos(r['position']) or 'SIN POSICIÓN'
                activities_32.append((
                    f'INVERTIR EL NEUMÁTICO {pos} DEL EQUIPO {eq}.',
                    f'EXT {fmt_tech(out)} / INT {fmt_tech(inn)} MM'
                ))

        # Agrupar P1-P4 por equipo para 3.3 y 3.4.
        eqmap={}
        for r in rows:
            if r['equipment_id'] is None: continue
            p=norm_pos(r['position'])
            if p not in ('P1','P2','P3','P4'): continue
            b=eqmap.setdefault(r['equipment_id'],{'code':r['equipment_code'] or f"Equipo {r['equipment_id']}",'pos':{}})
            b['pos'][p]=r

        for _,eq in sorted(eqmap.items(),key=lambda kv:str(kv[1]['code'])):
            pos=eq['pos']; code=eq['code']
            # 3.3: diferencia > 7.5 mm dentro del mismo eje.
            for pa,pb,axle_name in (('P1','P2','DELANTERO'),('P3','P4','POSTERIOR')):
                if pa in pos and pb in pos:
                    a=latest_avg(pos[pa]); b=latest_avg(pos[pb])
                    if a is not None and b is not None and abs(a-b) > 7.5:
                        activities_33.append((
                            f'NIVELACIÓN DEL EJE {axle_name} DEL EQUIPO {code}.',
                            f'{pa} {fmt_tech(a)} / {pb} {fmt_tech(b)} MM'
                        ))
            # 3.4: conserva el criterio vigente; se añade la información técnica de ambos ejes.
            if all(p in pos for p in ('P1','P2','P3','P4')):
                vals=[latest_avg(pos[p]) for p in ('P1','P2','P3','P4')]
                if all(v is not None for v in vals) and (max(vals)-min(vals)) > 7.5:
                    eje_del=(vals[0]+vals[1])/2.0
                    eje_post=(vals[2]+vals[3])/2.0
                    activities_34.append((
                        f'NIVELACIÓN DE EJES DEL EQUIPO {code}.',
                        f'EJE DEL. {fmt_tech(eje_del)} / EJE POST. {fmt_tech(eje_post)} MM'
                    ))

        # 3.5: nivelación de presión con los mismos parámetros del Módulo 2.
        # Solo genera actividad para >10 psi de diferencia o sobrepresión >20%.
        for r in rows:
            z=query("""SELECT pressure FROM occurrences
                       WHERE tire_id=? AND operation_id=? AND event_code IN ('INSP','INSC')
                       ORDER BY id DESC LIMIT 1""",(r['id'], op_id))
            act=z[0]['pressure'] if z and z[0]['pressure'] is not None else None
            rec=r['recommended_pressure']
            try:
                act=float(act); rec=float(rec)
                if rec <= 0:
                    continue
            except Exception:
                continue
            diff=abs(act-rec)
            if (act > rec * 1.20) or (diff > 10):
                eq=r['equipment_code'] or 'SIN EQUIPO'; pos=norm_pos(r['position']) or 'SIN POSICIÓN'
                activities_35.append((
                    f'NIVELAR PRESIÓN DEL NEUMÁTICO {pos} DEL EQUIPO {eq}.',
                    f'{pos} {fmt_tech(act)} PSI'
                ))

        def section(title, items):
            controls=[ft.Text(title,size=15,weight=ft.FontWeight.BOLD,color=TEXT_MAIN)]
            if items:
                for action,detail in items:
                    controls.append(ft.Row([
                        ft.Text(f'- {action}',size=12,color=TEXT_MAIN),
                        ft.Text(detail,size=12,color=ft.Colors.RED_700,weight=ft.FontWeight.BOLD),
                    ],spacing=5,wrap=True))
            else:
                controls.append(ft.Text('- SIN ACTIVIDADES DE MANTENIMIENTO PENDIENTES.',size=12,color=TEXT_MAIN))
            return ft.Column(controls,spacing=7)

        report_date=dt.datetime.now(dt.timezone(dt.timedelta(hours=-5))).strftime('%d/%m/%Y')

        def build_pdf_bytes():
            """Genera un PDF simple multipágina sin dependencias externas."""
            sections=[
                ('3.1 EVALUACIÓN DE REMANENTE',activities_31),
                ('3.2 DIFERENCIA DE RTD ENTRE HOMBROS',activities_32),
                ('3.3 DIFERENCIA DE RTD DEL MISMO EJE',activities_33),
                ('3.4 DIFERENCIA ENTRE EJES POR EQUIPO',activities_34),
                ('3.5 NIVELACIÓN DE PRESIÓN',activities_35),
            ]

            def pdf_escape(txt):
                txt=str(txt or '').replace('–','-').replace('—','-')
                raw=txt.encode('cp1252','replace').decode('latin-1')
                return raw.replace('\\','\\\\').replace('(','\\(').replace(')','\\)')

            pages=[]; cmds=[]; y=795
            def new_page():
                nonlocal cmds,y
                if cmds:
                    pages.append('\n'.join(cmds))
                cmds=[]; y=795
                cmds.append('BT /F1 17 Tf 0 0 0 rg 45 795 Td ('+pdf_escape('PROGRAMA DE MANTENIMIENTO DE NEUMÁTICOS')+') Tj ET')
                cmds.append('BT /F1 10 Tf 0 0 0 rg 45 774 Td ('+pdf_escape('Fecha de visualización: '+report_date)+') Tj ET')
                cmds.append('0.75 w 45 762 m 550 762 l S')
                y=742

            def ensure(space=28):
                nonlocal y
                if y-space < 45:
                    new_page()

            def put(text,size=10,bold=False,red=False,indent=0,leading=14):
                nonlocal y
                ensure(leading+5)
                font='/F2' if bold else '/F1'
                color='0.80 0.08 0.10 rg' if red else '0 0 0 rg'
                cmds.append(f'BT {font} {size} Tf {color} {45+indent} {y} Td ('+pdf_escape(text)+') Tj ET')
                y-=leading

            new_page()
            for title,items in sections:
                ensure(42)
                put(title,11,bold=True,leading=18)
                if not items:
                    put('- SIN ACTIVIDADES DE MANTENIMIENTO PENDIENTES.',9,leading=16)
                else:
                    for action,detail in items:
                        action_text='- '+action
                        wrapped=textwrap.wrap(action_text,width=88,break_long_words=False,break_on_hyphens=False) or [action_text]
                        for i,line in enumerate(wrapped):
                            put(line,9,indent=0 if i==0 else 10,leading=13)
                        put(detail,9,bold=True,red=True,indent=15,leading=15)
                y-=7
                ensure(15)
                cmds.append(f'0.85 G 45 {y+3} m 550 {y+3} l S')
                y-=8
            if cmds:
                pages.append('\n'.join(cmds))

            objects=[]
            objects.append(b'<< /Type /Catalog /Pages 2 0 R >>')
            # Pages object is filled after page objects are known.
            objects.append(b'')
            objects.append(b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>')
            objects.append(b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica-Bold /Encoding /WinAnsiEncoding >>')
            page_ids=[]
            for content_stream in pages:
                stream=content_stream.encode('latin-1','replace')
                content_id=len(objects)+1
                objects.append(b'<< /Length '+str(len(stream)).encode()+b' >>\nstream\n'+stream+b'\nendstream')
                page_id=len(objects)+1
                page_ids.append(page_id)
                objects.append((f'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] '
                                f'/Resources << /Font << /F1 3 0 R /F2 4 0 R >> >> '
                                f'/Contents {content_id} 0 R >>').encode('ascii'))
            kids=' '.join(f'{pid} 0 R' for pid in page_ids)
            objects[1]=f'<< /Type /Pages /Kids [{kids}] /Count {len(page_ids)} >>'.encode('ascii')

            out=bytearray(b'%PDF-1.4\n%\xe2\xe3\xcf\xd3\n')
            offsets=[0]
            for i,obj in enumerate(objects,1):
                offsets.append(len(out))
                out.extend(f'{i} 0 obj\n'.encode('ascii')); out.extend(obj); out.extend(b'\nendobj\n')
            xref=len(out)
            out.extend(f'xref\n0 {len(objects)+1}\n'.encode('ascii'))
            out.extend(b'0000000000 65535 f \n')
            for off in offsets[1:]:
                out.extend(f'{off:010d} 00000 n \n'.encode('ascii'))
            out.extend((f'trailer\n<< /Size {len(objects)+1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF').encode('ascii'))
            return bytes(out)

        async def download_pdf(e):
            try:
                pdf_bytes=build_pdf_bytes()
                file_name='Programa_Mantenimiento_'+report_date.replace('/','-')+'.pdf'
                # En Flet Web, FilePicker.save_file con src_bytes entrega el archivo
                # directamente al navegador; evita los data: URI que Chrome/Render
                # pueden bloquear.
                await ft.FilePicker().save_file(
                    file_name=file_name,
                    file_type=ft.FilePickerFileType.CUSTOM,
                    allowed_extensions=['pdf'],
                    src_bytes=pdf_bytes,
                )
                snack(f'PDF preparado: {file_name}')
            except Exception as ex:
                snack(f'No se pudo descargar el PDF: {ex}',True)

        content.content=ft.Column([
            page_title('PROGRAMA DE MANTENIMIENTO DE NEUMÁTICOS',
                       f'Fecha de visualización: {report_date} · Actividades generadas automáticamente a partir de condiciones de emergencia'),
            maintenance_nav_row('36'),
            ft.Row([ft.Button('DESCARGAR PDF',icon=ft.Icons.DOWNLOAD_OUTLINED,on_click=download_pdf)], spacing=10),
            card(ft.Column([
                section('3.1 EVALUACIÓN DE REMANENTE',activities_31),
                ft.Divider(height=18,color='#DDE5ED'),
                section('3.2 DIFERENCIA DE RTD ENTRE HOMBROS',activities_32),
                ft.Divider(height=18,color='#DDE5ED'),
                section('3.3 DIFERENCIA DE RTD DEL MISMO EJE',activities_33),
                ft.Divider(height=18,color='#DDE5ED'),
                section('3.4 DIFERENCIA ENTRE EJES POR EQUIPO',activities_34),
                ft.Divider(height=18,color='#DDE5ED'),
                section('3.5 NIVELACIÓN DE PRESIÓN',activities_35),
            ],spacing=10),padding=20),
        ],scroll=ft.ScrollMode.AUTO,spacing=16)
        page.update()

    def reports_view():
        op_id = active_operation_id()
        """Módulo 7 · Tablas y reportes.

        Replica la separación funcional observada en NEXA/FLT8000:
        9.1 formato llantas de baja-anual,
        9.2 retiro por equipo (RTEQ),
        9.4 formato reporte general,
        9.5 resumen por equipo,
        9.7 costo acumulado de llantas nuevas/reencauchadas instaladas,
        9.8 costo actual de llantas operativas en equipos y
        9.9 costo actual de llantas de repuesto (stand-by).
        """
        search=ft.TextField(label='Buscar código / marca / medida / equipo',prefix_icon=ft.Icons.SEARCH,width=330)
        body_91=ft.Column(spacing=0)
        body_92=ft.Column(spacing=0)
        body_94=ft.Column(spacing=0)
        body_95=ft.Column(spacing=0)
        body_97=ft.Column(spacing=0)
        body_98=ft.Column(spacing=0)
        body_99=ft.Column(spacing=0)
        total_91=ft.Text('',size=11,weight=ft.FontWeight.BOLD,color=TEXT_MAIN)
        total_92=ft.Text('',size=11,weight=ft.FontWeight.BOLD,color=TEXT_MAIN)
        total_94=ft.Text('',size=11,weight=ft.FontWeight.BOLD,color=TEXT_MAIN)
        total_95=ft.Text('',size=11,weight=ft.FontWeight.BOLD,color=TEXT_MAIN)
        total_97=ft.Text('',size=11,weight=ft.FontWeight.BOLD,color=TEXT_MAIN)
        total_98=ft.Text('',size=11,weight=ft.FontWeight.BOLD,color=TEXT_MAIN)
        total_99=ft.Text('',size=11,weight=ft.FontWeight.BOLD,color=TEXT_MAIN)

        def fnum(v,dec=0):
            if v in (None,''): return '—'
            try:
                x=float(v)
                return f'{x:,.{dec}f}'
            except Exception:
                return str(v)

        def money(v):
            try: return f'US$ {float(v or 0):,.2f}'
            except Exception: return 'US$ 0.00'

        def avg2(a,b):
            vals=[]
            for x in (a,b):
                try:
                    if x is not None: vals.append(float(x))
                except Exception: pass
            return (sum(vals)/len(vals)) if vals else None

        def install_info(tid, latest=False):
            order='DESC' if latest else 'ASC'
            q=query(f"""SELECT o.event_date,o.meter,o.equipment_id,o.position,e.code equipment_code
                         FROM occurrences o LEFT JOIN equipment e ON e.id=o.equipment_id
                         WHERE o.tire_id=? AND o.operation_id=? AND UPPER(TRIM(o.event_code))='INST'
                         ORDER BY o.id {order} LIMIT 1""",(tid,op_id))
            return q[0] if q else None

        def last_dins_info(tid):
            q=query("""SELECT o.event_date,o.meter,o.equipment_id,o.position,e.code equipment_code
                       FROM occurrences o LEFT JOIN equipment e ON e.id=o.equipment_id
                       WHERE o.tire_id=? AND o.operation_id=? AND UPPER(TRIM(o.event_code))='DINS'
                       ORDER BY o.id DESC LIMIT 1""",(tid,op_id))
            return q[0] if q else None

        def current_value(r):
            """Valorización por profundidad utilizable remanente.

            NEXA expone PROMEDIO, U$/mm y U$-Total. En SQLite se obtiene el
            equivalente con profundidad nueva, profundidad de retiro y RTD actual.
            """
            new_avg=avg2(r['new_tread_outer'],r['new_tread_inner'])
            if new_avg is None:
                try: new_avg=float(r['new_tread'])
                except Exception: new_avg=None
            cur_avg=avg2(r['tread_outer'],r['tread_inner'])
            try: retire=float(r['retirement_tread'] or 0)
            except Exception: retire=0.0
            try: cost=float(r['cost_usd'] or 0)
            except Exception: cost=0.0
            usable=(new_avg-retire) if new_avg is not None else 0.0
            rem=max(0.0,(cur_avg-retire)) if cur_avg is not None else 0.0
            usdmm=(cost/usable) if usable>0 else 0.0
            total=rem*usdmm
            return cur_avg,usdmm,total

        def make_table(columns, rows, total_text=None, width=None, highlight_last=False):
            # Estándar visual para reportes 9.x y futuros:
            # encabezado azul, filas alternadas y total general resaltado con
            # línea divisoria roja, fondo azul oscuro y texto blanco en negrita.
            def cell(value,w,header=False,total=False):
                return ft.Container(
                    width=w,
                    padding=ft.Padding(left=6,top=8,right=6,bottom=8),
                    content=ft.Text(
                        str(value if value not in (None,'') else '—'),size=10,
                        weight=ft.FontWeight.BOLD if (header or total) else ft.FontWeight.NORMAL,
                        color='#FFFFFF' if (header or total) else TEXT_MAIN,
                        no_wrap=True
                    )
                )
            hdr=ft.Container(bgcolor='#1F4E78',content=ft.Row([cell(label,w,True) for label,w in columns],spacing=0))
            controls=[hdr]
            for i,row in enumerate(rows):
                is_total=bool(highlight_last and i==len(rows)-1)
                controls.append(ft.Container(
                    bgcolor='#0B4A72' if is_total else ('#FFFFFF' if i%2==0 else '#F8FAFC'),
                    border=ft.Border(
                        top=ft.BorderSide(3,'#D32F2F') if is_total else ft.BorderSide(0,'#00000000'),
                        bottom=ft.BorderSide(1,'#E5E9EF')
                    ),
                    content=ft.Row([cell(row[j],columns[j][1],total=is_total) for j in range(len(columns))],spacing=0)
                ))
            if not rows:
                controls.append(ft.Container(padding=14,content=ft.Text('Sin registros para esta condición.',size=11,color=TEXT_MUTED)))
            if total_text is not None:
                total_text.color='#FFFFFF'
                total_text.weight=ft.FontWeight.BOLD
                controls.append(ft.Container(
                    bgcolor='#0B4A72',
                    border=ft.Border(top=ft.BorderSide(3,'#D32F2F')),
                    padding=ft.Padding(left=10,top=9,right=10,bottom=9),
                    content=total_text
                ))
            table=ft.Column(controls,spacing=0)
            return ft.Row([ft.Container(content=table,width=width or sum(w for _,w in columns))],scroll=ft.ScrollMode.ALWAYS)

        cols91=[('MEDIDA',90),('FECHA RETIRO',100),('SEC',115),('CÓDIGO',75),('MC',75),('MODELO',100),('V.U %',65),('REMA mm',75),('HORAS ACUM.',95),('US$/HR',80),('HsxMM',75),('COSTO',90),('COSTO PERDIDO',105),('OC N°',70),('EST',82),('MOTIVO',120),('EQ-I',75),('P',48),('EQ-F',75),('P',48)]
        cols92=[('CÓDIGO',78),('SERIE',115),('EQUIPO OUT',88),('POS.',55),('RTD NUEVO',82),('RTD RETIRO',82),('RTD ACTUAL',82),('% REM. ÚTIL',90),('HORAS ACUM.',95),('COSTO NEUM. US$',110),('COSTO ACUM. US$',110),('US$/NO UTILIZADO',125),('US$/H',78),('Hs/mm',75),('FECHA RETIRO',100),('MOTIVO',105)]
        cols94=[('EQ',70),('P',45),('COD',68),('MC',75),('MED',85),('MOD',90),('H.T.',65),('$/H',70),('CO',72),('EX',50),('IN',50),('REM',60),('DR',50),('DH',50),('P-Ac',60),('Rc',55),('COND',78),('T',45),('Hs/mm',65),('Proye',70),('FECHA',88),('HORO',70),('OC',58)]
        cols95=[('EQUIP',78),('LL/NEW',92),('LL/REE',92),('$CORTE',92),('$NO OPT',92),('$/HRS',78),('HRS/LL',82),('H/D',60),('Hr-Rod',82),('LL-UT',68),('MM$REE',78),('MM$BAJA',82),('$TOTAL',92),('DGT',55),('REP',55),('INV',55),('CTB',55),('CTL',55),('PSB',55),('XRE',55),('PRE',55),('SEP',55),('USA',55)]
        cols97=[('FECHA',95),('CÓDIGO',78),('MARCA',105),('MEDIDA',95),('DISEÑO',105),('EQUIPO',85),('POS.',58),('RTD EXT/INT',100),('COSTO US$',100),('PROVEEDOR',105),('CONDICIÓN',105)]
        cols98=[('EQUIPO',82),('POS.',55),('CÓDIGO',75),('MARCA',95),('MEDIDA',88),('DISEÑO',95),('H.T.',70),('$/H',70),('COND.',82),('EXT',55),('INT',55),('PROM.',65),('U$ COSTO',90),('U$/mm',75),('U$ TOTAL',95),('FECHA',95)]
        cols99=[('MEDIDA',90),('F. RETIRO',95),('ESTADO',90),('CÓDIGO',75),('MARCA',95),('DISEÑO',95),('COND.',82),('EXT/INT',82),('PROM.',65),('U$ COSTO',90),('U$/MM',75),('U$ TOTAL',95),('EQ-OUT',85)]

        def refresh(e=None):
            term=(search.value or '').strip().upper()

            # 9.1: FORMATO LLANTAS DE BAJA-ANUAL (NEXA FLT2015.FXP).
            # Historial de neumáticos dados de baja. Conserva la estructura del formato original.
            bajas=query("""SELECT t.* FROM tires t WHERE UPPER(TRIM(t.status))='BAJA' AND t.operation_id=? ORDER BY t.size,t.code""",(op_id,))
            rows91=[]; total91_cost=0.0; total91_lost=0.0
            for r in bajas:
                bq=query("""SELECT o.*,e.code equipment_code FROM occurrences o
                            LEFT JOIN equipment e ON e.id=o.equipment_id
                            WHERE o.tire_id=? AND o.operation_id=? AND UPPER(TRIM(o.event_code))='BAJA'
                            ORDER BY o.id DESC LIMIT 1""",(r['id'], op_id))
                baja=bq[0] if bq else None
                iq=query("""SELECT o.*,e.code equipment_code FROM occurrences o
                            LEFT JOIN equipment e ON e.id=o.equipment_id
                            WHERE o.tire_id=? AND o.operation_id=? AND UPPER(TRIM(o.event_code))='INST'
                            ORDER BY o.id ASC LIMIT 1""",(r['id'], op_id))
                ini=iq[0] if iq else None
                eqf=(baja['equipment_code'] if baja else '') or ''
                hay=' '.join(str(x or '') for x in (r['code'],r['serial'],r['brand'],r['size'],r['design'],eqf)).upper()
                if term and term not in hay: continue
                ext=(baja['tread_outer'] if baja and baja['tread_outer'] is not None else r['tread_outer'])
                inn=(baja['tread_inner'] if baja and baja['tread_inner'] is not None else r['tread_inner'])
                remavg=avg2(ext,inn)
                newavg=avg2(r['new_tread_outer'],r['new_tread_inner'])
                if newavg is None:
                    try: newavg=float(r['new_tread'])
                    except Exception: newavg=None
                vu=(remavg/newavg*100.0) if remavg is not None and newavg and newavg>0 else None
                events=query("""SELECT event_code,meter FROM occurrences WHERE tire_id=? AND operation_id=?
                                AND UPPER(TRIM(event_code)) IN ('INST','DINS','BAJA') ORDER BY id""",(r['id'], op_id))
                hrs=0.0; st=None
                for ev in events:
                    try: mv=float(ev['meter']) if ev['meter'] is not None else None
                    except Exception: mv=None
                    ec=str(ev['event_code'] or '').upper().strip()
                    if ec=='INST' and mv is not None: st=mv
                    elif ec in ('DINS','BAJA') and st is not None and mv is not None:
                        if mv>=st: hrs+=mv-st
                        st=None
                if hrs<=0:
                    try:
                        im=float(r['installation_meter']) if r['installation_meter'] is not None else None
                        bm=float(baja['meter']) if baja and baja['meter'] is not None else None
                        hrs=(bm-im) if im is not None and bm is not None and bm>=im else 0.0
                    except Exception: hrs=0.0
                cost=float(r['cost_usd'] or 0); total91_cost+=cost
                # Costo perdido estimado = parte del costo de compra consumida
                # según el porcentaje de vida útil remanente mostrado en V.U.%.
                costo_perdido=max(0.0, cost*(1.0-(vu/100.0))) if vu is not None else 0.0
                total91_lost+=costo_perdido
                cph=(cost/hrs) if hrs>0 else 0.0
                wear=(newavg-remavg) if newavg is not None and remavg is not None else 0.0
                hsmm=(hrs/wear) if wear and wear>0 else 0.0
                reason=(baja['reason'] if baja else '') or ''
                cond='REENC.' if 'REENC' in str(r['tire_condition'] or '').upper() else 'NUEVA'
                rows91.append([r['size'],format_date(baja['event_date']) if baja else '—',r['serial'],r['code'],r['brand'],r['design'],fnum(vu,1),fnum(remavg,1),fnum(hrs,0),f"${cph:.2f}",fnum(hsmm,1),money(cost),money(costo_perdido),'BAJA',cond,reason,(ini['equipment_code'] if ini else '—'),(ini['position'] if ini else '—'),eqf or '—',(baja['position'] if baja else '—')])
            total_91.value=f"TOTAL GENERAL   |   LLANTAS DE BAJA: {len(rows91)}   |   COSTO ACUMULADO: {money(total91_cost)}   |   COSTO PERDIDO: {money(total91_lost)}"
            body_91.controls=[make_table(cols91,rows91,total_91)]

            # 9.2: RETIRO POR EQUIPO (NEXA FLT2080.FXP / RTEQ).
            # Se incluyen solo BAJAS cuyo motivo identifica retiro del equipo hacia otra operación.
            # Costo acumulado = costo del neumático x porcentaje de profundidad útil consumida.
            # US$/NO UTILIZADO = costo del neumático - costo acumulado.
            rteq_rows=[]
            rteq_tot_cost=0.0; rteq_tot_acc=0.0; rteq_tot_unused=0.0; rteq_tot_hours=0.0
            all_bajas=query("""SELECT t.* FROM tires t WHERE UPPER(TRIM(t.status))='BAJA' AND t.operation_id=? ORDER BY t.code""",(op_id,))
            for r in all_bajas:
                bq=query("""SELECT o.*,e.code equipment_code FROM occurrences o
                            LEFT JOIN equipment e ON e.id=o.equipment_id
                            WHERE o.tire_id=? AND o.operation_id=? AND UPPER(TRIM(o.event_code))='BAJA'
                            ORDER BY o.id DESC LIMIT 1""",(r['id'], op_id))
                if not bq: continue
                baja=bq[0]
                motivo=str(baja['reason'] or '').strip()
                motup=motivo.upper()
                if not (motup=='RTEQ' or 'RETIRO' in motup and 'EQUIP' in motup):
                    continue
                eqout=baja['equipment_code'] or ''
                hay=' '.join(str(x or '') for x in (r['code'],r['serial'],r['brand'],r['size'],r['design'],eqout,motivo)).upper()
                if term and term not in hay: continue

                newavg=avg2(r['new_tread_outer'],r['new_tread_inner'])
                if newavg is None:
                    try: newavg=float(r['new_tread'])
                    except Exception: newavg=None
                actual=avg2(baja['tread_outer'],baja['tread_inner'])
                if actual is None: actual=avg2(r['tread_outer'],r['tread_inner'])
                try: retiro=float(r['retirement_tread'] or 0)
                except Exception: retiro=0.0
                try: costo=float(r['cost_usd'] or 0)
                except Exception: costo=0.0

                util_original=max(0.0,(newavg or 0)-retiro)
                util_actual=max(0.0,(actual or 0)-retiro)
                if util_original>0:
                    util_actual=min(util_actual,util_original)
                    pct_rem=util_actual/util_original*100.0
                    pct_consumido=1.0-(util_actual/util_original)
                else:
                    pct_rem=0.0; pct_consumido=0.0
                costo_acum=costo*pct_consumido
                no_utilizado=max(0.0,costo-costo_acum)

                # Horas reales acumuladas entre INST y DINS/BAJA.
                evs=query("""SELECT event_code,meter FROM occurrences WHERE tire_id=? AND operation_id=?
                              AND UPPER(TRIM(event_code)) IN ('INST','DINS','BAJA') ORDER BY id""",(r['id'], op_id))
                hrs=0.0; st=None
                for ev in evs:
                    ec=str(ev['event_code'] or '').upper().strip()
                    try: mv=float(ev['meter']) if ev['meter'] is not None else None
                    except Exception: mv=None
                    if ec=='INST' and mv is not None: st=mv
                    elif ec in ('DINS','BAJA') and st is not None and mv is not None:
                        if mv>=st: hrs+=mv-st
                        st=None
                if hrs<=0:
                    try:
                        im=float(r['installation_meter']) if r['installation_meter'] is not None else None
                        bm=float(baja['meter']) if baja['meter'] is not None else None
                        if im is not None and bm is not None and bm>=im: hrs=bm-im
                    except Exception: pass
                usd_h=(costo_acum/hrs) if hrs>0 else 0.0
                mm_consumidos=max(0.0,(newavg or 0)-(actual or 0))
                hs_mm=(hrs/mm_consumidos) if mm_consumidos>0 else 0.0

                rteq_rows.append([r['code'],r['serial'],eqout or '—',baja['position'],fnum(newavg,1),fnum(retiro,1),fnum(actual,1),f'{pct_rem:.1f}%',fnum(hrs,0),money(costo),money(costo_acum),money(no_utilizado),f'${usd_h:.2f}',fnum(hs_mm,1),format_date(baja['event_date']),motivo or 'RTEQ'])
                rteq_tot_cost+=costo; rteq_tot_acc+=costo_acum; rteq_tot_unused+=no_utilizado; rteq_tot_hours+=hrs
            if rteq_rows:
                rteq_rows.append(['TOTAL GENERAL','—','—','—','—','—','—','—',fnum(rteq_tot_hours,0),money(rteq_tot_cost),money(rteq_tot_acc),money(rteq_tot_unused),'—','—','—','—'])
            total_92.value=''
            body_92.controls=[make_table(cols92,rteq_rows,highlight_last=bool(rteq_rows))]

            # 9.4: FORMATO REPORTE GENERAL (NEXA DEMO13.FXP).
            # Equivalente a RELACION DE NEUMATICOS EN USO: solo neumáticos actualmente en servicio.
            active94=query("""SELECT t.*,e.code equipment_code FROM tires t LEFT JOIN equipment e ON e.id=t.equipment_id
                              WHERE UPPER(TRIM(t.status))='SERVICIO' AND t.equipment_id IS NOT NULL AND t.operation_id=?
                              ORDER BY e.code,t.position,t.code""",(op_id,))
            rows94=[]
            for r in active94:
                hay=' '.join(str(x or '') for x in (r['code'],r['brand'],r['size'],r['design'],r['equipment_code'])).upper()
                if term and term not in hay: continue
                inst=install_info(r['id'],latest=True)
                lq=query("""SELECT event_date,event_code,meter,tread_outer,tread_inner,pressure,notes
                            FROM occurrences WHERE tire_id=? AND operation_id=? ORDER BY id DESC LIMIT 1""",(r['id'], op_id))
                last=lq[0] if lq else None
                try:
                    cm=float(last['meter']) if last and last['meter'] is not None else (float(r['current_meter']) if r['current_meter'] is not None else None)
                    im=float(inst['meter']) if inst and inst['meter'] is not None else None
                    ht=max(0.0,cm-im) if cm is not None and im is not None else 0.0
                except Exception: ht=0.0; cm=None
                cost=float(r['cost_usd'] or 0); cph=(cost/ht) if ht>0 else 0.0
                ext=(last['tread_outer'] if last and last['tread_outer'] is not None else r['tread_outer'])
                inn=(last['tread_inner'] if last and last['tread_inner'] is not None else r['tread_inner'])
                curavg=avg2(ext,inn); newavg=avg2(r['new_tread_outer'],r['new_tread_inner'])
                if newavg is None:
                    try: newavg=float(r['new_tread'])
                    except Exception: newavg=None
                rem=(curavg/newavg*100.0) if curavg is not None and newavg and newavg>0 else None
                wear=(newavg-curavg) if newavg is not None and curavg is not None else 0.0
                hsmm=(ht/wear) if wear and wear>0 else 0.0
                try: proy=float(r['projected_life']) if r['projected_life'] is not None else None
                except Exception: proy=None
                pact=(last['pressure'] if last else None)
                notes=str(last['notes'] or '').upper() if last else ''
                tapa='SI' if ('TAPA' in notes or 'VALVULA' in notes or 'VÁLVULA' in notes) else 'NO'
                cond='REENC.' if 'REENC' in str(r['tire_condition'] or '').upper() else 'NUEVA'
                rows94.append([r['equipment_code'],r['position'],r['code'],r['brand'],r['size'],r['design'],fnum(ht,0),f"${cph:.2f}",r['construction_type'] or '—',fnum(ext,0),fnum(inn,0),fnum(rem,1),'—','—',fnum(pact,0),fnum(r['recommended_pressure'],0),cond,tapa,fnum(hsmm,1),fnum(proy,0),format_date(last['event_date']) if last else '—',fnum(cm,0),(last['event_code'] if last else '—')])
            total_94.value=f"TOTAL GENERAL   |   NEUMÁTICOS OPERATIVOS: {len(rows94)}"
            body_94.controls=[make_table(cols94,rows94,total_94)]

            # 9.5: RESUMEN POR EQUIPO (NEXA DEMO041.FXP / HISTORIAL POR EQUIPO).
            # H/D queda sin cálculo porque el FXP compilado no permite demostrar su fórmula.
            def _n(v):
                try: return float(v) if v is not None else None
                except Exception: return None

            def _tread_min(row):
                vals=[]
                for k in ('tread_outer','tread_inner'):
                    try: v=_n(row[k])
                    except Exception: v=None
                    if v is not None and v>=0: vals.append(v)
                return min(vals) if vals else None

            def _orig_min(row):
                vals=[]
                for k in ('new_tread_outer','new_tread_inner','new_tread'):
                    try: v=_n(row[k])
                    except Exception: v=None
                    if v is not None and v>0: vals.append(v)
                return min(vals) if vals else None

            def _is_ree(v):
                txt=str(v or '').strip().upper()
                return ('REENCAUCH' in txt) or ('REENC' in txt)

            def _is_cut(v):
                txt=str(v or '').strip().upper()
                return ('CORTE' in txt) or txt in ('CTL','CTB','CPB') or txt.startswith('CT')

            rows95=[]; grand95=0.0
            totals95={
                'll_new':0.0,'ll_ree':0.0,'cut_loss':0.0,'no_opt':0.0,
                'hr_rod':0.0,'ll_ut':0,'mm_ree':0.0,'mm_baja':0.0,'total':0.0,
                'DGT':0,'REP':0,'INV':0,'CTB':0,'CTL':0,'PSB':0,'XRE':0,'PRE':0,'SEP':0,'USA':0,
            }
            eqs=query("SELECT id,code FROM equipment WHERE operation_id=? ORDER BY code",(op_id,))
            for eq in eqs:
                occs=query("""SELECT o.id,o.tire_id,o.event_code,o.event_date,o.meter,o.tread_outer,o.tread_inner,o.reason,
                                      t.cost_usd,t.new_tread,t.new_tread_outer,t.new_tread_inner,t.retirement_tread,t.tire_condition
                               FROM occurrences o JOIN tires t ON t.id=o.tire_id
                               WHERE o.equipment_id=? AND o.operation_id=? AND t.operation_id=? ORDER BY o.id""",(eq['id'],op_id,op_id))
                if not occs: continue
                if term and term not in str(eq['code'] or '').upper(): continue
                meters=[_n(r['meter']) for r in occs if _n(r['meter']) is not None]
                hr_rod=max(0.0,(max(meters)-min(meters))) if len(meters)>=2 else 0.0
                by_tire={}
                for r in occs: by_tire.setdefault(int(r['tire_id']),[]).append(r)
                ll_new=ll_ree=cut_loss=no_opt=0.0
                mm_ree=mm_baja=0.0
                for tid,events in by_tire.items():
                    first=events[0]; original=_orig_min(first); cost=_n(first['cost_usd'])
                    if not original or not cost or original<=0: continue
                    pxmm=cost/original
                    start=None
                    for ev in events:
                        tm=_tread_min(ev)
                        if ev['event_code']=='INST' and tm is not None: start=tm; break
                    if start is None:
                        for ev in events:
                            tm=_tread_min(ev)
                            if tm is not None: start=tm; break
                    end=None
                    for ev in reversed(events):
                        tm=_tread_min(ev)
                        if tm is not None: end=tm; break
                    used_mm=max(0.0,start-end) if start is not None and end is not None else 0.0
                    used_value=used_mm*pxmm
                    if _is_ree(first['tire_condition']):
                        ll_ree+=used_value; mm_ree+=used_mm
                    else: ll_new+=used_value
                    baja=next((ev for ev in reversed(events) if ev['event_code']=='BAJA'),None)
                    if baja is not None:
                        bt=_tread_min(baja)
                        if bt is None: bt=end
                        if bt is not None:
                            mm_baja+=max(0.0,bt)
                            retirement=_n(first['retirement_tread']) or 0.0
                            if _is_cut(baja['reason']): cut_loss+=max(0.0,bt)*pxmm
                            else: no_opt+=max(0.0,bt-retirement)*pxmm
                ll_ut=len(by_tire)
                cph=((ll_new+ll_ree)/hr_rod) if hr_rod>0 else 0.0
                hrs_ll=(hr_rod/ll_ut) if ll_ut>0 else 0.0
                total=ll_new+ll_ree+cut_loss+no_opt; grand95+=total
                counters={k:0 for k in ('DGT','REP','INV','CTB','CTL','PSB','XRE','PRE','SEP','USA')}
                for ev in occs:
                    code=str(ev['event_code'] or '').upper().strip()
                    reason=str(ev['reason'] or '').upper().strip()
                    if code=='REPA': counters['REP']+=1
                    if code in ('INVE','INV'): counters['INV']+=1
                    token=reason or code
                    for k in counters:
                        if token==k or token.startswith(k): counters[k]+=1
                rows95.append([eq['code'],money(ll_new),money(ll_ree),money(cut_loss),money(no_opt),f"${cph:.2f}",fnum(hrs_ll,0),'—',fnum(hr_rod,0),ll_ut,fnum(mm_ree,1),fnum(mm_baja,1),money(total),
                               counters['DGT'],counters['REP'],counters['INV'],counters['CTB'],counters['CTL'],counters['PSB'],counters['XRE'],counters['PRE'],counters['SEP'],counters['USA']])
                totals95['ll_new']+=ll_new; totals95['ll_ree']+=ll_ree
                totals95['cut_loss']+=cut_loss; totals95['no_opt']+=no_opt
                totals95['hr_rod']+=hr_rod; totals95['ll_ut']+=ll_ut
                totals95['mm_ree']+=mm_ree; totals95['mm_baja']+=mm_baja; totals95['total']+=total
                for k in counters: totals95[k]+=counters[k]
            if rows95:
                total_cph=((totals95['ll_new']+totals95['ll_ree'])/totals95['hr_rod']) if totals95['hr_rod']>0 else 0.0
                total_hrs_ll=(totals95['hr_rod']/totals95['ll_ut']) if totals95['ll_ut']>0 else 0.0
                rows95.append([
                    'TOTAL GENERAL',money(totals95['ll_new']),money(totals95['ll_ree']),money(totals95['cut_loss']),money(totals95['no_opt']),
                    f"${total_cph:.2f}",fnum(total_hrs_ll,0),'—',fnum(totals95['hr_rod'],0),totals95['ll_ut'],fnum(totals95['mm_ree'],1),fnum(totals95['mm_baja'],1),money(totals95['total']),
                    totals95['DGT'],totals95['REP'],totals95['INV'],totals95['CTB'],totals95['CTL'],totals95['PSB'],totals95['XRE'],totals95['PRE'],totals95['SEP'],totals95['USA']
                ])
            total_95.value=''
            body_95.controls=[make_table(cols95,rows95,highlight_last=bool(rows95))]

            # 9.7: una inversión por neumático que haya sido instalado al menos una vez.
            # Se toma la primera INST para no duplicar el costo por reinstalaciones posteriores.
            tires=query("""SELECT t.*,e.code equipment_code FROM tires t
                           LEFT JOIN equipment e ON e.id=t.equipment_id WHERE t.operation_id=? ORDER BY t.code""",(op_id,))
            rows97=[]; sum_new=0.0; sum_ree=0.0
            for r in tires:
                inst=install_info(r['id'],latest=False)
                if not inst: continue
                eq=inst['equipment_code'] or ''
                hay=' '.join(str(x or '') for x in (r['code'],r['brand'],r['size'],r['design'],eq)).upper()
                if term and term not in hay: continue
                cond=str(r['tire_condition'] or 'Nueva').strip()
                is_ree='REENC' in cond.upper()
                cost=float(r['cost_usd'] or 0)
                if is_ree: sum_ree+=cost
                else: sum_new+=cost
                ext=r['new_tread_outer'] if r['new_tread_outer'] is not None else r['new_tread']
                inn=r['new_tread_inner'] if r['new_tread_inner'] is not None else r['new_tread']
                rows97.append([
                    format_date(inst['event_date']),r['code'],r['brand'],r['size'],r['design'],eq,
                    inst['position'],f"{fnum(ext,0)}/{fnum(inn,0)}",money(cost),r['supplier'],('REENC.' if is_ree else 'NUEVA')
                ])
            total_97.value=f"TOTAL GENERAL   |   LLANTAS NUEVAS: {money(sum_new)}   |   REENCAUCHADAS: {money(sum_ree)}   |   VALOR TOTAL: {money(sum_new+sum_ree)}"
            body_97.controls=[make_table(cols97,rows97,total_97)]

            # 9.8: equivalente a NEXA c_est_tire=0 y c_sit_tire=1 -> operativas en equipos.
            active=query("""SELECT t.*,e.code equipment_code FROM tires t
                            LEFT JOIN equipment e ON e.id=t.equipment_id
                            WHERE UPPER(TRIM(t.status))='SERVICIO' AND t.equipment_id IS NOT NULL AND t.operation_id=?
                            ORDER BY e.code,t.position,t.code""",(op_id,))
            rows98=[]; total98=0.0
            for r in active:
                hay=' '.join(str(x or '') for x in (r['code'],r['brand'],r['size'],r['design'],r['equipment_code'])).upper()
                if term and term not in hay: continue
                inst=install_info(r['id'],latest=True)
                try:
                    current=float(r['current_meter']) if r['current_meter'] is not None else None
                    im=float(inst['meter']) if inst and inst['meter'] is not None else None
                    ht=max(0.0,current-im) if current is not None and im is not None else 0.0
                except Exception: ht=0.0
                cost=float(r['cost_usd'] or 0)
                cph=(cost/ht) if ht>0 else 0.0
                prom,usdmm,val=current_value(r); total98+=val
                cond='REENC.' if 'REENC' in str(r['tire_condition'] or '').upper() else 'NUEVA'
                rows98.append([
                    r['equipment_code'],r['position'],r['code'],r['brand'],r['size'],r['design'],
                    fnum(ht,0),f"${cph:.2f}",cond,fnum(r['tread_outer'],0),fnum(r['tread_inner'],0),fnum(prom,1),
                    money(cost),f"${usdmm:.2f}",money(val),format_date(inst['event_date']) if inst else '—'
                ])
            total_98.value=f"TOTAL GENERAL   |   VALOR DE LLANTAS OPERATIVAS: {money(total98)}"
            body_98.controls=[make_table(cols98,rows98,total_98)]

            # 9.9: NEXA FLTRET08 = NEUMÁTICOS EN STAND BY / repuestos.
            standby=query("""SELECT t.* FROM tires t WHERE UPPER(TRIM(t.status)) IN ('STAND-BY','STAND BY','STANDBY') AND t.operation_id=? ORDER BY t.size,t.code""",(op_id,))
            rows99=[]; total99=0.0
            for r in standby:
                dins=last_dins_info(r['id'])
                eqout=dins['equipment_code'] if dins else ''
                hay=' '.join(str(x or '') for x in (r['code'],r['brand'],r['size'],r['design'],eqout)).upper()
                if term and term not in hay: continue
                prom,usdmm,val=current_value(r); total99+=val
                cond='REENC.' if 'REENC' in str(r['tire_condition'] or '').upper() else 'NUEVA'
                rows99.append([
                    r['size'],format_date(dins['event_date']) if dins else '—','STAND-BY',r['code'],r['brand'],r['design'],cond,
                    f"{fnum(r['tread_outer'],0)}/{fnum(r['tread_inner'],0)}",fnum(prom,1),money(r['cost_usd']),f"${usdmm:.2f}",money(val),eqout or '—'
                ])
            total_99.value=f"TOTAL GENERAL   |   VALOR DE LLANTAS DE REPUESTO: {money(total99)}"
            body_99.controls=[make_table(cols99,rows99,total_99)]
            page.update()

        search.on_change=refresh
        refresh()

        # Presentación tipo panel: los tres reportes se muestran como accesos visuales
        # y solamente se abre la tabla seleccionada. La lógica NEXA de 9.7/9.8/9.9
        # permanece intacta; este bloque modifica únicamente la navegación/presentación.
        report_detail=ft.Column(spacing=12)
        selected_report={'id':None}

        report_meta={
            '91':{
                'num':'7.1','icon':ft.Icons.DELETE_SWEEP_OUTLINED,'accent':'#B71C1C','soft':'#FFF1F1',
                'title':'FORMATO LLANTAS DE BAJA - ANUAL',
                'desc':'Historial anual de neumáticos dados de baja, rendimiento, remanente, costo y motivo de retiro.',
                'detail':'7.1 FORMATO LLANTAS DE BAJA - ANUAL','body':body_91,
            },
            '92':{
                'num':'7.2','icon':ft.Icons.EXIT_TO_APP_OUTLINED,'accent':'#C62828','soft':'#FFF4F4',
                'title':'RETIRO POR EQUIPO',
                'desc':'Neumáticos retirados por salida del equipo a otra operación (RTEQ), con costo utilizado y remanente económico no utilizado.',
                'detail':'7.2 REPORT - RETIRO POR EQUIPO','body':body_92,
            },
            '94':{
                'num':'7.3','icon':ft.Icons.ASSESSMENT_OUTLINED,'accent':'#00796B','soft':'#ECF8F6',
                'title':'FORMATO REPORTE GENERAL',
                'desc':'Relación técnica general de los neumáticos actualmente instalados y en uso.',
                'detail':'7.3 FORMATO REPORTE GENERAL','body':body_94,
            },
            '95':{
                'num':'7.4','icon':ft.Icons.SUMMARIZE_OUTLINED,'accent':'#6A1B9A','soft':'#F7F0FB',
                'title':'RESUMEN POR EQUIPO',
                'desc':'Historial consolidado por equipo: utilización, costos, pérdidas, rendimiento y eventos.',
                'detail':'7.4 RESUMEN POR EQUIPO','body':body_95,
            },
            '97':{
                'num':'7.5','icon':ft.Icons.RECEIPT_LONG_OUTLINED,'accent':'#1565C0','soft':'#EEF6FF',
                'title':'COSTO ACUMULADO\nLL/NUEVAS Y REENCAUCHADAS INSTALADAS',
                'desc':'Inversión histórica de neumáticos nuevos y reencauchados que registran instalación.',
                'detail':'7.5 COSTO ACUMUL. LL/NUEVAS Y REENC. INSTALADAS','body':body_97,
            },
            '98':{
                'num':'7.6','icon':ft.Icons.PRECISION_MANUFACTURING_OUTLINED,'accent':'#138A3D','soft':'#EEFAF2',
                'title':'COSTO ACTUAL\nLL/OPERATIVAS EN EQUIPOS',
                'desc':'Valorización económica actual de los neumáticos que se encuentran en servicio.',
                'detail':'7.6 COSTO ACTUAL LL/OPERATIVAS EN EQUIPOS','body':body_98,
            },
            '99':{
                'num':'7.7','icon':ft.Icons.INVENTORY_2_OUTLINED,'accent':'#EF6C00','soft':'#FFF5EA',
                'title':'COSTO ACTUAL\nLL/DE REPUESTO',
                'desc':'Valorización económica de los neumáticos disponibles actualmente en Stand-by.',
                'detail':'7.7 COSTO ACTUAL LL/DE REPUESTOS','body':body_99,
            },
        }

        def show_report(key):
            selected_report['id']=key
            m=report_meta[key]
            report_detail.controls=[
                ft.Row([
                    ft.OutlinedButton('VOLVER A REPORTES',icon=ft.Icons.ARROW_BACK,on_click=lambda e: close_report()),
                    ft.Container(expand=True),
                    ft.Container(
                        bgcolor=m['accent'],border_radius=8,padding=ft.Padding(10,5,10,5),
                        content=ft.Text(m['num'],color='#FFFFFF',weight=ft.FontWeight.BOLD,size=13)
                    ),
                ]),
                card(ft.Column([
                    ft.Row([
                        ft.Container(width=46,height=46,border_radius=12,bgcolor=m['soft'],alignment=ft.Alignment.CENTER,
                                     content=ft.Icon(m['icon'],color=m['accent'],size=26)),
                        ft.Column([
                            ft.Text(m['detail'],size=17,weight=ft.FontWeight.BOLD,color=TEXT_MAIN),
                            ft.Text(m['desc'],size=10.5,color=TEXT_MUTED),
                        ],spacing=2,expand=True),
                    ],spacing=12),
                    search,
                    m['body'],
                ],spacing=12),padding=16),
            ]
            report_cards.visible=False
            report_detail.visible=True
            page.update()

        def close_report():
            selected_report['id']=None
            report_detail.visible=False
            report_cards.visible=True
            page.update()

        def access_card(key):
            m=report_meta[key]
            return ft.Container(
                width=292,
                height=260,
                bgcolor=m['soft'],
                border=ft.Border.all(1,m['accent']+'55'),
                border_radius=14,
                padding=18,
                on_click=lambda e,k=key: show_report(k),
                ink=True,
                content=ft.Column([
                    ft.Row([
                        ft.Container(
                            bgcolor=m['accent'],border_radius=9,padding=ft.Padding(11,6,11,6),
                            content=ft.Text(m['num'],size=16,weight=ft.FontWeight.BOLD,color='#FFFFFF')
                        ),
                        ft.Container(expand=True),
                        ft.Container(width=54,height=54,border_radius=14,bgcolor='#FFFFFF',alignment=ft.Alignment.CENTER,
                                     content=ft.Icon(m['icon'],size=31,color=m['accent'])),
                    ]),
                    ft.Text(m['title'],size=16,weight=ft.FontWeight.BOLD,color=TEXT_MAIN,text_align=ft.TextAlign.CENTER),
                    ft.Text(m['desc'],size=11,color=TEXT_MUTED,text_align=ft.TextAlign.CENTER),
                    ft.Container(height=4),
                    ft.Container(
                        bgcolor=m['accent'],border_radius=9,padding=10,alignment=ft.Alignment.CENTER,
                        content=ft.Row([
                            ft.Icon(ft.Icons.BAR_CHART,color='#FFFFFF',size=20),
                            ft.Text('VER REPORTE',color='#FFFFFF',weight=ft.FontWeight.BOLD,size=13),
                            ft.Icon(ft.Icons.CHEVRON_RIGHT,color='#FFFFFF',size=20),
                        ],alignment=ft.MainAxisAlignment.CENTER,spacing=8)
                    ),
                ],spacing=12,horizontal_alignment=ft.CrossAxisAlignment.CENTER)
            )

        report_cards=ft.Column([
            card(ft.Column([
                ft.Row([
                    ft.Container(width=48,height=48,border_radius=24,bgcolor='#EAF2FF',alignment=ft.Alignment.CENTER,
                                 content=ft.Icon(ft.Icons.PAID_OUTLINED,color=NAV_ACCENT,size=27)),
                    ft.Column([
                        ft.Text('COSTOS ECONÓMICOS DE NEUMÁTICOS',size=18,weight=ft.FontWeight.BOLD,color=TEXT_MAIN),
                        ft.Text('Seleccione un reporte para visualizar el detalle.',size=11,color=TEXT_MUTED),
                    ],spacing=2),
                ],spacing=12),
                ft.Row([access_card('91'),access_card('92'),access_card('94'),access_card('95')],spacing=12),
                ft.Row([access_card('97'),access_card('98'),access_card('99'),ft.Container(width=292,height=260)],spacing=12),
            ],spacing=16),padding=16),
        ],spacing=12)

        report_detail.visible=False
        content.content=ft.Column([
            page_title('7. TABLAS Y REPORTES','Información técnica y económica para una mejor toma de decisiones'),
            report_cards,
            report_detail,
        ],scroll=ft.ScrollMode.AUTO,spacing=16)
        page.update()

    def standby_view():
        """4. Retén / Stand-by - menú visual con el mismo estándar de los módulos 3, 6 y 9."""
        def access_card():
            accent='#1565C0'; soft='#EEF5FF'
            return ft.Container(
                width=292,height=260,bgcolor=soft,
                border=ft.Border.all(1,accent+'55'),border_radius=14,padding=18,
                on_click=lambda e: standby_detail_view(),ink=True,
                content=ft.Column([
                    ft.Row([
                        ft.Container(bgcolor=accent,border_radius=9,padding=ft.Padding(11,6,11,6),
                                     content=ft.Text('4.1',size=16,weight=ft.FontWeight.BOLD,color='#FFFFFF')),
                        ft.Container(expand=True),
                        ft.Container(width=54,height=54,border_radius=14,bgcolor='#FFFFFF',alignment=ft.Alignment.CENTER,
                                     content=ft.Icon(ft.Icons.INVENTORY_2_OUTLINED,size=31,color=accent)),
                    ]),
                    ft.Text('NEUMÁTICOS EN\nRETÉN / STAND-BY',size=16,weight=ft.FontWeight.BOLD,color=TEXT_MAIN,text_align=ft.TextAlign.CENTER),
                    ft.Text('Control de neumáticos desmontados disponibles y sus acciones operativas.',size=11,color=TEXT_MUTED,text_align=ft.TextAlign.CENTER),
                    ft.Container(height=4),
                    ft.Container(bgcolor=accent,border_radius=9,padding=10,alignment=ft.Alignment.CENTER,
                                 content=ft.Row([
                                     ft.Icon(ft.Icons.BAR_CHART,color='#FFFFFF',size=20),
                                     ft.Text('VER REPORTE',color='#FFFFFF',weight=ft.FontWeight.BOLD,size=13),
                                     ft.Icon(ft.Icons.CHEVRON_RIGHT,color='#FFFFFF',size=20),
                                 ],alignment=ft.MainAxisAlignment.CENTER,spacing=8)),
                ],spacing=12,horizontal_alignment=ft.CrossAxisAlignment.CENTER)
            )
        content.content=ft.Column([
            page_title('4. NEUMÁTICOS EN STAND BY','Control operativo de neumáticos disponibles'),
            card(ft.Column([
                ft.Row([
                    ft.Container(width=48,height=48,border_radius=24,bgcolor='#EAF2FF',alignment=ft.Alignment.CENTER,
                                 content=ft.Icon(ft.Icons.INVENTORY_2_OUTLINED,color=NAV_ACCENT,size=27)),
                    ft.Column([
                        ft.Text('NEUMÁTICOS EN STAND BY',size=18,weight=ft.FontWeight.BOLD,color=TEXT_MAIN),
                        ft.Text('Seleccione el reporte para visualizar el detalle.',size=11,color=TEXT_MUTED),
                    ],spacing=2),
                ],spacing=12),
                ft.Row([access_card(),ft.Container(width=292,height=260),ft.Container(width=292,height=260),ft.Container(width=292,height=260)],spacing=12),
            ],spacing=16),padding=16),
        ],scroll=ft.ScrollMode.AUTO,spacing=16)
        page.update()

    def standby_detail_view():
        """Módulo 4.1 · Retén / Stand-by: control operativo de neumáticos disponibles."""
        search=ft.TextField(label='Buscar código / serie / marca / medida',prefix_icon=ft.Icons.SEARCH,width=330)
        rows_box=ft.Column(spacing=0)
        summary=ft.Text('',size=12,color=TEXT_MUTED)

        columns=[
            ('Código',85),('Serie',115),('Marca',105),('Medida',100),('Diseño',105),('TRA',90),
            ('Condición',105),('RTD EXT',75),('RTD INT',75),('RTD mín.',75),('% rem.',75),
            ('Horas acum.',95),('Costo/h',80),('Última fecha',100),('Último evento',90),('Ubicación',120),('Acciones',330)
        ]

        def cell(value,width,header=False):
            return ft.Container(
                content=ft.Text(str(value if value not in (None,'') else '—'),size=10.5,
                                weight=ft.FontWeight.BOLD if header else ft.FontWeight.NORMAL,
                                color=TEXT_MAIN,no_wrap=True),
                width=width,padding=ft.Padding(left=5,top=8,right=5,bottom=8))

        header=ft.Container(
            content=ft.Row([cell(a,b,True) for a,b in columns],spacing=0),
            bgcolor='#EEF2F7',border=ft.Border(bottom=ft.BorderSide(1,'#D5DCE5')))

        def standby_metric(title, value, subtitle, value_color=TEXT_MAIN):
            return ft.Container(
                width=225,
                padding=ft.Padding(left=18,top=14,right=18,bottom=14),
                bgcolor='#FFFFFF',
                border=ft.Border.all(1,'#E0E6EE'),
                border_radius=12,
                content=ft.Column([
                    ft.Text(title,size=10.5,weight=ft.FontWeight.BOLD,color=TEXT_MUTED),
                    ft.Text(str(value),size=25,weight=ft.FontWeight.BOLD,color=value_color),
                    ft.Text(subtitle,size=9.5,color=TEXT_MUTED),
                ],spacing=3)
            )

        def fmt(v,dec=1):
            if v is None: return '—'
            try:
                f=float(v)
                return str(int(f)) if f.is_integer() else f'{f:.{dec}f}'
            except Exception: return str(v)

        def go_action(tid,event_code):
            session['movement_tire_id']=str(tid)
            session['movement_event']=event_code
            movement_view()

        def refresh(e=None):
            # V12 MULTITALLER: el Módulo 4 solo puede leer el taller/operación activa.
            op_id = active_operation_id()
            term=(search.value or '').strip()
            sql="""SELECT t.* FROM tires t
                     WHERE t.status='STAND-BY' AND t.operation_id=?"""
            params=[op_id]
            if term:
                sql += " AND (t.code LIKE ? OR t.serial LIKE ? OR t.brand LIKE ? OR t.size LIKE ? OR t.design LIKE ?)"
                q=f'%{term}%'; params.extend([q,q,q,q,q])
            sql += " ORDER BY t.code"
            tires=query(sql,tuple(params))

            # Carga en bloque la última ocurrencia de los neumáticos visibles para evitar N+1 consultas.
            last_by_tire={}
            tire_ids=[int(r['id']) for r in tires]
            if tire_ids:
                marks=','.join('?' for _ in tire_ids)
                occs=query(f"""
                    SELECT o.event_date,o.event_code,o.meter,o.tread_outer,o.tread_inner,o.location,o.tire_id,o.id
                    FROM occurrences o
                    JOIN (
                        SELECT tire_id, MAX(id) max_id
                        FROM occurrences
                        WHERE operation_id=? AND tire_id IN ({marks})
                        GROUP BY tire_id
                    ) x ON x.max_id=o.id
                    WHERE o.operation_id=?
                """, (op_id,*tire_ids,op_id))
                last_by_tire={int(o['tire_id']):o for o in occs}

            rows_box.controls=[]
            apt=repair=low=0
            for idx,r in enumerate(tires):
                z=last_by_tire.get(int(r['id']))
                ext=(z['tread_outer'] if z and z['tread_outer'] is not None else r['tread_outer'])
                inn=(z['tread_inner'] if z and z['tread_inner'] is not None else r['tread_inner'])
                vals=[float(x) for x in (ext,inn) if x is not None]
                rtd=min(vals) if vals else None
                originals=[float(x) for x in (r['new_tread_outer'],r['new_tread_inner']) if x is not None]
                orig=min(originals) if originals else None
                rem=(rtd/orig*100.0) if rtd is not None and orig and orig>0 else None
                hours=z['meter'] if z and z['meter'] is not None else r['current_meter']
                try:
                    cph=float(r['cost_usd'])/float(hours) if r['cost_usd'] is not None and hours and float(hours)>0 else None
                except Exception: cph=None
                retirement=r['retirement_tread']
                try:
                    if rtd is not None and retirement is not None and rtd <= float(retirement): low+=1
                    else: apt+=1
                except Exception: apt+=1
                if z and z['event_code']=='REPA': repair+=1
                actions=ft.Row([
                    ft.OutlinedButton('INST',on_click=lambda e,tid=r['id']:go_action(tid,'INST')),
                    ft.OutlinedButton('INVE',on_click=lambda e,tid=r['id']:go_action(tid,'INVE')),
                    ft.OutlinedButton('REPA',on_click=lambda e,tid=r['id']:go_action(tid,'REPA')),
                    ft.OutlinedButton('BAJA',on_click=lambda e,tid=r['id']:go_action(tid,'BAJA')),
                ],spacing=4)
                values=[r['code'],r['serial'],r['brand'],r['size'],r['design'],r['compound'],r['tire_condition'],
                        fmt(ext),fmt(inn),fmt(rtd),('—' if rem is None else f'{rem:.1f}%'),fmt(hours,0),
                        ('—' if cph is None else f'{cph:.2f}'),format_date(z['event_date']) if z and z['event_date'] else '—',
                        z['event_code'] if z else '—',z['location'] if z else '—']
                controls=[cell(values[i],columns[i][1]) for i in range(len(values))]
                controls.append(ft.Container(content=actions,width=columns[-1][1],padding=ft.Padding(left=3,top=3,right=3,bottom=3)))
                rows_box.controls.append(ft.Container(content=ft.Row(controls,spacing=0),
                    bgcolor='#FFFFFF' if idx%2==0 else '#F8FAFC',border=ft.Border(bottom=ft.BorderSide(1,'#E5E9EF'))))
            total=len(tires)
            summary.value=f'{total} neumático(s) en Retén / Stand-by'
            kpis.controls=[
                standby_metric('EN STAND-BY',total,'Neumáticos disponibles / retén'),
                standby_metric('APTOS PARA INSTALAR',apt,'Con remanente sobre retiro','#2E9B45'),
                standby_metric('ÚLTIMO EVENTO REPA',repair,'Reparación como último evento','#C47A00'),
                standby_metric('EVALUAR PARA BAJA',low,'RTD en profundidad de retiro','#C81D2A'),
            ]
            page.update()

        search.on_change=refresh
        kpis=ft.Row([],wrap=True,spacing=12,run_spacing=12)
        refresh()
        content.content=ft.Column([
            page_title('4. RETÉN / STAND-BY · 4.1 NEUMÁTICOS EN RETÉN / STAND-BY','Neumáticos desmontados disponibles para instalación, inversión, reparación o baja'),
            ft.Row([ft.OutlinedButton('VOLVER A RETÉN / STAND-BY',icon=ft.Icons.ARROW_BACK,on_click=lambda e: standby_view())]),
            kpis,
            card(ft.Column([
                ft.Row([ft.Text('NEUMÁTICOS EN RETÉN / STAND-BY',size=16,weight=ft.FontWeight.BOLD,color=TEXT_MAIN),ft.Container(expand=True),search]),
                summary,
                ft.Row([ft.Column([header,rows_box],spacing=0)],scroll=ft.ScrollMode.AUTO),
                ft.Text('Acciones habilitadas: INST · INVE · REPA · BAJA. Cada acción abre Movimiento de neumáticos con el neumático y evento seleccionados.',size=10,color=TEXT_MUTED,italic=True),
            ],spacing=10))
        ],scroll=ft.ScrollMode.AUTO,spacing=16)
        page.update()

    def baja_history_view():
        """Módulo 5 · Neumáticos fuera de servicio."""
        search=ft.TextField(
            label='Buscar código / serie / marca / medida / equipo / motivo',
            prefix_icon=ft.Icons.SEARCH,
            width=420
        )
        summary=ft.Text('',size=12,color=TEXT_MUTED)

        table=ft.DataTable(
            heading_row_height=38,
            data_row_min_height=38,
            data_row_max_height=42,
            column_spacing=22,
            columns=[ft.DataColumn(ft.Text(x,size=10.5,weight=ft.FontWeight.BOLD)) for x in [
                'Código','Serie','Marca','Medida','Diseño','TRA','Condición','Equipo proc.','Pos.',
                'Fecha baja','Motivo','RTD EXT','RTD INT','RTD mín.','Horas acum.','Costo/h','Observaciones'
            ]],
            rows=[]
        )

        def baja_metric(title,value,subtitle,value_color=TEXT_MAIN):
            return ft.Container(
                width=225,
                padding=ft.Padding(left=18,top=14,right=18,bottom=14),
                bgcolor='#FFFFFF',
                border=ft.Border.all(1,'#E0E6EE'),
                border_radius=12,
                content=ft.Column([
                    ft.Text(title,size=10.5,weight=ft.FontWeight.BOLD,color=TEXT_MUTED),
                    ft.Text(str(value),size=25,weight=ft.FontWeight.BOLD,color=value_color),
                    ft.Text(subtitle,size=9.5,color=TEXT_MUTED),
                ],spacing=3)
            )

        def fmt(v,dec=1):
            if v is None: return '—'
            try:
                f=float(v)
                return str(int(f)) if f.is_integer() else f'{f:.{dec}f}'
            except Exception:
                return str(v)

        def refresh(e=None):
            term=(search.value or '').strip()
            sql="""
                SELECT t.*
                FROM tires t
                WHERE t.status='BAJA' AND t.operation_id=?
            """
            params=[op_id]
            if term:
                q=f'%{term}%'
                sql += """ AND (
                    t.code LIKE ? OR t.serial LIKE ? OR t.brand LIKE ? OR t.size LIKE ? OR
                    t.design LIKE ? OR EXISTS(
                        SELECT 1 FROM occurrences ox LEFT JOIN equipment ex ON ex.id=ox.equipment_id
                        WHERE ox.tire_id=t.id AND ox.operation_id=? AND ox.event_code='BAJA'
                          AND (COALESCE(ex.code,'') LIKE ? OR COALESCE(ox.reason,'') LIKE ?)
                    )
                )"""
                params=[op_id,q,q,q,q,q,op_id,q,q]
            sql += ' ORDER BY t.code'
            tires=query(sql,tuple(params))
            table.rows=[]

            year_now=(dt.datetime.utcnow()-dt.timedelta(hours=5)).year
            bajas_year=0
            meters=[]
            reason_counts={}

            for r in tires:
                baja=query("""
                    SELECT o.event_date,o.equipment_id,o.position,o.meter,
                           o.tread_outer,o.tread_inner,o.reason,o.notes,e.code equipment_code
                    FROM occurrences o
                    LEFT JOIN equipment e ON e.id=o.equipment_id
                    WHERE o.tire_id=? AND o.operation_id=? AND o.event_code='BAJA'
                    ORDER BY o.id DESC LIMIT 1
                """,(r['id'],op_id))
                z=baja[0] if baja else None

                ext=(z['tread_outer'] if z and z['tread_outer'] is not None else r['tread_outer'])
                inn=(z['tread_inner'] if z and z['tread_inner'] is not None else r['tread_inner'])
                vals=[]
                for x in (ext,inn):
                    try:
                        if x is not None: vals.append(float(x))
                    except Exception:
                        pass
                rtd=min(vals) if vals else None

                meter=(z['meter'] if z and z['meter'] is not None else r['current_meter'])
                try:
                    m=float(meter) if meter is not None else None
                    if m is not None: meters.append(m)
                except Exception:
                    m=None
                try:
                    cph=float(r['cost_usd'])/m if r['cost_usd'] is not None and m and m>0 else None
                except Exception:
                    cph=None

                date_text=format_date(z['event_date']) if z and z['event_date'] else '—'
                if z and z['event_date']:
                    raw=str(z['event_date'])
                    parsed=None
                    for f in ('%Y-%m-%d','%d/%m/%Y'):
                        try:
                            parsed=dt.datetime.strptime(raw[:10],f)
                            break
                        except Exception:
                            pass
                    if parsed and parsed.year==year_now:
                        bajas_year+=1

                reason=(str(z['reason']).strip().upper() if z and z['reason'] else 'SIN MOTIVO')
                reason_counts[reason]=reason_counts.get(reason,0)+1

                values=[
                    r['code'],r['serial'],r['brand'],r['size'],r['design'],r['compound'],r['tire_condition'],
                    z['equipment_code'] if z else '—',z['position'] if z else '—',date_text,reason,
                    fmt(ext),fmt(inn),fmt(rtd),fmt(m,0),('—' if cph is None else f'{cph:.2f}'),
                    z['notes'] if z and z['notes'] else '—'
                ]
                table.rows.append(ft.DataRow(cells=[
                    ft.DataCell(ft.Text(str(v if v not in (None,'') else '—'),size=10.2,no_wrap=True,tooltip=str(v if v not in (None,'') else '—')))
                    for v in values
                ]))

            total=len(tires)
            avg_hours=(sum(meters)/len(meters)) if meters else 0
            top_reason='—'
            if reason_counts:
                top_reason=max(reason_counts.items(),key=lambda kv:(kv[1],kv[0]))[0]

            summary.value=f'{total} neumático(s) registrados en la tabla tires con estado BAJA'
            kpis.controls=[
                baja_metric('TOTAL DADOS DE BAJA',total,'Histórico registrado','#C81D2A'),
                baja_metric('BAJAS DEL AÑO',bajas_year,f'Año {year_now}','#C81D2A'),
                baja_metric('HORAS PROMEDIO AL RETIRO',f'{avg_hours:.0f}' if meters else '—','Promedio de lectura al evento BAJA'),
                baja_metric('MOTIVO MÁS FRECUENTE',top_reason,'Según último evento BAJA','#C47A00'),
            ]
            page.update()

        search.on_change=refresh
        kpis=ft.Row([],wrap=True,spacing=12,run_spacing=12)
        refresh()
        content.content=ft.Column([
            page_title('5. NEUMÁTICOS FUERA DE SERVICIO','Neumáticos retirados definitivamente de operación'),
            kpis,
        ],scroll=ft.ScrollMode.AUTO,spacing=14)
        page.update()

    def nfu_view():
        """5. NFU / Bajas - menú visual con el mismo estándar de los módulos 3, 6 y 9."""
        def access_card():
            accent='#C62828'; soft='#FFF0F0'
            return ft.Container(
                width=292,height=260,bgcolor=soft,
                border=ft.Border.all(1,accent+'55'),border_radius=14,padding=18,
                on_click=lambda e: nfu_detail_view(),ink=True,
                content=ft.Column([
                    ft.Row([
                        ft.Container(bgcolor=accent,border_radius=9,padding=ft.Padding(11,6,11,6),
                                     content=ft.Text('5.1',size=16,weight=ft.FontWeight.BOLD,color='#FFFFFF')),
                        ft.Container(expand=True),
                        ft.Container(width=54,height=54,border_radius=14,bgcolor='#FFFFFF',alignment=ft.Alignment.CENTER,
                                     content=ft.Icon(ft.Icons.DELETE_FOREVER_OUTLINED,size=31,color=accent)),
                    ]),
                    ft.Text('NEUMÁTICOS\nNFU / BAJAS',size=16,weight=ft.FontWeight.BOLD,color=TEXT_MAIN,text_align=ft.TextAlign.CENTER),
                    ft.Text('Control técnico e histórico de neumáticos retirados definitivamente de operación.',size=11,color=TEXT_MUTED,text_align=ft.TextAlign.CENTER),
                    ft.Container(height=4),
                    ft.Container(bgcolor=accent,border_radius=9,padding=10,alignment=ft.Alignment.CENTER,
                                 content=ft.Row([
                                     ft.Icon(ft.Icons.BAR_CHART,color='#FFFFFF',size=20),
                                     ft.Text('VER REPORTE',color='#FFFFFF',weight=ft.FontWeight.BOLD,size=13),
                                     ft.Icon(ft.Icons.CHEVRON_RIGHT,color='#FFFFFF',size=20),
                                 ],alignment=ft.MainAxisAlignment.CENTER,spacing=8)),
                ],spacing=12,horizontal_alignment=ft.CrossAxisAlignment.CENTER)
            )
        content.content=ft.Column([
            page_title('5. NEUMÁTICOS DE BAJA','Control de neumáticos fuera de servicio'),
            card(ft.Column([
                ft.Row([
                    ft.Container(width=48,height=48,border_radius=24,bgcolor='#EAF2FF',alignment=ft.Alignment.CENTER,
                                 content=ft.Icon(ft.Icons.DELETE_FOREVER_OUTLINED,color=NAV_ACCENT,size=27)),
                    ft.Column([
                        ft.Text('NFU / BAJAS',size=18,weight=ft.FontWeight.BOLD,color=TEXT_MAIN),
                        ft.Text('Seleccione el reporte para visualizar el detalle.',size=11,color=TEXT_MUTED),
                    ],spacing=2),
                ],spacing=12),
                ft.Row([access_card(),ft.Container(width=292,height=260),ft.Container(width=292,height=260),ft.Container(width=292,height=260)],spacing=12),
            ],spacing=16),padding=16),
        ],scroll=ft.ScrollMode.AUTO,spacing=16)
        page.update()

    def nfu_detail_view():
        """Módulo 5.1 · NFU / BAJA: cuadro independiente de neumáticos dados de baja."""
        op_id = active_operation_id()
        search=ft.TextField(
            label='Buscar código / serie / marca / medida / equipo',
            prefix_icon=ft.Icons.SEARCH,
            width=390,
            dense=True,
        )
        summary=ft.Text('',size=11,color=TEXT_MUTED)
        rows_box=ft.Column(spacing=0)
        kpis=ft.Row([],wrap=True,spacing=12,run_spacing=12)

        columns=[
            ('Código',78),('Serie',120),('Marca',100),('Medida',105),('Diseño',100),('TRA',80),
            ('Condición',90),('RTD EXT',72),('RTD INT',72),('% REM',76),('Horas recorrido',105),
            ('Costo x hrs',90),('Hs/mm',80),('Fecha de baja',105),('Equipo',82),('Posición',70)
        ]

        def cell(value,width,header=False):
            value='—' if value in (None,'') else value
            return ft.Container(
                width=width,
                padding=ft.Padding(left=5,top=8,right=5,bottom=8),
                content=ft.Text(
                    str(value),size=10,
                    weight=ft.FontWeight.BOLD if header else ft.FontWeight.NORMAL,
                    color=ft.Colors.WHITE if header else TEXT_MAIN,
                    no_wrap=True,tooltip=str(value)
                )
            )

        header=ft.Container(
            bgcolor=NAV_BG,
            border_radius=ft.BorderRadius.only(top_left=8,top_right=8),
            content=ft.Row([cell(a,b,True) for a,b in columns],spacing=0)
        )

        def metric(title,value,subtitle='',value_color=TEXT_MAIN):
            return ft.Container(
                width=225,
                padding=ft.Padding(left=18,top=14,right=18,bottom=14),
                bgcolor='#FFFFFF',
                border=ft.Border.all(1,'#E0E6EE'),
                border_radius=12,
                content=ft.Column([
                    ft.Text(title,size=10.5,weight=ft.FontWeight.BOLD,color=TEXT_MUTED),
                    ft.Text(str(value),size=25,weight=ft.FontWeight.BOLD,color=value_color),
                    ft.Text(subtitle,size=9.5,color=TEXT_MUTED),
                ],spacing=3)
            )

        def num(v):
            try:
                return float(v) if v is not None else None
            except Exception:
                return None

        def fmt_mm(v):
            v=num(v)
            if v is None: return '—'
            return f'{v:.0f}' if float(v).is_integer() else f'{v:.1f}'

        def accumulated_tire_hours(tire_id):
            """Suma horas reales de uso entre INST y DINS/BAJA; no usa el horómetro BAJA como vida total."""
            events=query("""SELECT id,event_code,meter FROM occurrences
                            WHERE tire_id=? AND operation_id=? AND event_code IN ('INST','DINS','BAJA')
                            ORDER BY id""",(tire_id,op_id))
            total=0.0
            start=None
            for ev in events:
                code=(ev['event_code'] or '').upper()
                meter=num(ev['meter'])
                if code=='INST':
                    if meter is not None:
                        start=meter
                elif code in ('DINS','BAJA') and start is not None and meter is not None:
                    diff=meter-start
                    if diff>=0:
                        total += diff
                    start=None
            if total>0: return total
            z=query('SELECT installation_meter FROM tires WHERE id=? AND operation_id=?',(tire_id,op_id)); b=query("SELECT meter FROM occurrences WHERE tire_id=? AND operation_id=? AND event_code='BAJA' ORDER BY id DESC LIMIT 1",(tire_id,op_id))
            im=num(z[0]['installation_meter']) if z else None; bm=num(b[0]['meter']) if b else None
            return (bm-im) if im is not None and bm is not None and bm>=im else None

        def parse_year(raw):
            if not raw: return None
            s=str(raw)
            for f in ('%Y-%m-%d','%d/%m/%Y'):
                try:
                    return dt.datetime.strptime(s[:10],f).year
                except Exception:
                    pass
            return None

        def refresh(e=None):
            term=(search.value or '').strip()
            sql="""SELECT t.* FROM tires t WHERE t.status='BAJA' AND t.operation_id=?"""
            params=[op_id]
            if term:
                q=f'%{term}%'
                sql += """ AND (
                    t.code LIKE ? OR t.serial LIKE ? OR t.brand LIKE ? OR t.size LIKE ? OR t.design LIKE ? OR
                    EXISTS(SELECT 1 FROM occurrences ox LEFT JOIN equipment ex ON ex.id=ox.equipment_id
                           WHERE ox.tire_id=t.id AND ox.operation_id=? AND ox.event_code='BAJA' AND COALESCE(ex.code,'') LIKE ?)
                )"""
                params.extend([q,q,q,q,q,op_id,q])
            sql += ' ORDER BY t.code'
            tires=query(sql,tuple(params))

            rows_box.controls=[]
            year_now=(dt.datetime.utcnow()-dt.timedelta(hours=5)).year
            bajas_year=0
            hour_values=[]
            reason_counts={}

            for idx,r in enumerate(tires):
                last=query("""SELECT o.event_date,o.equipment_id,o.position,o.meter,
                                     o.tread_outer,o.tread_inner,o.reason,e.code equipment_code
                              FROM occurrences o
                              LEFT JOIN equipment e ON e.id=o.equipment_id
                              WHERE o.tire_id=? AND o.operation_id=? AND o.event_code='BAJA'
                              ORDER BY o.id DESC LIMIT 1""",(r['id'],op_id))
                z=last[0] if last else None

                ext=(z['tread_outer'] if z and z['tread_outer'] is not None else r['tread_outer'])
                inn=(z['tread_inner'] if z and z['tread_inner'] is not None else r['tread_inner'])
                ext_n=num(ext); inn_n=num(inn)
                current_vals=[v for v in (ext_n,inn_n) if v is not None]
                current_min=min(current_vals) if current_vals else None

                original_vals=[]
                for key in ('new_tread_outer','new_tread_inner','new_tread'):
                    try:
                        v=num(r[key])
                    except Exception:
                        v=None
                    if v is not None and v>0:
                        original_vals.append(v)
                original_min=min(original_vals) if original_vals else None

                rem=(current_min/original_min*100.0) if current_min is not None and original_min else None
                if rem is not None:
                    rem=max(0.0,min(100.0,rem))

                hours=accumulated_tire_hours(r['id'])
                if hours is not None:
                    hour_values.append(hours)
                cost=num(r['cost_usd']) if 'cost_usd' in r.keys() else None
                cph=(cost/hours) if cost is not None and hours and hours>0 else None
                wear=(original_min-current_min) if original_min is not None and current_min is not None else None
                hsmm=(hours/wear) if hours and wear is not None and wear>0 else None

                baja_date=format_date(z['event_date']) if z and z['event_date'] else '—'
                if z and parse_year(z['event_date'])==year_now:
                    bajas_year+=1
                reason=(str(z['reason']).strip().upper() if z and z['reason'] else 'SIN MOTIVO')
                reason_counts[reason]=reason_counts.get(reason,0)+1

                values=[
                    r['code'],r['serial'],r['brand'],r['size'],r['design'],r['compound'],r['tire_condition'],
                    fmt_mm(ext_n),fmt_mm(inn_n),('—' if rem is None else f'{rem:.1f}%'),
                    ('—' if hours is None else f'{hours:,.0f} h'),
                    ('—' if cph is None else f'${cph:.2f}/h'),
                    ('—' if hsmm is None else f'{hsmm:.1f} h/mm'),
                    baja_date,
                    (z['equipment_code'] if z and z['equipment_code'] else '—'),
                    (z['position'] if z and z['position'] else '—'),
                ]
                row_controls=[cell(values[i],columns[i][1]) for i in range(len(columns))]
                rows_box.controls.append(ft.Container(
                    bgcolor='#FFFFFF' if idx%2==0 else '#F8FAFC',
                    border=ft.Border(bottom=ft.BorderSide(1,'#E5E9EF')),
                    content=ft.Row(row_controls,spacing=0)
                ))

            total=len(tires)
            avg=(sum(hour_values)/len(hour_values)) if hour_values else None
            top_reason='—'
            if reason_counts:
                top_reason=max(reason_counts.items(),key=lambda kv:(kv[1],kv[0]))[0]

            kpis.controls=[
                metric('TOTAL DE BAJAS',total,'Neumáticos NFU registrados','#C81D2A'),
                metric('BAJAS DEL AÑO',bajas_year,f'Año {year_now}','#C81D2A'),
                metric('HORAS PROMEDIO AL RETIRO','—' if avg is None else f'{avg:,.0f} h','Horas reales acumuladas'),
                metric('MOTIVO MÁS FRECUENTE',top_reason,'Según último evento BAJA','#C47A00'),
            ]
            summary.value=f'Total de registros: {total}'
            page.update()

        search.on_change=refresh
        refresh()
        content.content=ft.Column([
            page_title('10. NFU / BAJA','Neumáticos Fuera de Uso · registro histórico de bajas'),
            ft.Row([search,ft.Container(expand=True),summary],vertical_alignment=ft.CrossAxisAlignment.CENTER),
            kpis,
            card(ft.Column([
                ft.Text('LISTADO DE NEUMÁTICOS DADOS DE BAJA',size=16,weight=ft.FontWeight.BOLD,color=TEXT_MAIN),
                ft.Row([ft.Column([header,rows_box],spacing=0)],scroll=ft.ScrollMode.ALWAYS),
            ],spacing=10))
        ],scroll=ft.ScrollMode.AUTO,spacing=14)
        page.update()

    def inventory_consumption_view():
        op_id = active_operation_id()
        """6. Análisis de operación -> 6.1 Utilización y pérdida por equipo.

        Criterio evaluado:
        - Hr-Rod: MAX(horómetro) - MIN(horómetro) registrado por equipo.
        - LL/NEW / LL/REE: valor del caucho efectivamente consumido en cada equipo,
          según mm usados y costo por mm del neumático.
        - $CORTE: valor remanente perdido en bajas cuyo motivo contiene CORTE/CTL/CTB.
        - $NO OPT: valor remanente por encima de la profundidad de retiro en otras bajas.
        """

        def n(v):
            try:
                return float(v) if v is not None else None
            except Exception:
                return None

        def tread_min(row):
            vals=[]
            for key in ('tread_outer','tread_inner'):
                try:
                    v=n(row[key])
                except Exception:
                    v=None
                if v is not None and v >= 0:
                    vals.append(v)
            return min(vals) if vals else None

        def tire_original_min(t):
            vals=[]
            for key in ('new_tread_outer','new_tread_inner','new_tread'):
                if key in t.keys():
                    v=n(t[key])
                    if v is not None and v > 0:
                        vals.append(v)
            return min(vals) if vals else None

        def is_reencauchada(value):
            txt=str(value or '').strip().upper()
            return ('REENCAUCH' in txt) or ('REENC' in txt)

        def is_cut_reason(value):
            txt=str(value or '').strip().upper()
            return ('CORTE' in txt) or txt in ('CTL','CTB','CPB') or txt.startswith('CT')

        eq_rows=query("SELECT id,code,brand,model,vehicle_type,active FROM equipment WHERE operation_id=? ORDER BY code",(op_id,))
        result=[]

        for eq in eq_rows:
            eid=int(eq['id'])
            occs=query("""
                SELECT o.id,o.tire_id,o.event_code,o.event_date,o.meter,
                       o.tread_outer,o.tread_inner,o.reason,
                       t.cost_usd,t.new_tread,t.new_tread_outer,t.new_tread_inner,
                       t.retirement_tread,t.tire_condition
                FROM occurrences o
                JOIN tires t ON t.id=o.tire_id
                WHERE o.equipment_id=? AND o.operation_id=? AND t.operation_id=?
                ORDER BY o.id
            """,(eid,op_id,op_id))

            if not occs:
                continue

            meters=[n(r['meter']) for r in occs if n(r['meter']) is not None]
            hr_rod=(max(meters)-min(meters)) if len(meters)>=2 else 0.0
            if hr_rod < 0:
                hr_rod=0.0

            by_tire={}
            for r in occs:
                by_tire.setdefault(int(r['tire_id']),[]).append(r)

            ll_new=0.0
            ll_ree=0.0
            cut_loss=0.0
            no_opt=0.0

            for tid,events in by_tire.items():
                first=events[0]
                original=tire_original_min(first)
                cost=n(first['cost_usd'])
                if not original or not cost or original <= 0:
                    continue
                pxmm=cost/original

                # Inicio: preferir cocada del evento INST en este equipo; si no existe,
                # usar la primera cocada conocida. Fin: última cocada conocida.
                start_tread=None
                for ev in events:
                    tm=tread_min(ev)
                    if ev['event_code']=='INST' and tm is not None:
                        start_tread=tm
                        break
                if start_tread is None:
                    for ev in events:
                        tm=tread_min(ev)
                        if tm is not None:
                            start_tread=tm
                            break

                end_tread=None
                for ev in reversed(events):
                    tm=tread_min(ev)
                    if tm is not None:
                        end_tread=tm
                        break

                used_mm=0.0
                if start_tread is not None and end_tread is not None:
                    used_mm=max(0.0,start_tread-end_tread)
                used_value=used_mm*pxmm

                if is_reencauchada(first['tire_condition']):
                    ll_ree += used_value
                else:
                    ll_new += used_value

                # Pérdidas: solo si existe una BAJA registrada en este equipo.
                baja=None
                for ev in reversed(events):
                    if ev['event_code']=='BAJA':
                        baja=ev
                        break
                if baja is not None:
                    baja_tread=tread_min(baja)
                    if baja_tread is None:
                        baja_tread=end_tread
                    if baja_tread is not None:
                        reason=baja['reason']
                        retirement=n(first['retirement_tread']) or 0.0
                        if is_cut_reason(reason):
                            cut_loss += max(0.0,baja_tread)*pxmm
                        else:
                            no_opt += max(0.0,baja_tread-retirement)*pxmm

            result.append({
                'code':eq['code'],
                'll_new':ll_new,
                'll_ree':ll_ree,
                'cut':cut_loss,
                'no_opt':no_opt,
                'hr_rod':hr_rod,
            })

        def money(v):
            return f"${v:,.2f}"

        columns=[
            ft.DataColumn(ft.Text('EQUIP',weight=ft.FontWeight.BOLD,color=ft.Colors.WHITE)),
            ft.DataColumn(ft.Text('Hr-Rod',weight=ft.FontWeight.BOLD,color=ft.Colors.WHITE),numeric=True),
            ft.DataColumn(ft.Text('LL/NEW',weight=ft.FontWeight.BOLD,color=ft.Colors.WHITE),numeric=True),
            ft.DataColumn(ft.Text('LL/REE',weight=ft.FontWeight.BOLD,color=ft.Colors.WHITE),numeric=True),
            ft.DataColumn(ft.Text('$CORTE',weight=ft.FontWeight.BOLD,color=ft.Colors.WHITE),numeric=True),
            ft.DataColumn(ft.Text('$NO OPT',weight=ft.FontWeight.BOLD,color=ft.Colors.WHITE),numeric=True),
        ]
        rows=[]
        for r in result:
            rows.append(ft.DataRow(cells=[
                ft.DataCell(ft.Text(str(r['code']),weight=ft.FontWeight.BOLD,color=TEXT_MAIN)),
                ft.DataCell(ft.Text(f"{r['hr_rod']:,.0f} h",color=TEXT_MAIN)),
                ft.DataCell(ft.Text(money(r['ll_new']),color=TEXT_MAIN)),
                ft.DataCell(ft.Text(money(r['ll_ree']),color=TEXT_MAIN)),
                ft.DataCell(ft.Text(money(r['cut']),color=TEXT_MAIN)),
                ft.DataCell(ft.Text(money(r['no_opt']),color=TEXT_MAIN)),
            ]))

        if not rows:
            rows=[ft.DataRow(cells=[
                ft.DataCell(ft.Text('Sin movimientos registrados',color=TEXT_MUTED)),
                ft.DataCell(ft.Text('—')),ft.DataCell(ft.Text('—')),
                ft.DataCell(ft.Text('—')),ft.DataCell(ft.Text('—')),ft.DataCell(ft.Text('—')),
            ])]

        table=ft.DataTable(
            columns=columns,
            rows=rows,
            heading_row_color=NAV_BG,
            heading_row_height=44,
            data_row_min_height=42,
            data_row_max_height=46,
            horizontal_margin=16,
            column_spacing=38,
            border=ft.Border.all(1,'#DCE4EC'),
            divider_thickness=1,
        )

        # Gráfico Power BI: utilización y pérdidas monetarias por equipo.
        # Hr-Rod se presenta junto al equipo y no forma parte de la barra monetaria.
        C_NEW='#118DFF'       # Llantas nuevas - azul Power BI
        C_REE='#E66C37'       # Llantas reencauchadas - naranja Power BI
        C_CORTE='#D64545'     # Llantas accidentadas - rojo
        C_NOOPT='#1AAB40'     # Vida útil no optimizada - verde
        chart_width=500
        max_total=max([r['ll_new']+r['ll_ree']+r['cut']+r['no_opt'] for r in result] or [1.0])
        if max_total <= 0:
            max_total=1.0

        # Escala monetaria tipo Power BI: redondear el máximo a un intervalo limpio.
        import math
        axis_step=5000.0 if max_total > 5000 else 1000.0
        axis_max=max(axis_step,math.ceil(max_total/axis_step)*axis_step)

        def legend_item(color,label,total):
            return ft.Row([
                ft.Container(width=11,height=11,bgcolor=color,border_radius=2),
                ft.Text(label,size=11,color=TEXT_MAIN),
                ft.Text(money(total),size=11,weight=ft.FontWeight.BOLD,color=TEXT_MAIN),
            ],spacing=6,tight=True)

        chart_rows=[]
        for r in sorted(result,key=lambda x:(x['ll_new']+x['ll_ree']+x['cut']+x['no_opt']),reverse=True):
            total=r['ll_new']+r['ll_ree']+r['cut']+r['no_opt']
            losses=r['cut']+r['no_opt']
            segs=[]
            for value,color in ((r['ll_new'],C_NEW),(r['ll_ree'],C_REE),(r['cut'],C_CORTE),(r['no_opt'],C_NOOPT)):
                if value > 0:
                    w=max(3,chart_width*(value/axis_max))
                    segs.append(ft.Container(
                        width=w,height=30,bgcolor=color,
                        alignment=ft.Alignment.CENTER,
                        border=ft.Border.all(1,'#FFFFFF'),
                        content=ft.Text(money(value),size=10,weight=ft.FontWeight.BOLD,color=ft.Colors.WHITE)
                                if w >= 78 else None,
                    ))
            if not segs:
                segs=[ft.Container(width=2,height=30,bgcolor='#DCE4EC')]

            # Equipo y horas rodadas en una sola línea horizontal.
            team_label=ft.Row([
                ft.Container(width=8,height=8,bgcolor=C_CORTE,border_radius=4) if losses > 0 else ft.Container(width=8,height=8),
                ft.Text(str(r['code']),size=12,weight=ft.FontWeight.BOLD,color=TEXT_MAIN),
                ft.Text('/',size=11,color=TEXT_MUTED),
                ft.Text(f"{r['hr_rod']:,.0f} h",size=11,color=TEXT_MUTED),
            ],spacing=4,tight=True)

            chart_rows.append(ft.Row([
                ft.Container(width=130,content=team_label),
                ft.Container(width=chart_width,content=ft.Row(segs,spacing=0)),
            ],spacing=10,vertical_alignment=ft.CrossAxisAlignment.CENTER))

        totals={
            'new':sum(r['ll_new'] for r in result),
            'ree':sum(r['ll_ree'] for r in result),
            'cut':sum(r['cut'] for r in result),
            'noopt':sum(r['no_opt'] for r in result),
        }
        chart_panel=ft.Container(
            bgcolor='#FFFFFF',
            border=ft.Border.all(1,'#DCE4EC'),
            border_radius=12,
            padding=18,
            content=ft.Column([
                ft.Row([
                    ft.Column([
                        ft.Text('UTILIZACIÓN Y PÉRDIDA POR EQUIPO',size=16,weight=ft.FontWeight.BOLD,color=TEXT_MAIN),
                        ft.Text('EQUIPO / HRS. TRABAJADAS · valores monetarios',size=11,color=TEXT_MUTED),
                    ],spacing=2),
                    ft.Container(expand=True),
                    ft.Column([
                        legend_item(C_NEW,'LLANTAS NUEVAS',totals['new']),
                        legend_item(C_REE,'LLANTAS REENCAUCHADAS',totals['ree']),
                        legend_item(C_CORTE,'LLANTAS ACCIDENTADAS',totals['cut']),
                        legend_item(C_NOOPT,'VIDA ÚTIL NO OPTIMIZADA',totals['noopt']),
                    ],spacing=4),
                ],vertical_alignment=ft.CrossAxisAlignment.START),
                ft.Divider(height=12,color='#E7EDF3'),
                ft.Column(chart_rows,spacing=8,scroll=ft.ScrollMode.AUTO),
                # Eje monetario visible debajo de las barras.
                ft.Row([
                    ft.Container(width=130),
                    ft.Container(width=chart_width,content=ft.Column([
                        ft.Container(height=1,bgcolor='#C9D4DF'),
                        ft.Row([
                            ft.Text(money(axis_max*i/4),size=9,color=TEXT_MUTED)
                            for i in range(5)
                        ],alignment=ft.MainAxisAlignment.SPACE_BETWEEN),
                    ],spacing=3)),
                ],spacing=10),
                ft.Row([
                    ft.Container(width=8,height=8,bgcolor=C_CORTE,border_radius=4),
                    ft.Text('Equipo con pérdidas ($CORTE o $NO OPT)',size=10,color=TEXT_MUTED),
                ],spacing=6,tight=True),
            ],spacing=10),
        )

        # 6.2 EQUIPOS - COSTO POR HORA
        # Vista actual basada en los mismos datos consolidados de 6.1.
        # El histórico mensual se incorporará cuando existan cierres INSC suficientes.
        cost_rows=[]
        for r in result:
            tire_cost=(r['ll_new'] or 0.0)+(r['ll_ree'] or 0.0)
            hr=r['hr_rod'] or 0.0
            cost_hour=(tire_cost/hr) if hr>0 else 0.0
            cost_rows.append({
                'code':r['code'],'hr_rod':hr,'ll_new':r['ll_new'] or 0.0,
                'll_ree':r['ll_ree'] or 0.0,'cost_total':tire_cost,'cost_hour':cost_hour
            })
        cost_rows=sorted(cost_rows,key=lambda x:x['cost_hour'],reverse=True)

        # Power BI: azul principal para la serie actual.
        C_COST='#118DFF'
        ymax=max([r['cost_hour'] for r in cost_rows] or [1.0])
        if ymax<=0:
            ymax=1.0
        if ymax<=2:
            step=0.25
        elif ymax<=6:
            step=1.0
        else:
            step=2.0
        axis_max=max(step,math.ceil(ymax/step)*step)
        chart_h=260
        bar_w=42

        groups=[]
        for r in cost_rows:
            val=r['cost_hour']
            bh=max(2,chart_h*(val/axis_max)) if val>0 else 2
            groups.append(ft.Column([
                ft.Container(
                    height=chart_h+30,
                    alignment=ft.Alignment.BOTTOM_CENTER,
                    content=ft.Column([
                        ft.Text(f"${val:.2f}",size=10,weight=ft.FontWeight.BOLD,color=TEXT_MAIN),
                        ft.Container(
                            width=bar_w,height=bh,bgcolor=C_COST,
                            border_radius=ft.BorderRadius.only(top_left=4,top_right=4),
                        ),
                    ],spacing=3,horizontal_alignment=ft.CrossAxisAlignment.CENTER,
                      alignment=ft.MainAxisAlignment.END),
                ),
                ft.Container(width=70,alignment=ft.Alignment.CENTER,
                             content=ft.Text(str(r['code']),size=11,weight=ft.FontWeight.BOLD,color=TEXT_MAIN,text_align=ft.TextAlign.CENTER)),
            ],spacing=5,horizontal_alignment=ft.CrossAxisAlignment.CENTER))

        # Líneas/etiquetas del eje Y, conservando el estilo del gráfico de referencia.
        y_labels=[]
        ticks=5
        for i in range(ticks,-1,-1):
            y_labels.append(ft.Text(f"${axis_max*i/ticks:.2f}",size=9,color=TEXT_MUTED))

        cost_chart=ft.Container(
            bgcolor='#FFFFFF',border=ft.Border.all(1,'#DCE4EC'),border_radius=12,padding=18,
            content=ft.Column([
                ft.Row([
                    ft.Column([
                        ft.Text('COSTO US$ / HORA - EQUIPOS',size=17,weight=ft.FontWeight.BOLD,color=TEXT_MAIN),
                        ft.Text('Costo por hora acumulado según los datos actuales de 6.1',size=11,color=TEXT_MUTED),
                    ],spacing=2),
                    ft.Container(expand=True),
                    ft.Row([
                        ft.Container(width=10,height=10,bgcolor=C_COST,border_radius=2),
                        ft.Text('Costo US$/Hr',size=10,color=TEXT_MAIN),
                    ],spacing=5,tight=True),
                ],vertical_alignment=ft.CrossAxisAlignment.START),
                ft.Divider(height=12,color='#E7EDF3'),
                ft.Row([
                    ft.Container(width=58,height=chart_h+30,
                                 content=ft.Column(y_labels,alignment=ft.MainAxisAlignment.SPACE_BETWEEN)),
                    ft.Container(
                        expand=True,
                        border=ft.Border.only(bottom=ft.BorderSide(1,'#AEBAC6')),
                        content=ft.Row(groups,spacing=18,scroll=ft.ScrollMode.AUTO,
                                       vertical_alignment=ft.CrossAxisAlignment.END),
                    ),
                ],vertical_alignment=ft.CrossAxisAlignment.END),
                ft.Container(alignment=ft.Alignment.CENTER,content=ft.Text('EQUIPO',size=11,weight=ft.FontWeight.BOLD,color=TEXT_MUTED)),
                ft.Text('Costo US$/hora = (LL/NEW + LL/REE) / Hr-Rod.',size=10,color=TEXT_MUTED),
            ],spacing=8),
        )

        cost_columns=[
            ft.DataColumn(ft.Text('EQUIPO',weight=ft.FontWeight.BOLD,color=ft.Colors.WHITE)),
            ft.DataColumn(ft.Text('Hr-Rod',weight=ft.FontWeight.BOLD,color=ft.Colors.WHITE),numeric=True),
            ft.DataColumn(ft.Text('LL/NEW',weight=ft.FontWeight.BOLD,color=ft.Colors.WHITE),numeric=True),
            ft.DataColumn(ft.Text('LL/REE',weight=ft.FontWeight.BOLD,color=ft.Colors.WHITE),numeric=True),
            ft.DataColumn(ft.Text('COSTO TOTAL',weight=ft.FontWeight.BOLD,color=ft.Colors.WHITE),numeric=True),
            ft.DataColumn(ft.Text('$/HRS',weight=ft.FontWeight.BOLD,color=ft.Colors.WHITE),numeric=True),
        ]
        cost_table_rows=[]
        for r in cost_rows:
            cost_table_rows.append(ft.DataRow(cells=[
                ft.DataCell(ft.Text(str(r['code']),weight=ft.FontWeight.BOLD,color=TEXT_MAIN)),
                ft.DataCell(ft.Text(f"{r['hr_rod']:,.0f} h",color=TEXT_MAIN)),
                ft.DataCell(ft.Text(money(r['ll_new']),color=TEXT_MAIN)),
                ft.DataCell(ft.Text(money(r['ll_ree']),color=TEXT_MAIN)),
                ft.DataCell(ft.Text(money(r['cost_total']),color=TEXT_MAIN)),
                ft.DataCell(ft.Text(f"${r['cost_hour']:.2f}/h",weight=ft.FontWeight.BOLD,color=TEXT_MAIN)),
            ]))
        if not cost_table_rows:
            cost_table_rows=[ft.DataRow(cells=[ft.DataCell(ft.Text('Sin datos',color=TEXT_MUTED))]+[ft.DataCell(ft.Text('—')) for _ in range(5)])]
        cost_table=ft.DataTable(
            columns=cost_columns,rows=cost_table_rows,heading_row_color=NAV_BG,heading_row_height=44,
            data_row_min_height=40,data_row_max_height=44,horizontal_margin=14,column_spacing=30,
            border=ft.Border.all(1,'#DCE4EC'),divider_thickness=1,
        )

        section_62=ft.Column([
            ft.Column([
                ft.Text('6.2 EQUIPOS COSTO POR HORA',size=17,weight=ft.FontWeight.BOLD,color=TEXT_MAIN),
                ft.Text('Comparativo por equipo con los datos actualmente obtenidos',size=12,color=TEXT_MUTED),
            ],spacing=2),
            cost_chart,
        ],spacing=12)

        # 6.3 · Análisis estadístico/económico de neumáticos dados de BAJA (NFU).
        # Se alimenta exclusivamente de tires.status='BAJA' y del último evento BAJA.
        baja_tires=query("SELECT * FROM tires WHERE status='BAJA' AND operation_id=? ORDER BY code",(op_id,))
        baja_groups={}

        def baja_hours(tire_id):
            evs=query("""SELECT event_code,meter FROM occurrences
                         WHERE tire_id=? AND operation_id=? AND event_code IN ('INST','DINS','BAJA') ORDER BY id""",(tire_id,op_id))
            total=0.0; start=None
            for ev in evs:
                code=str(ev['event_code'] or '').upper(); meter=n(ev['meter'])
                if code=='INST' and meter is not None:
                    start=meter
                elif code in ('DINS','BAJA') and start is not None and meter is not None:
                    if meter>=start: total += meter-start
                    start=None
            if total>0: return total
            z=query('SELECT installation_meter FROM tires WHERE id=? AND operation_id=?',(tire_id,op_id)); b=query("SELECT meter FROM occurrences WHERE tire_id=? AND operation_id=? AND event_code='BAJA' ORDER BY id DESC LIMIT 1",(tire_id,op_id))
            im=n(z[0]['installation_meter']) if z else None; bm=n(b[0]['meter']) if b else None
            return (bm-im) if im is not None and bm is not None and bm>=im else None

        for t in baja_tires:
            last=query("""SELECT o.reason
                          FROM occurrences o
                          WHERE o.tire_id=? AND o.operation_id=? AND o.event_code='BAJA' ORDER BY o.id DESC LIMIT 1""",(t['id'],op_id))
            z=last[0] if last else None
            brand=str(t['brand'] or '—')
            size=str(t['size'] or '—')
            design=str(t['design'] or '—')
            tra=str(t['compound'] or '—')
            # 6.3 se consolida por medida y luego por marca; el equipo no interviene.
            key=(size,brand,design,tra)
            reason=(z['reason'] if z else '')
            cond=t['tire_condition'] if 'tire_condition' in t.keys() else ''
            category='Cortes' if is_cut_reason(reason) else ('Reencauche' if is_reencauchada(cond) else 'Desgaste Regular')
            hrs=baja_hours(t['id']); cost=n(t['cost_usd']) if 'cost_usd' in t.keys() else None
            cph=(cost/hrs) if cost is not None and hrs and hrs>0 else None
            baja_groups.setdefault(key,[]).append({'category':category,'hours':hrs,'cph':cph})

        # Estética tipo reporte comparativo (referencia histórica): encabezado gris,
        # fila de Desgaste Regular resaltada en amarillo y acumulados en verde.
        baja_columns=[
            ft.DataColumn(ft.Text('MEDIDA',weight=ft.FontWeight.BOLD,color=ft.Colors.WHITE,size=12)),
            ft.DataColumn(ft.Text('MARCA',weight=ft.FontWeight.BOLD,color=ft.Colors.WHITE,size=12)),
            ft.DataColumn(ft.Text('DISEÑO / TRA',weight=ft.FontWeight.BOLD,color=ft.Colors.WHITE,size=12)),
            ft.DataColumn(ft.Text('CONCEPTO DE RETIRO',weight=ft.FontWeight.BOLD,color=ft.Colors.WHITE,size=12)),
            ft.DataColumn(ft.Text('# LLANTAS\nEVALUADAS',weight=ft.FontWeight.BOLD,color=ft.Colors.WHITE,size=11,text_align=ft.TextAlign.CENTER),numeric=True),
            ft.DataColumn(ft.Text('HORAS\nPROMEDIO',weight=ft.FontWeight.BOLD,color=ft.Colors.WHITE,size=11,text_align=ft.TextAlign.CENTER),numeric=True),
            ft.DataColumn(ft.Text('COSTO / HORA\nPROMEDIO',weight=ft.FontWeight.BOLD,color=ft.Colors.WHITE,size=11,text_align=ft.TextAlign.CENTER),numeric=True),
            ft.DataColumn(ft.Text('VARIACIÓN RESPECTO\nAL DESG. REG.',weight=ft.FontWeight.BOLD,color=ft.Colors.WHITE,size=11,text_align=ft.TextAlign.CENTER),numeric=True),
        ]
        baja_table_rows=[]

        def stat(items):
            hv=[x['hours'] for x in items if x['hours'] is not None]
            cv=[x['cph'] for x in items if x['cph'] is not None]
            return len(items),(sum(hv)/len(hv) if hv else None),(sum(cv)/len(cv) if cv else None)

        for key,items in sorted(baja_groups.items(), key=lambda kv: (kv[0][0], kv[0][1], kv[0][2], kv[0][3])):
            size,brand,design,tra=key
            regular=[x for x in items if x['category']=='Desgaste Regular']
            reenc=[x for x in items if x['category']=='Reencauche']
            cuts=[x for x in items if x['category']=='Cortes']
            reg_cph=stat(regular)[2] if regular else None
            concepts=[('Desgaste Regular',regular,'#FFF89A'),('Reencauche',reenc,'#FFFFFF'),('Cortes',cuts,'#FFFFFF'),
                      ('Desg. Reg. + Reenc.',regular+reenc,'#FFFFFF'),('Desg. Reg. + Reenc. + Cortes',regular+reenc+cuts,'#FFFFFF')]
            first=True
            for label,subset,bg in concepts:
                # Mantener siempre las cinco filas del análisis. Si una categoría aún no
                # tiene neumáticos (por ejemplo Reencauche), sus indicadores calculados
                # se muestran en cero y la variación queda sin aplicar.
                cnt,avgh,avgc=stat(subset)
                variation=((avgc/reg_cph)-1.0)*100.0 if subset and avgc is not None and reg_cph not in (None,0) and label!='Desgaste Regular' else None
                var_txt='—' if variation is None else f'{variation:+.0f}%'
                var_color=('#D64545' if variation is not None and variation>0 else '#1AAB40' if variation is not None and variation<0 else TEXT_MAIN)
                concept_color = '#1F4E79' if label in ('Desgaste Regular','Reencauche','Cortes') else '#548235'
                value_color = '#548235' if label in ('Desg. Reg. + Reenc.','Desg. Reg. + Reenc. + Cortes') else TEXT_MAIN
                baja_table_rows.append(ft.DataRow(color=bg,cells=[
                    ft.DataCell(ft.Text(size if first else '',weight=ft.FontWeight.BOLD,color=TEXT_MAIN,size=12)),
                    ft.DataCell(ft.Text(brand if first else '',weight=ft.FontWeight.BOLD,color=TEXT_MAIN,size=12)),
                    ft.DataCell(ft.Text((f'{design} | {tra}') if first else '',weight=ft.FontWeight.BOLD,color=TEXT_MAIN,size=12)),
                    ft.DataCell(ft.Text(label,weight=ft.FontWeight.BOLD if label in ('Desgaste Regular','Desg. Reg. + Reenc.','Desg. Reg. + Reenc. + Cortes') else ft.FontWeight.NORMAL,color=concept_color,size=12)),
                    ft.DataCell(ft.Text(str(cnt),color=value_color,weight=ft.FontWeight.BOLD if label.startswith('Desg. Reg. +') else ft.FontWeight.NORMAL,size=12)),
                    ft.DataCell(ft.Text('0 h' if avgh is None else f'{avgh:,.0f} h',color=value_color,weight=ft.FontWeight.BOLD if label.startswith('Desg. Reg. +') else ft.FontWeight.NORMAL,size=12)),
                    ft.DataCell(ft.Text('$0.00/h' if avgc is None else f'${avgc:.2f}/h',color=value_color,weight=ft.FontWeight.BOLD if label.startswith('Desg. Reg. +') else ft.FontWeight.NORMAL,size=12)),
                    ft.DataCell(ft.Text(var_txt,weight=ft.FontWeight.BOLD,color=var_color,size=12)),
                ]))
                first=False

        if not baja_table_rows:
            baja_table_rows=[ft.DataRow(cells=[ft.DataCell(ft.Text('Sin neumáticos dados de baja',color=TEXT_MUTED))]+[ft.DataCell(ft.Text('—')) for _ in range(7)])]

        baja_table=ft.DataTable(columns=baja_columns,rows=baja_table_rows,heading_row_color='#A6A6A6',heading_row_height=58,
                                data_row_min_height=34,data_row_max_height=42,horizontal_margin=10,column_spacing=18,
                                border=ft.Border.all(1,'#8A8A8A'),divider_thickness=0.8)
        section_63=ft.Column([
            ft.Container(
                bgcolor='#C00000',padding=ft.Padding.symmetric(horizontal=14,vertical=10),
                content=ft.Row([
                    ft.Text('6.3  COSTO POR HORA COMPARATIVO',size=17,weight=ft.FontWeight.BOLD,color=ft.Colors.WHITE),
                    ft.Container(expand=True),
                    ft.Text('ANÁLISIS DE NEUMÁTICOS DADOS DE BAJA',size=12,weight=ft.FontWeight.BOLD,color=ft.Colors.WHITE),
                ],vertical_alignment=ft.CrossAxisAlignment.CENTER)
            ),
            ft.Text('Rendimiento y costo por concepto de retiro · fuente: NFU / BAJAS',size=11,color=TEXT_MUTED),
            ft.Container(content=ft.Row([baja_table],scroll=ft.ScrollMode.ALWAYS)),
            ft.Container(bgcolor='#FFFDEB',border=ft.Border.all(1,'#D9D2A8'),padding=9,
                         content=ft.Text('Costo/hora del neumático = costo registrado / horas reales acumuladas. La variación se compara con Desgaste Regular dentro de la misma medida, marca y diseño.',size=10,color=TEXT_MUTED)),
        ],spacing=8)

        # 6.4 · BALANCE GENERAL – UTILIZACIÓN Y PÉRDIDA DE NEUMÁTICOS OTR.
        # Ecuación aprobada por el usuario:
        # Inventario inicial + Ingresos = Salidas + Saldo.
        # Mientras no exista inventario histórico de apertura:
        #   - Apertura repuestos = 0
        #   - Apertura operativas = (Salidas + Saldo) - Ingresos
        # De esta manera el balance queda conciliado por construcción.
        def _avg_vals(a,b):
            vals=[]
            for x in (a,b):
                v=n(x)
                if v is not None:
                    vals.append(v)
            return (sum(vals)/len(vals)) if vals else None

        def _current_residual_value(t):
            new_avg=_avg_vals(t['new_tread_outer'],t['new_tread_inner'])
            if new_avg is None:
                new_avg=n(t['new_tread'])
            cur_avg=_avg_vals(t['tread_outer'],t['tread_inner'])
            retirement=n(t['retirement_tread']) or 0.0
            cost=n(t['cost_usd']) or 0.0
            usable=(new_avg-retirement) if new_avg is not None else 0.0
            rem=max(0.0,(cur_avg-retirement)) if cur_avg is not None else 0.0
            return (cost/usable)*rem if usable>0 else 0.0

        # INGRESOS · fuente 9.7. Una sola inversión por neumático instalado.
        income_new=0.0
        income_reenc=0.0
        seen_inst=set()
        inst_rows=query("""SELECT o.tire_id,t.cost_usd,t.tire_condition
                           FROM occurrences o JOIN tires t ON t.id=o.tire_id
                           WHERE UPPER(TRIM(o.event_code))='INST' AND o.operation_id=? AND t.operation_id=? ORDER BY o.id""",(op_id,op_id))
        for r in inst_rows:
            tid=int(r['tire_id'])
            if tid in seen_inst:
                continue
            seen_inst.add(tid)
            cost=n(r['cost_usd']) or 0.0
            if is_reencauchada(r['tire_condition']):
                income_reenc += cost
            else:
                income_new += cost

        # No existe aún un indicador/campo específico para identificar neumáticos
        # que llegaron montados con un equipo al ingresar a la operación.
        income_equipment_arrival=0.0

        # SALIDAS · fuente 9.5: consolidar los mismos conceptos económicos.
        out_new=0.0
        out_reenc=0.0
        out_cut=0.0
        out_noopt=0.0
        out_mmree_value=0.0
        for eq in eq_rows:
            occs=query("""
                SELECT o.id,o.tire_id,o.event_code,o.event_date,o.meter,
                       o.tread_outer,o.tread_inner,o.reason,
                       t.cost_usd,t.new_tread,t.new_tread_outer,t.new_tread_inner,
                       t.retirement_tread,t.tire_condition
                FROM occurrences o JOIN tires t ON t.id=o.tire_id
                WHERE o.equipment_id=? AND o.operation_id=? AND t.operation_id=? ORDER BY o.id
            """,(int(eq['id']),op_id,op_id))
            by_tire={}
            for r in occs:
                by_tire.setdefault(int(r['tire_id']),[]).append(r)
            for events in by_tire.values():
                first=events[0]
                original=tire_original_min(first)
                cost=n(first['cost_usd'])
                if not original or not cost or original<=0:
                    continue
                pxmm=cost/original
                start_tread=None
                for ev in events:
                    tm=tread_min(ev)
                    if str(ev['event_code'] or '').upper().strip()=='INST' and tm is not None:
                        start_tread=tm
                        break
                if start_tread is None:
                    for ev in events:
                        tm=tread_min(ev)
                        if tm is not None:
                            start_tread=tm
                            break
                end_tread=None
                for ev in reversed(events):
                    tm=tread_min(ev)
                    if tm is not None:
                        end_tread=tm
                        break
                used_mm=max(0.0,start_tread-end_tread) if start_tread is not None and end_tread is not None else 0.0
                used_value=used_mm*pxmm
                if is_reencauchada(first['tire_condition']):
                    out_reenc += used_value
                    # En 9.5 MM$REE representa el consumo asociado al reencauche.
                    out_mmree_value += used_value
                else:
                    out_new += used_value
                baja=next((ev for ev in reversed(events) if str(ev['event_code'] or '').upper().strip()=='BAJA'),None)
                if baja is not None:
                    bt=tread_min(baja)
                    if bt is None:
                        bt=end_tread
                    if bt is not None:
                        retirement=n(first['retirement_tread']) or 0.0
                        if is_cut_reason(baja['reason']):
                            out_cut += max(0.0,bt)*pxmm
                        else:
                            out_noopt += max(0.0,bt-retirement)*pxmm

        # SALIDAS · fuente 9.2: valor del remanente no utilizado en bajas RTEQ.
        out_rteq=0.0
        rteq_rows=query("""SELECT t.*,o.tread_outer baja_outer,o.tread_inner baja_inner,o.reason
                           FROM occurrences o JOIN tires t ON t.id=o.tire_id
                           WHERE UPPER(TRIM(o.event_code))='BAJA' AND o.operation_id=? AND t.operation_id=?
                           ORDER BY o.id""",(op_id,op_id))
        latest_rteq={}
        for r in rteq_rows:
            reason=str(r['reason'] or '').upper().strip()
            if reason=='RTEQ' or 'RETIRO' in reason and 'EQUIP' in reason:
                latest_rteq[int(r['id']) if 'id' in r.keys() else len(latest_rteq)]=r
        for r in latest_rteq.values():
            new_avg=_avg_vals(r['new_tread_outer'],r['new_tread_inner'])
            if new_avg is None:
                new_avg=n(r['new_tread'])
            cur_avg=_avg_vals(r['baja_outer'],r['baja_inner'])
            retirement=n(r['retirement_tread']) or 0.0
            cost=n(r['cost_usd']) or 0.0
            usable=(new_avg-retirement) if new_avg is not None else 0.0
            consumed=max(0.0,(new_avg-cur_avg)) if new_avg is not None and cur_avg is not None else 0.0
            consumed_pct=min(1.0,consumed/usable) if usable>0 else 0.0
            cost_acum=cost*consumed_pct
            out_rteq += max(0.0,cost-cost_acum)

        # SALDO · fuentes 9.8 y 9.9: valor residual actual por RTD útil.
        close_oper=0.0
        for t in query("""SELECT * FROM tires
                          WHERE UPPER(TRIM(status))='SERVICIO' AND equipment_id IS NOT NULL AND operation_id=?""",(op_id,)):
            close_oper += _current_residual_value(t)
        close_spare=0.0
        for t in query("""SELECT * FROM tires
                          WHERE UPPER(TRIM(status)) IN ('STAND-BY','STAND BY','STANDBY') AND operation_id=?""",(op_id,)):
            close_spare += _current_residual_value(t)

        total_income=income_new+income_equipment_arrival+income_reenc
        total_out=out_new+out_reenc+out_cut+out_noopt+out_mmree_value+out_rteq
        total_close=close_oper+close_spare
        opening_spare=0.0
        opening_oper=(total_out+total_close)-total_income
        total_opening=opening_oper+opening_spare
        total_available=total_opening+total_income
        balance_final=(total_opening+total_income)-(total_out+total_close)

        def _pct(v):
            return (100.0*v/total_available) if abs(total_available)>1e-9 else 0.0

        balance_rows=[
            ('INVENTARIO INICIAL','Inventario Apertura Llantas Operativas en Equipos',opening_oper,_pct(opening_oper)),
            ('','Inventario Apertura Llantas de Repuesto',opening_spare,_pct(opening_spare)),
            ('INGRESOS','Llantas Nuevas Instaladas (compras)',income_new,_pct(income_new)),
            ('','Equipos que ingresaron con llantas a la operación',income_equipment_arrival,_pct(income_equipment_arrival)),
            ('','Llantas Reencauchadas Instaladas (compras)',income_reenc,_pct(income_reenc)),
            ('SALIDAS','Neumáticos Originales - Retiros x Desgaste Regular',out_new,_pct(out_new)),
            ('','Neumáticos Reencauchados - Retiros x Desgaste',out_reenc,_pct(out_reenc)),
            ('','Neumáticos Retirados x Cortes',out_cut,_pct(out_cut)),
            ('','Remanente No utilizado (Retiro de llanta x Seguridad)',out_noopt,_pct(out_noopt)),
            ('','Remanente utilizado para el Reencauche',out_mmree_value,_pct(out_mmree_value)),
            ('','Remanente No utilizado (Retiro de llanta x Equipo de baja)',out_rteq,_pct(out_rteq)),
            ('SALDO','Inventario Cierre Llantas Operativas en Equipos',close_oper,_pct(close_oper)),
            ('','Inventario Cierre Llantas de Repuesto',close_spare,_pct(close_spare)),
        ]

        # Presentación 6.4: réplica del balance histórico tipo Excel aprobado.
        # Se muestran subtotales numéricos por bloque sin etiquetas adicionales.
        def balance_cell(text,width,bold=False,align=ft.TextAlign.LEFT,color=TEXT_MAIN,size=11):
            return ft.Container(
                width=width,
                padding=ft.Padding(left=4,top=4,right=4,bottom=4),
                content=ft.Text(
                    str(text),size=size,
                    weight=ft.FontWeight.BOLD if bold else ft.FontWeight.NORMAL,
                    color=color,text_align=align,no_wrap=True
                )
            )

        month_names={1:'ENE',2:'FEB',3:'MAR',4:'ABR',5:'MAY',6:'JUN',7:'JUL',8:'AGO',9:'SEP',10:'OCT',11:'NOV',12:'DIC'}
        _today=dt.datetime.now()
        period_label=f"{month_names.get(_today.month,'')}-{_today.year} $"

        # Anchos fijos para conservar simetría y alineación del formato original.
        W_BLOCK=112
        W_LABEL=515
        W_VALUE=132
        W_PCT=68

        def balance_data_row(block,label,value,pct=None,bold_value=False):
            pct_text='' if pct is None else f'{pct:.1f}%'
            return ft.Row([
                balance_cell(block,W_BLOCK),
                balance_cell(label,W_LABEL),
                balance_cell(f'{value:,.2f}',W_VALUE,bold=bold_value,align=ft.TextAlign.RIGHT),
                balance_cell(pct_text,W_PCT,bold=(pct is not None),align=ft.TextAlign.RIGHT),
            ],spacing=0,vertical_alignment=ft.CrossAxisAlignment.CENTER)

        def balance_subtotal(value,pct_text=''):
            return ft.Row([
                balance_cell('',W_BLOCK),
                balance_cell('',W_LABEL),
                ft.Container(
                    width=W_VALUE,
                    border=ft.Border(top=ft.BorderSide(1.2,'#1F1F1F')),
                    padding=ft.Padding(left=4,top=4,right=4,bottom=5),
                    content=ft.Text(f'{value:,.2f}',size=11,weight=ft.FontWeight.BOLD,text_align=ft.TextAlign.RIGHT,color=TEXT_MAIN,no_wrap=True),
                ),
                balance_cell(pct_text,W_PCT,bold=bool(pct_text),align=ft.TextAlign.RIGHT),
            ],spacing=0)

        # El % conserva la lógica del formato: participación respecto al total disponible
        # (Inventario inicial + Ingresos).
        bal_controls=[
            # Encabezado de columnas monetarias.
            ft.Row([
                balance_cell('',W_BLOCK),
                balance_cell('',W_LABEL),
                balance_cell(period_label,W_VALUE,align=ft.TextAlign.RIGHT,size=11),
                balance_cell('%',W_PCT,align=ft.TextAlign.RIGHT,size=11),
            ],spacing=0),
            ft.Container(height=8),

            # INVENTARIO INICIAL (sin porcentaje individual, como el formato fuente).
            balance_data_row('INVENTARIO','Inventario Apertura Llantas Operativas en Equipos',opening_oper,None),
            balance_data_row('INICIAL','Inventario Apertura Llantas de Repuesto',opening_spare,None),
            ft.Container(height=12),

            # INGRESOS.
            balance_data_row('INGRESOS','Llantas Nuevas Instaladas  (compras)',income_new,_pct(income_new)),
            balance_data_row('','Equipos que ingresaron con llantas a la operación',income_equipment_arrival,_pct(income_equipment_arrival)),
            balance_data_row('','Llantas Reencauchadas Instaladas (compras)',income_reenc,_pct(income_reenc)),
            # Subtotal INVENTARIO + INGRESOS, sin texto adicional.
            balance_subtotal(total_available,'100%'),
            ft.Container(height=14),
            ft.Divider(height=1,thickness=1.2,color='#1F1F1F'),
            ft.Container(height=12),

            # SALIDAS.
            balance_data_row('SALIDAS','Neumáticos Originales - Retiros x Desgaste Regular',out_new,_pct(out_new)),
            balance_data_row('','Neumáticos Reencauchados - Retiros x Desgaste',out_reenc,_pct(out_reenc)),
            balance_data_row('','Neumáticos Retirados x Cortes',out_cut,_pct(out_cut)),
            balance_data_row('','Remanente No utilizado (Retiro de llanta x Seguridad)',out_noopt,_pct(out_noopt)),
            balance_data_row('','Remanente utilizado para el Reencauche',out_mmree_value,_pct(out_mmree_value)),
            balance_data_row('','Remanente No utilizado (Retiro de llanta x Equipo de baja)',out_rteq,_pct(out_rteq)),
            # Subtotal SALIDAS, sin etiqueta.
            balance_subtotal(total_out),
            ft.Container(height=14),
            ft.Divider(height=1,thickness=1.2,color='#1F1F1F'),
            ft.Container(height=12),

            # SALDO.
            balance_data_row('SALDO','Inventario Cierre Operativas en Equipos',close_oper,_pct(close_oper)),
            balance_data_row('','Inventario Cierre Llantas de Repuesto',close_spare,_pct(close_spare)),
            # Subtotal SALDO, sin etiqueta.
            balance_subtotal(total_close),
            ft.Container(height=12),
            ft.Divider(height=1,thickness=1.2,color='#1F1F1F'),
        ]

        section_64=ft.Column([
            # Título rojo, centrado, como el formato Excel.
            ft.Container(
                bgcolor='#F00000',height=48,alignment=ft.Alignment.CENTER,
                content=ft.Text('BALANCE GENERAL',size=20,weight=ft.FontWeight.BOLD,color=ft.Colors.WHITE,text_align=ft.TextAlign.CENTER)
            ),
            # Subtítulo gris independiente.
            ft.Container(
                bgcolor='#D9D9D9',height=45,alignment=ft.Alignment.CENTER,
                content=ft.Text('UTILIZACIÓN Y PERDIDA DE NEUMÁTICOS OTR',size=15,weight=ft.FontWeight.BOLD,color='#000000',text_align=ft.TextAlign.CENTER)
            ),
            ft.Container(height=12),
            ft.Container(
                bgcolor='#FFFFFF',
                padding=ft.Padding(left=10,top=4,right=10,bottom=4),
                content=ft.Column(bal_controls,spacing=0)
            ),
        ],spacing=0)

        note=ft.Container(
            bgcolor='#F7FAFC',
            border=ft.Border.all(1,'#DCE4EC'),
            border_radius=8,
            padding=12,
            content=ft.Text(
                'Criterio Hr-Rod: diferencia entre el mayor y el menor horómetro registrado para cada equipo. '
                'LL/NEW y LL/REE valorizan los mm consumidos; $CORTE y $NO OPT muestran pérdida de remanente.',
                size=11,color=TEXT_MUTED
            )
        )

        # Presentación de accesos 6.x: mismo estándar visual aprobado para el módulo 9.
        section_61=ft.Column([
            ft.Text('6.1 UTILIZACIÓN Y PÉRDIDA POR EQUIPO',size=17,weight=ft.FontWeight.BOLD,color=TEXT_MAIN),
            ft.Text('Resumen económico y horas rodadas por equipo',size=12,color=TEXT_MUTED),
            ft.Row([
                ft.Container(width=520,content=ft.Row([table],scroll=ft.ScrollMode.ALWAYS)),
                ft.Container(expand=True,content=chart_panel),
            ],spacing=14,vertical_alignment=ft.CrossAxisAlignment.START,scroll=ft.ScrollMode.AUTO),
            note,
        ],spacing=12)

        inventory_meta={
            '61':{
                'num':'6.1','icon':ft.Icons.ACCOUNT_BALANCE_WALLET_OUTLINED,'accent':'#1976D2','soft':'#EEF5FF',
                'title':'UTILIZACIÓN Y PÉRDIDA\nPOR EQUIPO',
                'desc':'Resumen económico de utilización, pérdidas y horas rodadas por equipo.',
                'body':section_61,
            },
            '62':{
                'num':'6.2','icon':ft.Icons.SPEED_OUTLINED,'accent':'#138A3D','soft':'#EEFAF2',
                'title':'EQUIPOS\nCOSTO POR HORA',
                'desc':'Comparativo del costo por hora acumulado para cada equipo.',
                'body':section_62,
            },
            '63':{
                'num':'6.3','icon':ft.Icons.ANALYTICS_OUTLINED,'accent':'#EF6C00','soft':'#FFF5EA',
                'title':'ANÁLISIS DE NEUMÁTICOS\nDE BAJA',
                'desc':'Rendimiento y costo comparativo según concepto de retiro.',
                'body':section_63,
            },
            '64':{
                'num':'6.4','icon':ft.Icons.BALANCE_OUTLINED,'accent':'#C62828','soft':'#FFF0F0',
                'title':'BALANCE GENERAL',
                'desc':'Utilización y pérdida de neumáticos OTR, ingresos, salidas y saldo.',
                'body':section_64,
            },
        }

        inventory_cards=ft.Column(spacing=12)
        inventory_detail=ft.Column(spacing=12,visible=False)

        def close_inventory_section():
            inventory_detail.visible=False
            inventory_cards.visible=True
            page.update()

        def show_inventory_section(key):
            m=inventory_meta[key]
            inventory_detail.controls=[
                ft.Row([
                    ft.OutlinedButton('VOLVER A ANÁLISIS DE OPERACIÓN',icon=ft.Icons.ARROW_BACK,on_click=lambda e: close_inventory_section()),
                    ft.Container(expand=True),
                    ft.Container(bgcolor=m['accent'],border_radius=8,padding=ft.Padding(10,5,10,5),
                                 content=ft.Text(m['num'],color='#FFFFFF',weight=ft.FontWeight.BOLD,size=13)),
                ]),
                card(ft.Column([
                    ft.Row([
                        ft.Container(width=46,height=46,border_radius=12,bgcolor=m['soft'],alignment=ft.Alignment.CENTER,
                                     content=ft.Icon(m['icon'],color=m['accent'],size=26)),
                        ft.Column([
                            ft.Text(m['title'].replace('\n',' '),size=17,weight=ft.FontWeight.BOLD,color=TEXT_MAIN),
                            ft.Text(m['desc'],size=10.5,color=TEXT_MUTED),
                        ],spacing=2,expand=True),
                    ],spacing=12),
                    m['body'],
                ],spacing=12),padding=16),
            ]
            inventory_cards.visible=False
            inventory_detail.visible=True
            page.update()

        def inventory_access_card(key):
            m=inventory_meta[key]
            return ft.Container(
                width=292,height=260,bgcolor=m['soft'],
                border=ft.Border.all(1,m['accent']+'55'),border_radius=14,padding=18,
                on_click=lambda e,k=key: show_inventory_section(k),ink=True,
                content=ft.Column([
                    ft.Row([
                        ft.Container(bgcolor=m['accent'],border_radius=9,padding=ft.Padding(11,6,11,6),
                                     content=ft.Text(m['num'],size=16,weight=ft.FontWeight.BOLD,color='#FFFFFF')),
                        ft.Container(expand=True),
                        ft.Container(width=54,height=54,border_radius=14,bgcolor='#FFFFFF',alignment=ft.Alignment.CENTER,
                                     content=ft.Icon(m['icon'],size=31,color=m['accent'])),
                    ]),
                    ft.Text(m['title'],size=16,weight=ft.FontWeight.BOLD,color=TEXT_MAIN,text_align=ft.TextAlign.CENTER),
                    ft.Text(m['desc'],size=11,color=TEXT_MUTED,text_align=ft.TextAlign.CENTER),
                    ft.Container(height=4),
                    ft.Container(bgcolor=m['accent'],border_radius=9,padding=10,alignment=ft.Alignment.CENTER,
                                 content=ft.Row([
                                     ft.Icon(ft.Icons.BAR_CHART,color='#FFFFFF',size=20),
                                     ft.Text('VER REPORTE',color='#FFFFFF',weight=ft.FontWeight.BOLD,size=13),
                                     ft.Icon(ft.Icons.CHEVRON_RIGHT,color='#FFFFFF',size=20),
                                 ],alignment=ft.MainAxisAlignment.CENTER,spacing=8)),
                ],spacing=12,horizontal_alignment=ft.CrossAxisAlignment.CENTER)
            )

        inventory_cards.controls=[
            card(ft.Column([
                ft.Row([
                    ft.Container(width=48,height=48,border_radius=24,bgcolor='#EAF2FF',alignment=ft.Alignment.CENTER,
                                 content=ft.Icon(ft.Icons.WAREHOUSE_OUTLINED,color=NAV_ACCENT,size=27)),
                    ft.Column([
                        ft.Text('ANÁLISIS DE OPERACIÓN',size=18,weight=ft.FontWeight.BOLD,color=TEXT_MAIN),
                        ft.Text('Seleccione un reporte para visualizar el detalle.',size=11,color=TEXT_MUTED),
                    ],spacing=2),
                ],spacing=12),
                ft.Row([inventory_access_card('61'),inventory_access_card('62'),inventory_access_card('63'),inventory_access_card('64')],spacing=12),
            ],spacing=16),padding=16),
        ]

        content.content=ft.Column([
            page_title('6. ANÁLISIS DE OPERACIÓN','Indicadores de utilización, costos, consumos y remanentes'),
            inventory_cards,
            inventory_detail,
        ],scroll=ft.ScrollMode.AUTO,spacing=16)
        page.update()

    def administration_view():
        """8. Administración - menú visual para equipos y neumáticos."""
        admin_meta={
            '81':{
                'num':'8.1','icon':ft.Icons.PRECISION_MANUFACTURING_OUTLINED,'accent':'#1565C0','soft':'#EEF5FF',
                'title':'EQUIPOS',
                'desc':'Administración de equipos: registro, consulta y actualización de la flota.',
                'action':equipment_view,
            },
            '82':{
                'num':'8.2','icon':ft.Icons.TIRE_REPAIR,'accent':'#138A3D','soft':'#EEFAF2',
                'title':'NEUMÁTICOS',
                'desc':'Registro maestro de neumáticos: datos técnicos, costos y parámetros de control.',
                'action':lambda: tires_view(),
            },
        }

        def admin_access_card(key):
            m=admin_meta[key]
            card_click = lambda e,k=key: admin_meta[k]['action']()
            return ft.Container(
                width=292,height=260,bgcolor=m['soft'],
                border=ft.Border.all(1,m['accent']+'55'),border_radius=14,padding=18,
                on_click=card_click,ink=True,
                content=ft.Column([
                    ft.Row([
                        ft.Container(bgcolor=m['accent'],border_radius=9,padding=ft.Padding(11,6,11,6),
                                     content=ft.Text(m['num'],size=16,weight=ft.FontWeight.BOLD,color='#FFFFFF')),
                        ft.Container(expand=True),
                        ft.Container(width=54,height=54,border_radius=14,bgcolor='#FFFFFF',alignment=ft.Alignment.CENTER,
                                     content=ft.Icon(m['icon'],size=31,color=m['accent'])),
                    ]),
                    ft.Text(m['title'],size=16,weight=ft.FontWeight.BOLD,color=TEXT_MAIN,text_align=ft.TextAlign.CENTER),
                    ft.Text(m['desc'],size=11,color=TEXT_MUTED,text_align=ft.TextAlign.CENTER),
                    ft.Container(height=4),
                    ft.Container(
                        bgcolor=m['accent'],border_radius=9,padding=10,alignment=ft.Alignment.CENTER,
                        content=ft.Row([
                            ft.Icon(ft.Icons.LOGIN,color='#FFFFFF',size=20),
                            ft.Text('INGRESAR',color='#FFFFFF',weight=ft.FontWeight.BOLD,size=13),
                            ft.Icon(ft.Icons.CHEVRON_RIGHT,color='#FFFFFF',size=20),
                        ],alignment=ft.MainAxisAlignment.CENTER,spacing=8)
                    ),
                ],spacing=12,horizontal_alignment=ft.CrossAxisAlignment.CENTER)
            )

        content.content=ft.Column([
            page_title('8. ADMINISTRACIÓN','Gestión maestra de equipos y neumáticos'),
            card(ft.Column([
                ft.Row([
                    ft.Container(width=48,height=48,border_radius=24,bgcolor='#EAF2FF',alignment=ft.Alignment.CENTER,
                                 content=ft.Icon(ft.Icons.ADMIN_PANEL_SETTINGS_OUTLINED,color=NAV_ACCENT,size=27)),
                    ft.Column([
                        ft.Text('ADMINISTRACIÓN',size=18,weight=ft.FontWeight.BOLD,color=TEXT_MAIN),
                        ft.Text('Seleccione una opción para ingresar al registro maestro.',size=11,color=TEXT_MUTED),
                    ],spacing=2),
                ],spacing=12),
                ft.Row([
                    admin_access_card('81'),admin_access_card('82'),
                    ft.Container(width=292,height=260),ft.Container(width=292,height=260)
                ],spacing=12),
            ],spacing=16),padding=16),
        ],scroll=ft.ScrollMode.AUTO,spacing=16)
        page.update()

    def placeholder(title,desc):
        content.content=ft.Column([
            page_title(title,desc),
            card(ft.Column([
                ft.Icon(ft.Icons.CONSTRUCTION_OUTLINED,size=42,color=NAV_ACCENT),
                ft.Text('Módulo preparado para la siguiente etapa.',size=16,weight=ft.FontWeight.BOLD,color=TEXT_MAIN),
                ft.Text('La navegación y estructura ya están integradas en la versión Web.',color=TEXT_MUTED)
            ],horizontal_alignment=ft.CrossAxisAlignment.CENTER),width=520)
        ],scroll=ft.ScrollMode.AUTO)
        page.update()

    def select(idx):
        # El manejo de ESC pertenece solo a Neumáticos en servicio.
        if idx != 2:
            page.on_keyboard_event = None
        if idx==0: dashboard()
        elif idx==1: movement_view()
        elif idx==2: service_menu_view()
        elif idx==3: maintenance_menu_view()
        elif idx==4: standby_view()
        elif idx==5: nfu_view()
        elif idx==6: inventory_consumption_view()
        elif idx==7: reports_view()
        elif idx==8: administration_view()

    def build_shell():
        nonlocal nav
        user=session['user']
        nav=ft.NavigationRail(
            selected_index=0,
            extended=True,
            min_extended_width=290,
            bgcolor=NAV_BG,
            indicator_color='#234A73',
            leading=ft.Container(
                padding=16,
                content=ft.Column([
                    ft.Row([ft.Icon(ft.Icons.TIRE_REPAIR,color=ft.Colors.WHITE,size=28),ft.Text('MegaSoftire',size=22,weight=ft.FontWeight.BOLD,color=ft.Colors.WHITE)]),
                    ft.Text('Web 2026',size=11,color='#B9C9D9')
                ],spacing=4)
            ),
            destinations=[
                ft.NavigationRailDestination(icon=ft.Icon(ft.Icons.DASHBOARD_OUTLINED,color='#D7E3EF'),selected_icon=ft.Icon(ft.Icons.DASHBOARD,color=ft.Colors.WHITE),label=ft.Text('Panel principal',color=ft.Colors.WHITE))
            ] + [
                ft.NavigationRailDestination(icon=ft.Icon(icon,color='#D7E3EF'),selected_icon=ft.Icon(icon,color=ft.Colors.WHITE),label=ft.Text(f'{i+1}. {name}',color=ft.Colors.WHITE))
                for i,(name,icon) in enumerate(MODULES)
            ]
        )
        nav.on_change=lambda e: select(e.control.selected_index)

        active_op_label = ft.Text(
            f"TALLER ACTIVO: {active_operation_row()['name']}",
            size=12, weight=ft.FontWeight.BOLD, color=NAV_ACCENT
        )

        userbar=ft.Container(
            height=62,
            bgcolor=ft.Colors.WHITE,
            border=ft.Border(bottom=ft.BorderSide(1,'#E4EAF0')),
            padding=ft.Padding.symmetric(horizontal=20),
            content=ft.Row([
                ft.Text('Gestión integral de neumáticos OTR',size=13,color=TEXT_MUTED),
                ft.Container(expand=True),
                active_op_label,
                ft.Container(width=18),
                ft.Icon(ft.Icons.ACCOUNT_CIRCLE_OUTLINED,color=NAV_ACCENT),
                ft.Column([ft.Text(user['full_name'] or user['username'],size=12,weight=ft.FontWeight.BOLD,color=TEXT_MAIN),ft.Text('ADMINISTRADOR GENERAL' if str(user['role']).upper()=='ADMIN' else user['role'],size=10,color=TEXT_MUTED)],spacing=0),
                ft.TextButton('Salir',icon=ft.Icons.LOGOUT,on_click=lambda e: logout())
            ])
        )
        shell=ft.Row([
            nav,
            ft.VerticalDivider(width=1,color='#D7E0E8'),
            ft.Column([userbar,content],expand=True,spacing=0)
        ],expand=True,spacing=0)
        app_host.content=shell
        dashboard()
        adapt()

    def logout():
        session['user']=None
        show_login()

    def show_login():
        # Portada Software-ETM limpia, sin imágenes.
        # Se conserva íntegramente la autenticación existente.
        expected_user = 'admin'
        expected_password_hash = '04445e6487736590d1ef50186b414e737e0164683cbbec64e00e73c000fd3bef'

        user_field = ft.TextField(
            value='', width=390, height=56, autofocus=True,
            text_size=16, color='#172B3A',
            border_color='#B7C5D0', focused_border_color='#FF6338',
            cursor_color='#FF6338',
            hint_text='Usuario',
            hint_style=ft.TextStyle(color='#7B8B97'),
            bgcolor='#FFFFFF',
            content_padding=ft.Padding(left=46, top=8, right=14, bottom=8),
            prefix_icon=ft.Icons.PERSON_OUTLINE,
        )

        password_field = ft.TextField(
            value='', width=390, height=56, password=True, can_reveal_password=False,
            text_size=16, color='#172B3A',
            border_color='#B7C5D0', focused_border_color='#FF6338',
            cursor_color='#FF6338',
            hint_text='Contraseña',
            hint_style=ft.TextStyle(color='#7B8B97'),
            bgcolor='#FFFFFF',
            content_padding=ft.Padding(left=46, top=8, right=48, bottom=8),
            prefix_icon=ft.Icons.LOCK_OUTLINE,
        )

        # Ojo: permite mostrar/ocultar la contraseña.
        password_eye = ft.IconButton(
            icon=ft.Icons.VISIBILITY_OFF,
            icon_color='#617381',
            tooltip='Mostrar / ocultar contraseña'
        )

        def toggle_password_visibility(e=None):
            password_field.password = not password_field.password
            password_eye.icon = (
                ft.Icons.VISIBILITY_OFF
                if password_field.password
                else ft.Icons.VISIBILITY
            )
            page.update()

        password_eye.on_click = toggle_password_visibility

        password_control = ft.Stack(
            width=390, height=56,
            controls=[
                password_field,
                ft.Container(
                    content=password_eye,
                    right=2, top=4, width=46, height=46,
                    alignment=ft.Alignment.CENTER
                ),
            ],
        )

        login_message = ft.Text(
            '', size=12, color='#D84315',
            text_align=ft.TextAlign.CENTER
        )

        def enter_system(e=None):
            username = (user_field.value or '').strip()
            password = password_field.value or ''
            db_user = authenticate(username, password)
            legacy_admin_ok = (
                username == expected_user and
                hashlib.sha256(password.encode('utf-8')).hexdigest() == expected_password_hash
            )

            if db_user:
                urow = query(
                    'SELECT id,username,full_name,role,operation_id FROM users WHERE id=?',
                    (db_user['id'],)
                )[0]
                session['user'] = dict(urow)
            elif legacy_admin_ok:
                session['user'] = {
                    'username': 'admin',
                    'full_name': 'Administrador',
                    'role': 'ADMIN',
                    'operation_id': 1
                }
            else:
                login_message.value = 'USUARIO O CLAVE INCORRECTOS'
                password_field.value = ''
                page.update()
                return

            role_now = str(session['user'].get('role') or '').upper()
            if role_now not in ('ADMIN', 'ADMINISTRADOR GENERAL', 'SUPERADMIN'):
                assigned = session['user'].get('operation_id')
                if not assigned:
                    session['user'] = None
                    login_message.value = 'USUARIO SIN TALLER ASIGNADO'
                    page.update()
                    return
                operation_state['id'] = int(assigned)

            try:
                build_shell()
                page.update()
            except Exception as ex:
                session['user'] = None
                login_message.value = f'ERROR AL INGRESAR: {str(ex)[:120]}'
                page.update()

        user_field.on_submit = lambda e: password_field.focus()
        password_field.on_submit = enter_system

        login_panel = ft.Container(
            width=500,
            bgcolor='#FFFFFF',
            border_radius=18,
            border=ft.Border(
                left=ft.BorderSide(1, '#D7E0E7'),
                top=ft.BorderSide(1, '#D7E0E7'),
                right=ft.BorderSide(1, '#D7E0E7'),
                bottom=ft.BorderSide(1, '#D7E0E7'),
            ),
            padding=44,
            shadow=ft.BoxShadow(
                blur_radius=24,
                color='#26000000',
                offset=ft.Offset(0, 8)
            ),
            content=ft.Column(
                [
                    ft.Text(
                        'Software-',
                        size=40,
                        weight=ft.FontWeight.BOLD,
                        color='#193247'
                    ),
                    ft.Text(
                        'ETM',
                        size=50,
                        weight=ft.FontWeight.BOLD,
                        color='#FF6338'
                    ),
                    ft.Text(
                        'Enterprise Tire Management',
                        size=17,
                        italic=True,
                        color='#607684'
                    ),
                    ft.Container(height=10),
                    ft.Divider(color='#D8E0E6', height=1),
                    ft.Text(
                        'GESTIÓN INTEGRAL DE NEUMÁTICOS OTR',
                        size=11,
                        weight=ft.FontWeight.BOLD,
                        color='#71838F'
                    ),
                    ft.Container(height=18),
                    user_field,
                    password_control,
                    login_message,
                    ft.Container(height=4),
                    ft.Button(
                        'INICIAR SESIÓN',
                        on_click=enter_system,
                        width=390,
                        height=54,
                        bgcolor='#FF6338',
                        color=ft.Colors.WHITE,
                        style=ft.ButtonStyle(
                            shape=ft.RoundedRectangleBorder(radius=7),
                            text_style=ft.TextStyle(
                                size=16,
                                weight=ft.FontWeight.BOLD
                            )
                        ),
                    ),
                    ft.Container(height=2),
                    ft.TextButton(
                        'Recuperar Contraseña',
                        style=ft.ButtonStyle(color='#607684')
                    ),
                ],
                horizontal_alignment=ft.CrossAxisAlignment.CENTER,
                spacing=10
            ),
        )

        # Fondo neutro: no depende de archivos externos ni de assets.
        background = ft.Container(
            expand=True,
            bgcolor='#EEF2F5',
        )

        app_host.content = ft.Stack(
            expand=True,
            controls=[
                background,
                ft.Container(
                    expand=True,
                    alignment=ft.Alignment.CENTER,
                    content=login_panel
                ),
            ],
        )
        page.update()

    def adapt(e=None):
        if not nav: return
        w=page.width or 1280
        nav.extended = w >= 1120
        content.padding = 14 if w < 800 else 24
        page.update()

    page.on_resized=adapt
    page.add(app_host)
    show_login()


if __name__=='__main__':
    if os.environ.get('PORT'):
        os.environ.setdefault('FLET_SERVER_PORT', os.environ['PORT'])
        os.environ.setdefault('FLET_SERVER_IP', '0.0.0.0')
        os.environ.setdefault('FLET_FORCE_WEB_SERVER', 'true')
    ft.run(main, assets_dir='assets')
